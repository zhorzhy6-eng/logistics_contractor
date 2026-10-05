#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сборка эталонного пустого бланка Хавалов (ЭТАП 3.1.E.A.1).

Создаёт ОДИН файл:

    templates/shablon_havaly.xlsx — пустой бланк заявки
    «ЗАЯВКА на перевозку автомобилей», один лист TDSheet, 32 колонки,
    10 пустых строк данных (максимум машин) и нижний блок сторон.

Запуск:  python tools/make_havaly_template.py

═══════════════════════════════════════════════════════════════════════
РЕЗУЛЬТАТ РАЗБОРА ОБРАЗЦА (ЭТАП 3.1.E.A.0)
═══════════════════════════════════════════════════════════════════════
Источник структуры — templates/Хавалы.xlsx (присланный заказчиком файл,
копия для репозитория: templates/Хавалы_образец.xlsx, SHA256
106da4ec323d8b850dbdcef2e089043e7de0f889d41f1a466137258c796d1280).
Образец только читается и НЕ изменяется (см. tests/test_havaly_template.py).

Что показал openpyxl (openpyxl==3.1.5):

  * листов РОВНО ОДИН, имя 'TDSheet', лист видимый; скрытых листов,
    строк и колонок нет; закрепления областей, автофильтра, областей
    печати, примечаний, гиперссылок, таблиц, проверок данных нет;
  * ws.max_row = 19, ws.max_column = 32 (A..AF), ws.min_row = 2;
  * объединённые ячейки — ровно две: 'B2:C2' (заголовок) и 'AF7:AF12'
    (ставка, объединённая на 6 машин образца) и больше ничего;
  * определенённых имён нет, формулы отсутствуют — только текст.

Геометрия строк и точные координаты:

    R1    пусто                     (высота 14,25)
    R2    B2:C2 'ЗАЯВКА\\n на перевозку автомобилей'  (высота 23,25)
    R3    пусто                     (высота 23,25)
    R4    пусто                     (высота по умолчанию)
    R5    A5 'Дата заявки:' + B5 (дата, в образце '16.09.2026'); высота 15,0
    R6    ШАПКА ТАБЛИЦЫ, ровно 32 колонки A..AF      (высота 23,25)
    R7-R12  данные: 6 строк образца, все с рамкой    (высота 11,25)
    R13   пусто                     (высота 11,25)
    R14   пусто                     (высота 12,75)
    R15   C15 'Заказчик'   | J15 'Перевозчик'        (высота 12,75)
    R16   C16 'Сюрлогистик' | J16 'ООО "ТЕХНОЛОГИСТИКА"' (высота 15,75)
    R17   пусто                     (высота 11,25)
    R18   пусто                     (высота 11,25)
    R19   C19 'ФИО, подпись, печать' | J19 то же самое   (высота 11,25)

  ВАЖНО про строку начала данных: разделительной пустой строки между
  шапкой (R6) и первой строкой данных НЕТ — данные начинаются сразу
  со строки R7 (у R7 полная рамка, как у R6). Гипотеза ТЗ о «пустой
  строке-разделителе» на этом файле не подтвердилась, расхождения нет.

  В образце 6 строк данных (R7-R12) — по числу машин в присланном файле.
  В эталонном бланке их 10 (R7-R16, максимум машин); лишние строки позже
  обрезает генератор.

Тексты шапки (R6), A..AF — дословно из образца:

    A  Номер Лота                   Q  Фамилия
    B  VIN                          R  Имя
    C  Марка                        S  Отчество
    D  Модель                       T  Номер В/У
    E  Город погрузки               U  Дата выдачи В/У
    F  Пункт погрузки               V  Серия Паспорта
    G  Город доставки               W  Номер Паспорта
    H  Пункт разгрузки              X  Кем выдан паспорт
    I  Дилер                        Y  Дата выдачи паспорта
    J  Код дилера                   Z  Гражданство
    K  Наименование транспортной    AA Дата Рождения
       компании                    AB Прописка
    L  Марка Автовоза               AC Планируемая дата\\n погрузки
    M  Цвет кабины                  AD Время погрузки
    N  Номер автовоза               AE № телефона водителя
    O  Марка прицепа                AF Ставка с НДС
    P  Номер Прицепа

  У AC6 в образце перенос строки внутри текста ('Планируемая дата' + '\\n'
  + пробел + 'погрузки') и включён wrap_text.

Оформление образца:

  * заливки НЕТ вообще: в xl/styles.xml только patternFill none и gray125,
    ни одной цветной заливки на листе (ТЗ предполагало заливку шапки —
    в файле её нет, поэтому бланк собирается без заливки);
  * границы — только тонкие чёрные (thin, FF000000), других цветов нет;
  * шрифты: шапка R6 — Arial 10 bold по центру; данные R7-R12 — Arial 8;
    'Заказчик'/'Перевозчик' (C15/J15) — Arial 10 bold слева; значения
    нижнего блока (C16, J16, C19, J19) — Calibri 8 с цветом FF000000
    (шрифт по умолчанию Excel у заказчика);
  * wrap_text включён ровно у четырёх ячеек шапки: E6, G6, AC6, AE6;
  * числовые форматы: AD7:AD12 — 'h:mm' (время погрузки),
    AF7 — '#,##0" р."' (ставка с НДС);
  * ширина колонок (ws.column_dimensions) — ниже в COLUMN_WIDTHS;
    bestFit=1 у колонок B, H, I;
  * высоты строк 20-100 в образце равны 7,5 pt — это следы работы
    в файле заказчика, на вид формы не влияют и в бланк не переносятся;
  * в образце есть xl/drawings/drawing1.xml — две горизонтальные чёрные
    линии (подписи) под ячейками C19/J19. openpyxl фигуры создавать
    не умеет, поэтому в бланке линий нет (доложено в отчёте по этапу).

Чего в образце НЕТ (проверено по всему архиву):

  * ни одного зелёного цвета: перебор всех 6-значных HEX по всем частям
    архива не нашёл ни одного зелёного значения, в xl/styles.xml
    цвет только один — FF000000. «Зелёная обводка вокруг "Город погрузки"»
    из ТЗ в файле отсутствует (вероятно, это было выделение ячейки
    в Excel на скриншоте) — воспроизводить нечего, и бланк её не содержит
    (это закреплено тестом).

ПДн: в ячейках образца только шапка, пример даты ('16.09.2026') и блок
сторон; ФИО водителей, паспортов, телефонов, адресов, ИНН, VIN и номеров
нет, скрытых листов тоже нет. В метаданных: docProps/core.xml →
lastModifiedBy 'Сергей Жерженов' (creator пуст), xl/workbook.xml →
absPath 'D:\\Рабочий стол\\Данные договоров\\'. Это данные владельца
файла, а не третьих лиц, и то же lastModifiedBy уже лежит в репозитории
в templates/shablon_ooo.docx и templates/shablon_ip_*.docx, поэтому по
решению владельца образец коммитится как есть (см. отчёт по этапу).

═══════════════════════════════════════════════════════════════════════
СОЗНАТЕЛЬНЫЕ РЕШЕНИЯ ПО БЛАНКУ (согласованы перед сборкой)
═══════════════════════════════════════════════════════════════════════
  1. Нижний блок сторон стоит в колонках C и J — как в образце
     (в ТЗ было указано C и I; в файле заказчика 'Перевозчик' стоит
     именно в J, под 'Код дилера').
  2. Разрывы строк взяты из образца: после таблицы две пустые строки,
     затем 'Заказчик/Перевозчик' и 'Сюрлогистик/ООО "ТЕХНОЛОГИСТИКА"',
     затем ещё две пустые строки под подписи и 'ФИО, подпись, печать'.
     При 10 строках данных это R17-R18 пусто, R19, R20, R21-R22 пусто,
     R23 — ровно как в образце при 6 строках данных (там R13-R14 пусто,
     R15, R16, R17-R18 пусто, R19).
  3. Заливки нет — как в образце (см. выше).
  4. Две чёрные линии-подписи из xl/drawings/drawing1.xml не
     воспроизводятся: openpyxl не создаёт фигуры.
  5. Границы ячеек таблицы (R6-R16) — тонкая чёрная рамка со всех четырёх
     сторон у каждой ячейки A..AF. В образце Excel применяет сокращённую
     запись (общие рёбра соседних ячеек не дублируются), визуально это
     ровно та же сплошная сетка, но она не рассыпается при обрезке строк
     генератором.
  6. Шрифт всех ячеек данных (A..AF) — Arial 8: в образце AF8:AF12 имеет
     Calibri 8 (след ручной правки заказчика), в бланке колонка ставки
     набрана тем же шрифтом, что и остальная таблица.
  7. Числовые форматы 'h:mm' (AD) и '#,##0" р."' (AF) протянуты на все
     10 строк данных, а не только на заполненные строки образца.

ИДЕМПОТЕНТНОСТЬ И SHA256. openpyxl дважды пишет в файл текущее время:
в метку времени каждой записи zip и в docProps/core.xml (dcterms:created
и dcterms:modified — последнее openpyxl перезаписывает прямо в
save_workbook). Из-за этого байты одного и того же бланка отличались бы
от запуска к запуску, а SHA256 «плыл». Чтобы тест мог закрепить SHA256
шаблона (tests/test_havaly_template.py, TEMPLATE_SHA256), готовый архив
нормализуется: всем записям проставляется фиксированная дата
1980-01-01 00:00:00, а в docProps/core.xml — фиксированные created и
modified. Порядок записей и всё остальное содержимое не меняются.
Поэтому повторный запуск скрипта даёт побайтово тот же файл, а сам
скрипт не перезаписывает уже актуальный бланк (см. build_template).

ВНИМАНИЕ: шаблон закреплён по SHA256 в tests/test_havaly_template.py
(TEMPLATE_SHA256). После осознанной пересборки обнови эту константу,
иначе тест test_template_matches_pinned_sha упадёт.
"""

import hashlib
import re
import sys
import zipfile
from pathlib import Path
from typing import Dict, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

PROJECT_ROOT = Path(__file__).resolve().parent.parent
TEMPLATES_DIR = PROJECT_ROOT / "templates"

#: Имя собираемого бланка.
TEMPLATE_NAME = "shablon_havaly.xlsx"

#: Имя образца-источника в templates/ (копия присланного заказчиком файла).
SAMPLE_NAME = "Хавалы_образец.xlsx"

#: Лист всегда ОДИН и всегда с этим именем (требование заказчика).
SHEET_NAME = "TDSheet"

#: Заголовок бланка: две строки в одной ячейке B2 (объединение B2:C2).
TITLE = "ЗАЯВКА\n на перевозку автомобилей"

#: Подпись поля даты. Импорт ищет её по ТЕКСТУ, а не по номеру строки.
DATE_LABEL = "Дата заявки:"

#: Шапка таблицы — ровно 32 колонки A..AF, порядок и тексты как в образце.
HEADERS: Tuple[str, ...] = (
    "Номер Лота",                       # A
    "VIN",                              # B
    "Марка",                            # C
    "Модель",                           # D
    "Город погрузки",                   # E
    "Пункт погрузки",                   # F
    "Город доставки",                   # G
    "Пункт разгрузки",                  # H
    "Дилер",                            # I
    "Код дилера",                       # J
    "Наименование транспортной компании",  # K
    "Марка Автовоза",                   # L
    "Цвет кабины",                      # M
    "Номер автовоза",                   # N
    "Марка прицепа",                    # O
    "Номер Прицепа",                    # P
    "Фамилия",                          # Q
    "Имя",                              # R
    "Отчество",                         # S
    "Номер В/У",                        # T
    "Дата выдачи В/У",                  # U
    "Серия Паспорта",                   # V
    "Номер Паспорта",                   # W
    "Кем выдан паспорт",                # X
    "Дата выдачи паспорта",             # Y
    "Гражданство",                      # Z
    "Дата Рождения",                    # AA
    "Прописка",                         # AB
    "Планируемая дата\n погрузки",      # AC — перенос строки как в образце
    "Время погрузки",                   # AD
    "№ телефона водителя",              # AE
    "Ставка с НДС",                     # AF
)

#: Ячейки шапки с включённым переносом текста — ровно как в образце.
HEADER_WRAP_COLUMNS = ("E", "G", "AC", "AE")

#: Строки бланка.
TITLE_ROW = 2
DATE_ROW = 5
HEADER_ROW = 6
FIRST_DATA_ROW = 7

#: Максимум машин = максимум строк данных эталонного бланка.
DATA_ROWS = 10
LAST_DATA_ROW = FIRST_DATA_ROW + DATA_ROWS - 1

#: Нижний блок: две пустые строки, подписи сторон, две пустые строки
#: под собственноручные подписи и строка «ФИО, подпись, печать».
PARTY_ROW = LAST_DATA_ROW + 3          # 19: Заказчик / Перевозчик
PARTY_NAME_ROW = PARTY_ROW + 1         # 20: Сюрлогистик / ООО «ТЕХНОЛОГИСТИКА»
SIGNATURE_ROW = PARTY_ROW + 4          # 23: ФИО, подпись, печать
SPACER_ROWS = (
    LAST_DATA_ROW + 1, LAST_DATA_ROW + 2,   # 17, 18 — перед блоком сторон
    SIGNATURE_ROW - 2, SIGNATURE_ROW - 1,   # 21, 22 — под подписи
)

#: Стороны: Сюрлогистик — Заказчик, ООО «ТЕХНОЛОГИСТИКА» — Перевозчик.
#: В бланке зафиксированы, от заявки к заявке не меняются.
CUSTOMER = "Сюрлогистик"
CARRIER = 'ООО "ТЕХНОЛОГИСТИКА"'
PARTY_LABELS = ("Заказчик", "Перевозчик")
SIGNATURE_TEXT = "ФИО, подпись, печать"

#: Колонки нижнего блока — как в образце: слева C, справа J.
PARTY_COLUMNS = ("C", "J")

#: Ширины колонок A..AF — значения ws.column_dimensions образца.
COLUMN_WIDTHS: Dict[str, float] = {
    "A": 18.83203125, "B": 19.33203125, "C": 15.0, "D": 14.33203125,
    "E": 19.1640625, "F": 19.1640625, "G": 19.1640625, "H": 66.33203125,
    "I": 23.0, "J": 18.0, "K": 44.33203125, "L": 22.1640625,
    "M": 19.6640625, "N": 19.6640625, "O": 21.83203125, "P": 21.6640625,
    "Q": 10.5, "R": 10.5, "S": 10.5, "T": 18.33203125, "U": 21.0,
    "V": 26.5, "W": 18.0, "X": 27.1640625, "Y": 24.1640625,
    "Z": 28.1640625, "AA": 26.83203125, "AB": 25.6640625,
    "AC": 28.1640625, "AD": 26.33203125, "AE": 23.33203125,
    "AF": 23.33203125,
}

#: Колонки образца с bestFit=1 (Excel подгоняет ширину под содержимое).
COLUMN_BEST_FIT = ("B", "H", "I")

#: Высоты строк бланка (значения ws.row_dimensions образца).
ROW_HEIGHTS: Dict[int, float] = {
    1: 14.25,
    TITLE_ROW: 23.25,
    3: 23.25,
    DATE_ROW: 15.0,
    HEADER_ROW: 23.25,
    **{row: 11.25 for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1)},
    LAST_DATA_ROW + 1: 11.25,   # 17 — пусто перед блоком сторон
    LAST_DATA_ROW + 2: 12.75,   # 18 — пусто перед блоком сторон
    PARTY_ROW: 12.75,           # 19 — Заказчик / Перевозчик
    PARTY_NAME_ROW: 15.75,      # 20 — Сюрлогистик / ООО «ТЕХНОЛОГИСТИКА»
    SIGNATURE_ROW - 2: 11.25,   # 21 — пусто под подписи
    SIGNATURE_ROW - 1: 11.25,   # 22 — пусто под подписи
    SIGNATURE_ROW: 11.25,       # 23 — ФИО, подпись, печать
}

#: Реквизиты оформления образца.
FONT_FAMILY = "Arial"
#: Шрифт значений нижнего блока: в образце это Calibri 8 (шрифт по
#: умолчанию Excel у заказчика) с явно заданным чёрным цветом.
VALUE_FONT_FAMILY = "Calibri"
VALUE_FONT_COLOR = "FF000000"

SHEET_FONT_SIZE = 10.0
DATA_FONT_SIZE = 8.0

#: Тонкая чёрная граница ячеек таблицы (единственный стиль границ образца).
THIN_BLACK = Side(style="thin", color="FF000000")

#: Числовые форматы колонок данных: время погрузки и ставка с НДС.
TIME_FORMAT = "h:mm"
MONEY_FORMAT = '#,##0" р."'

#: Формат ячейки даты заявки (B5) — в образце шрифт Arial 8 по центру.
DATE_CELL_FONT_SIZE = 8.0

#: Фиксированная метка времени для записей zip: иначе SHA256 бланка
#: «плывёт» от запуска к запуску (см. docstring).
ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)

#: Часть архива с метаданными документа и фиксированные даты в ней.
CORE_PROPERTIES_PART = "docProps/core.xml"
DOCUMENT_TIMESTAMP = "2026-01-01T00:00:00Z"

#: dcterms:created и dcterms:modified в docProps/core.xml.
CORE_TIMESTAMP_RE = re.compile(
    r"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)"
)

# ─────────────────────────────────────────────────────────────
# Примитивы оформления
# ─────────────────────────────────────────────────────────────

#: Тонкая чёрная рамка со всех четырёх сторон — рамка ячеек таблицы.
BOX_BORDER = Border(left=THIN_BLACK, right=THIN_BLACK,
                    top=THIN_BLACK, bottom=THIN_BLACK)

#: Шрифты бланка.
TITLE_FONT = Font(name=FONT_FAMILY, size=SHEET_FONT_SIZE, bold=True)
LABEL_FONT = Font(name=FONT_FAMILY, size=SHEET_FONT_SIZE)
DATE_FONT = Font(name=FONT_FAMILY, size=DATE_CELL_FONT_SIZE)
HEADER_FONT = Font(name=FONT_FAMILY, size=SHEET_FONT_SIZE, bold=True)
DATA_FONT = Font(name=FONT_FAMILY, size=DATA_FONT_SIZE)
PARTY_FONT = Font(name=FONT_FAMILY, size=SHEET_FONT_SIZE, bold=True)
VALUE_FONT = Font(name=VALUE_FONT_FAMILY, size=DATA_FONT_SIZE,
                  color=VALUE_FONT_COLOR)

TITLE_ALIGNMENT = Alignment(horizontal="center", vertical="center",
                            wrap_text=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
HEADER_WRAP_ALIGNMENT = Alignment(horizontal="center", vertical="center",
                                  wrap_text=True)
LABEL_ALIGNMENT = Alignment(horizontal="left", vertical="center")
DATE_ALIGNMENT = Alignment(horizontal="center")


def write_cell(cell, value=None, *, font=None, alignment=None, border=None,
               number_format=None) -> None:
    """Единая запись ячейки: значение + шрифт + выравнивание + рамка + формат."""
    if value is not None:
        cell.value = value
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    if border is not None:
        cell.border = border
    if number_format is not None:
        cell.number_format = number_format


def configure_sheet(ws: Worksheet) -> None:
    """Параметры листа: имя и формат по умолчанию — как в образце."""
    ws.title = SHEET_NAME
    ws.sheet_format.baseColWidth = 8
    ws.sheet_format.defaultColWidth = 16.83203125
    ws.sheet_format.defaultRowHeight = 15.0
    ws.sheet_format.customHeight = True


def build_title(ws: Worksheet) -> None:
    """Строки 1-4: пустая строка, заголовок в объединении B2:C2, две пустые."""
    ws.merge_cells("B2:C2")
    write_cell(ws["B2"], TITLE, font=TITLE_FONT, alignment=TITLE_ALIGNMENT)


def build_date_row(ws: Worksheet) -> None:
    """
    Строка 5: подпись 'Дата заявки:' в колонке A и пустая ячейка B5.

    Дата из образца ('16.09.2026') в бланк НЕ переносится: ячейка B5
    остаётся пустой, но сохраняет оформление образца, чтобы её можно
    было заполнить руками или импортом.
    """
    write_cell(ws.cell(row=DATE_ROW, column=1), DATE_LABEL,
               font=LABEL_FONT, alignment=LABEL_ALIGNMENT)
    write_cell(ws.cell(row=DATE_ROW, column=2), None,
               font=DATE_FONT, alignment=DATE_ALIGNMENT)


def build_header_row(ws: Worksheet) -> None:
    """Строка 6: ровно 32 заголовка A..AF с тонкой рамкой и жирным шрифтом."""
    for index, title in enumerate(HEADERS, start=1):
        column = get_column_letter(index)
        alignment = (HEADER_WRAP_ALIGNMENT if column in HEADER_WRAP_COLUMNS
                     else HEADER_ALIGNMENT)
        write_cell(ws.cell(row=HEADER_ROW, column=index), title,
                   font=HEADER_FONT, alignment=alignment, border=BOX_BORDER)


def build_data_rows(ws: Worksheet) -> None:
    """
    Строки 7-16: 10 пустых строк данных (максимум машин).

    Строки пустые, но с рамкой, шрифтом и числовыми форматами: генератор
    позже обрежет лишние и заполнит фактические.
    """
    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        for index in range(1, len(HEADERS) + 1):
            column = get_column_letter(index)
            if column == "AD":
                number_format = TIME_FORMAT
            elif column == "AF":
                number_format = MONEY_FORMAT
            else:
                number_format = None
            write_cell(ws.cell(row=row, column=index), None,
                       font=DATA_FONT, border=BOX_BORDER,
                       number_format=number_format)


def build_parties_block(ws: Worksheet) -> None:
    """Строки 17-23: две пустые строки, стороны, две пустые, подпись."""
    left, right = PARTY_COLUMNS

    write_cell(ws[f"{left}{PARTY_ROW}"], PARTY_LABELS[0],
               font=PARTY_FONT, alignment=LABEL_ALIGNMENT)
    write_cell(ws[f"{right}{PARTY_ROW}"], PARTY_LABELS[1],
               font=PARTY_FONT, alignment=LABEL_ALIGNMENT)

    write_cell(ws[f"{left}{PARTY_NAME_ROW}"], CUSTOMER, font=VALUE_FONT)
    write_cell(ws[f"{right}{PARTY_NAME_ROW}"], CARRIER, font=VALUE_FONT)

    write_cell(ws[f"{left}{SIGNATURE_ROW}"], SIGNATURE_TEXT, font=VALUE_FONT)
    write_cell(ws[f"{right}{SIGNATURE_ROW}"], SIGNATURE_TEXT, font=VALUE_FONT)


def apply_column_widths(ws: Worksheet) -> None:
    """Ширины колонок A..AF — как в образце (плюс bestFit у B, H, I)."""
    for column, width in COLUMN_WIDTHS.items():
        dimension = ws.column_dimensions[column]
        dimension.width = width
        dimension.bestFit = column in COLUMN_BEST_FIT


def apply_row_heights(ws: Worksheet) -> None:
    """Высоты строк — как в образце."""
    for row, height in ROW_HEIGHTS.items():
        ws.row_dimensions[row].height = height


# ─────────────────────────────────────────────────────────────
# Сборка файла
# ─────────────────────────────────────────────────────────────

def template_path() -> Path:
    """Путь собираемого бланка: templates/shablon_havaly.xlsx."""
    return TEMPLATES_DIR / TEMPLATE_NAME


def normalize_core_properties(payload: bytes) -> bytes:
    """
    Ставит фиксированные created и modified в docProps/core.xml.

    openpyxl в save_workbook всегда перезаписывает dcterms:modified
    текущим временем, поэтому без этой правки SHA256 бланка не
    воспроизводим. Ожидаются ровно две даты: created и modified.
    """
    text = payload.decode("utf-8")
    text, replaced = CORE_TIMESTAMP_RE.subn(
        rf"\g<1>{DOCUMENT_TIMESTAMP}\g<2>", text
    )
    if replaced != 2:
        raise RuntimeError(
            "в docProps/core.xml ожидались dcterms:created и dcterms:modified, "
            f"заменено: {replaced}. Проверь версию openpyxl: формат "
            "метаданных изменился, SHA256 бланка перестанет быть стабильным."
        )
    return text.encode("utf-8")


def normalize_archive(path: Path) -> None:
    """
    Фиксирует метки времени архива и метаданных, не меняя содержимое.

    openpyxl записывает в zip текущее время, поэтому SHA256 одного и того
    же бланка отличался бы от запуска к запуску. После нормализации
    сборка воспроизводима побайтово, и тест может закрепить SHA256.
    """
    with zipfile.ZipFile(path) as source:
        entries = [(info, source.read(info.filename))
                   for info in source.infolist()]

    temporary = path.with_name(path.name + ".normalized")
    with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for info, payload in entries:
            info.date_time = ZIP_TIMESTAMP
            if info.filename == CORE_PROPERTIES_PART:
                payload = normalize_core_properties(payload)
            target.writestr(info, payload)
    temporary.replace(path)


def build_workbook() -> Workbook:
    """Собирает книгу бланка целиком (без записи на диск)."""
    workbook = Workbook()
    ws = workbook.active

    configure_sheet(ws)
    build_title(ws)
    build_date_row(ws)
    build_header_row(ws)
    build_data_rows(ws)
    build_parties_block(ws)
    apply_column_widths(ws)
    apply_row_heights(ws)

    return workbook


def build_template(path: Optional[Path] = None) -> Path:
    """
    Собирает бланк и возвращает путь к нему.

    Идемпотентность: готовый файл сравнивается по байтам, и если он уже
    совпадает со свежей сборкой, он не перезаписывается (не меняются
    ни содержимое, ни время изменения).
    """
    target = Path(path) if path is not None else template_path()
    target.parent.mkdir(parents=True, exist_ok=True)

    temporary = target.with_name(target.name + ".building")
    build_workbook().save(str(temporary))
    normalize_archive(temporary)

    if target.exists() and target.read_bytes() == temporary.read_bytes():
        temporary.unlink()
        return target

    temporary.replace(target)
    return target


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except AttributeError:  # pragma: no cover - экзотические потоки вывода
        pass

    existed_before = template_path().exists()
    sha_before = (hashlib.sha256(template_path().read_bytes()).hexdigest()
                  if existed_before else None)

    path = build_template()
    digest = hashlib.sha256(path.read_bytes()).hexdigest()

    print(f"Бланк Хавалов: {path}")
    print(f"  лист: {SHEET_NAME}, колонок: {len(HEADERS)}, "
          f"строк данных: {DATA_ROWS} (R{FIRST_DATA_ROW}-R{LAST_DATA_ROW})")
    print(f"  нижний блок: R{PARTY_ROW}, R{PARTY_NAME_ROW}, R{SIGNATURE_ROW}, "
          f"колонки {PARTY_COLUMNS[0]}/{PARTY_COLUMNS[1]}")
    print(f"  размер: {path.stat().st_size} байт, SHA256: {digest}")

    if not existed_before:
        print("  файл создан")
    elif sha_before == digest:
        print("  файл уже был актуален — не изменён (идемпотентно)")
    else:
        print("  файл пересобран: содержимое изменилось")

    return 0


if __name__ == "__main__":
    sys.exit(main())
