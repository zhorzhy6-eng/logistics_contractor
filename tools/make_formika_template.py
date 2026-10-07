#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сборка шаблона договора-заявки «Формика» (ЭТАП 3.1.A.1).

Создаёт templates/shablon_formika.docx — ПУСТОЙ БЛАНК с плейсхолдерами
{{...}} для docxtpl. Данные из образца
(templates/Заявка_ТЛ_447 Формика_Технологистика.docx) сюда не переносятся:
ни марок, ни VIN, ни ФИО, ни госномеров, ни сумм. Образец читается только
как источник структуры, формулировок и оформления.

Что воспроизводится по образцу:
  * A4, поля 2,0 / 1,5 / 2,0 / 2,0 см;
  * шапка: заголовок, две строки под ним, строка «г. Москва … дата»;
  * 8 разделов: груз, маршрут, исполнитель, стоимость, особые условия,
    ответственность, прочие условия, подписи;
  * таблица груза: шапка + 12 строк, ширины колонок 1000/4770/4253 twips,
    заливка шапки E5E5E5, границы, текст по центру;
  * таблица подписей: 1×2, колонки по 4500 twips.

Что исправлено по сравнению с образцом (опечатки/неряшливость):
  * «№ 1  от «28 » ноября 2025 г.» → «№ 1 от «28» ноября 2025 г.»;
  * «2. МАРШРУТ ПЕРЕВОЗКИ :» → «2. МАРШРУТ ПЕРЕВОЗКИ:»;
  * двойные пробелы в тексте и марках — ушли вместе с данными образца;
  * «Дата и время:» оформлена полужирной, как соседние метки
    («Пункт погрузки:», «Пункт выгрузки:», «Срок доставки:»);
  * абзац с тремя табуляциями и хвост из пустых абзацев не переносятся;
  * п. 4 «Порядок оплаты» печатает `{{payment_days}}
    ({{payment_days_words}}) банковских дней` вместо константы
    «3 (трех) банковских дней» — срок оплаты задаётся в интерфейсе
    (вкладка «Стоимость»), как в перевозке и в аренде ТС.

Шрифт — Times New Roman (заголовок 18 pt bold, остальное 11 pt), как в
рабочем шаблоне shablon_ooo.docx. Плейсхолдер всегда лежит в ОДНОМ run:
docxtpl(Jinja) не склеивает разорванные по runs переменные.

Запуск:  python tools/make_formika_template.py
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = PROJECT_ROOT / "templates" / "shablon_formika.docx"

FONT = "Times New Roman"
TITLE_SIZE = 18
TEXT_SIZE = 11

#: Геометрия страницы образца: A4, поля 2,0 / 1,5 / 2,0 / 2,0 см.
PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7
MARGIN_LEFT_CM = 2.0
MARGIN_RIGHT_CM = 1.5
MARGIN_TOP_CM = 2.0
MARGIN_BOTTOM_CM = 2.0

#: Ширина текстового блока — правый таб-стоп строки «г. Москва … дата».
TEXT_WIDTH_CM = PAGE_WIDTH_CM - MARGIN_LEFT_CM - MARGIN_RIGHT_CM

#: Ширины колонок таблиц в twips (1/1440 дюйма) — ровно как в образце.
CARGO_COLUMN_WIDTHS = (1000, 4770, 4253)
SIGN_COLUMN_WIDTHS = (4500, 4500)

#: Заливка шапки таблицы груза из образца.
HEADER_FILL = "E5E5E5"

#: Число строк таблицы груза: бланк рассчитан на 12 машин; лишние строки
#: удаляет постобработка генератора (RemoveEmptyVehicleRowsStep).
CAR_ROWS = 12


def twips(value: int) -> Emu:
    """twips → EMU (1 twip = 635 EMU)."""
    return Emu(int(value) * 635)


# ─────────────────────────────────────────────────────────────
# Примитивы оформления
# ─────────────────────────────────────────────────────────────

def _set_run_font(run, size: int, bold: bool, underline: bool) -> None:
    """Шрифт run: имя для всех алфавитов (ascii/hAnsi/cs/eastAsia)."""
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.underline = underline

    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT)


def add_paragraph(doc, parts, *, align=None, size=TEXT_SIZE,
                  space_before=0, space_after=0, tab_stop_right_cm=None):
    """
    Абзац из списка частей.

    Часть — строка (обычный текст) либо кортеж (текст, {bold/underline/size}).
    Каждая часть становится отдельным run, поэтому плейсхолдер целиком
    остаётся в одном run.
    """
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)

    if tab_stop_right_cm is not None:
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Emu(int(tab_stop_right_cm * 360000)), WD_TAB_ALIGNMENT.RIGHT
        )

    for part in parts:
        text, opts = (part, {}) if isinstance(part, str) else part
        run = paragraph.add_run(text)
        _set_run_font(
            run,
            size=opts.get("size", size),
            bold=opts.get("bold", False),
            underline=opts.get("underline", False),
        )
    return paragraph


def shade_cell(cell, fill: str) -> None:
    """Заливка ячейки (w:shd). Ставится после tcW — порядок тегов соблюдён."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_borders(cell, size_eighth_pt: int) -> None:
    """Границы ячейки в 1/8 pt (8 = 1 pt). Вставляются сразу после tcW."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size_eighth_pt))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        borders.append(element)

    tcW = tcPr.find(qn("w:tcW"))
    if tcW is not None:
        tcW.addnext(borders)
    else:
        tcPr.insert(0, borders)


def write_cell(cell, parts, *, align=WD_ALIGN_PARAGRAPH.CENTER, size=TEXT_SIZE):
    """Текст в первую (единственную) строку ячейки."""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    for part in parts:
        text, opts = (part, {}) if isinstance(part, str) else part
        run = paragraph.add_run(text)
        _set_run_font(
            run,
            size=opts.get("size", size),
            bold=opts.get("bold", False),
            underline=opts.get("underline", False),
        )


def set_table_widths(table, widths_twips) -> None:
    """Фиксированные ширины колонок: tblW + gridCol + tcW."""
    table.autofit = False

    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.insert(0, tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(sum(widths_twips)))

    for column, width in zip(table.columns, widths_twips):
        column.width = twips(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_twips):
            cell.width = twips(width)


# ─────────────────────────────────────────────────────────────
# Сборка документа
# ─────────────────────────────────────────────────────────────

def configure_document(doc) -> None:
    """Страница и базовый шрифт документа."""
    section = doc.sections[0]
    section.page_width = Emu(int(PAGE_WIDTH_CM * 360000))
    section.page_height = Emu(int(PAGE_HEIGHT_CM * 360000))
    section.left_margin = Emu(int(MARGIN_LEFT_CM * 360000))
    section.right_margin = Emu(int(MARGIN_RIGHT_CM * 360000))
    section.top_margin = Emu(int(MARGIN_TOP_CM * 360000))
    section.bottom_margin = Emu(int(MARGIN_BOTTOM_CM * 360000))

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(TEXT_SIZE)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT)


def build_header(doc) -> None:
    """Шапка: заголовок, две строки под ним, город/дата, преамбула."""
    add_paragraph(
        doc,
        [
            ("ДОГОВОР-ЗАЯВКА № ", {"bold": True, "size": TITLE_SIZE}),
            ("{{contract_number}}", {"bold": True, "size": TITLE_SIZE}),
        ],
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )

    add_paragraph(
        doc,
        ["к Генеральному договору на перевозку грузов автотранспортом"],
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_paragraph(
        doc,
        ["№ 1 от «28» ноября 2025 г."],
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    add_paragraph(doc, [""])
    add_paragraph(
        doc,
        ["г. Москва\t", "«{{contract_date}}» {{contract_month}} {{contract_year}} г."],
        tab_stop_right_cm=TEXT_WIDTH_CM,
    )
    add_paragraph(doc, [""])

    add_paragraph(
        doc,
        ["ООО «ТЕХНОЛОГИСТИКА» (Экспедитор) и ООО «Формика» (Заказчик), "
         "действующие на основании Генерального договора № 1 от «28» ноября "
         "2025 г., согласовали следующие условия перевозки:"],
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_paragraph(doc, [""])


def build_cargo_table(doc):
    """Таблица груза: шапка + 12 строк с плейсхолдерами."""
    table = doc.add_table(rows=1 + CAR_ROWS, cols=3)
    table.style = "Table Grid"
    set_table_widths(table, CARGO_COLUMN_WIDTHS)

    headers = ("№", "Марка, модель", "VIN-номер")
    for cell, title in zip(table.rows[0].cells, headers):
        write_cell(cell, [(title, {"bold": True})])
        shade_cell(cell, HEADER_FILL)

    for number in range(1, CAR_ROWS + 1):
        row = table.rows[number]
        write_cell(row.cells[0], [str(number)])
        write_cell(row.cells[1], [f"{{{{car_{number}_brand}}}}"])
        write_cell(row.cells[2], [f"{{{{car_{number}_vin}}}}"])

    return table


def build_route_block(doc) -> None:
    """Раздел 2 «Маршрут перевозки»."""
    add_paragraph(
        doc,
        [
            ("2. МАРШРУТ ПЕРЕВОЗКИ: ", {"bold": True}),
            ("{{route}}", {"underline": True}),
        ],
        space_before=6,
    )
    add_paragraph(
        doc,
        [("Пункт погрузки: ", {"bold": True}), "{{loading_address}}"],
    )
    add_paragraph(
        doc,
        [
            ("Дата и время: ", {"bold": True}),
            "{{loading_plan_date}} г. с {{loading_plan_time_from}} "
            "до {{loading_plan_time_to}}",
        ],
    )
    add_paragraph(
        doc,
        [("Пункт выгрузки: ", {"bold": True}), "{{unloading_address}}"],
    )
    add_paragraph(
        doc,
        [("Срок доставки: ", {"bold": True}),
         "3 календарных дня с момента погрузки"],
    )
    add_paragraph(doc, [""])


def build_executor_block(doc) -> None:
    """Раздел 3 «Информация об исполнителе»: водитель и ТС."""
    add_paragraph(doc, [("Водитель:", {"bold": True})])
    add_paragraph(doc, ["ФИО: {{driver_name}}"])
    add_paragraph(doc, ["Дата рождения: {{driver_birth_date}}"])
    add_paragraph(doc, ["Паспорт: {{driver_passport}}"])
    add_paragraph(doc, ["Выдан: {{driver_passport_issuer}}"])
    add_paragraph(doc, ["Дата выдачи: {{driver_passport_date}}"])
    add_paragraph(doc, ["Водительское удостоверение: {{driver_license}}"])
    add_paragraph(doc, ["Адрес регистрации: {{driver_address}}"])
    add_paragraph(doc, ["Телефон: {{driver_phone}}"])

    add_paragraph(doc, [("Транспортное средство:", {"bold": True})])
    add_paragraph(
        doc,
        ["Тягач: {{tractor_brand}} гос. номер: {{tractor_plate}}"],
    )
    add_paragraph(doc, ["Тип ТС: {{tractor_type}}"])
    add_paragraph(
        doc,
        ["Прицеп: {{trailer_brand}} гос. номер: {{trailer_plate}}"],
    )
    add_paragraph(doc, [""])


def build_cost_block(doc) -> None:
    """Раздел 4 «Стоимость и порядок оплаты»."""
    add_paragraph(
        doc,
        ["Стоимость перевозки составляет: {{sum_total}} руб. "
         "({{sum_total_words}}), включая НДС {{vat_rate}}."],
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    # Срок оплаты — плейсхолдеры, а не константа: число банковских дней
    # задаётся на вкладке «Стоимость». Формулировка — как в шаблоне
    # перевозки (п. 4.4) и в аренде ТС (п. 4.5): цифры, затем прописью
    # в скобках, затем «банковских дней».
    add_paragraph(
        doc,
        ["Порядок оплаты: в течение {{payment_days}} "
         "({{payment_days_words}}) банковских дней после получения "
         "Заказчиком оригиналов товарных накладных (ТН) и актов "
         "приема-передачи."],
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_paragraph(doc, [""])


def build_terms_block(doc, title: str, items) -> None:
    """Фиксированный раздел: заголовок + пункты."""
    add_paragraph(doc, [(title, {"bold": True})], space_before=6)
    for item in items:
        add_paragraph(doc, [item], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_paragraph(doc, [""])


def build_signatures(doc) -> None:
    """Раздел 8 «Подписи сторон»: таблица 1×2."""
    add_paragraph(doc, [("8. ПОДПИСИ СТОРОН", {"bold": True})], space_before=6)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_widths(table, SIGN_COLUMN_WIDTHS)

    left, right = table.rows[0].cells
    for cell in (left, right):
        set_cell_borders(cell, 8)

    def fill_column(cell, title: str, company: str, sign_line: str) -> None:
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(title)
        _set_run_font(run, size=TEXT_SIZE, bold=True, underline=False)

        company_par = cell.add_paragraph()
        company_par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        company_par.paragraph_format.space_after = Pt(0)
        run = company_par.add_run(company)
        _set_run_font(run, size=TEXT_SIZE, bold=True, underline=False)

        blank_par = cell.add_paragraph()
        blank_par.paragraph_format.space_after = Pt(0)

        sign_par = cell.add_paragraph()
        sign_par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        sign_par.paragraph_format.space_after = Pt(0)
        run = sign_par.add_run(sign_line)
        _set_run_font(run, size=TEXT_SIZE, bold=False, underline=False)

    fill_column(left, "Заказчик:", "ООО «Формика»",
                "Генеральный директор: ___________________ А.В. Калашников")
    fill_column(right, "Экспедитор:", "ООО «ТЕХНОЛОГИСТИКА»",
                "Генеральный директор: ______________________ Т.А. Ахмедов")


def build_template() -> Path:
    """Собирает бланк и сохраняет его в templates/shablon_formika.docx."""
    doc = Document()
    configure_document(doc)

    build_header(doc)

    add_paragraph(doc, [("1. ИНФОРМАЦИЯ О ГРУЗЕ", {"bold": True})], space_before=6)
    build_cargo_table(doc)
    add_paragraph(doc, [""])

    build_route_block(doc)

    add_paragraph(doc, [("3. ИНФОРМАЦИЯ ОБ ИСПОЛНИТЕЛЕ", {"bold": True})],
                   space_before=6)
    build_executor_block(doc)

    add_paragraph(doc, [("4. СТОИМОСТЬ И ПОРЯДОК ОПЛАТЫ", {"bold": True})],
                   space_before=6)
    build_cost_block(doc)

    build_terms_block(doc, "5. ОСОБЫЕ УСЛОВИЯ", [
        "Страхование груза осуществляется Экспедитором на полную стоимость, "
        "указанную в п.1.",
        "При приемке и сдаче груза составляются Акты приема-передачи "
        "с фотофиксацией.",
        "Перевозка осуществляется специализированным автовозом, пригодным "
        "для перевозки электромобилей.",
        "Водитель имеет все необходимые документы для осуществления "
        "международных перевозок.",
    ])

    build_terms_block(doc, "6. ОТВЕТСТВЕННОСТЬ", [
        "За нарушение сроков доставки Экспедитор уплачивает пеню 0,5% "
        "от стоимости перевозки за каждый день просрочки, но не более 10% "
        "от стоимости перевозки.",
        "За срыв погрузки по вине Экспедитора – штраф 20% от стоимости "
        "перевозки.",
        "За задержку погрузки/выгрузки по вине Заказчика – пеня 5000 рублей "
        "за каждые сутки.",
        "Ответственность Экспедитора ограничивается стоимостью перевозки, "
        "за исключением случаев утраты груза.",
        "В остальном Стороны руководствуются условиями Генерального договора.",
    ])

    build_terms_block(doc, "7. ПРОЧИЕ УСЛОВИЯ", [
        "Настоящий Договор-заявка является неотъемлемой частью Генерального "
        "договора.",
        "Договор-заявка вступает в силу с момента подписания и действует "
        "до полного исполнения обязательств.",
        "Изменения возможны только по письменному соглашению Сторон.",
    ])

    build_signatures(doc)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT_PATH))
    return OUTPUT_PATH


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except AttributeError:
        pass

    path = build_template()

    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    texts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    placeholders = re.findall(r"\{\{[^{}]*\}\}", "\n".join(texts))

    print(f"[OK] {path}")
    print(f"     абзацев={len(doc.paragraphs)} таблиц={len(doc.tables)} "
          f"строк таблицы груза={len(doc.tables[0].rows)}")
    print(f"     плейсхолдеров={len(placeholders)} "
          f"уникальных={len(set(placeholders))}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
