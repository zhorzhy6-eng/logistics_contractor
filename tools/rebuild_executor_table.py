#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Перестраивает пункт 3.5 «Информация об исполнителе» в таблицу.

Было: два столбца текста друг под другом — «Водитель:» и «Транспортное
средство:». Стало: одна таблица из двух колонок, где слева данные
водителя, справа — данные транспортного средства (тягач и полуприцеп).
Оформление берётся из таблицы машин (блок 3.1): та же шапка, заливка,
границы и шрифт.

Скрипт идемпотентен: если раздела 3.5 в виде абзацев уже нет (данные
перенесены в таблицу), файл не меняется.

Запуск:
    python tools/rebuild_executor_table.py            # все три шаблона
    python tools/rebuild_executor_table.py --check    # только показать план
"""

import argparse
import shutil
import sys
from copy import deepcopy
from datetime import datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document
from docx.shared import Cm

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

TEMPLATES = (
    "shablon_ooo.docx",
    "shablon_ip_with_vat.docx",
    "shablon_ip_without_vat.docx",
)

SECTION_START = "3.5."
SECTION_END = "3.6."
DRIVER_HEADING = "Водитель:"
VEHICLE_HEADING = "Транспортное средство:"

#: Ширина колонок таблицы 3.5 (полоса набора — 16,5 см).
COLUMN_WIDTH_CM = 8.25


def find_paragraph(doc: Document, prefix: str):
    """Первый абзац верхнего уровня, текст которого начинается с prefix."""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    return None


def collect_section(doc: Document):
    """
    Абзацы раздела 3.5 между заголовком и следующим разделом.

    Возвращает (заголовок_раздела, [абзацы], [абзацы водителя], [абзацы ТС])
    или None, если раздел уже перестроен (абзацев «Водитель:» нет).
    """
    paragraphs = list(doc.paragraphs)
    start = end = None
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if start is None and text.startswith(SECTION_START):
            start = index
        elif start is not None and text.startswith(SECTION_END):
            end = index
            break

    if start is None or end is None:
        return None

    body = [p for p in paragraphs[start + 1:end] if p.text.strip()]
    driver = []
    vehicle = []
    target = None
    for paragraph in body:
        text = paragraph.text.strip()
        if text == DRIVER_HEADING:
            target = driver
            continue
        if text == VEHICLE_HEADING:
            target = vehicle
            continue
        if target is not None:
            target.append(text)

    if not driver and not vehicle:
        return None

    return paragraphs[start], body, driver, vehicle


def sample_formats(doc: Document):
    """
    Образцы оформления из таблицы машин (блок 3.1): tcPr и rPr шапки/данных.
    """
    for table in doc.tables:
        head = [c.text.strip() for c in table.rows[0].cells] if table.rows else []
        if "VIN" not in "".join(head) or len(table.rows) < 2:
            continue
        return table.rows[0].cells[0], table.rows[1].cells[0]
    return None, None


def copy_cell_format(cell, sample_cell) -> None:
    """Переносит оформление ячейки-образца (заливку, границы, отступы)."""
    if sample_cell is None:
        return
    props = sample_cell._tc.find(W + "tcPr")
    if props is None:
        return

    copied = deepcopy(props)
    for tag in ("tcW", "gridSpan", "vMerge", "hMerge"):
        node = copied.find(W + tag)
        if node is not None:
            copied.remove(node)

    existing = cell._tc.find(W + "tcPr")
    if existing is not None:
        cell._tc.remove(existing)
    cell._tc.insert(0, copied)


def write_cell(cell, lines, sample_cell, bold: bool, center: bool = False) -> None:
    """Заполняет ячейку строками текста с оформлением ячейки-образца."""
    copy_cell_format(cell, sample_cell)

    sample_paragraph = (
        sample_cell.paragraphs[0]
        if (sample_cell is not None and sample_cell.paragraphs)
        else None
    )
    sample_run = (
        sample_paragraph.runs[0]._r
        if (sample_paragraph is not None and sample_paragraph.runs)
        else None
    )
    sample_ppr = sample_paragraph._p.find(W + "pPr") if sample_paragraph is not None else None
    sample_rpr = sample_run.find(W + "rPr") if sample_run is not None else None

    # Ячейка новой таблицы уже содержит один пустой абзац: чистим его и
    # удаляем лишние. cell.text = "" не годится — он оставляет пустой run.
    paragraphs = cell.paragraphs
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    paragraphs[0].clear()

    paragraphs = [cell.paragraphs[0]]
    for _ in range(max(0, len(lines) - 1)):
        paragraphs.append(cell.add_paragraph())

    for paragraph, line in zip(paragraphs, lines):
        if sample_ppr is not None:
            existing = paragraph._p.find(W + "pPr")
            if existing is not None:
                paragraph._p.remove(existing)
            paragraph._p.insert(0, deepcopy(sample_ppr))
        run = paragraph.add_run(line)
        if sample_rpr is not None:
            run._r.insert(0, deepcopy(sample_rpr))
        run.bold = bold
        if center:
            paragraph.alignment = 1  # WD_ALIGN_PARAGRAPH.CENTER


def usable_width_cm(doc: Document) -> float:
    """Ширина полосы набора первой секции, см."""
    from docx.shared import Emu

    section = doc.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    return Emu(width).cm


def set_table_widths(table, column_cm: float) -> None:
    """
    Фиксирует ширину колонок: tblGrid + tblW + tcW.

    Одного cell.width мало: Word считает колонки по tblGrid, а python-docx
    при создании таблицы раздаёт им ширину полосы набора целиком — без
    правки tblGrid таблица вылезает за поля страницы.
    """
    from docx.shared import Cm, Emu

    width = Cm(column_cm)
    total = Emu(width * len(table.columns))

    props = table._tbl.tblPr
    props.get_or_add_tblLayout().type = "fixed"

    tbl_width = props.find(W + "tblW")
    if tbl_width is None:
        tbl_width = props.makeelement(W + "tblW", {})
        props.insert(0, tbl_width)
    tbl_width.set(W + "type", "dxa")
    tbl_width.set(W + "w", str(int(total.twips)))

    for grid_col in table._tbl.tblGrid:
        grid_col.set(W + "w", str(int(width.twips)))

    for row in table.rows:
        for cell in row.cells:
            cell.width = width


def rebuild_template(path: Path, apply: bool) -> str:
    """Перестраивает раздел 3.5 в одном шаблоне."""
    doc = Document(str(path))

    section = collect_section(doc)
    if section is None:
        return "уже в таблице — пропуск"

    heading_paragraph, body, driver, vehicle = section
    if not driver or not vehicle:
        return f"не найдены данные (водитель: {len(driver)}, ТС: {len(vehicle)}) — пропуск"

    header_sample, data_sample = sample_formats(doc)

    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.autofit = False

    write_cell(table.rows[0].cells[0], ["Водитель"], header_sample, bold=True, center=True)
    write_cell(table.rows[0].cells[1], ["Транспортное средство"], header_sample,
               bold=True, center=True)
    write_cell(table.rows[1].cells[0], driver, data_sample, bold=False)
    write_cell(table.rows[1].cells[1], vehicle, data_sample, bold=False)

    for row in table.rows:
        for cell in row.cells:
            cell.width = Cm(COLUMN_WIDTH_CM)

    if not apply:
        return (f"будет заменено: абзацев {len(body)} -> таблица 2x2 "
                f"(водитель {len(driver)} строк, ТС {len(vehicle)} строк)")

    # Ширина колонок — половина полосы набора с небольшим запасом,
    # чтобы таблица гарантированно не вылезала за поля страницы.
    column_cm = max(4.0, (usable_width_cm(doc) - 0.1) / 2)
    set_table_widths(table, column_cm)

    # Переносим таблицу на место удаляемых абзацев
    anchor = heading_paragraph._p
    anchor.addnext(table._tbl)

    parent = heading_paragraph._p.getparent()
    for paragraph in body:
        parent.remove(paragraph._p)

    # Пустой абзац после таблицы: без него следующий раздел липнет к таблице
    spacer = doc.add_paragraph()
    table._tbl.addnext(spacer._p)

    doc.save(str(path))
    return (f"готово: таблица 2x2 (колонки {column_cm:.2f} см), "
            f"водитель {len(driver)} строк, ТС {len(vehicle)} строк")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true",
                        help="только показать, что будет сделано")
    parser.add_argument("--templates", default=None,
                        help="папка с шаблонами (по умолчанию templates/)")
    parser.add_argument("--no-backup", action="store_true",
                        help="не делать резервную копию шаблонов")
    args = parser.parse_args()

    templates_dir = Path(args.templates) if args.templates else PROJECT_ROOT / "templates"

    if not args.check and not args.no_backup:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_dir = PROJECT_ROOT / "backup" / f"templates_before_35_{stamp}"
        backup_dir.mkdir(parents=True, exist_ok=True)
        for name in TEMPLATES:
            source = templates_dir / name
            if source.exists():
                shutil.copy2(source, backup_dir / name)
        print(f"Резервная копия: {backup_dir}\n")

    for name in TEMPLATES:
        path = templates_dir / name
        if not path.exists():
            print(f"{name}: НЕ НАЙДЕН")
            continue
        result = rebuild_template(path, apply=not args.check)
        print(f"{name}: {result}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
