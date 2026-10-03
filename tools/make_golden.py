#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Генератор золотых эталонов договоров (Шаг 0 рефакторинга архитектуры контрактов).

Создаёт tests/data/golden/:
  <сценарий>.docx — вывод core.contract_generator.ContractGenerator на
                    облегчённой копии реального шаблона (без встроенных шрифтов);
  <сценарий>.json — структурный отпечаток документа (абзацы, таблицы,
                    стили, ширины колонок, порядок элементов тела, хеши),
                    по которому после рефакторинга сверяется, что старые
                    договоры генерируются без изменений.

Сценарии:
  ooo            — шаблон ООО, перевозчик «ООО (с НДС)»;
  ip_with_vat    — шаблон ИП с НДС, перевозчик «ИП с НДС»;
  ip_without_vat — шаблон ИП без НДС, перевозчик «ИП без НДС»;
  gap            — шаблон ООО, маршрут с «дыркой» (пустой адрес точки
                   в середине), машины привязаны к исходным номерам точек;
  formika_sample — шаблон Формики, 4 машины (ЭТАП 3.1.A, свой генератор).

Все данные синтетические (как в tests/conftest.py), реальных ПДн нет.

Перегенерация:  python tools/make_golden.py
Один сценарий:  python tools/make_golden.py formika_sample
(без аргументов перегенерируются ВСЕ эталоны; чтобы не переписывать
эталоны перевозки, указывайте нужный сценарий явно).
"""

import hashlib
import json
import re
import shutil
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document  # noqa: E402

from core.contract_generator import ContractGenerator  # noqa: E402
from core.contracts.factory import GeneratorFactory  # noqa: E402
from core.contracts.registry import ContractTypeRegistry  # noqa: E402

PROJECT_ROOT = Path(__file__).resolve().parent.parent
GOLDEN_DIR = PROJECT_ROOT / "tests" / "data" / "golden"

# ─────────────────────────────────────────────────────────────
# Синтетические данные (зеркалят tests/conftest.py)
# ─────────────────────────────────────────────────────────────

DRIVER = {
    "full_name": "Иванов Иван Иванович",
    "birth_date": "1980-01-01",
    "birth_place": "г. Москва",
    "passport_series": "18 22",
    "passport_number": "926830",
    "passport_issue_date": "2023-01-30",
    "passport_issuer": "Отделом УФМС России по г. Москве",
    "passport_code": "500-123",
    "registration_address": "г. Москва, ул. Тестовая, д. 1",
    "license_series": "99 36",
    "license_number": "123456",
    "license_issue_date": "2020-01-01",
    "license_expiry_date": "2030-01-01",
    "license_categories": "B, C, E",
    "phone": "+7 (999) 123-45-67",
}

ORGANIZATION = {
    "full_name": "ООО «Ромашка»",
    "short_name": "ООО «Ромашка»",
    "inn": "7701234567",
    "kpp": "770101001",
    "ogrn": "1027700132195",
    "legal_address": "г. Москва, ул. Тестовая, д. 1",
    "actual_address": "г. Москва, ул. Тестовая, д. 1",
    "bank_account": "40702810000000000001",
    "bik": "044525225",
    "correspondent_account": "30101810400000000225",
    "bank_name": "ПАО Сбербанк",
    "director_name": "Петров Пётр Петрович",
    "director_position": "Генеральный директор",
    "phone": "+7 (495) 123-45-67",
    "email": "info@example.ru",
}

TRACTOR = {"brand_model": "Foton Auman", "plate_number": "O844XY196",
           "color": "Белый", "year": 2023}
TRAILER = {"brand_model": "YANGMINDA", "plate_number": "71ABF18",
           "color": "Серый", "year": 2020}


def base_payload(carrier_type: str, vat_rate: str, vat_rate_num) -> dict:
    carrier = dict(ORGANIZATION, entity_type="ООО")
    customer = dict(ORGANIZATION)
    customer["full_name"] = "ООО «Заказчик»"
    customer["short_name"] = "ООО «Заказчик»"

    return {
        "driver": DRIVER,
        "carrier": carrier,
        "customer": customer,
        "vehicles": [
            {"vin": "EC3TEUMB0T0002608", "brand_model": "JETOUR T2",
             "plate_number": "А123ВС77", "year": 2024, "color": "Белый",
             "vehicle_type": "Легковой автомобиль"},
        ],
        "tractor": TRACTOR,
        "trailer": TRAILER,
        "contract": {
            "number": "23092026-74", "date": "2026-09-23",
            "route": "Мурманск - Пятигорск", "carrier_type": carrier_type,
            "vat_rate": vat_rate, "vat_rate_num": vat_rate_num,
            "price_without_vat": 180300.0, "price_with_vat": 219966.0,
            "payment_days": 10,
        },
        "loadings": [
            {"address": "183052, г.Мурманск, пр.Кольский, д.53",
             "date": "2026-09-24", "time_window": "09:00-18:00"},
        ],
        "unloadings": [
            {"address": "г. Пятигорск, Бештаугорское шоссе 17",
             "date": "2026-09-27", "time_window": ""},
        ],
    }


def gap_payload() -> dict:
    """Маршрут с «дыркой»: пустая точка в середине (как tests/test_contract_generator.py)."""
    payload = base_payload("ООО (с НДС)", "22%", 22)
    payload["contract"] = dict(payload["contract"], number="GAP-2026-1")

    payload["loadings"] = [
        {"address": "Точка 1", "date": "2026-09-30", "time_window": "08:00"},
        {"address": "", "date": "2026-09-30", "time_window": ""},
        {"address": "Точка 3", "date": "2026-09-30", "time_window": ""},
    ]
    payload["unloadings"] = [
        {"address": "Выгрузка 1", "date": "2026-10-02", "time_window": ""},
        {"address": "Выгрузка 2", "date": "2026-10-03", "time_window": ""},
        {"address": "", "date": "2026-10-04", "time_window": ""},
        {"address": "Выгрузка 4", "date": "2026-10-05", "time_window": ""},
    ]
    payload["vehicles"] = [
        {"brand_model": "МОДЕЛЬ 1", "vin": "EC3TEUMB0T0000001",
         "vehicle_type": "Легковой автомобиль",
         "loading_index": 1, "unloading_index": 1},
        {"brand_model": "МОДЕЛЬ 2", "vin": "EC3TEUMB0T0000002",
         "vehicle_type": "Легковой автомобиль",
         "loading_index": 3, "unloading_index": 2},
        {"brand_model": "МОДЕЛЬ 3", "vin": "EC3TEUMB0T0000003",
         "vehicle_type": "Легковой автомобиль",
         "loading_index": 3, "unloading_index": 4},
    ]
    return payload


SCENARIOS = {
    "ooo": ("shablon_ooo.docx",
            base_payload("ООО (с НДС)", "22%", 22)),
    "ip_with_vat": ("shablon_ip_with_vat.docx",
                    base_payload("ИП с НДС", "22%", 22)),
    "ip_without_vat": ("shablon_ip_without_vat.docx",
                       base_payload("ИП без НДС", "0%", 0)),
    "gap": ("shablon_ooo.docx", gap_payload()),
}


# ─────────────────────────────────────────────────────────────
# Сценарии Формики (ЭТАП 3.1.A): свой генератор, свой шаблон.
# Отдельный словарь, потому что golden-тест перевозки жёстко
# использует ContractGenerator, а Формика рендерится FormikaGenerator.
# ─────────────────────────────────────────────────────────────

def formika_payload() -> dict:
    """
    Данные договора-заявки «Формика»: 4 перевозимые машины.

    Тягач и полуприцеп лежат отдельными блоками и в таблицу груза не
    попадают; лишние (пустые) строки таблицы удаляет постобработка.
    """
    return {
        "driver": DRIVER,
        "carrier": dict(ORGANIZATION, entity_type="ООО"),
        "customer": dict(ORGANIZATION, full_name="ООО «Заказчик»",
                         short_name="ООО «Заказчик»"),
        "vehicles": [
            {"vin": "EC3TEUMB0T0000001", "brand_model": "МОДЕЛЬ 1",
             "vehicle_type": "Легковой автомобиль"},
            {"vin": "EC3TEUMB0T0000002", "brand_model": "МОДЕЛЬ 2",
             "vehicle_type": "Легковой автомобиль"},
            {"vin": "EC3TEUMB0T0000003", "brand_model": "МОДЕЛЬ 3",
             "vehicle_type": "Легковой автомобиль"},
            {"vin": "EC3TEUMB0T0000004", "brand_model": "МОДЕЛЬ 4",
             "vehicle_type": "Легковой автомобиль"},
            {"vin": "", "brand_model": "", "vehicle_type": ""},
        ],
        "tractor": TRACTOR,
        "trailer": TRAILER,
        "contract": {
            "number": "ФМ-2026-1", "date": "2026-07-24",
            "route": "г. Воронеж - г. Москва", "carrier_type": "ООО (с НДС)",
            "vat_rate": "22%", "vat_rate_num": 22,
            "price_without_vat": 180300.0, "price_with_vat": 219966.0,
            "loading_plan_date": "2026-07-27",
            "loading_plan_time_from": "", "loading_plan_time_to": "",
        },
        "loadings": [
            {"address": "г. Воронеж, ул. Остужева 52Б",
             "date": "2026-07-27", "time_window": "09:00-15:00"},
        ],
        "unloadings": [
            {"address": "г. Москва, Перерва 19 стр 3",
             "date": "2026-07-30", "time_window": ""},
        ],
    }


FORMIKA_SCENARIOS = {
    "formika_sample": ("shablon_formika.docx", formika_payload()),
}


# ─────────────────────────────────────────────────────────────
# Облегчённая копия шаблона (как в tests/test_contract_generator.py)
# ─────────────────────────────────────────────────────────────

def strip_embedded_fonts(source: Path, target: Path) -> Path:
    """Копия шаблона без встроенных шрифтов и ссылок на них."""
    with zipfile.ZipFile(str(source)) as archive:
        items = [(info, archive.read(info.filename)) for info in archive.infolist()]

    with zipfile.ZipFile(str(target), "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in items:
            name = info.filename
            if name.startswith("word/fonts/"):
                continue
            if name.endswith(".rels"):
                data = re.sub(
                    rb"<Relationship\b[^>]*Target=\"fonts/[^\"]*\"[^>]*/>", b"", data
                )
            if name == "word/settings.xml":
                data = re.sub(rb"<w:embedTrueTypeFonts[^>]*/>", b"", data)
                data = re.sub(rb"<w:saveSubsetFonts[^>]*/>", b"", data)
            out.writestr(info, data)
    return target


# ─────────────────────────────────────────────────────────────
# Структурный отпечаток документа
# ─────────────────────────────────────────────────────────────

def fingerprint(docx_path: Path) -> dict:
    """Детерминированное описание документа для сверки после рефакторинга."""
    doc = Document(str(docx_path))

    tables = []
    for table in doc.tables:
        tables.append({
            "style": table.style.name if table.style is not None else None,
            "rows": [[cell.text for cell in row.cells] for row in table.rows],
            "grid_twips": [int(c.get(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w"
            ) or 0) for c in table._tbl.tblGrid],
        })

    # Порядок элементов тела: абзацы и таблицы вперемешку.
    body = []
    table_index = 0
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p":
            text = "".join(node.text or "" for node in child.iter()
                           if node.tag.endswith("}t"))
            body.append({"kind": "p", "text": text})
        elif tag == "tbl":
            body.append({"kind": "tbl", "index": table_index})
            table_index += 1

    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )

    raw = docx_path.read_bytes()
    with zipfile.ZipFile(str(docx_path)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8", "ignore")
    text_without_tags = re.sub(r"<[^>]+>", "", document_xml)
    placeholders = sorted(set(re.findall(r"\{\{[^{}]*\}\}", text_without_tags)))

    return {
        "paragraphs": [p.text for p in doc.paragraphs],
        "tables": tables,
        "body": body,
        "placeholders_left": placeholders,
        # Справочное значение: docx-контейнер НЕ детерминирован между
        # прогонами (проверено эмпирически, отличается даже у старого
        # кода), поэтому в golden-тесте sha256_docx не сравнивается.
        "sha256_docx": hashlib.sha256(raw).hexdigest(),
        "sha256_text": hashlib.sha256(text.encode("utf-8")).hexdigest(),
    }


# ─────────────────────────────────────────────────────────────
# Основной прогон
# ─────────────────────────────────────────────────────────────

def _make_generator(kind: str, templates_dir: str):
    """
    Генератор сценария: перевозка (по умолчанию) или Формика.

    Реестр типов загружается явно (как в main.py при старте приложения):
    без load_builtin() фабрика не знает про formika и в нестрогом режиме
    молча отдаёт генератор перевозки — эталон тогда снимается не с того
    типа. strict=True превращает такую ошибку в исключение.
    """
    if kind == "formika":
        ContractTypeRegistry.load_builtin()
        return GeneratorFactory.get_generator(
            "formika", templates_dir=templates_dir, strict=True
        )
    return ContractGenerator(templates_dir=templates_dir)


def _render_scenario(scenario, template_name, payload, kind, work) -> dict:
    """Рендерит один сценарий в tests/data/golden/<scenario>.{docx,json}."""
    templates_dir = work / scenario
    templates_dir.mkdir(parents=True, exist_ok=True)
    strip_embedded_fonts(
        PROJECT_ROOT / "templates" / template_name,
        templates_dir / template_name,
    )

    generator = _make_generator(kind, str(templates_dir))
    docx_path = GOLDEN_DIR / f"{scenario}.docx"
    if docx_path.exists():
        docx_path.unlink()

    generator.generate_docx(payload, str(docx_path))

    fingerprint_path = GOLDEN_DIR / f"{scenario}.json"
    fingerprint_path.write_text(
        json.dumps(fingerprint(docx_path), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    size = docx_path.stat().st_size
    result = json.loads(fingerprint_path.read_text(encoding="utf-8"))
    print(f"[OK] {scenario}: {docx_path.name} ({size} байт), "
          f"sha256_text={result['sha256_text'][:16]}..., "
          f"плейсхолдеров осталось: {len(result['placeholders_left'])}")
    return result


def main(argv=None) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except AttributeError:  # старые Python без reconfigure
        pass

    selected = {arg for arg in (sys.argv[1:] if argv is None else argv) if arg}
    known = set(SCENARIOS) | set(FORMIKA_SCENARIOS)
    unknown = selected - known
    if unknown:
        print(f"Неизвестные сценарии: {sorted(unknown)}")
        print(f"Известны: {sorted(known)}")
        return 2

    GOLDEN_DIR.mkdir(parents=True, exist_ok=True)

    work = Path(tempfile.mkdtemp(prefix="make_golden_"))
    try:
        for scenario, (template_name, payload) in SCENARIOS.items():
            if selected and scenario not in selected:
                continue
            _render_scenario(scenario, template_name, payload, "perevozka", work)

        for scenario, (template_name, payload) in FORMIKA_SCENARIOS.items():
            if selected and scenario not in selected:
                continue
            _render_scenario(scenario, template_name, payload, "formika", work)

        readme = GOLDEN_DIR / "README.md"
        readme.write_text(
            "# Золотые эталоны договоров\n\n"
            "Выводы генераторов (Шаг 0 рефакторинга архитектуры контрактов) "
            "на облегчённых копиях реальных шаблонов (без встроенных шрифтов).\n\n"
            "Сценарии перевозки (`ContractGenerator`): `ooo`, `ip_with_vat`, "
            "`ip_without_vat`, `gap` (маршрут с пустой точкой).\n"
            "Сценарий Формики (`FormikaGenerator`, ЭТАП 3.1.A): "
            "`formika_sample` — 4 машины в таблице груза.\n\n"
            "Для каждого сценария: `<сценарий>.docx` — документ, "
            "`<сценарий>.json` — структурный отпечаток для сверки в "
            "`tests/test_contract_generator_golden.py` (перевозка) и "
            "`tests/test_formika_generator.py` (Формика).\n\n"
            "Перегенерация: `python tools/make_golden.py [сценарий ...]`.\n"
            "Без аргументов перегенерируются ВСЕ эталоны.\n\n"
            "Все данные синтетические, реальных ПДн нет.\n",
            encoding="utf-8",
        )
    finally:
        shutil.rmtree(work, ignore_errors=True)

    print(f"\nЭталоны: {GOLDEN_DIR}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
