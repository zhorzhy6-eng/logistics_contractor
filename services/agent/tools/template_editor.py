#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Полуавтоматическая правка шаблонов: предложить → подтвердить → применить.

Главное правило
---------------
Шаблон НИКОГДА не меняется молча. Порядок всегда один:

  1. propose_template_fix() — находит, что можно заменить, и НЕ трогает файл;
  2. оператор видит диалог «было / стало» (ui/template_diff_dialog.py) и
     отмечает, что применять;
  3. apply_template_changes() — делает .bak, меняет ТОЛЬКО отмеченное и
     проверяет результат;
  4. если проверка не прошла — файл восстанавливается из .bak,
     а оператор получает причину.

Почему только runs
------------------
python-docx не умеет менять текст абзаца, не теряя оформление: новый
абзац создаётся с нуля и стили предыдущего не наследует. Поэтому текст
меняется ТОЛЬКО в существующем run (куске текста с одинаковым
оформлением). Если нужный текст разбит на несколько runs, правка
помечается небезопасной и не предлагается: собрать её обратно в один run
можно только переformatировав абзац.

Почему плейсхолдеры неприкосновенны
------------------------------------
Абзац, в котором есть {{…}} или {%…%}, правкой не трогается вообще. Одна
лишняя буква внутри плейсхолдера — и договор не соберётся (или, хуже,
подставит не то поле). Такие абзацы помечаются safe: False с причиной:
их правят руками и через диалог подтверждения с юристом, а не кнопкой.
"""

import logging
import shutil
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from services.agent.tools import template_analyzer as analyzer

logger = logging.getLogger("services.agent.tools.template_editor")

#: Расширение резервной копии: шаблон.docx → шаблон.docx.bak.
BACKUP_SUFFIX = ".bak"

#: Известные замены устаревших текстов: что менять и на что.
#: Ключ — текст как он записан в шаблоне (сравнение без учёта регистра и
#: лишних пробелов), значение — новое написание.
KNOWN_FIXES: Dict[str, str] = {
    "НДС 20%": "НДС 22%",
    "НДС 20 %": "НДС 22%",
    "НДС 18%": "НДС 22%",
    "НДС 18 %": "НДС 22%",
}

#: Сколько правок максимум предлагать за раз: больше — это уже переписывание.
MAX_CHANGES = 50


# ─────────────────────────────────────────────────────────────
# Общее
# ─────────────────────────────────────────────────────────────

def _normalize(text: str) -> str:
    """Текст без лишних пробелов — для сравнения, а не для записи."""
    return " ".join(str(text or "").split())


def backup_path(template_path: Any) -> Path:
    """Путь резервной копии шаблона (шаблон.docx.bak)."""
    path = Path(str(template_path or ""))
    return path.with_name(path.name + BACKUP_SUFFIX)


def make_backup(template_path: Any) -> Optional[Path]:
    """
    Делает .bak рядом с шаблоном.

    Если .bak уже есть (правка идёт второй раз за сеанс), он НЕ
    перезаписывается: первая копия — это состояние ДО всех правок, и именно
    она нужна для отката.

    :return: путь копии или None, если копию сделать не удалось.
    """
    source = Path(str(template_path or ""))
    if not source.is_file():
        return None
    target = backup_path(source)
    if target.is_file():
        return target
    try:
        shutil.copy2(str(source), str(target))
    except OSError as error:
        logger.error("Резервная копия не создана (%s): %s",
                     source.name, type(error).__name__)
        return None
    logger.info("Резервная копия шаблона: %s", target.name)
    return target


def restore_backup(template_path: Any) -> bool:
    """
    Восстанавливает шаблон из .bak.

    :return: True — файл восстановлен; False — копии нет или она не легла.
    """
    source = Path(str(template_path or ""))
    target = backup_path(source)
    if not target.is_file():
        return False
    try:
        shutil.copy2(str(target), str(source))
    except OSError as error:
        logger.error("Откат не удался (%s): %s", source.name,
                     type(error).__name__)
        return False
    logger.warning("Шаблон восстановлен из резервной копии: %s", source.name)
    return True


# ─────────────────────────────────────────────────────────────
# Чтение абзацев DOCX
# ─────────────────────────────────────────────────────────────

def iter_docx_paragraphs(document: Any) -> Iterable[Tuple[int, str, Any]]:
    """
    Абзацы документа по порядку: (номер, место, объект абзаца).

    Порядок тот же, что у анализатора: сначала тело, потом таблицы. Без
    совпадения порядка номер абзаца в диалоге не совпадал бы с номером в
    анализе, и оператор искал бы правку не там.
    """
    index = 0
    for paragraph in document.paragraphs:
        yield index, "body", paragraph
        index += 1
    for table_number, table in enumerate(document.tables, start=1):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield index, f"table {table_number}", paragraph
                    index += 1


def _has_placeholder(text: str) -> bool:
    """Есть ли в абзаце плейсхолдер или тег docxtpl."""
    return bool(analyzer.PLACEHOLDER_RE.search(text or "")
                or analyzer.TAG_RE.search(text or ""))


def _run_with_text(paragraph: Any, old: str) -> Optional[int]:
    """
    Номер run, в котором целиком лежит искомый текст (None — не лежит).

    Именно этот поиск решает, можно ли править безопасно: текст, разбитый
    на несколько runs, собрать обратно без потери оформления нельзя.
    """
    for index, run in enumerate(paragraph.runs):
        if old.casefold() in str(run.text or "").casefold():
            return index
    return None


# ─────────────────────────────────────────────────────────────
# Предложение правки (файл НЕ меняется)
# ─────────────────────────────────────────────────────────────

def _resolve_replacement(issue: str, old: Optional[str],
                         new: Optional[str]) -> Tuple[str, str]:
    """
    Что на что менять: явные старый/новый тексты или известная замена.

    :return: (старый текст, новый текст); новый пустой — замену не знаем.
    """
    if old is not None and new is not None:
        return str(old), str(new)

    key = _normalize(issue)
    for known_old, known_new in KNOWN_FIXES.items():
        if _normalize(known_old).casefold() == key.casefold():
            return known_old, known_new
    return key, ""


def propose_template_fix(template_path: Any, issue: str = "", *,
                         old: Optional[str] = None,
                         new: Optional[str] = None,
                         limit: int = MAX_CHANGES) -> Dict[str, Any]:
    """
    Предлагает правку шаблона, НЕ применяя её.

    :param template_path: файл .docx или .xlsx;
    :param issue: что не так («НДС 20%», «устаревшая ставка»);
    :param old: заменяемый текст (если известен точно);
    :param new: новый текст;
    :param limit: сколько правок максимум предлагать.
    :return: словарь вида::

        {
          "ok": True, "template": "templates/shablon_ooo.docx",
          "issue": "НДС 20%",
          "changes": [
            {"type": "replace", "paragraph_index": 42, "location": "body",
             "run_index": 3, "old": "НДС 20%", "new": "НДС 22%",
             "has_placeholder": False, "safe": True, "reason": ""}
          ],
          "warnings": ["…"],
        }

    Небезопасные места (абзац с плейсхолдером, текст в нескольких runs,
    файл не открылся) остаются в списке с safe: False и причиной — оператор
    должен видеть, что есть и почему это правится руками.
    """
    path = Path(str(template_path or ""))
    old_text, new_text = _resolve_replacement(issue, old, new)
    warnings: List[str] = []
    changes: List[Dict[str, Any]] = []

    if not old_text:
        return {
            "ok": False, "template": str(path), "issue": str(issue),
            "changes": [], "warnings": ["не указано, что менять"],
        }
    if not new_text:
        warnings.append(f"нет известной замены для «{old_text}»: укажите "
                         "новый текст (new)")

    result = analyzer.analyze_template(path)
    if not result.get("ok"):
        return {
            "ok": False, "template": str(path), "issue": str(issue),
            "changes": [], "warnings": list(result.get("warnings") or []),
        }

    suffix = path.suffix.casefold()
    if suffix == ".docx":
        changes = _propose_docx(path, old_text, new_text, limit, warnings)
    elif suffix == ".xlsx":
        changes = _propose_xlsx(path, old_text, new_text, limit, warnings)
    else:
        warnings.append("правка поддерживается только для .docx и .xlsx")

    if not changes:
        warnings.append(f"«{old_text}» в шаблоне не найдено — менять нечего")

    logger.info("Предложена правка шаблона %s: изменений %s (из них "
                "безопасных %s)", path.name, len(changes),
                sum(1 for item in changes if item["safe"]))
    return {
        "ok": bool(changes),
        "template": str(path),
        "issue": str(issue),
        "old": old_text,
        "new": new_text,
        "changes": changes,
        "warnings": warnings,
    }


def _propose_docx(path: Path, old: str, new: str, limit: int,
                  warnings: List[str]) -> List[Dict[str, Any]]:
    """Ищет замены по абзацам DOCX (файл не меняется)."""
    from docx import Document

    document = Document(str(path))
    changes: List[Dict[str, Any]] = []
    for index, location, paragraph in iter_docx_paragraphs(document):
        text = paragraph.text or ""
        if old.casefold() not in text.casefold():
            continue
        if len(changes) >= limit:
            warnings.append(f"правок больше {limit}: остальные не показаны")
            break

        placeholder = _has_placeholder(text)
        run_index = _run_with_text(paragraph, old)
        reason = ""
        if not new:
            # Замена не задана: стирать текст вместо замены нельзя.
            reason = "не указана замена (new)"
        elif placeholder:
            reason = ("в абзаце есть плейсхолдер или тег docxtpl — "
                      "такой абзац правят руками")
        elif run_index is None:
            reason = "текст разбит на несколько фрагментов оформления"
        safe = not reason
        changes.append({
            "type": "replace",
            "paragraph_index": index,
            "location": location,
            "run_index": run_index if run_index is not None else -1,
            "old": old,
            "new": new,
            "paragraph_text": _normalize(text)[:300],
            "has_placeholder": placeholder,
            "safe": safe,
            "reason": reason,
        })
    return changes


def _propose_xlsx(path: Path, old: str, new: str, limit: int,
                  warnings: List[str]) -> List[Dict[str, Any]]:
    """Ищет замены по ячейкам XLSX (файл не меняется)."""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=False)
    changes: List[Dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str):
                        continue
                    if old.casefold() not in value.casefold():
                        continue
                    if len(changes) >= limit:
                        warnings.append(f"правок больше {limit}: остальные не "
                                        "показаны")
                        break
                    placeholder = _has_placeholder(value)
                    if not new:
                        # Замена не задана: стирать текст вместо замены нельзя.
                        reason = "не указана замена (new)"
                    elif placeholder:
                        reason = "в ячейке есть плейсхолдер — правят руками"
                    else:
                        reason = ""
                    changes.append({
                        "type": "replace",
                        "sheet": sheet.title,
                        "cell": cell.coordinate,
                        "old": old,
                        "new": new,
                        "paragraph_text": _normalize(value)[:300],
                        "has_placeholder": placeholder,
                        "safe": not reason,
                        "reason": reason,
                    })
    finally:
        workbook.close()
    return changes


# ─────────────────────────────────────────────────────────────
# Применение правки (только после подтверждения)
# ─────────────────────────────────────────────────────────────

def selected_changes(changes: Sequence[Dict[str, Any]],
                     selected: Optional[Sequence[int]] = None) -> List[Dict[str, Any]]:
    """
    Какие правки применять: отмеченные оператором и только безопасные.

    selected=None означает «все безопасные»: так диалог может применить
    предложение целиком, не перечисляя номера. Небезопасные правки не
    применяются НИКОГДА — даже если оператор их отметил: они помечены
    такими не зря (плейсхолдер, несколько runs).

    :param changes: список предложенных правок;
    :param selected: номера отмеченных правок (None — все безопасные).
    """
    chosen: List[Dict[str, Any]] = []
    for index, change in enumerate(changes or []):
        if not isinstance(change, dict) or not change.get("safe"):
            continue
        if selected is not None and index not in set(selected):
            continue
        chosen.append(dict(change))
    return chosen


def _replace_first(text: str, old: str, new: str) -> Tuple[str, bool]:
    """
    Заменяет ПЕРВОЕ вхождение без учёта регистра.

    Регистр ищется без учёта (в шаблоне бывает «ндс 20%» строчными), а
    заменяется ровно найденный кусок: остальной текст остаётся как был.

    :return: (новый текст, была ли замена).
    """
    source = str(text or "")
    probe = str(old or "")
    if not probe:
        return source, False
    position = source.casefold().find(probe.casefold())
    if position < 0:
        return source, False
    return (source[:position] + str(new) + source[position + len(probe):],
            True)


def _apply_docx(path: Path, chosen: Sequence[Dict[str, Any]],
                ) -> Tuple[int, List[Dict[str, Any]]]:
    """Применяет правки к DOCX: run за run, форматирование сохраняется."""
    from docx import Document

    document = Document(str(path))
    paragraphs = list(iter_docx_paragraphs(document))
    applied = 0
    skipped: List[Dict[str, Any]] = []

    for change in chosen:
        index = int(change.get("paragraph_index", -1))
        if not 0 <= index < len(paragraphs):
            skipped.append({**change, "reason": "абзац не найден"})
            continue
        _, location, paragraph = paragraphs[index]
        text = paragraph.text or ""
        if str(change.get("old", "")).casefold() not in text.casefold():
            skipped.append({**change, "reason": "текста в абзаце уже нет"})
            continue
        if _has_placeholder(text):
            skipped.append({**change, "reason": "в абзаце плейсхолдер"})
            continue

        run_index = int(change.get("run_index", -1))
        runs = paragraph.runs
        if not 0 <= run_index < len(runs):
            skipped.append({**change, "reason": "фрагмент оформления не найден"})
            continue
        run = runs[run_index]
        replaced, done = _replace_first(run.text, change.get("old", ""),
                                        change.get("new", ""))
        if not done:
            skipped.append({**change, "reason": "фрагмент не тот"})
            continue
        run.text = replaced
        applied += 1
        logger.info("Шаблон: замена в %s, абзац %s", location, index + 1)

    if applied:
        document.save(str(path))
    return applied, skipped


def _apply_xlsx(path: Path, chosen: Sequence[Dict[str, Any]],
                ) -> Tuple[int, List[Dict[str, Any]]]:
    """Применяет правки к XLSX: значение ячейки подменяется целиком."""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=False)
    applied = 0
    skipped: List[Dict[str, Any]] = []
    try:
        for change in chosen:
            sheet_name = str(change.get("sheet") or "")
            coordinate = str(change.get("cell") or "")
            if sheet_name not in workbook.sheetnames or not coordinate:
                skipped.append({**change, "reason": "ячейка не найдена"})
                continue
            cell = workbook[sheet_name][coordinate]
            value = cell.value
            if not isinstance(value, str):
                skipped.append({**change, "reason": "значение не текст"})
                continue
            if _has_placeholder(value):
                skipped.append({**change, "reason": "в ячейке плейсхолдер"})
                continue
            replaced, done = _replace_first(value, change.get("old", ""),
                                            change.get("new", ""))
            if not done:
                skipped.append({**change, "reason": "текста в ячейке уже нет"})
                continue
            cell.value = replaced
            applied += 1
        if applied:
            workbook.save(str(path))
    finally:
        workbook.close()
    return applied, skipped


def apply_template_changes(template_path: Any,
                           changes: Sequence[Dict[str, Any]],
                           selected: Optional[Sequence[int]] = None) -> Dict[str, Any]:
    """
    Применяет правки к шаблону: .bak → запись → проверка → (откат).

    Порядок обязателен и не настраивается:

      1. анализ «до» (чтобы знать, что должно остаться);
      2. .bak рядом с шаблоном;
      3. запись только выбранных безопасных правок;
      4. анализ «после» и сравнение (regressions);
      5. любая поломка — восстановление из .bak.

    :param template_path: файл шаблона;
    :param changes: предложенные правки (из propose_template_fix);
    :param selected: номера отмеченных правок (None — все безопасные).
    :return: {"ok", "applied", "skipped", "problems", "rolled_back",
        "backup", "before", "after"}.
    """
    path = Path(str(template_path or ""))
    if not path.is_file():
        return {
            "ok": False, "applied": 0, "skipped": [], "rolled_back": False,
            "backup": "", "problems": [f"шаблона нет: {path}"],
        }

    chosen = selected_changes(changes, selected)
    if not chosen:
        return {
            "ok": False, "applied": 0, "skipped": [], "rolled_back": False,
            "backup": "", "problems": ["не выбрано ни одной безопасной правки"],
        }

    before = analyzer.analyze_template(path)
    backup = make_backup(path)
    if backup is None:
        return {
            "ok": False, "applied": 0, "skipped": [], "rolled_back": False,
            "backup": "",
            "problems": ["резервную копию создать не удалось — правка отменена"],
        }

    try:
        if path.suffix.casefold() == ".docx":
            applied, skipped = _apply_docx(path, chosen)
        elif path.suffix.casefold() == ".xlsx":
            applied, skipped = _apply_xlsx(path, chosen)
        else:
            applied, skipped = 0, list(chosen)
    except Exception as error:  # noqa: BLE001 — файл занят, нет прав и т.п.
        rolled_back = restore_backup(path)
        logger.exception("Правка шаблона не выполнена: %s", error)
        return {
            "ok": False, "applied": 0, "skipped": skipped,
            "rolled_back": rolled_back, "backup": str(backup),
            "problems": [f"запись в шаблон не удалась: {type(error).__name__}"],
        }

    after = analyzer.analyze_template(path)
    problems = analyzer.regressions(before, after)

    if problems:
        rolled_back = restore_backup(path)
        logger.error("Правка шаблона отменена: %s", "; ".join(problems))
        return {
            "ok": False, "applied": 0, "skipped": skipped,
            "rolled_back": rolled_back, "backup": str(backup),
            "problems": problems, "before": before, "after": after,
        }

    logger.info("Шаблон %s изменён: правок %s, резервная копия %s",
                path.name, applied, backup.name)
    return {
        "ok": True, "applied": applied, "skipped": skipped,
        "rolled_back": False, "backup": str(backup), "problems": [],
        "before": before, "after": after,
        "message": (f"Применено правок: {applied}. Резервная копия: "
                    f"{backup.name}."),
    }


def apply_plan(plan: Dict[str, Any],
               selected: Optional[Sequence[int]] = None) -> Dict[str, Any]:
    """
    Применяет предложение целиком (plan из propose_template_fix).

    Удобно диалогу: у него на руках словарь предложения, а не путь с
    отдельным списком правок.
    """
    plan = plan or {}
    return apply_template_changes(plan.get("template", ""),
                                  plan.get("changes") or [], selected)


def format_proposal(plan: Dict[str, Any]) -> str:
    """
    Текст предложения для панели: что меняем, сколько и что придётся руками.

    Оператор должен видеть не только «правка найдена», но и сколько мест
    останется нетронутыми и почему.
    """
    plan = plan or {}
    changes = list(plan.get("changes") or [])
    safe = [item for item in changes if item.get("safe")]
    manual = [item for item in changes if not item.get("safe")]
    lines: List[str] = []
    title = Path(str(plan.get("template") or "")).name
    lines.append(f"Шаблон: {title}")
    old_text = plan.get("old", "") or ""
    new_text = plan.get("new", "") or ""
    if old_text or new_text:
        lines.append(f"Замена: «{old_text}» → «{new_text}»")
    lines.append(f"Найдено мест: {len(changes)} (можно применить "
                  f"автоматически: {len(safe)}).")
    for item in manual:
        place = item.get("paragraph_index", item.get("cell", ""))
        reason = item.get("reason", "") or ""
        lines.append(f"• руками (место {place}): {reason}")
    for warning in plan.get("warnings") or []:
        lines.append(f"! {warning}")
    return "\n".join(lines)


__all__ = [
    "BACKUP_SUFFIX",
    "KNOWN_FIXES",
    "MAX_CHANGES",
    "apply_plan",
    "apply_template_changes",
    "backup_path",
    "format_proposal",
    "iter_docx_paragraphs",
    "make_backup",
    "propose_template_fix",
    "restore_backup",
    "selected_changes",
]


