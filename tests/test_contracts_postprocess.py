# -*- coding: utf-8 -*-
"""
Тесты шагов постобработки (Шаг 5 рефакторинга архитектуры контрактов).

Каждый шаг проверяется отдельно на программно собранном Document() —
без тяжёлых шаблонов templates/*.docx.
"""

import gc

import pytest
from docx import Document

from core.contract_data import ContractData
from core.contracts.base_generator import BaseContractGenerator, ConvertNewlinesStep
from core.contracts.perevozka.generator import PerevozkaGenerator
from core.contracts.perevozka.postprocess import (
    RemoveEmptyVehicleRowsStep,
    RouteTablesStep,
)

NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@pytest.fixture(autouse=True)
def _release_documents():
    yield
    gc.collect()


@pytest.fixture
def generator(tmp_path) -> PerevozkaGenerator:
    """Генератор без реальных шаблонов: шаги их не используют."""
    return PerevozkaGenerator(templates_dir=str(tmp_path))


def _save_doc(doc, path):
    doc.save(str(path))
    return Document(str(path))


# ─────────────────────────────────────────────────────────────
# ConvertNewlinesStep
# ─────────────────────────────────────────────────────────────

def test_convert_newlines_step_turns_newline_into_break(tmp_path):
    gen = PerevozkaGenerator(templates_dir=str(tmp_path))
    doc = Document()
    paragraph = doc.add_paragraph("первая строка\nвторая строка")

    ConvertNewlinesStep(gen).apply(doc, None)

    assert "вторая строка" in paragraph.text
    assert paragraph._p.findall(f".//{NS}br"), "ожидался <w:br/>"


def test_convert_newlines_step_keeps_plain_text(tmp_path):
    gen = PerevozkaGenerator(templates_dir=str(tmp_path))
    doc = Document()
    paragraph = doc.add_paragraph("без переносов")

    ConvertNewlinesStep(gen).apply(doc, None)

    assert not paragraph._p.findall(f".//{NS}br")


# ─────────────────────────────────────────────────────────────
# RouteTablesStep
# ─────────────────────────────────────────────────────────────

def _route_doc_with_marker():
    """Документ с абзацем-меткой LOADING_TABLE_HERE (как после docxtpl)."""
    doc = Document()
    doc.add_paragraph("3.2. Погрузка:")
    doc.add_paragraph("LOADING_TABLE_HERE")
    doc.add_paragraph("3.3. Выгрузка:")
    return doc


def test_route_tables_step_replaces_marker_with_heading_and_table(
    generator, tmp_path
):
    doc = _route_doc_with_marker()
    data = ContractData(
        vehicles=[
            {"brand_model": "JETOUR T2", "vin": "EC3TEUMB0T0002608",
             "vehicle_type": "Легковой автомобиль",
             "loading_index": 1, "unloading_index": 0},
        ],
        loadings=[{"address": "г. Мурманск, ул. Тестовая, д. 1",
                   "date": "2026-09-24", "time_window": "09:00"}],
    )

    RouteTablesStep(generator).apply(doc, data)

    paragraphs = [p.text for p in doc.paragraphs]
    assert "Погрузка 1: г. Мурманск, ул. Тестовая, д. 1" in paragraphs
    assert "LOADING_TABLE_HERE" not in paragraphs

    route_table = doc.tables[0]
    headers = [cell.text for cell in route_table.rows[0].cells]
    assert headers == ["№", "Марка/Модель", "VIN-номер"]
    data_row = [cell.text for cell in route_table.rows[1].cells]
    assert data_row == ["1", "JETOUR T2", "EC3TEUMB0T0002608"]


def test_route_tables_step_without_marker_is_noop(generator):
    doc = Document()
    doc.add_paragraph("обычный текст без меток")

    RouteTablesStep(generator).apply(doc, ContractData())

    assert not doc.tables
    assert [p.text for p in doc.paragraphs] == ["обычный текст без меток"]


# ─────────────────────────────────────────────────────────────
# RemoveEmptyVehicleRowsStep
# ─────────────────────────────────────────────────────────────

def _vehicle_table_doc():
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    for cell, text in zip(table.rows[0].cells, ("№", "Марка/Модель", "VIN-номер")):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ("1", "JETOUR T2", "EC3TEUMB0T0002608")):
        cell.text = text
    # третья строка — пустая (как незаполненная строка шаблона)
    return doc


def test_remove_empty_vehicle_rows_step_removes_empty_row(generator, tmp_path):
    doc = _vehicle_table_doc()
    assert len(doc.tables[0].rows) == 3

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    table = doc.tables[0]
    assert len(table.rows) == 2
    assert [cell.text for cell in table.rows[1].cells] == [
        "1", "JETOUR T2", "EC3TEUMB0T0002608",
    ]


def test_remove_empty_vehicle_rows_step_ignores_foreign_tables(generator, tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "что-то другое"
    table.rows[1].cells[0].text = "данные"

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 2  # чужая таблица не тронута


# ─────────────────────────────────────────────────────────────
# Состав конвейера
# ─────────────────────────────────────────────────────────────

def test_perevozka_steps_with_data(generator):
    steps = generator.postprocess_steps(ContractData())
    assert [step.name for step in steps] == [
        "convert_newlines", "route_tables", "remove_empty_vehicle_rows",
        "normalize_spaces",
    ]


def test_perevozka_steps_without_data_keeps_old_behaviour(generator):
    """Без данных таблицы маршрута не строятся — как в старом коде."""
    steps = generator.postprocess_steps(None)
    assert [step.name for step in steps] == [
        "convert_newlines", "remove_empty_vehicle_rows", "normalize_spaces",
    ]


def test_base_generator_default_steps(tmp_path):
    base = BaseContractGenerator(templates_dir=str(tmp_path))
    steps = base.postprocess_steps(ContractData())
    assert [step.name for step in steps] == ["convert_newlines"]
