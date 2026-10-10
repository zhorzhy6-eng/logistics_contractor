#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты ставок НДС (шаг «Ставки НДС в UI + форма ООО/ИП»).

Что проверяется:

  * список ставок — «Без НДС», «0%», «5%», «7%», «10%», «22%», по умолчанию
    базовая 22 %;
  * пересчёт сумм от базы без НДС для КАЖДОЙ ставки (1000 ₽: 22 % → 220 и
    1220, 5 % → 50 и 1050, «Без НДС» и «0%» → 0 и 1000);
  * `compute_carrier_type`: форма и ставка → вид перевозчика для бланка;
  * загрузка старых записей: есть только `carrier_type` (без `entity_type`
    и `vat_rate`) — форма и ставка восстанавливаются из него;
  * вкладка «Договор»: значения по умолчанию, список из шести ставок,
    пересчёт при смене ставки и форма в данных;
  * генератор: вид перевозчика вычисляется, если поля `carrier_type` в
    данных нет, и «Без НДС» печатается словами, а «0%» — ставкой.

Числа ставок — по Федеральному закону от 28.11.2025 № 425-ФЗ (с 01.01.2026).

Данные синтетические, ПДн нет; шаблоны только читаются.
"""

import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest

from core.vat import (
    DEFAULT_VAT_RATE, ENTITY_TYPES, VAT_FREE, VAT_RATES,
    compute_carrier_type, compute_vat, is_vat_free, normalize_entity_type,
    normalize_vat_rate, split_carrier_type, vat_rate_label, vat_rate_number,
)

PROJECT_ROOT = Path(__file__).resolve().parent.parent

#: Требуемый набор и порядок пунктов списка ставок.
EXPECTED_RATES = ("Без НДС", "0%", "5%", "7%", "10%", "22%")

#: 1000 ₽ без НДС и ожидаемые (НДС, итог) для каждой ставки.
VAT_CASES = (
    (VAT_FREE, 0.0, 1000.0),
    ("0%", 0.0, 1000.0),
    ("5%", 50.0, 1050.0),
    ("7%", 70.0, 1070.0),
    ("10%", 100.0, 1100.0),
    ("22%", 220.0, 1220.0),
)


# ─────────────────────────────────────────────────────────────
# Ядро: список ставок и пересчёт
# ─────────────────────────────────────────────────────────────

def test_rate_list_matches_law():
    """Список ставок — ровно шесть пунктов в порядке от «без налога»."""
    assert VAT_RATES == EXPECTED_RATES
    assert list(VAT_RATES) == list(EXPECTED_RATES)


def test_default_rate_is_base_22():
    """По умолчанию — базовая ставка 22 %."""
    assert DEFAULT_VAT_RATE == "22%"
    assert DEFAULT_VAT_RATE in VAT_RATES


@pytest.mark.parametrize("rate, nds, total", VAT_CASES)
def test_compute_vat_for_every_rate(rate, nds, total):
    """1000 ₽ без НДС: налог и итог для каждой ставки списка."""
    assert compute_vat(1000, rate) == (nds, total)


def test_vat_free_is_not_zero_rate():
    """«Без НДС» и «0%» — разные пункты: налог нулевой у обоих, смысл разный."""
    assert VAT_FREE != "0%"
    assert is_vat_free(VAT_FREE) is True
    assert is_vat_free("0%") is False
    assert vat_rate_number(VAT_FREE) == 0.0
    assert vat_rate_number("0%") == 0.0
    assert vat_rate_label("0%") == "0%"


@pytest.mark.parametrize("value, expected", [
    ("22%", 22.0), ("22", 22.0), (" 5 %", 5.0), ("0%", 0.0),
    (VAT_FREE, 0.0), ("без ндс", 0.0), ("10", 10.0),
])
def test_rate_number_from_any_spelling(value, expected):
    """Ставка числом принимается и со знаком процента, и без него."""
    assert vat_rate_number(value) == expected


@pytest.mark.parametrize("value, expected", [
    ("22%", "22%"), ("22", "22%"), ("5 %", "5%"), (VAT_FREE, VAT_FREE),
    ("НДС не облагается", VAT_FREE), ("20%", ""), ("мусор", ""), ("", ""),
])
def test_rate_normalization_for_the_form(value, expected):
    """Ставка из данных приводится к пункту списка; чужое — пустая строка."""
    assert normalize_vat_rate(value) == expected


def test_unknown_rate_does_not_become_zero():
    """Мусор в ставке не превращается в ноль: берётся базовая ставка."""
    assert vat_rate_number("мусор") == 22.0
    assert vat_rate_number("") == 22.0
    assert vat_rate_label("мусор") == "22%"


# ─────────────────────────────────────────────────────────────
# Вид перевозчика: форма + ставка
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("entity_type, rate, expected", [
    ("ООО", "22%", "ООО (с НДС)"),
    ("ООО", "5%", "ООО (с НДС)"),
    ("ООО", "0%", "ООО (с НДС)"),
    ("ООО", VAT_FREE, "ООО (без НДС)"),
    ("ИП", "22%", "ИП с НДС"),
    ("ИП", "7%", "ИП с НДС"),
    ("ИП", "0%", "ИП с НДС"),
    ("ИП", VAT_FREE, "ИП без НДС"),
])
def test_compute_carrier_type(entity_type, rate, expected):
    """Вид перевозчика: ООО без НДС и ИП без НДС — разные значения."""
    assert compute_carrier_type(entity_type, rate) == expected


def test_ooo_without_vat_is_a_real_case():
    """ООО на УСН с доходом до 20 млн ₽ — «ООО (без НДС)», а не ошибка."""
    assert compute_carrier_type("ООО", VAT_FREE) == "ООО (без НДС)"
    assert compute_carrier_type("ООО", "5%") == "ООО (с НДС)"


def test_empty_entity_type_is_treated_as_ooo():
    """Форма не пришла — вид считается как у ООО (прежнее поведение)."""
    assert compute_carrier_type("", "22%") == "ООО (с НДС)"
    assert compute_carrier_type(None, VAT_FREE) == "ООО (без НДС)"


@pytest.mark.parametrize("value, expected", [
    ("ООО", "ООО"), ("ИП", "ИП"), ("ООО (с НДС)", "ООО"),
    ("ИП без НДС", "ИП"), ("Индивидуальный предприниматель", "ИП"),
    ("мусор", ""), ("", ""),
])
def test_normalize_entity_type(value, expected):
    """Форма стороны распознаётся и в записи вида «ИП без НДС»."""
    assert normalize_entity_type(value) == expected
    assert expected in ENTITY_TYPES or expected == ""


# ─────────────────────────────────────────────────────────────
# Старые записи: только carrier_type
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("carrier_type, entity_type, rate", [
    ("ООО (с НДС)", "ООО", "22%"),
    ("ИП с НДС", "ИП", "22%"),
    ("ИП без НДС", "ИП", VAT_FREE),
    ("ООО (без НДС)", "ООО", VAT_FREE),
])
def test_legacy_carrier_type_restores_form_and_rate(carrier_type, entity_type, rate):
    """Старая запись без entity_type/vat_rate восстанавливается из carrier_type."""
    assert split_carrier_type(carrier_type) == (entity_type, rate)


@pytest.mark.parametrize("carrier_type", ["", "   ", "Прочее"])
def test_unknown_carrier_type_gives_nothing(carrier_type):
    """Непонятный вид не подставляет ни форму, ни ставку."""
    assert split_carrier_type(carrier_type) == ("", "")


@pytest.mark.parametrize("carrier_type", [
    "ООО (с НДС)", "ООО (без НДС)", "ИП с НДС", "ИП без НДС",
])
def test_legacy_round_trip(carrier_type):
    """Вид → форма+ставка → вид: преобразование обратимо."""
    entity_type, rate = split_carrier_type(carrier_type)
    assert compute_carrier_type(entity_type, rate) == carrier_type


# ─────────────────────────────────────────────────────────────
# Вкладка «Договор»
# ─────────────────────────────────────────────────────────────

pytest.importorskip("PyQt5")

from PyQt5.QtWidgets import QApplication  # noqa: E402

from ui.tabs.contract_tab import ContractTab  # noqa: E402


@pytest.fixture(scope="module")
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture
def tab(qapp, isolated_db) -> ContractTab:
    return ContractTab()


def test_tab_has_six_rates_in_order(tab):
    """Выпадающий список ставок — те же шесть пунктов и в том же порядке."""
    rates = [tab.vat_rate.itemText(i) for i in range(tab.vat_rate.count())]

    assert rates == list(EXPECTED_RATES)


def test_tab_defaults_are_ooo_and_22(tab):
    """По умолчанию: форма ООО, ставка 22 %."""
    assert tab.entity_type.currentText() == "ООО"
    assert tab.vat_rate.currentText() == DEFAULT_VAT_RATE


@pytest.mark.parametrize("rate, nds, total", VAT_CASES)
def test_tab_recalculates_for_every_rate(tab, rate, nds, total):
    """Смена ставки пересчитывает НДС и итог (база — 1000 ₽)."""
    tab.price_input.setValue(1000)
    tab.vat_rate.setCurrentText(rate)

    assert tab.vat_amount.text() == f"{nds:.2f} ₽"
    assert tab.price_with_vat.text() == f"{total:.2f} ₽"


def test_tab_get_data_carries_form_rate_and_computed_type(tab):
    """get_data(): форма, ставка и вычисленный по ним вид перевозчика."""
    tab.entity_type.setCurrentText("ИП")
    tab.vat_rate.setCurrentText(VAT_FREE)

    data = tab.get_data()

    assert data["entity_type"] == "ИП"
    assert data["vat_rate"] == VAT_FREE
    assert data["vat_rate_num"] == 0.0
    assert data["carrier_type"] == "ИП без НДС"


@pytest.mark.parametrize("rate, carrier_type", [
    ("22%", "ООО (с НДС)"),
    ("5%", "ООО (с НДС)"),
    ("0%", "ООО (с НДС)"),
    (VAT_FREE, "ООО (без НДС)"),
])
def test_tab_carrier_type_for_ooo(tab, rate, carrier_type):
    """Форма ООО: вид меняется только на «без НДС»."""
    tab.vat_rate.setCurrentText(rate)

    assert tab.get_data()["carrier_type"] == carrier_type


def test_tab_accepts_legacy_carrier_type(tab):
    """Старая запись (только carrier_type) заполняет форму и ставку."""
    tab.fill_data({"carrier_type": "ИП без НДС"})

    assert tab.entity_type.currentText() == "ИП"
    assert tab.vat_rate.currentText() == VAT_FREE
    assert tab.get_data()["carrier_type"] == "ИП без НДС"


def test_tab_explicit_fields_win_over_legacy(tab):
    """Явные entity_type и vat_rate важнее старого carrier_type."""
    tab.fill_data({"carrier_type": "ИП без НДС", "entity_type": "ООО",
                   "vat_rate": "5%"})

    assert tab.entity_type.currentText() == "ООО"
    assert tab.vat_rate.currentText() == "5%"
    assert tab.get_data()["carrier_type"] == "ООО (с НДС)"


def test_tab_stores_price_without_vat(tab):
    """«Стоимость» — база без НДС: в данных она равна введённому числу."""
    tab.price_input.setValue(1000)

    data = tab.get_data()

    assert data["price_without_vat"] == 1000.0
    assert data["price_with_vat"] == 1220.0
    assert data["vat_amount"] == 220.0


def test_tab_clear_returns_defaults(tab):
    """Очистка возвращает форму ООО и ставку 22 %."""
    tab.entity_type.setCurrentText("ИП")
    tab.vat_rate.setCurrentText("10%")

    tab.clear()

    assert tab.entity_type.currentText() == "ООО"
    assert tab.vat_rate.currentText() == DEFAULT_VAT_RATE


def test_tab_apply_entity_type_from_the_reference(tab):
    """Вид из справочника ставится в «Форму» и не трогает ставку."""
    tab.vat_rate.setCurrentText("5%")

    assert tab.apply_entity_type("ИП") is True
    assert tab.entity_type.currentText() == "ИП"
    assert tab.vat_rate.currentText() == "5%"
    assert tab.apply_entity_type("мусор") is False


# ─────────────────────────────────────────────────────────────
# Генератор: вид вычисляется, суммы совпадают с формой
# ─────────────────────────────────────────────────────────────

from core.contract_data import ContractData  # noqa: E402
from core.contract_generator import ContractGenerator  # noqa: E402

TEMPLATES = PROJECT_ROOT / "templates"

CARRIER_OOO = {
    "full_name": "ООО «Ромашка»", "short_name": "ООО «Ромашка»",
    "inn": "7701234567", "kpp": "770101001", "ogrn": "1027700132195",
    "entity_type": "ООО",
}
CUSTOMER = {"full_name": "ООО «Заказчик»", "short_name": "ООО «Заказчик»"}
CONTRACT_BASE = {
    "number": "НДС-ТЕСТ", "date": "2026-10-10", "route": "Москва - Тверь",
    "price_without_vat": 1000.0, "payment_days": 10,
}


@pytest.fixture(scope="module")
def generator() -> ContractGenerator:
    return ContractGenerator(templates_dir=str(TEMPLATES))


def _payload(**contract) -> dict:
    return {
        "carrier": dict(CARRIER_OOO),
        "customer": dict(CUSTOMER),
        "contract": dict(CONTRACT_BASE, **contract),
    }


TEMPLATE_CASES = (
    ("ООО (с НДС)", "22%", "shablon_ooo.docx"),
    ("ООО (без НДС)", VAT_FREE, "shablon_ooo.docx"),
    ("ИП с НДС", "5%", "shablon_ip_with_vat.docx"),
    ("ИП без НДС", VAT_FREE, "shablon_ip_without_vat.docx"),
)


@pytest.mark.parametrize("carrier_type, rate, filename", TEMPLATE_CASES)
def test_template_follows_carrier_type(generator, carrier_type, rate, filename):
    """Бланк выбирается по виду перевозчика, вид — по форме и ставке."""
    payload = _payload(carrier_type=carrier_type, vat_rate=rate)

    assert generator._get_template_path(carrier_type).endswith(filename)
    assert generator.get_template_path(
        ContractData.coerce(payload)
    ).endswith(filename)


@pytest.mark.parametrize("rate, nds, total", VAT_CASES)
def test_replacements_match_the_form(generator, rate, nds, total):
    """Суммы в карте замен — те же, что считает вкладка для этой ставки."""
    payload = _payload(
        carrier_type=compute_carrier_type("ООО", rate), vat_rate=rate,
    )

    replacements = generator._build_replacements_map(payload)

    assert replacements["sum_wo_nds"] == "1000.00"
    assert replacements["sum_total"] == f"{total:.2f}"
    if nds:
        assert replacements["sum_nds"] == f"{nds:.2f}"
        assert replacements["vat_rate"] == rate
        assert replacements["is_vat_free"] is False
    else:
        assert replacements["sum_total"] == replacements["sum_wo_nds"]


def test_replacements_mark_vat_free_only_for_the_free_rate(generator):
    """Флаг условного блока: True — только у «Без НДС», не у «0%»."""
    free = generator._build_replacements_map(
        _payload(carrier_type="ООО (без НДС)", vat_rate=VAT_FREE)
    )
    zero = generator._build_replacements_map(_payload(vat_rate="0%"))

    assert free["is_vat_free"] is True
    assert free["nds_text"] == "НДС не облагается"
    assert free["sum_nds"] == ""

    assert zero["is_vat_free"] is False
    assert zero["sum_nds"] == "0.00"
    assert zero["vat_rate"] == "0%"
    assert "0%" in zero["nds_text"]


@pytest.mark.parametrize("entity_type, rate, expected", [
    ("ИП", "7%", "ИП с НДС"),
    ("ООО", VAT_FREE, "ООО (без НДС)"),
])
def test_carrier_type_is_computed_when_field_is_missing(
    generator, entity_type, rate, expected
):
    """Без поля carrier_type вид вычисляется из формы и ставки."""
    payload = _payload(entity_type=entity_type, vat_rate=rate)
    carrier = dict(CARRIER_OOO)
    if entity_type == "ИП":
        carrier["full_name"] = "Индивидуальный предприниматель Тестов Т. Т."
        carrier["entity_type"] = "ИП"
    payload["carrier"] = carrier

    contract_data = ContractData.coerce(payload)

    assert generator._carrier_type_of(contract_data) == expected


def test_ooo_without_vat_uses_ooo_template_and_free_text(generator, work_file):
    """ООО без НДС: бланк ООО, в п. 4.1 — «НДС не облагается»."""
    payload = _payload(carrier_type="ООО (без НДС)", vat_rate=VAT_FREE)
    out = work_file("vat_free.docx")

    generator.generate_docx(payload, str(out))

    text = _document_text(out)
    assert "НДС не облагается (упрощённая система налогообложения)." in text
    assert "НДС по ставке, действующей на дату оказания услуг" not in text


def test_ooo_with_vat_keeps_the_rate_line(generator, work_file):
    """ООО с НДС 5 %: в п. 4.1 печатается ставка и сумма налога."""
    payload = _payload(carrier_type="ООО (с НДС)", vat_rate="5%")
    out = work_file("vat_5.docx")

    generator.generate_docx(payload, str(out))

    text = _document_text(out)
    assert "НДС по ставке, действующей на дату оказания услуг" in text
    assert "(в настоящее время 5%)" in text
    assert "НДС не облагается" not in text


def _document_text(path) -> str:
    """Весь текст документа одной строкой (абзацы + таблицы)."""
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)
