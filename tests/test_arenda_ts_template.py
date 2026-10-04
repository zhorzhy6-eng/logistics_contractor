#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты шаблонов договора аренды ТС с экипажем (ЭТАП 3.1.D.A.1).

Проверяют, что все три шаблона —

    templates/shablon_arenda_ts_ooo.docx             (Арендатор ООО)
    templates/shablon_arenda_ts_ip_with_vat.docx     (Арендатор ИП, с НДС)
    templates/shablon_arenda_ts_ip_without_vat.docx  (Арендатор ИП, без НДС)

— именно ПУСТЫЕ БЛАНКИ с плейсхолдерами docxtpl, а не копии образца
с данными:

  * файл существует и открывается python-docx;
  * все ключевые плейсхолдеры на месте, набор имён ровно ожидаемый,
    ни один плейсхолдер не разорван между runs;
  * таблица автомобилей п. 3.1: 5 колонок («№ / Марка, модель /
    VIN-номер / Точка погрузки / Точка выгрузки») и 12 строк с
    плейсхолдерами (всего 13 строк с шапкой);
  * до 10 точек погрузки (адрес, дата, время с/по) и до 10 точек
    выгрузки (адрес, дата);
  * у ООО и ИП с НДС три суммы (sum_wo_vat / vat_rate + sum_vat /
    sum_total) и все три прописью; у ИП без НДС только sum_total
    и sum_total_words, а плейсхолдеров НДС нет;
  * КПП есть только в ООО-варианте (раздел 9), у ИП его нет;
  * Приложение № 1 (Акт приема-передачи и возврата ТС) есть во всех трёх
    шаблонах: шапка с номером и датой договора, таблицы «Передача ТС»
    и «Возврат ТС», таблица подписей;
  * в бланках нет данных образца (марок, VIN, ФИО, госномеров, адресов,
    сумм, ИНН/ОГРН, банковских реквизитов, номера ТЛ-574);
  * оформление повторяет образец: Letter, поля 2,0/1,7 см,
    Times New Roman (10,5 pt основной текст);
  * образец-источник не изменён (сверка по SHA256);
  * docxtpl рендерит шаблон без остатка плейсхолдеров;
  * сборщик tools/make_arenda_ts_template.py даёт ровно тот же текст.

Образец
(templates/Договор_аренды_ТС_с_экипажем_ТЛ-574_Технологистика_ЛЦ_обновленный.docx)
— только источник структуры, формулировок и оформления: тест
test_sample_document_untouched намеренно падает, если его отредактировали
или пересобрали.
"""

import gc
import hashlib
import re

import pytest
from docx import Document

#: Три варианта арендатора и их файлы — те же ключи, что у перевозки.
TEMPLATE_NAMES = {
    "ООО": "shablon_arenda_ts_ooo.docx",
    "ИП с НДС": "shablon_arenda_ts_ip_with_vat.docx",
    "ИП без НДС": "shablon_arenda_ts_ip_without_vat.docx",
}

CARRIER_TYPES = tuple(TEMPLATE_NAMES)

#: Варианты, в которых арендная плата облагается НДС (три суммы).
VAT_VARIANTS = ("ООО", "ИП с НДС")

#: SHA256 собранных шаблонов (ЭТАП 3.1.D.A.1).
#: Если шаблон пересобрали осознанно (например, поменяли формулировку),
#: значения нужно обновить — тест ловит ручную правку .docx мимо сборщика.
TEMPLATE_SHA256 = {
    "ООО": "aa9e8580f5c7b97770d19ff7f7378da6ad257ef4595b3f22202a1472e16f5bba",
    "ИП с НДС": "d4227d7156b104c4e31eb1e12a6ae57f0df5832393c190f5a8d66089d00ecbd4",
    "ИП без НДС": "4f6936c981dfcd68b170c334378b784e027d470cec7713307545f10f54f9075a",
}

#: Образец-источник и его SHA256 на момент сборки шаблонов.
SAMPLE_NAME = (
    "Договор_аренды_ТС_с_экипажем_ТЛ-574_Технологистика_ЛЦ_обновленный.docx"
)
SAMPLE_SHA256 = (
    "0d9f25c9e589fe8a931efeb25ef8cd16e03f7d54196fb9db6ee26bed8fde10a9"
)

#: Заголовки таблицы автомобилей п. 3.1 — как в задании к этапу.
CAR_HEADERS = ("№", "Марка, модель", "VIN-номер", "Точка погрузки",
               "Точка выгрузки")

#: Ширины колонок таблицы автомобилей в twips — как в образце.
CAR_COLUMN_WIDTHS = ("453", "1360", "2382", "1984", "2552")

CAR_ROWS = 12
MAX_POINTS = 10

#: Плейсхолдеры, общие для всех трёх вариантов.
COMMON_PLACEHOLDERS = (
    "contract_number",
    "contract_date",
    "lessee_full_name",
    "lessee_short_name",
    "lessee_director_position",
    "lessee_director_name",
    "lessee_basis",
    "lessor_full_name",
    "lessor_short_name",
    "lessor_inn",
    "lessor_ogrn",
    "lessor_director_position",
    "lessor_director_name",
    "lessor_basis",
    "tractor_brand",
    "tractor_plate",
    "tractor_type",
    "trailer_brand",
    "trailer_plate",
    "lease_start_date",
    "lease_end_date",
    "cargo_count",
    "route",
    "driver_full_name",
    "driver_birth_date",
    "driver_passport",
    "driver_passport_issuer",
    "driver_passport_issue_date",
    "driver_license",
    "driver_license_issue_date",
    "driver_address",
    "driver_phone",
    "sum_total",
    "sum_total_words",
    "lessee_inn",
    "lessee_ogrn_label",
    "lessee_ogrn",
    "lessee_address",
    "lessee_account",
    "lessee_bank",
    "lessee_bik",
    "lessee_corr_account",
    "lessee_email",
    "lessee_edo",
    "lessee_director_short",
    "lessor_ogrn_label",
    "lessor_address",
    "lessor_account",
    "lessor_bank",
    "lessor_bik",
    "lessor_corr_account",
    "lessor_email",
    "lessor_edo",
    "lessor_director_short",
)

#: Плейсхолдеры, которых нет у ИП без НДС (арендная плата без НДС).
VAT_PLACEHOLDERS = (
    "sum_wo_vat", "sum_wo_vat_words", "vat_rate", "sum_vat", "sum_vat_words",
)

#: Данные образца, которых в бланках быть не должно.
FORBIDDEN_FRAGMENTS = (
    # марки тягача, прицепа и перевозимых автомобилей
    "SITRAK", "LUXUDA", "Haval", "M6",
    # VIN из образца
    "EC2EF4A5",
    # ФИО (водитель, директора)
    "Шамин", "Сергей Александрович", "Ахмедов", "Тимур", "Чаговец", "Иван",
    # наименования сторон образца
    "ТЕХНОЛОГИСТИКА", "Логистический Центр", "ООО «ЛЦ»",
    # государственные регистрационные знаки
    "Р081ХО", "АО268770",
    # адреса и точки маршрута образца
    "Росва", "ПСМА", "Калуж", "Чехов", "Магнитогорск", "Челябинск",
    "Екатеринбург", "Космонавтов", "Копейское", "УРАЛ БЭСТ", "ЦС-Моторс",
    "Леваневского",
    # суммы и даты образца
    "188 524,59", "41 475,41", "230 000,00", "41 475",
    "17.09.2026", "28.09.2026", "26.09.2026", "16.09.2026", "15.09.2026",
    # номер договора образца
    "ТЛ-574",
    # ИНН и ОГРН образца
    "2310238790", "1242300061068", "9709112631", "1247700454792",
    # банковские реквизиты и контакты образца
    "40702810610001632504", "40702810610001720490", "044525974",
    "30101810145250000974", "logistika.tehnologistika", "centrallogo23",
    "2AE41A857C4", "2AEFCFA386A",
    # документы водителя образца
    "03 20 771606", "99 07 349327", "670-69-35",
)

#: Плейсхолдеры, которые обязаны встретиться в Приложении № 1.
APPENDIX_PLACEHOLDERS = (
    "contract_number",
    "contract_date",
    "tractor_brand",
    "tractor_plate",
    "trailer_brand",
    "trailer_plate",
    "driver_full_name",
    "lessee_short_name",
    "lessee_director_short",
    "lessor_short_name",
    "lessor_director_short",
)

PLACEHOLDER_RE = re.compile(r"\{\{[^{}]*\}\}")

#: Пространство имён WordprocessingML (для чтения w:tblGrid и w:shd).
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@pytest.fixture(autouse=True)
def _release_documents():
    """Word-файлы на Windows освобождаем до удаления временных копий."""
    yield
    gc.collect()


@pytest.fixture(params=CARRIER_TYPES)
def variant(request) -> str:
    """Вариант арендатора: «ООО» / «ИП с НДС» / «ИП без НДС»."""
    return request.param


@pytest.fixture
def template_path(templates_dir, variant):
    path = templates_dir / TEMPLATE_NAMES[variant]
    assert path.exists(), f"нет шаблона аренды ТС: {path}"
    return path


@pytest.fixture(scope="module")
def templates(templates_dir):
    """Все три шаблона: {вариант: Document}."""
    return {
        key: Document(str(templates_dir / name))
        for key, name in TEMPLATE_NAMES.items()
    }


@pytest.fixture
def template_doc(templates, variant):
    return templates[variant]


# ─────────────────────────────────────────────────────────────
# Вспомогательные функции
# ─────────────────────────────────────────────────────────────

def _expected_placeholders(variant: str) -> set:
    """Полный набор имён плейсхолдеров шаблона этого варианта."""
    names = set(COMMON_PLACEHOLDERS)
    names |= {
        f"car_{number}_{field}"
        for number in range(1, CAR_ROWS + 1)
        for field in ("brand", "vin", "loading_point", "unloading_point")
    }
    names |= {
        f"loading_{number}_{field}"
        for number in range(1, MAX_POINTS + 1)
        for field in ("address", "date", "time_from", "time_to")
    }
    names |= {
        f"unloading_{number}_{field}"
        for number in range(1, MAX_POINTS + 1)
        for field in ("address", "date")
    }
    if variant == "ООО":
        names.add("lessee_kpp")
    if variant in VAT_VARIANTS:
        names |= set(VAT_PLACEHOLDERS)
    return names


def _document_text(doc) -> str:
    """Весь текст документа: абзацы, таблицы (включая вложенные), колонтитулы."""
    parts = [p.text for p in doc.paragraphs]

    def table_text(table):
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
                for nested in cell.tables:
                    table_text(nested)

    for table in doc.tables:
        table_text(table)

    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)

    return "\n".join(parts)


def _all_paragraphs(doc):
    """Все абзацы документа, включая абзацы таблиц и колонтитулов."""
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _body_texts(doc):
    """Тексты абзацев верхнего уровня (без таблиц), без крайних пробелов."""
    return [p.text.strip() for p in doc.paragraphs]


def _document_placeholders(doc) -> list:
    """Плейсхолдеры документа в порядке появления."""
    return PLACEHOLDER_RE.findall(_document_text(doc))


def _table_by_headers(doc, headers):
    """Таблица по заголовкам первой строки, а не по индексу."""
    for table in doc.tables:
        if not table.rows:
            continue
        if tuple(cell.text.strip() for cell in table.rows[0].cells) == headers:
            return table
    return None


def _car_table(doc):
    """Таблица автомобилей п. 3.1."""
    return _table_by_headers(doc, CAR_HEADERS)


def _act_table(doc, label: str):
    """Таблица Акта по тексту первой ячейки (метка поля)."""
    for table in doc.tables:
        if not table.rows:
            continue
        if table.rows[0].cells[0].text.strip() == label:
            return table
    return None


def _two_column_row_table(doc, left_prefix, right_prefix):
    """Таблица 1×2, у которой левая ячейка начинается с left_prefix."""
    for table in doc.tables:
        if len(table.rows) != 1 or len(table.columns) != 2:
            continue
        left, right = table.rows[0].cells
        if left.text.startswith(left_prefix) and right.text.startswith(
                right_prefix):
            return table
    return None


def _requisites_table(doc):
    """Таблица раздела 9 с блоками реквизитов сторон."""
    return _two_column_row_table(doc, "АРЕНДАТОР:", "АРЕНДОДАТЕЛЬ:")


def _appendix_signature_table(doc):
    """Таблица подписей Акта: 1×2 с кратким именем Арендатора в левой ячейке."""
    for table in doc.tables:
        if len(table.rows) != 1 or len(table.columns) != 2:
            continue
        left, right = table.rows[0].cells
        if ("{{lessee_short_name}}" in left.text
                and "{{lessor_short_name}}" in right.text
                and "{{lessee_inn}}" not in left.text):
            return table
    return None


def _index_of_starting_with(texts, prefix: str):
    """Индекс первого абзаца, начинающегося с prefix (или None)."""
    for index, text in enumerate(texts):
        if text.startswith(prefix):
            return index
    return None


# ─────────────────────────────────────────────────────────────
# Файл и плейсхолдеры
# ─────────────────────────────────────────────────────────────

def test_template_file_exists_and_opens(template_path, template_doc):
    assert template_path.stat().st_size > 0
    assert template_doc.paragraphs, "в бланке нет ни одного абзаца"
    assert template_doc.tables, "в бланке нет ни одной таблицы"


def test_common_placeholders_present(template_doc, variant):
    text = _document_text(template_doc)
    missing = [name for name in COMMON_PLACEHOLDERS
               if "{{" + name + "}}" not in text]
    assert not missing, f"в шаблоне {variant} нет плейсхолдеров: {missing}"


def test_car_placeholders_present(template_doc, variant):
    """Все 12 машин: марка, VIN, точка погрузки и точка выгрузки."""
    text = _document_text(template_doc)
    missing = [
        f"{{{{car_{number}_{field}}}}}"
        for number in range(1, CAR_ROWS + 1)
        for field in ("brand", "vin", "loading_point", "unloading_point")
        if f"{{{{car_{number}_{field}}}}}" not in text
    ]
    assert not missing, f"нет плейсхолдеров машин: {missing}"


def test_loading_point_placeholders_present(template_doc, variant):
    """Все 10 точек погрузки: адрес, дата и окно времени подачи ТС."""
    text = _document_text(template_doc)
    missing = [
        f"{{{{loading_{number}_{field}}}}}"
        for number in range(1, MAX_POINTS + 1)
        for field in ("address", "date", "time_from", "time_to")
        if f"{{{{loading_{number}_{field}}}}}" not in text
    ]
    assert not missing, f"нет плейсхолдеров точек погрузки: {missing}"


def test_unloading_point_placeholders_present(template_doc, variant):
    """Все 10 точек выгрузки: адрес и плановая дата завершения."""
    text = _document_text(template_doc)
    missing = [
        f"{{{{unloading_{number}_{field}}}}}"
        for number in range(1, MAX_POINTS + 1)
        for field in ("address", "date")
        if f"{{{{unloading_{number}_{field}}}}}" not in text
    ]
    assert not missing, f"нет плейсхолдеров точек выгрузки: {missing}"


def test_placeholders_are_not_split_across_runs(template_doc, variant):
    """
    Плейсхолдер обязан лежать целиком в одном run.

    Если Word разорвал «{{contract_number}}» на несколько runs, docxtpl
    (Jinja) его не найдёт и в документе останется сырой текст.
    """
    joined_placeholders = set()
    for paragraph in _all_paragraphs(template_doc):
        joined = "".join(run.text for run in paragraph.runs)
        for placeholder in PLACEHOLDER_RE.findall(joined):
            if not any(placeholder in run.text for run in paragraph.runs):
                joined_placeholders.add(placeholder)

    assert not joined_placeholders, (
        f"плейсхолдеры разорваны между runs: {sorted(joined_placeholders)}"
    )


def test_no_run_contains_partial_placeholder(template_doc, variant):
    """Внутри одного run не бывает «обрывка» плейсхолдера."""
    for paragraph in _all_paragraphs(template_doc):
        for run in paragraph.runs:
            rest = PLACEHOLDER_RE.sub("", run.text)
            assert "{{" not in rest and "}}" not in rest, (
                f"незакрытый плейсхолдер в run: {run.text!r}"
            )


def test_no_unclosed_placeholders(template_doc, variant):
    text = _document_text(template_doc)
    opening = text.count("{{")
    closing = text.count("}}")

    assert opening == closing, (
        f"несбалансированные скобки: «{{{{» = {opening}, «}}}}» = {closing}"
    )
    assert len(PLACEHOLDER_RE.findall(text)) == opening, (
        "есть «{{» вне корректно закрытого плейсхолдера"
    )


def test_placeholder_names_are_known(template_doc, variant):
    """В бланке нет опечаток в именах: каждое имя — из ожидаемого набора."""
    expected = _expected_placeholders(variant)
    found = {
        match.strip("{} ")
        for match in _document_placeholders(template_doc)
    }
    assert found == expected, (
        f"лишние/недостающие плейсхолдеры ({variant}): {found ^ expected}"
    )


def test_no_spaces_inside_placeholder_braces(template_doc, variant):
    """Jinja понимает «{{ name }}», но шаблоны проекта пишутся без пробелов."""
    text = _document_text(template_doc)
    for name in _expected_placeholders(variant):
        assert "{{ " + name not in text
        assert name + " }}" not in text


# ─────────────────────────────────────────────────────────────
# Вариант: НДС, КПП, формулировки
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("variant", VAT_VARIANTS)
def test_vat_variants_have_three_sums(templates, variant):
    """ООО и ИП с НДС: без НДС / НДС по ставке / итого, каждая с прописью."""
    texts = _body_texts(templates[variant])
    assert ("– {{sum_wo_vat}} руб. ({{sum_wo_vat_words}}) — стоимость "
            "без НДС;") in texts
    assert "– НДС {{vat_rate}} — {{sum_vat}} руб. ({{sum_vat_words}});" in texts
    assert "Итого с НДС: {{sum_total}} руб. ({{sum_total_words}})." in texts


@pytest.mark.parametrize("variant", VAT_VARIANTS)
def test_vat_variant_states_vat_payer(templates, variant):
    """В вариантах с НДС Арендодатель подтверждает общую систему."""
    text = _document_text(templates[variant])
    assert "применяет общую систему налогообложения" in text
    assert "является плательщиком НДС" in text
    assert "НДС не облагается" not in text


def test_ip_without_vat_has_single_sum(templates):
    """ИП без НДС: одна сумма без НДС и оговорка про УСН."""
    doc = templates["ИП без НДС"]
    texts = _body_texts(doc)

    assert "{{sum_total}} руб. ({{sum_total_words}})." in texts
    assert ("НДС не облагается (упрощённая система налогообложения)."
            in texts)

    text = _document_text(doc)
    for name in VAT_PLACEHOLDERS:
        assert "{{" + name + "}}" not in text, (
            f"в шаблоне «ИП без НДС» лишний плейсхолдер {name!r}"
        )
    assert text.count("{{sum_total}}") == 1

    assert "применяет упрощённую систему налогообложения" in text
    assert "не является плательщиком НДС" in text
    assert "применяет общую систему налогообложения" not in text


def test_kpp_only_in_ooo_template(templates, variant):
    """КПП — только у Арендатора-ООО; у ИП его нет."""
    text = _document_text(templates[variant])
    if variant == "ООО":
        assert "КПП {{lessee_kpp}}" in text
    else:
        assert "{{lessee_kpp}}" not in text
        assert "КПП" not in text


def test_lessee_wording_matches_variant(templates, variant):
    """ООО — «именуемое», ИП — «именуемый»."""
    texts = _body_texts(templates[variant])
    line = next(t for t in texts if t.startswith("1.1. Арендатор:"))

    if variant == "ООО":
        assert "именуемое в дальнейшем «Арендатор»" in line
    else:
        assert "именуемый в дальнейшем «Арендатор»" in line


def test_driver_block_is_complete(template_doc, variant):
    """П. 3.5: девять строк экипажа — ровно как в задании."""
    texts = _body_texts(template_doc)
    expected = (
        "ФИО: {{driver_full_name}}",
        "Дата рождения: {{driver_birth_date}}",
        "Паспорт: {{driver_passport}}",
        "Выдан: {{driver_passport_issuer}}",
        "Дата выдачи: {{driver_passport_issue_date}}",
        "Водительское удостоверение: {{driver_license}}",
        "Дата выдачи: {{driver_license_issue_date}}",
        "Адрес регистрации: {{driver_address}}",
        "Телефон: {{driver_phone}}",
    )
    for line in expected:
        assert line in texts, f"в п. 3.5 нет строки «{line}»"
    assert "3.5. Член экипажа Арендодателя (водитель):" in texts


def test_lease_period_line(template_doc, variant):
    """П. 2.5: плановый период аренды — два плейсхолдера."""
    texts = _body_texts(template_doc)
    line = next(t for t in texts if t.startswith("2.5. Плановый период аренды:"))
    assert "с {{lease_start_date}} г. по {{lease_end_date}} г. включительно" \
        in line


def test_tractor_and_trailer_lines(template_doc, variant):
    """П. 2.1: тягач с типом ТС и прицеп/полуприцеп."""
    texts = _body_texts(template_doc)
    assert ("– тягач: {{tractor_brand}}, государственный регистрационный "
            "знак {{tractor_plate}}, тип ТС — {{tractor_type}};") in texts
    assert ("– прицеп/полуприцеп: {{trailer_brand}}, государственный "
            "регистрационный знак {{trailer_plate}}.") in texts


# ─────────────────────────────────────────────────────────────
# Таблица автомобилей п. 3.1
# ─────────────────────────────────────────────────────────────

def test_car_table_is_thirteen_by_five(template_doc, variant):
    """Таблица машин: 5 колонок и 13 строк (шапка + 12 машин)."""
    table = _car_table(template_doc)
    assert table is not None, "таблица автомобилей не найдена по заголовкам"
    assert len(table.rows) == 1 + CAR_ROWS
    assert len(table.columns) == 5


def test_car_rows_are_numbered_and_placeholder_driven(template_doc, variant):
    table = _car_table(template_doc)
    for number in range(1, CAR_ROWS + 1):
        cells = [cell.text.strip() for cell in table.rows[number].cells]
        assert cells == [
            str(number),
            f"{{{{car_{number}_brand}}}}",
            f"{{{{car_{number}_vin}}}}",
            f"{{{{car_{number}_loading_point}}}}",
            f"{{{{car_{number}_unloading_point}}}}",
        ], f"строка {number} таблицы автомобилей: {cells!r}"


def test_car_table_header_is_bold_and_shaded(template_doc, variant):
    table = _car_table(template_doc)
    fills = []
    for cell in table.rows[0].cells:
        paragraph = cell.paragraphs[0]
        assert str(paragraph.alignment) == "CENTER (1)"
        assert all(run.bold for run in paragraph.runs), "шапка не полужирная"

        shd = cell._tc.tcPr.find(W_NS + "shd")
        fills.append(shd.get(W_NS + "fill") if shd is not None else None)

    assert fills == ["E5E5E5"] * 5


def test_car_column_widths_match_sample(template_doc, variant):
    """Ширины колонок таблицы автомобилей — как в образце (twips)."""
    table = _car_table(template_doc)
    widths = [column.get(W_NS + "w") for column in table._tbl.tblGrid]
    assert widths == list(CAR_COLUMN_WIDTHS)


def test_car_table_header_repeats_on_new_page(template_doc, variant):
    """Шапка таблицы повторяется на новой странице (w:tblHeader)."""
    table = _car_table(template_doc)
    trPr = table.rows[0]._tr.find(W_NS + "trPr")
    assert trPr is not None
    assert trPr.find(W_NS + "tblHeader") is not None


def test_car_table_has_sample_style_headers(template_doc, variant):
    table = _car_table(template_doc)
    headers = tuple(cell.text.strip() for cell in table.rows[0].cells)
    assert headers == CAR_HEADERS


def test_cargo_count_line_present(template_doc, variant):
    assert "Общее количество: {{cargo_count}} шт." in _body_texts(template_doc)


def test_route_points_are_numbered(template_doc, variant):
    """Точки погрузки и выгрузки пронумерованы 3.2.N / 3.3.N до 10."""
    texts = _body_texts(template_doc)
    for number in range(1, MAX_POINTS + 1):
        loading = (f"3.2.{number}. Точка погрузки № {number} — "
                   f"{{{{loading_{number}_address}}}}.")
        unloading = (f"3.3.{number}. Точка выгрузки № {number} — "
                     f"{{{{unloading_{number}_address}}}}.")
        assert any(t.startswith(loading) for t in texts), f"нет п. 3.2.{number}"
        assert any(t.startswith(unloading) for t in texts), \
            f"нет п. 3.3.{number}"

    assert "3.2. Согласованные точки погрузки:" in texts
    assert "3.3. Согласованные точки выгрузки:" in texts
    assert "3.4. Согласованный маршрут: {{route}}." in texts


def test_loading_point_line_format(template_doc, variant):
    """Строка точки погрузки: адрес, дата и окно подачи ТС."""
    texts = _body_texts(template_doc)
    line = next(t for t in texts if t.startswith("3.2.1. Точка погрузки № 1"))
    assert line == (
        "3.2.1. Точка погрузки № 1 — {{loading_1_address}}. Плановая дата "
        "и время подачи ТС: {{loading_1_date}} г., с "
        "{{loading_1_time_from}} до {{loading_1_time_to}}."
    )


def test_unloading_point_line_format(template_doc, variant):
    """Строка точки выгрузки: адрес и плановая дата завершения."""
    texts = _body_texts(template_doc)
    line = next(t for t in texts if t.startswith("3.3.1. Точка выгрузки № 1"))
    assert line == (
        "3.3.1. Точка выгрузки № 1 — {{unloading_1_address}}. Плановая дата "
        "завершения: {{unloading_1_date}} г."
    )


# ─────────────────────────────────────────────────────────────
# Приложение № 1 (Акт приема-передачи и возврата ТС)
# ─────────────────────────────────────────────────────────────

def test_appendix_is_present_in_every_template(template_doc, variant):
    text = _document_text(template_doc)
    assert "Приложение № 1" in text
    assert "АКТ ПРИЕМА-ПЕРЕДАЧИ И ВОЗВРАТА" in text
    assert "ТРАНСПОРТНОГО СРЕДСТВА С ЭКИПАЖЕМ" in text
    assert "1. ПЕРЕДАЧА ТС В АРЕНДУ" in text
    assert "2. ВОЗВРАТ ТС" in text


def test_appendix_keeps_contract_and_vehicle_placeholders(template_doc, variant):
    """В Акте — номер и дата договора, тягач, прицеп и экипаж."""
    text = _document_text(template_doc)
    for name in APPENDIX_PLACEHOLDERS:
        assert "{{" + name + "}}" in text, f"в Акте нет {name!r}"

    assert "с экипажем № {{contract_number}} от {{contract_date}} г." in text
    assert ("Настоящий Акт составлен во исполнение Договора аренды "
            "транспортного средства с экипажем № {{contract_number}} от "
            "{{contract_date}} г.") in text


def test_appendix_starts_on_new_page(template_doc, variant):
    paragraph = next(p for p in template_doc.paragraphs
                     if p.text.startswith("Приложение № 1"))
    assert paragraph.paragraph_format.page_break_before is True
    assert str(paragraph.alignment) == "RIGHT (2)"


def test_appendix_transfer_table(template_doc, variant):
    """Таблица «Передача ТС в аренду»: 8 строк, данные — из договора."""
    table = _act_table(template_doc, "Место передачи")
    assert table is not None, "нет таблицы «Передача ТС в аренду»"
    assert len(table.columns) == 2

    labels = [row.cells[0].text.strip() for row in table.rows]
    assert labels == [
        "Место передачи",
        "Фактические дата и время передачи",
        "Тягач",
        "Прицеп/полуприцеп",
        "Пробег на момент передачи",
        "Внешнее состояние / замечания",
        "Переданные документы",
        "Экипаж",
    ]

    values = [row.cells[1].text.strip() for row in table.rows]
    assert values == [
        "",
        "",
        "{{tractor_brand}}, гос. номер {{tractor_plate}}",
        "{{trailer_brand}}, гос. номер {{trailer_plate}}",
        "",
        "",
        "СТС на тягач и прицеп/полуприцеп; ОСАГО; иные:",
        "{{driver_full_name}}",
    ]


def test_appendix_return_table(template_doc, variant):
    """Таблица «Возврат ТС»: 5 строк для ручного заполнения."""
    table = _act_table(template_doc, "Место возврата")
    assert table is not None, "нет таблицы «Возврат ТС»"
    assert len(table.columns) == 2

    labels = [row.cells[0].text.strip() for row in table.rows]
    assert labels == [
        "Место возврата",
        "Фактические дата и время возврата",
        "Пробег на момент возврата",
        "Состояние ТС / замечания",
        "Иные отметки",
    ]
    values = [row.cells[1].text.strip() for row in table.rows]
    assert values == [""] * 5


def test_appendix_has_explanatory_paragraphs(template_doc, variant):
    """Вводный абзац Акта и оговорка о моменте возврата ТС."""
    texts = _body_texts(template_doc)
    assert any(t.startswith("Настоящий Акт составлен во исполнение Договора")
               for t in texts)
    assert ("С момента, указанного в строке «Фактические дата и время "
            "возврата», ТС считается возвращенным Арендодателю, если "
            "в настоящем Акте не указано иное.") in texts


def test_appendix_signatures(template_doc, variant):
    """Подписи Акта: наименования сторон, фамилии с инициалами и М.П."""
    table = _appendix_signature_table(template_doc)
    assert table is not None, "нет таблицы подписей Акта"

    lessee, lessor = table.rows[0].cells
    assert lessee.text.split("\n") == [
        "АРЕНДАТОР:",
        "{{lessee_short_name}}",
        "______________/ {{lessee_director_short}} /",
        "М.П.",
    ]
    assert lessor.text.split("\n") == [
        "АРЕНДОДАТЕЛЬ:",
        "{{lessor_short_name}}",
        "______________/ {{lessor_director_short}} /",
        "М.П.",
    ]


def test_appendix_is_last_part_of_document(template_doc, variant):
    """Акт идёт после раздела 9 — Приложение в том же файле."""
    texts = _body_texts(template_doc)
    requisites = texts.index("9. РЕКВИЗИТЫ И ПОДПИСИ СТОРОН")
    appendix = _index_of_starting_with(texts, "Приложение № 1")
    assert appendix is not None
    assert appendix > requisites


# ─────────────────────────────────────────────────────────────
# Отсутствие данных образца
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("fragment", FORBIDDEN_FRAGMENTS)
def test_template_has_no_sample_data(templates, variant, fragment):
    assert fragment not in _document_text(templates[variant]), (
        f"в шаблоне {variant} найдены данные образца: «{fragment}»"
    )


def test_car_table_has_no_sample_brands_or_vins(template_doc, variant):
    table = _car_table(template_doc)
    text = "\n".join(cell.text for row in table.rows for cell in row.cells)
    for fragment in ("Haval", "SITRAK", "LUXUDA", "EC2EF4A5"):
        assert fragment not in text, f"в таблице машин осталось «{fragment}»"


def test_no_free_text_looks_like_data(template_doc, variant):
    """Все изменяемые данные — только плейсхолдеры, без «рыбы» из образца."""
    texts = _body_texts(template_doc)
    for text in texts:
        assert "____" not in text, f"в бланке остались пропуски: {text[:60]}"
        assert "ТЛ-" not in text


# ─────────────────────────────────────────────────────────────
# Оформление
# ─────────────────────────────────────────────────────────────

def test_body_font_is_times_new_roman(template_doc, variant):
    """Основной текст — Times New Roman 10,5 pt, как Normal образца."""
    normal = template_doc.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size.pt == 10.5


def test_all_runs_are_times_new_roman(template_doc, variant):
    """Во всём бланке нет ни одного run другого шрифта."""
    allowed_sizes = {8.5, 9.0, 10.5, 11.0, 12.0, 13.0}
    wrong = []
    for paragraph in _all_paragraphs(template_doc):
        for run in paragraph.runs:
            if not run.text:
                continue
            size = run.font.size.pt if run.font.size else None
            if run.font.name != "Times New Roman" or size not in allowed_sizes:
                wrong.append((run.text[:30], run.font.name, size))

    assert not wrong, f"runs не Times New Roman: {wrong[:5]}"


def test_page_geometry_matches_sample(template_doc, variant):
    """Геометрия образца: Letter, поля 2,0 см по бокам и 1,7 см сверху/снизу."""
    section = template_doc.sections[0]
    assert round(section.page_width.cm, 2) == 21.59
    assert round(section.page_height.cm, 2) == 27.94
    assert round(section.left_margin.cm, 2) == 2.0
    assert round(section.right_margin.cm, 2) == 2.0
    assert round(section.top_margin.cm, 2) == 1.7
    assert round(section.bottom_margin.cm, 2) == 1.7


def test_title_is_bold_centered(template_doc, variant):
    """Заголовок договора, «(на один рейс)» и строка «г. Москва … дата»."""
    texts = _body_texts(template_doc)

    assert texts[0].startswith("ДОГОВОР АРЕНДЫ ТРАНСПОРТНОГО СРЕДСТВА")
    assert "С ЭКИПАЖЕМ № {{contract_number}}" in texts[0]
    assert texts[1] == "(на один рейс)"

    title = template_doc.paragraphs[0]
    assert str(title.alignment) == "CENTER (1)"
    for run in title.runs:
        assert run.bold is True
        assert run.font.size.pt == 13

    assert "{{contract_date}} г." in _document_text(template_doc)


def test_sections_go_in_order(template_doc, variant):
    """Разделы 1–9 и Приложение № 1 идут в том же порядке, что в образце."""
    texts = _body_texts(template_doc)
    expected = [
        "1. СТОРОНЫ, ПРАВОВАЯ ПРИРОДА И СТАТУС",
        "2. ПРЕДМЕТ ДОГОВОРА И ОБЪЕКТ АРЕНДЫ",
        "3. УСЛОВИЯ КОММЕРЧЕСКОГО ИСПОЛЬЗОВАНИЯ НА РЕЙС",
        "4. АРЕНДНАЯ ПЛАТА, НДС И ПОРЯДОК ОПЛАТЫ",
        "5. ПЕРЕДАЧА, ЭКСПЛУАТАЦИЯ, СОДЕРЖАНИЕ И ВОЗВРАТ ТС",
        "6. СТРАХОВАНИЕ И ОТВЕТСТВЕННОСТЬ СТОРОН",
        "7. СЕРВИС «ТЕХНОЩИТ» И ПЕРСОНАЛЬНЫЕ ДАННЫЕ",
        "8. СРОК ДЕЙСТВИЯ, ФОРС-МАЖОР, СПОРЫ И ЗАКЛЮЧИТЕЛЬНЫЕ ПОЛОЖЕНИЯ",
        "9. РЕКВИЗИТЫ И ПОДПИСИ СТОРОН",
        "Приложение № 1",
    ]
    positions = []
    for marker in expected:
        index = _index_of_starting_with(texts, marker)
        assert index is not None, f"в шаблоне нет раздела «{marker}»"
        positions.append(index)

    assert positions == sorted(positions), (
        f"разделы идут не по порядку: {list(zip(expected, positions))}"
    )


def test_requisites_table_has_both_parties(template_doc, variant):
    """Раздел 9: таблица 1×2 с блоками АРЕНДАТОР и АРЕНДОДАТЕЛЬ."""
    table = _requisites_table(template_doc)
    assert table is not None, "нет таблицы реквизитов"
    assert len(table.rows) == 1
    assert len(table.columns) == 2

    lessee = table.rows[0].cells[0].text
    lessor = table.rows[0].cells[1].text
    for name in ("{{lessee_full_name}}", "{{lessee_inn}}",
                 "{{lessee_ogrn_label}} {{lessee_ogrn}}",
                 "{{lessee_address}}", "{{lessee_account}}",
                 "{{lessee_bank}}", "{{lessee_bik}}",
                 "{{lessee_corr_account}}", "{{lessee_email}}",
                 "{{lessee_edo}}", "{{lessee_director_position}}",
                 "{{lessee_director_short}}"):
        assert name in lessee, f"в блоке Арендатора нет {name!r}"
    for name in ("{{lessor_full_name}}", "{{lessor_inn}}",
                 "{{lessor_ogrn_label}} {{lessor_ogrn}}",
                 "{{lessor_address}}", "{{lessor_account}}",
                 "{{lessor_bank}}", "{{lessor_bik}}",
                 "{{lessor_corr_account}}", "{{lessor_email}}",
                 "{{lessor_edo}}", "{{lessor_director_position}}",
                 "{{lessor_director_short}}"):
        assert name in lessor, f"в блоке Арендодателя нет {name!r}"

    assert "М.П." in lessee and "М.П." in lessor


def test_numbered_clauses_match_sample(template_doc, variant):
    """П. 5.1–5.8, 6.1–6.12, 7.1–7.4, 8.1–8.7 — как в образце."""
    text = _document_text(template_doc)
    for section, last in (("5", 8), ("6", 12), ("7", 4), ("8", 7)):
        for number in range(1, last + 1):
            marker = f"{section}.{number}."
            assert marker in text, f"нет пункта {marker}"


# ─────────────────────────────────────────────────────────────
# Образец не изменён и шаблон рендерится
# ─────────────────────────────────────────────────────────────

def test_sample_document_untouched(templates_dir):
    """
    Образец — источник структуры, только чтение.

    Если этот тест упал, значит образец отредактировали (или пересобрали).
    Вернуть файл из git:
        git checkout -- "templates/Договор_аренды_ТС_с_экипажем_ТЛ-574_…docx"
    """
    sample = templates_dir / SAMPLE_NAME
    assert sample.exists(), f"пропал образец договора аренды: {sample}"

    digest = hashlib.sha256(sample.read_bytes()).hexdigest()
    assert digest == SAMPLE_SHA256, (
        f"образец {SAMPLE_NAME!r} изменён: ожидался {SAMPLE_SHA256}, "
        f"получен {digest}"
    )


@pytest.mark.parametrize("variant", CARRIER_TYPES)
def test_templates_match_pinned_sha(templates_dir, variant):
    """
    Шаблон на диске — ровно то, что собрал сборщик (ЭТАП 3.1.D.A.1).

    Если этот тест упал после осознанной правки шаблона — пересобери его
    (python tools/make_arenda_ts_template.py) и обнови TEMPLATE_SHA256;
    если правки не было — .docx кто-то отредактировал вручную.
    """
    path = templates_dir / TEMPLATE_NAMES[variant]
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    assert digest == TEMPLATE_SHA256[variant], (
        f"шаблон {TEMPLATE_NAMES[variant]!r} не совпадает со сборкой: "
        f"ожидался {TEMPLATE_SHA256[variant]}, получен {digest}"
    )


def test_docxtpl_renders_without_leftovers(template_path, variant, work_file):
    """Шаблон совместим с docxtpl: после рендера плейсхолдеров не остаётся."""
    docxtpl = pytest.importorskip("docxtpl")

    text_before = _document_text(Document(str(template_path)))
    values = {
        name.strip("{} "): "ЗНАЧЕНИЕ"
        for name in PLACEHOLDER_RE.findall(text_before)
    }
    assert set(values) == _expected_placeholders(variant)

    template = docxtpl.DocxTemplate(str(template_path))
    template.render(values, autoescape=True)

    output = work_file(f"arenda_ts_{variant}_rendered.docx")
    template.save(str(output))

    text_after = _document_text(Document(str(output)))
    assert "{{" not in text_after
    assert "}}" not in text_after
    assert "ЗНАЧЕНИЕ" in text_after
    # Акт рендерится вместе с договором — он в том же файле.
    assert "АКТ ПРИЕМА-ПЕРЕДАЧИ И ВОЗВРАТА" in text_after


def test_builder_reproduces_same_document(templates_dir, variant, work_dir,
                                          monkeypatch):
    """
    Сборщик и лежащий в templates/ шаблон не разъехались.

    Байты сравнивать нельзя: python-docx пишет в zip текущее время. Поэтому
    шаблон пересобирается во временную папку и сравнивается текст документа.
    """
    import tools.make_arenda_ts_template as builder

    monkeypatch.setattr(builder, "TEMPLATES_DIR", work_dir)
    rebuilt = builder.build_template(variant)

    expected = Document(str(templates_dir / TEMPLATE_NAMES[variant]))
    actual = Document(str(rebuilt))

    assert _document_text(actual) == _document_text(expected)
    assert len(actual.paragraphs) == len(expected.paragraphs)
    assert len(actual.tables) == len(expected.tables)
    assert (len(_car_table(actual).rows)
            == len(_car_table(expected).rows) == 1 + CAR_ROWS)


def test_builder_resolves_variants():
    """Тип арендатора → файл шаблона: правило совпадает с перевозкой."""
    import tools.make_arenda_ts_template as builder

    assert builder.resolve_variant("ООО") == "ООО"
    assert builder.resolve_variant("ООО (с НДС)") == "ООО"
    assert builder.resolve_variant("ИП с НДС") == "ИП с НДС"
    assert builder.resolve_variant("ИП без НДС") == "ИП без НДС"
    assert builder.resolve_variant("что-то иное") == "ООО"

    for variant, filename in TEMPLATE_NAMES.items():
        assert builder.VARIANTS[variant]["filename"] == filename
