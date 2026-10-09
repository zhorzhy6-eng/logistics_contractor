#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Заказчик-ИП в договоре перевозки (ШАГ «Фикс сохранения + динамические стороны
в шаблонах», часть B–C).

Что было
--------
П. 1.1 бланка печатал КОНСТАНТУ: «…в лице Генерального директора Ахмедова
Тимура Артуровича, действующего на основании Устава». Константа одна на все
договоры, поэтому у заказчика-ИП (женщины, без директора) в договоре стояло
чужое имя. П. 9 печатал такую же константу в подписи: «Генеральный директор
/Т.А. Ахмедов /».

Что стало
---------
Ветвь выбирает шаблон по ключу `is_client_ip`:

  * ИП — «{{client_legal_form}} {{client_full_name}}, {{client_pronoun}}
    в дальнейшем «Заказчик», {{client_acting}} на основании свидетельства
    о государственной регистрации» («именуемая/именуемый» и
    «действующая/действующий» — по роду ФИО), подпись — «Индивидуальный
    предприниматель ________ /ФИО/»;
  * ООО — прежняя формулировка, но с ДАННЫМИ заказчика.

Данные синтетические, ПДн нет (имена выдуманы для теста).
"""

import re

import pytest
from docx import Document

from core.contract_generator import ContractGenerator

#: Синтетические заказчики.
IP_CLIENT_FEMALE = {
    "full_name": "Индивидуальный предприниматель Смирнова Елена Владимировна",
    "short_name": "ИП Смирнова Е.В.",
    "inn": "770123456789", "kpp": "", "ogrn": "315770000000012",
    "legal_address": "г. Волгоград", "actual_address": "",
    "bank_account": "40702810123456789012", "bik": "044525225",
    "correspondent_account": "30101810400000000225", "bank_name": "ПАО Сбербанк",
    "director_name": "", "director_position": "Генеральный директор",
    "phone": "", "email": "ip@example.ru",
}

IP_CLIENT_MALE = dict(
    IP_CLIENT_FEMALE,
    full_name="Индивидуальный предприниматель Смирнов Сергей Сергеевич",
)

OOO_CLIENT = {
    "full_name": "ООО «Заказчик»",
    "short_name": "ООО «Заказчик»",
    "inn": "7707654321", "kpp": "770701001", "ogrn": "1027700132195",
    "legal_address": "г. Москва, ул. Тестовая, д. 1",
    "actual_address": "", "bank_account": "40702810000000000001",
    "bik": "044525225", "correspondent_account": "30101810400000000225",
    "bank_name": "ПАО Сбербанк",
    "director_name": "Петров Пётр Петрович",
    "director_position": "Генеральный директор",
    "phone": "", "email": "info@example.ru",
}

#: Перевозчик — ООО: проверяем ЗАКАЗЧИКА, поэтому вторая сторона простая.
CARRIER = {
    "full_name": "ООО «Перевозчик»", "short_name": "ООО «Перевозчик»",
    "inn": "7701234567", "kpp": "770101001", "ogrn": "1027700132195",
    "legal_address": "г. Москва", "actual_address": "",
    "bank_account": "40702810000000000002", "bik": "044525225",
    "correspondent_account": "30101810400000000225", "bank_name": "ПАО Сбербанк",
    "director_name": "Сидоров Сидор Сидорович",
    "director_position": "Генеральный директор",
    "phone": "", "email": "", "entity_type": "ООО",
}

DRIVER = {
    "full_name": "Иванов Иван Иванович", "birth_date": "1980-01-01",
    "birth_place": "г. Москва", "passport_series": "18 22",
    "passport_number": "926830", "passport_issue_date": "2023-01-30",
    "passport_issuer": "ОВД", "passport_code": "500-123",
    "registration_address": "г. Москва, ул. Тестовая, д. 1",
    "license_series": "99 36", "license_number": "123456",
    "license_issue_date": "2020-01-01", "license_expiry_date": "2030-01-01",
    "license_categories": "B, C, E", "phone": "+7 (999) 123-45-67",
}


@pytest.fixture
def generator(templates_dir) -> ContractGenerator:
    return ContractGenerator(templates_dir=str(templates_dir))


def _payload(customer) -> dict:
    return {
        "driver": DRIVER, "carrier": CARRIER, "customer": customer,
        "vehicles": [{"vin": "EC3TEUMB0T0002608", "brand_model": "JETOUR T2",
                      "vehicle_type": "Легковой автомобиль"}],
        "tractor": {"brand_model": "Foton Auman", "plate_number": "O844XY196",
                    "color": "Белый", "year": 2023},
        "trailer": {"brand_model": "YANGMINDA", "plate_number": "71ABF18",
                    "color": "Серый", "year": 2020},
        "contract": {
            "number": "CLIENT-IP-1", "date": "2026-10-09",
            "carrier_type": "ООО (с НДС)", "vat_rate": "22%", "vat_rate_num": 22,
            "price_without_vat": 180300.0, "payment_days": 10,
        },
        "loadings": [{"address": "г. Воронеж, ул. Остужева 52Б",
                      "date": "2026-10-10", "time_window": "09:00-15:00"}],
        "unloadings": [{"address": "г. Москва, Перерва 19 стр 3",
                        "date": "2026-10-13", "time_window": ""}],
    }


def _render(generator, customer, output) -> Document:
    generator.generate_docx(_payload(customer), str(output))
    assert output.exists(), f"файл не создан: {output}"
    return Document(str(output))


def _flatten(text: str) -> str:
    """Текст одной строкой: пробелы (в том числе неразрывные) по одному."""
    return re.sub(r"\s+", " ", text.replace("\u00a0", " ")).strip()


def _document_text(doc) -> str:
    """Весь текст документа: абзацы и таблицы (включая вложенные)."""
    parts = [_flatten(p.text) for p in doc.paragraphs]

    def table_text(table):
        for row in table.rows:
            for cell in row.cells:
                parts.append(_flatten(cell.text))
                for nested in cell.tables:
                    table_text(nested)

    for table in doc.tables:
        table_text(table)
    return "\n".join(parts)


def _clause_text(doc, label: str) -> str:
    """
    Абзац, идущий сразу за строкой-меткой («1.1. Заказчик:»).

    Пункт ищется по метке, а не по номеру абзаца: номер в бланках разный,
    и жёсткий индекс молча попал бы в чужую строку.
    """
    texts = [p.text for p in doc.paragraphs]
    for index, text in enumerate(texts):
        if text.strip().startswith(label):
            for following in texts[index + 1:]:
                if following.strip():
                    return _flatten(following)
    raise AssertionError(f"абзац {label!r} не найден")


# ─────────────────────────────────────────────────────────────
# Заказчик-ИП: п. 1.1
# ─────────────────────────────────────────────────────────────

def test_ip_client_no_akhmedov_in_1_1(generator, work_file):
    """Заказчик-ИП — в п. 1.1 нет чужого директора из константы бланка."""
    doc = _render(generator, IP_CLIENT_FEMALE, work_file("client_ip_1_1.docx"))

    clause = _clause_text(doc, "1.1. Заказчик:")
    assert "Ахмедов" not in clause
    assert "Ахмедов" not in _document_text(doc), "чужое ФИО осталось в договоре"


def test_ip_client_female_uses_imenuemaya(generator, work_file):
    """Женщина-ИП: «именуемая» и «действующая» — род по отчеству."""
    doc = _render(generator, IP_CLIENT_FEMALE, work_file("client_ip_f.docx"))

    clause = _clause_text(doc, "1.1. Заказчик:")
    assert clause == (
        "Индивидуальный предприниматель Смирнова Елена Владимировна, "
        "именуемая в дальнейшем «Заказчик», действующая на основании "
        "свидетельства о государственной регистрации."
    )


def test_ip_client_male_uses_imenuemyj(generator, work_file):
    """Мужчина-ИП: «именуемый» и «действующий»."""
    doc = _render(generator, IP_CLIENT_MALE, work_file("client_ip_m.docx"))

    clause = _clause_text(doc, "1.1. Заказчик:")
    assert clause == (
        "Индивидуальный предприниматель Смирнов Сергей Сергеевич, "
        "именуемый в дальнейшем «Заказчик», действующий на основании "
        "свидетельства о государственной регистрации."
    )


def test_ip_client_no_director_block(generator, work_file):
    """У ИП нет «в лице» и нет должности директора в п. 1.1."""
    doc = _render(generator, IP_CLIENT_FEMALE, work_file("client_ip_dir.docx"))

    clause = _clause_text(doc, "1.1. Заказчик:")
    assert "в лице" not in clause
    assert "Генеральный директор" not in clause
    assert "Устава" not in clause, "основание у ИП — свидетельство"


def test_ip_client_signature_is_ip(generator, work_file):
    """П. 9: подпись заказчика-ИП — «Индивидуальный предприниматель / ФИО /»."""
    doc = _render(generator, IP_CLIENT_FEMALE, work_file("client_ip_sign.docx"))

    text = _document_text(doc)
    assert (
        "Индивидуальный предприниматель ________ "
        "/Смирнова Елена Владимировна /" in text
    )
    assert "Генеральный директор ________ /Т.А. Ахмедов /" not in text


def test_ip_client_ogrn_label_is_ogrnip(generator, work_file):
    """У заказчика-ИП в реквизитах ОГРНИП, а не ОГРН."""
    doc = _render(generator, IP_CLIENT_FEMALE, work_file("client_ip_ogrn.docx"))

    text = _document_text(doc)
    assert "ОГРНИП 315770000000012" in text
    assert "ОГРН 315770000000012" not in text


# ─────────────────────────────────────────────────────────────
# Заказчик-ООО: прежняя формулировка, но с данными
# ─────────────────────────────────────────────────────────────

def test_ooo_client_uses_director_data(generator, work_file):
    """ООО-заказчик: в п. 1.1 подставляются должность и ФИО из данных."""
    doc = _render(generator, OOO_CLIENT, work_file("client_ooo_1_1.docx"))

    clause = _clause_text(doc, "1.1. Заказчик:")
    assert clause == (
        "ООО «Заказчик», именуемое в дальнейшем «Заказчик», в лице "
        "Генеральный директор Петров Пётр Петрович, действующего на "
        "основании Устава."
    )


def test_ooo_client_signature_is_director(generator, work_file):
    """П. 9: подпись ООО-заказчика — должность из данных и краткое ФИО."""
    doc = _render(generator, OOO_CLIENT, work_file("client_ooo_sign.docx"))

    text = _document_text(doc)
    assert "Генеральный директор ________ /П.П. Петров /" in text
    assert "Т.А. Ахмедов" not in text


def test_ooo_client_does_not_print_akhmedov_unless_in_data(generator, work_file):
    """Чужого ФИО нет; своё ФИО печатается — оно приходит из данных."""
    doc = _render(generator, OOO_CLIENT, work_file("client_ooo_other.docx"))
    assert "Ахмедов" not in _document_text(doc)

    own = dict(OOO_CLIENT, director_name="Ахмедов Тимур Артурович")
    doc = _render(generator, own, work_file("client_ooo_own.docx"))
    text = _document_text(doc)
    assert "в лице Генеральный директор Ахмедов Тимур Артурович" in text
    assert "Генеральный директор ________ /Т.А. Ахмедов /" in text
