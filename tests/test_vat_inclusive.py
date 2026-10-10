#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сквозные тесты «НДС в том числе» (шаг «Калькулятор „НДС в том числе“»).

Проверяется ГЛАВНОЕ правило шага на живом контуре «вкладка → данные →
генератор → документ»:

  * оператор вводит ОДНУ сумму — «Стоимость», и это ИТОГ договора;
  * база без НДС и налог ВЫНИМАЮТСЯ из итога (а не прибавляются к нему);
  * для всех шести ставок: 250 000 ₽ дают ровно те суммы, что в задании
    (22 % → 204 918.03 и 45 081.97; 5 % → 238 095.24 и 11 904.76; «Без НДС»
    и «0%» → 250 000 и 0);
  * в документе эти же суммы и печатаются, а при «Без НДС» — «НДС не
    облагается»; бланк выбирается по форме и ставке, как и раньше;
  * запись, прошедшая через базу, читается обратно и даёт те же суммы;
  * СТАРАЯ запись (в базе только база без НДС) пересобирается в прежние
    суммы: 1000 ₽ + 22 % печатались как 1000 / 220 / 1220 — такими и
    остаются.

Данные синтетические, ПДн нет; бланки только читаются (не правятся).
"""

import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest

from core.vat import VAT_FREE, compute_vat, total_from_base

PROJECT_ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = PROJECT_ROOT / "templates"

#: Итог, который вводит оператор во всех сквозных проверках.
TOTAL = 250000.0

#: (ставка, форма, база без НДС, НДС, бланк) — числа из задания шага.
CASES = (
    ("22%", "ООО", 204918.03, 45081.97, "shablon_ooo.docx"),
    ("10%", "ООО", 227272.73, 22727.27, "shablon_ooo.docx"),
    ("7%", "ИП", 233644.86, 16355.14, "shablon_ip_with_vat.docx"),
    ("5%", "ИП", 238095.24, 11904.76, "shablon_ip_with_vat.docx"),
    ("0%", "ООО", 250000.00, 0.00, "shablon_ooo.docx"),
    (VAT_FREE, "ООО", 250000.00, 0.00, "shablon_ooo.docx"),
)

CARRIER_OOO = {
    "full_name": "ООО «Ромашка»", "short_name": "ООО «Ромашка»",
    "inn": "7701234567", "kpp": "770101001", "ogrn": "1027700132195",
    "entity_type": "ООО",
}
CARRIER_IP = {
    "full_name": "Индивидуальный предприниматель Тестов Тимур Тимурович",
    "short_name": "ИП Тестов Т. Т.",
    "inn": "770123456789", "ogrn": "310770000000012", "entity_type": "ИП",
}
CUSTOMER = {"full_name": "ООО «Заказчик»", "short_name": "ООО «Заказчик»"}


@pytest.fixture(scope="module")
def qapp():
    from PyQt5.QtWidgets import QApplication

    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture
def tab(qapp, isolated_db):
    from ui.tabs.contract_tab import ContractTab

    return ContractTab()


@pytest.fixture(scope="module")
def generator():
    from core.contract_generator import ContractGenerator

    return ContractGenerator(templates_dir=str(TEMPLATES))


def _fill_tab(tab, form: str, rate: str, total: float = TOTAL) -> None:
    """Оператор выбирает форму и ставку и вводит одну сумму — итог."""
    tab.entity_type.setCurrentText(form)
    tab.vat_rate.setCurrentText(rate)
    tab.price_input.setValue(total)


def _payload_for(form: str, rate: str, data: dict) -> dict:
    """Данные договора: стороны + то, что отдала вкладка."""
    carrier = CARRIER_IP if form == "ИП" else CARRIER_OOO
    return {
        "carrier": dict(carrier),
        "customer": dict(CUSTOMER),
        "contract": {
            "number": "НДС-ВКЛ-1", "date": "2026-10-10",
            "route": "Москва - Тверь", "payment_days": 10,
            "entity_type": form, "vat_rate": rate,
            "carrier_type": data["carrier_type"],
            "price_with_vat": data["price_with_vat"],
            "price_without_vat": data["price_without_vat"],
            "vat_amount": data["vat_amount"],
        },
    }


def _document_text(path) -> str:
    """Весь текст документа одной строкой (абзацы + таблицы)."""
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


# ─────────────────────────────────────────────────────────────
# Вкладка: одна сумма на входе, три на выходе
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("rate, form, base, nds, template", CASES)
def test_tab_shows_base_and_tax_from_the_total(tab, rate, form, base, nds, template):
    """250 000 ₽ — итог: база и налог на экране и в данных вкладки."""
    _fill_tab(tab, form, rate)

    data = tab.get_data()

    assert data["price_with_vat"] == TOTAL
    assert data["price_without_vat"] == base
    assert data["vat_amount"] == nds
    assert tab.price_without_vat.text() == f"{base:.2f} ₽"
    assert tab.vat_amount.text() == f"{nds:.2f} ₽"


@pytest.mark.parametrize("rate, form, base, nds, template", CASES)
def test_document_keeps_the_entered_total(tab, generator, work_file,
                                          rate, form, base, nds, template):
    """В договоре — тот же итог 250 000 ₽ и вынутый из него налог."""
    _fill_tab(tab, form, rate)
    payload = _payload_for(form, rate, tab.get_data())
    out = work_file(f"vat_inclusive_{rate.replace('%', 'pct')}_{form}.docx")

    generator.generate_docx(payload, str(out))
    text = _document_text(out)

    assert "250000.00" in text
    assert f"{base:.2f}" in text
    if nds:
        assert f"{nds:.2f}" in text


@pytest.mark.parametrize("rate, form, base, nds, template", CASES)
def test_template_still_follows_form_and_rate(generator, tab,
                                              rate, form, base, nds, template):
    """Бланк выбирается по форме и ставке — правило шага не изменилось."""
    from core.contract_data import ContractData

    _fill_tab(tab, form, rate)
    data = tab.get_data()
    contract_data = ContractData.coerce(_payload_for(form, rate, data))

    assert generator.get_template_path(contract_data).endswith(template)


def test_vat_free_document_says_not_taxed(tab, generator, work_file):
    """«Без НДС»: 250 000 ₽ и «НДС не облагается» вместо ставки."""
    _fill_tab(tab, "ООО", VAT_FREE)
    payload = _payload_for("ООО", VAT_FREE, tab.get_data())
    out = work_file("vat_inclusive_free.docx")

    generator.generate_docx(payload, str(out))
    text = _document_text(out)

    assert "НДС не облагается" in text
    assert "250000.00" in text
    assert "НДС по ставке, действующей на дату оказания услуг" not in text


def test_zero_rate_document_prints_the_rate(tab, generator, work_file):
    """«0%»: налог нулевой, но печатается ставка, а не «не облагается»."""
    _fill_tab(tab, "ООО", "0%")
    payload = _payload_for("ООО", "0%", tab.get_data())
    out = work_file("vat_inclusive_zero.docx")

    generator.generate_docx(payload, str(out))
    text = _document_text(out)

    assert "(в настоящее время 0%)" in text
    assert "НДС не облагается" not in text


# ─────────────────────────────────────────────────────────────
# База: запись → чтение → те же суммы
# ─────────────────────────────────────────────────────────────

def test_saved_contract_reads_back_with_the_same_sums(tab, qapp):
    """Суммы, записанные в базу, читаются обратно теми же числами."""
    from db.crud.contracts import save_contract
    from db.database import get_connection
    from ui.tabs.contract_tab import ContractTab

    _fill_tab(tab, "ООО", "22%")
    data = tab.get_data()

    contract_id = save_contract({
        "number": "НДС-ВКЛ-БД", "date": "2026-10-10", "route": "Москва - Тверь",
        "vat_rate": data["vat_rate"],
        "price_with_vat": data["price_with_vat"],
        "price_without_vat": data["price_without_vat"],
    })

    conn = get_connection()
    try:
        row = conn.execute(
            "SELECT price_with_vat, price_without_vat, vat_rate FROM contracts "
            "WHERE id = ?", (contract_id,)
        ).fetchone()
    finally:
        conn.close()

    price_with_vat, price_without_vat, vat_rate = row

    assert price_with_vat == 250000.0
    assert price_without_vat == 204918.03
    assert vat_rate == "22%"

    restored = ContractTab()
    try:
        restored.fill_data({
            "vat_rate": vat_rate,
            "price_with_vat": price_with_vat,
            "price_without_vat": price_without_vat,
        })
        back = restored.get_data()
    finally:
        restored.deleteLater()

    assert back["price_with_vat"] == 250000.0
    assert back["price_without_vat"] == 204918.03
    assert back["vat_amount"] == 45081.97


def test_old_record_regenerates_the_same_sums(tab, generator, work_file):
    """
    Старая запись (в базе только база без НДС) — суммы те же, что печатались.

    До шага договор «1000 ₽ без НДС, 22 %» печатался как 1000 / 220 / 1220.
    Вкладка восстанавливает итог (1000 × 1.22 = 1220) и вынимает налог
    обратно — в документе снова ровно 1000 / 220 / 1220.
    """
    tab.fill_data({"vat_rate": "22%", "price_without_vat": 1000.0})
    data = tab.get_data()

    assert tab.price_input.value() == 1220.0
    assert data["price_without_vat"] == 1000.0
    assert data["vat_amount"] == 220.0
    assert data["price_with_vat"] == 1220.0

    payload = _payload_for("ООО", "22%", data)
    out = work_file("vat_inclusive_legacy.docx")

    generator.generate_docx(payload, str(out))
    text = _document_text(out)

    assert "1220.00" in text
    assert "1000.00" in text
    assert "220.00" in text


def test_old_record_formula_is_the_previous_one():
    """Восстановление итога из базы — ровно прежняя формула «сверху»."""
    for rate, nds in (("5%", 50.0), ("7%", 70.0), ("10%", 100.0), ("22%", 220.0)):
        total = total_from_base(1000.0, rate)
        result = compute_vat(total, rate)

        assert total == 1000.0 + nds, rate
        assert result["sum_wo_nds"] == 1000.0, rate
        assert result["sum_nds"] == nds, rate
