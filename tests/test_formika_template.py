#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты шаблона Формики (ЭТАП 3.1.A.1).

Проверяют, что templates/shablon_formika.docx — именно ПУСТОЙ БЛАНК
с плейсхолдерами docxtpl, а не копия образца с данными:

  * файл существует и открывается python-docx;
  * все ключевые плейсхолдеры на месте и ни один не разорван между runs
    (разорванный docxtpl не подставит);
  * таблица груза: шапка + ровно 12 строк, в каждой строке — плейсхолдеры
    марки и VIN;
  * в шаблоне нет данных образца (марки, VIN, ФИО, паспорт, госномера,
    сумма, адреса) и нет незакрытых «{{»;
  * оформление повторяет образец: A4, те же поля, Times New Roman,
    заголовок 18 pt полужирный по центру;
  * п. «Порядок оплаты» печатает срок оплаты плейсхолдерами, а не
    константой «3 (трех) банковских дней» (шаг FIX-2.4);
  * бланк закреплён по SHA256 (TEMPLATE_SHA256);
  * образец-источник не изменён (сверка по SHA256);
  * docxtpl рендерит шаблон без остатка плейсхолдеров.

Образец (templates/Заявка_ТЛ_447 Формика_Технологистика.docx) — только
источник структуры: тест test_sample_document_untouched намеренно падает,
если его отредактировали или пересобрали.
"""

import gc
import hashlib
import re

import pytest
from docx import Document

TEMPLATE_NAME = "shablon_formika.docx"
SAMPLE_NAME = "Заявка_ТЛ_447 Формика_Технологистика.docx"

#: SHA256 образца на момент создания шаблона (ЭТАП 3.1.A.1).
SAMPLE_SHA256 = "d5441240c8cdaad730a039ae07c7673c6c1fc72fbc4a5b68a4dd6080264e85b8"

#: SHA256 собранного бланка (обновлён на шаге FIX-2.4: п. «Порядок оплаты»
#: печатает {{payment_days}} + {{payment_days_words}} вместо константы
#: «3 (трех) банковских дней»).
#: Если тест упал после правки tools/make_formika_template.py — пересобери
#: бланк (python tools/make_formika_template.py) и обнови константу.
TEMPLATE_SHA256 = "3436cffee2e443c561f67cc3d2dba8d0db2d1c5c91ad220a75ad2bb194c85345"

#: Заголовки таблицы груза — по ним таблица ищется в документе
#: (индекс таблицы не используем: в бланке их две).
CARGO_HEADERS = ("№", "Марка, модель", "VIN-номер")

CARGO_ROWS = 12

#: Плейсхолдеры, которые обязаны быть в бланке.
REQUIRED_PLACEHOLDERS = (
    "contract_number",
    "contract_date",
    "contract_month",
    "contract_year",
    "car_1_brand",
    "car_1_vin",
    "car_12_brand",
    "car_12_vin",
    "route",
    "loading_address",
    "loading_plan_date",
    "loading_plan_time_from",
    "loading_plan_time_to",
    "unloading_address",
    "driver_name",
    "driver_birth_date",
    "driver_passport",
    "driver_passport_issuer",
    "driver_passport_date",
    "driver_license",
    "driver_address",
    "driver_phone",
    "tractor_brand",
    "tractor_plate",
    "tractor_type",
    "trailer_brand",
    "trailer_plate",
    "sum_total",
    "sum_total_words",
    "vat_rate",
    "payment_days",
    "payment_days_words",
)

#: Данные образца, которых в бланке быть не должно.
FORBIDDEN_FRAGMENTS = (
    "Haval", "Geely", "Chery", "ВАЗ", "LADA", "XRAY", "JOLION",
    "Coolr", "TIGGO", "Vin по факту",
    "Дмитренко", "350179", "768037",
    "FAW", "LOHR", "Н 658", "ЕР 3166",
    "75 000", "АВТОХАУС", "Воронеж", "ТЛ-447", "Остужева",
)

PLACEHOLDER_RE = re.compile(r"\{\{[^{}]*\}\}")


@pytest.fixture(autouse=True)
def _release_documents():
    """Word-файлы на Windows освобождаем до удаления временных копий."""
    yield
    gc.collect()


@pytest.fixture(scope="module")
def template_path(templates_dir):
    path = templates_dir / TEMPLATE_NAME
    assert path.exists(), f"нет шаблона Формики: {path}"
    return path


@pytest.fixture(scope="module")
def template_doc(template_path):
    return Document(str(template_path))


def _document_text(doc) -> str:
    """Весь текст документа: абзацы, таблицы (включая вложенные), колонтитулы."""
    parts = [p.text for p in doc.paragraphs]

    def table_text(table):
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
                for nested in cell.tables:
                    table_text(nested)

    for table in doc.tables:
        table_text(table)

    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)

    return "\n".join(parts)


def _all_paragraphs(doc):
    """Все абзацы документа, включая абзацы таблиц и колонтитулов."""
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _cargo_table(doc):
    """Таблица груза — по заголовкам колонок, а не по индексу."""
    for table in doc.tables:
        if not table.rows:
            continue
        headers = tuple(cell.text.strip() for cell in table.rows[0].cells)
        if headers == CARGO_HEADERS:
            return table
    return None


# ─────────────────────────────────────────────────────────────
# Файл и плейсхолдеры
# ─────────────────────────────────────────────────────────────

def test_template_file_exists_and_opens(template_path, template_doc):
    assert template_path.stat().st_size > 0
    assert template_doc.tables, "в бланке нет ни одной таблицы"


def test_required_placeholders_present(template_doc):
    text = _document_text(template_doc)
    missing = [name for name in REQUIRED_PLACEHOLDERS
               if "{{" + name + "}}" not in text]
    assert not missing, f"в шаблоне нет плейсхолдеров: {missing}"


def test_all_car_placeholders_present(template_doc):
    """Все 12 машин: и марка, и VIN."""
    text = _document_text(template_doc)
    missing = [
        f"{{{{car_{number}_{field}}}}}"
        for number in range(1, CARGO_ROWS + 1)
        for field in ("brand", "vin")
        if f"{{{{car_{number}_{field}}}}}" not in text
    ]
    assert not missing, f"нет плейсхолдеров машин: {missing}"


def test_placeholders_are_not_split_across_runs(template_doc):
    """
    Плейсхолдер обязан лежать целиком в одном run.

    Если Word разорвал «{{contract_number}}» на несколько runs, docxtpl
    (Jinja) его не найдёт и в документе останется сырой текст.
    """
    joined_placeholders = set()
    for paragraph in _all_paragraphs(template_doc):
        joined = "".join(run.text for run in paragraph.runs)
        for placeholder in PLACEHOLDER_RE.findall(joined):
            if not any(placeholder in run.text for run in paragraph.runs):
                joined_placeholders.add(placeholder)

    assert not joined_placeholders, (
        f"плейсхолдеры разорваны между runs: {sorted(joined_placeholders)}"
    )


def test_no_unclosed_placeholders(template_doc):
    text = _document_text(template_doc)
    opening = text.count("{{")
    closing = text.count("}}")

    assert opening == closing, (
        f"несбалансированные скобки: «{{{{» = {opening}, «}}}}» = {closing}"
    )
    assert len(PLACEHOLDER_RE.findall(text)) == opening, (
        "есть «{{» вне корректно закрытого плейсхолдера"
    )


def test_placeholder_names_are_known(template_doc):
    """В бланке нет опечаток в именах: каждое имя — из ожидаемого набора."""
    expected = set(REQUIRED_PLACEHOLDERS)
    expected |= {
        f"car_{number}_{field}"
        for number in range(1, CARGO_ROWS + 1)
        for field in ("brand", "vin")
    }

    found = {
        match.strip("{} ")
        for match in PLACEHOLDER_RE.findall(_document_text(template_doc))
    }
    assert found == expected, f"лишние/недостающие плейсхолдеры: {found ^ expected}"


# ─────────────────────────────────────────────────────────────
# Таблица груза
# ─────────────────────────────────────────────────────────────

def test_cargo_table_has_header_and_12_data_rows(template_doc):
    table = _cargo_table(template_doc)
    assert table is not None, "таблица груза не найдена по заголовкам"
    assert len(table.rows) == 1 + CARGO_ROWS, (
        f"строк в таблице груза {len(table.rows)}, ожидалось {1 + CARGO_ROWS}"
    )
    assert len(table.columns) == 3


def test_cargo_rows_are_numbered_and_placeholder_driven(template_doc):
    table = _cargo_table(template_doc)
    for number in range(1, CARGO_ROWS + 1):
        cells = [cell.text.strip() for cell in table.rows[number].cells]
        assert cells == [
            str(number),
            f"{{{{car_{number}_brand}}}}",
            f"{{{{car_{number}_vin}}}}",
        ], f"строка {number} таблицы груза: {cells!r}"


def test_cargo_table_has_no_sample_brands_or_vins(template_doc):
    table = _cargo_table(template_doc)
    text = "\n".join(cell.text for row in table.rows for cell in row.cells)
    for fragment in ("Haval", "Geely", "Chery", "ВАЗ", "LADA",
                     "XRAY", "JOLION", "Coolr", "TIGGO", "Vin"):
        assert fragment not in text, f"в таблице груза осталось «{fragment}»"


def test_cargo_header_is_bold_and_centered(template_doc):
    table = _cargo_table(template_doc)
    for cell in table.rows[0].cells:
        paragraph = cell.paragraphs[0]
        assert paragraph.alignment is not None
        assert str(paragraph.alignment) == "CENTER (1)"
        assert all(run.bold for run in paragraph.runs), "шапка не полужирная"


def test_payment_clause_uses_placeholders_not_constant(template_doc):
    """
    П. 4 «Порядок оплаты» — плейсхолдеры, а не константа «3 (трех)».

    Формулировка та же, что в перевозке (п. 4.4) и в аренде ТС (п. 4.5):
    цифры, затем прописью в скобках, затем «банковских дней». Срок оплаты
    задаётся на вкладке «Стоимость», поэтому числа в бланке быть не должно.
    """
    line = next(
        p.text for p in template_doc.paragraphs if p.text.startswith("Порядок оплаты:")
    )

    assert line.startswith(
        "Порядок оплаты: в течение {{payment_days}} "
        "({{payment_days_words}}) банковских дней после получения"
    )
    assert "3 (трех)" not in line


def test_payment_placeholders_are_single_run(template_doc):
    """Оба плейсхолдера обязаны лежать целиком в одном run (иначе docxtpl молчит)."""
    paragraph = next(
        p for p in template_doc.paragraphs if p.text.startswith("Порядок оплаты:")
    )
    runs = [run.text for run in paragraph.runs]

    assert any("{{payment_days}}" in text for text in runs)
    assert any("{{payment_days_words}}" in text for text in runs)


# ─────────────────────────────────────────────────────────────
# Отсутствие данных образца
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("fragment", FORBIDDEN_FRAGMENTS)
def test_template_has_no_sample_data(template_doc, fragment):
    assert fragment not in _document_text(template_doc), (
        f"в шаблоне найдены данные образца: «{fragment}»"
    )


# ─────────────────────────────────────────────────────────────
# Оформление
# ─────────────────────────────────────────────────────────────

def test_title_is_times_new_roman_18_bold_centered(template_doc):
    title = template_doc.paragraphs[0]
    assert title.text.startswith("ДОГОВОР-ЗАЯВКА № {{contract_number}}")

    for run in title.runs:
        assert run.font.name == "Times New Roman"
        assert run.font.size.pt == 18
        assert run.bold is True

    assert str(title.alignment) == "CENTER (1)"


def test_page_geometry_matches_sample(template_doc):
    """A4 и поля образца: 2,0 / 1,5 / 2,0 / 2,0 см."""
    section = template_doc.sections[0]
    assert round(section.page_width.cm, 2) == 21.0
    assert round(section.page_height.cm, 2) == 29.7
    assert round(section.left_margin.cm, 2) == 2.0
    assert round(section.right_margin.cm, 2) == 1.5
    assert round(section.top_margin.cm, 2) == 2.0
    assert round(section.bottom_margin.cm, 2) == 2.0


def test_cargo_column_widths_match_sample(template_doc):
    """Ширины колонок таблицы груза — как в образце (twips)."""
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    table = _cargo_table(template_doc)
    widths = [column.get(ns + "w") for column in table._tbl.tblGrid]
    assert widths == ["1000", "4770", "4253"]


def test_body_font_is_times_new_roman(template_doc):
    """Основной текст — Times New Roman 11 pt."""
    paragraph = next(
        p for p in template_doc.paragraphs if p.text.startswith("Пункт погрузки:")
    )
    for run in paragraph.runs:
        assert run.font.name == "Times New Roman"
        assert run.font.size.pt == 11


# ─────────────────────────────────────────────────────────────
# Образец не изменён и шаблон рендерится
# ─────────────────────────────────────────────────────────────

def test_sample_document_untouched(templates_dir):
    """
    Образец — источник структуры, только чтение.

    Если этот тест упал, значит templates/Заявка_ТЛ_447 Формика_
    Технологистика.docx отредактировали (или пересобрали). Вернуть файл
    из git: git checkout -- "templates/Заявка_ТЛ_447 Формика_Технологистика.docx"
    """
    sample = templates_dir / SAMPLE_NAME
    assert sample.exists(), f"пропал образец Формики: {sample}"

    digest = hashlib.sha256(sample.read_bytes()).hexdigest()
    assert digest == SAMPLE_SHA256, (
        "образец Формики изменён: "
        f"ожидался {SAMPLE_SHA256}, получен {digest}"
    )


def test_template_sha256_unchanged(template_path):
    """
    Бланк закреплён по SHA256 (шаг FIX-2.4).

    Тест ловит ЛЮБУЮ правку templates/shablon_formika.docx, сделанную в
    обход сборщика. Если бланк пересобран осознанно — сначала
    python tools/make_formika_template.py, затем обнови TEMPLATE_SHA256.
    """
    digest = hashlib.sha256(template_path.read_bytes()).hexdigest()
    assert digest == TEMPLATE_SHA256, (
        "бланк Формики изменился: "
        f"ожидался {TEMPLATE_SHA256}, получен {digest}. "
        "Пересобери его сборщиком и обнови TEMPLATE_SHA256."
    )


def test_docxtpl_renders_without_leftovers(template_path, work_file):
    """Шаблон совместим с docxtpl: после рендера плейсхолдеров не остаётся."""
    docxtpl = pytest.importorskip("docxtpl")

    text_before = _document_text(Document(str(template_path)))
    values = {
        name.strip("{} "): "ЗНАЧЕНИЕ"
        for name in PLACEHOLDER_RE.findall(text_before)
    }

    expected = set(REQUIRED_PLACEHOLDERS) | {
        f"car_{number}_{field}"
        for number in range(1, CARGO_ROWS + 1)
        for field in ("brand", "vin")
    }
    assert set(values) == expected

    template = docxtpl.DocxTemplate(str(template_path))
    template.render(values, autoescape=True)

    output = work_file("formika_template_rendered.docx")
    template.save(str(output))

    text_after = _document_text(Document(str(output)))
    assert "{{" not in text_after
    assert "}}" not in text_after
    assert "ЗНАЧЕНИЕ" in text_after
