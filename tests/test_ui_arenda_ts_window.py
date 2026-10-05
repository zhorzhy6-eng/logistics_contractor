#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты окна типа «Разовая аренда» на реальных вкладках (ЭТАП 3.1.D.B.3).

Проверяют то, что появилось на этом шаге: окно собирается на семи настоящих
вкладках (а не на заглушках «в разработке»), сигналы вкладок подключены к
слотам окна, данные всех вкладок уходят в один ContractData, «Очистить форму»
чистит только свою вкладку, «Создать договор» проходит путь валидатор →
диалог → генератор → диалог успеха, а распознавание раскладывает ответ модели
по вкладкам, не стирая ручной ввод.

Отдельная группа тестов — сама раскладка блоков ответа по вкладкам
(_lessee_tab_data, _lessor_tab_data, _vehicle_tab_data, _route_tab_data,
_cargo_tab_data, _crew_tab_data, _price_tab_data): промпт аренды
(core/prompts/arenda_ts.py) отдаёт блоки lessee / lessor / tractor / trailer /
vehicles / loadings / unloadings / driver / contract, а часть полей — route,
lease_start_date, lease_end_date — кладёт в КОРЕНЬ ответа, тогда как вкладки
ждут СВОИ имена полей. Проверять это через окно неудобно, поэтому карта
ключей проверяется напрямую — без поднятия интерфейса.

Qt — в offscreen-режиме. Сеть и системное хранилище ключей не трогаются:
клиент GigaChat подменяется заглушкой, а промпт типа проверяется по
аргументам, с которыми окно позвало recognize_text.

Все данные синтетические, реальных ПДн нет.
"""

import gc
import inspect
import os
import re
import weakref
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from PyQt5.QtWidgets import QApplication, QMessageBox  # noqa: E402

import ui.windows.arenda_ts.window as arenda_window_module  # noqa: E402
from core.contract_data import ContractData  # noqa: E402
from core.contracts.arenda_ts.generator import ArendaTsGenerator  # noqa: E402
from core.contracts.arenda_ts.validator import ArendaTsValidator  # noqa: E402
from core.contracts.factory import GeneratorFactory  # noqa: E402
from core.contracts.registry import ContractTypeRegistry  # noqa: E402
from core.prompts import get_prompt  # noqa: E402
from ui.windows.arenda_ts import ArendaTsWindow  # noqa: E402
from ui.windows.arenda_ts.tabs import (  # noqa: E402
    CargoTab, CrewTab, LesseeTab, LessorTab, PriceTab, RouteTab, VehicleTab,
)
from ui.windows.arenda_ts.window import RecognitionTask  # noqa: E402

# ─────────────────────────────────────────────────────────────
# Константы тестовых данных (синтетика)
# ─────────────────────────────────────────────────────────────

#: Сколько машин помещается в таблицу п. 3.1 бланка (см. CargoTab.MAX_CARS).
MAX_CARS = 12

#: Ожидаемый состав вкладок: (заголовок, класс) — порядок как в окне.
TAB_SPECS = (
    ("Арендатор", LesseeTab),
    ("Арендодатель", LessorTab),
    ("ТС", VehicleTab),
    ("Маршрут", RouteTab),
    ("Груз", CargoTab),
    ("Экипаж", CrewTab),
    ("Стоимость", PriceTab),
)

#: Шапка договора и срок аренды (п. 2.5).
CONTRACT_NUMBER = "ТЛ-574"
CONTRACT_DATE = "2026-09-19"
LEASE_START = "2026-09-21"
LEASE_END = "2026-09-27"
ROUTE = "г. Москва — г. Калуга"

#: Арендатор-ООО — наша сторона в варианте по умолчанию.
LESSEE_NAME = "ООО «Арендатор-Тест»"
LESSEE_SHORT = "ООО «АТ»"
LESSEE_INN = "7701234567"
LESSEE_KPP = "770101001"
LESSEE_OGRN = "1027700132195"
LESSEE_ADDRESS = "г. Москва, ул. Арендаторская, д. 1"
LESSEE_ACCOUNT = "40702810000000000001"
LESSEE_BANK = "ПАО Сбербанк"
LESSEE_BIK = "044525225"
LESSEE_CORR = "30101810400000000225"
LESSEE_EMAIL = "arenda@example.ru"
LESSEE_EDO = "2AE-7F31-4C50"
LESSEE_DIRECTOR = "Петров Пётр Петрович"

#: Арендодатель — вторая сторона; во всех бланках это ООО.
LESSOR_NAME = "ООО «Арендодатель-Тест»"
LESSOR_SHORT = "ООО «АДТ»"
LESSOR_INN = "7709876543"
LESSOR_OGRN = "1027700132196"
LESSOR_ADDRESS = "г. Москва, ул. Арендодательская, д. 2"
LESSOR_ACCOUNT = "40702810000000000002"
LESSOR_BANK = "АО «Банк Второй»"
LESSOR_BIK = "044525226"
LESSOR_CORR = "30101810400000000226"
LESSOR_EMAIL = "lessor@example.ru"
LESSOR_EDO = "3CD-8A42-5D60"
LESSOR_DIRECTOR = "Сидоров Сидор Сидорович"

#: Объект аренды (п. 2.1) — тягач и прицеп.
TRACTOR_BRAND = "Тягач-Модель 5440"
TRACTOR_PLATE = "А001АА01"
TRACTOR_TYPE = "грузовой тягач седельный"
TRAILER_BRAND = "Прицеп-Модель 9"
TRAILER_PLATE = "Б002ББ02"

#: Перевозимые машины (таблица п. 3.1): VIN — 17 символов без букв I, O, Q.
VIN_1 = "XTC651150N0001001"

#: Точки маршрута (п. 3.2 и п. 3.3).
LOADING_NAME = "ООО «Склад Север»"
LOADING_ADDRESS = "г. Москва, ул. Складская, д. 1"
LOADING_DATE = "2026-09-21"
LOADING_TIME_FROM = "08:00"
LOADING_TIME_TO = "18:00"
UNLOADING_NAME = "ООО «Приёмка»"
UNLOADING_ADDRESS = "г. Чехов, ул. Приёмная, д. 9"
UNLOADING_DATE = "2026-09-27"

#: Экипаж (п. 3.5): паспорт и удостоверение — одной строкой, как у промпта.
DRIVER_NAME = "Иванов Иван Иванович"
DRIVER_BIRTH = "1980-01-01"
DRIVER_PASSPORT = "18 22 926830"
DRIVER_PASSPORT_ISSUER = "Отделом УФМС России по г. Москве"
DRIVER_PASSPORT_DATE = "2023-01-30"
DRIVER_LICENSE = "99 36 123456"
DRIVER_LICENSE_DATE = "2020-01-01"
DRIVER_ADDRESS = "г. Москва, ул. Водительская, д. 3"
DRIVER_PHONE = "+7 (999) 123-45-67"

#: Суммы п. 4.1: 221 099,18 + НДС 22% (48 641,82) = 269 741,00.
FORM_AMOUNT = 221099.18
VAT_AMOUNT = 48641.82
TOTAL_AMOUNT = 269741.00
VAT_RATE_TEXT = "22%"

SPECIAL_CONDITIONS = "Простой не более 24 часов"


# ─────────────────────────────────────────────────────────────
# Фикстуры и помощники
# ─────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def qt_app():
    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture(autouse=True)
def _release_documents():
    """Word-файлы на Windows освобождаем до удаления временных копий."""
    yield
    gc.collect()


@pytest.fixture
def quiet_messages(monkeypatch):
    """
    Глушит диалоги QMessageBox и запоминает их тексты.

    :return: словарь со списками текстов information / warning / critical.
    """
    seen = {"information": [], "warning": [], "critical": []}

    def recorder(kind):
        def _record(parent, title, text, *args, **kwargs):
            seen[kind].append(text)
            return QMessageBox.Ok
        return staticmethod(_record)

    for kind in seen:
        monkeypatch.setattr(QMessageBox, kind, recorder(kind))
    return seen


@pytest.fixture
def window(qt_app, quiet_messages):
    """
    Свежее окно «Разовой аренды»; по завершении теста закрывается по-настоящему.

    Перед закрытием ждём пул распознавания: поток, переживший окно, роняет
    процесс при разрушении QThreadPool (access violation на выходе).
    """
    win = ArendaTsWindow()
    yield win
    win.thread_pool.waitForDone(5000)
    qt_app.processEvents()
    win.force_close()


class FakeClient:
    """Заглушка клиента GigaChat: в сеть не ходит, помнит аргументы вызова."""

    def __init__(self, answer=None):
        self.answer = answer if answer is not None else {}
        self.calls = []

    def recognize_text(self, text, prompt=None):
        self.calls.append({"text": text, "prompt": prompt})
        return self.answer


def _fill_all_tabs(win, cars=1):
    """Заполняет все семь вкладок данными, которых хватает валидатору."""
    win.lessee_tab.fill_data({
        "carrier_type": "ООО",
        "full_name": LESSEE_NAME,
        "short_name": LESSEE_SHORT,
        "inn": LESSEE_INN,
        "kpp": LESSEE_KPP,
        "ogrn": LESSEE_OGRN,
        "address": LESSEE_ADDRESS,
        "actual_address": LESSEE_ADDRESS,
        "account": LESSEE_ACCOUNT,
        "bank": LESSEE_BANK,
        "bik": LESSEE_BIK,
        "corr_account": LESSEE_CORR,
        "email": LESSEE_EMAIL,
        "edo": LESSEE_EDO,
        "director_position": "Генеральный директор",
        "director_name": LESSEE_DIRECTOR,
    })
    win.lessor_tab.fill_data({
        "full_name": LESSOR_NAME,
        "short_name": LESSOR_SHORT,
        "inn": LESSOR_INN,
        "ogrn": LESSOR_OGRN,
        "address": LESSOR_ADDRESS,
        "actual_address": LESSOR_ADDRESS,
        "account": LESSOR_ACCOUNT,
        "bank": LESSOR_BANK,
        "bik": LESSOR_BIK,
        "corr_account": LESSOR_CORR,
        "email": LESSOR_EMAIL,
        "edo": LESSOR_EDO,
        "director_position": "Директор",
        "director_name": LESSOR_DIRECTOR,
    })
    win.vehicle_tab.fill_data({
        "contract_number": CONTRACT_NUMBER,
        "contract_date": CONTRACT_DATE,
        "lease_start_date": LEASE_START,
        "lease_end_date": LEASE_END,
        "tractor_brand": TRACTOR_BRAND,
        "tractor_plate": TRACTOR_PLATE,
        "tractor_type": TRACTOR_TYPE,
        "trailer_brand": TRAILER_BRAND,
        "trailer_plate": TRAILER_PLATE,
    })
    win.route_tab.fill_data({
        "route": ROUTE,
        "loadings": [{"name": LOADING_NAME, "address": LOADING_ADDRESS,
                      "date": LOADING_DATE, "time_from": LOADING_TIME_FROM,
                      "time_to": LOADING_TIME_TO}],
        "unloadings": [{"name": UNLOADING_NAME, "address": UNLOADING_ADDRESS,
                        "date": UNLOADING_DATE}],
    })
    win.cargo_tab.fill_data({"vehicles": [
        {"brand_model": f"МОДЕЛЬ {number}",
         "vin": f"XTC651150N0001{number:03d}",
         "loading_point": LOADING_ADDRESS,
         "unloading_point": UNLOADING_ADDRESS}
        for number in range(1, cars + 1)
    ]})
    win.crew_tab.fill_data({
        "driver_full_name": DRIVER_NAME,
        "driver_birth_date": DRIVER_BIRTH,
        "driver_passport": DRIVER_PASSPORT,
        "driver_passport_issuer": DRIVER_PASSPORT_ISSUER,
        "driver_passport_issue_date": DRIVER_PASSPORT_DATE,
        "driver_license": DRIVER_LICENSE,
        "driver_license_issue_date": DRIVER_LICENSE_DATE,
        "driver_registration_address": DRIVER_ADDRESS,
        "driver_phone": DRIVER_PHONE,
    })
    win.price_tab.fill_data({
        "sum_wo_vat": FORM_AMOUNT,
        "sum_vat": VAT_AMOUNT,
        "sum_total": TOTAL_AMOUNT,
        "vat_rate": VAT_RATE_TEXT,
        "vat_rate_num": 22.0,
        "special_conditions": SPECIAL_CONDITIONS,
    })


def _assert_other_tabs_intact(win, cleared_index: int) -> None:
    """Соседние вкладки после «Очистить форму» остались заполненными."""
    if cleared_index != 0:
        assert win.lessee_tab.get_data()["full_name"] == LESSEE_NAME
    if cleared_index != 1:
        assert win.lessor_tab.get_data()["full_name"] == LESSOR_NAME
    if cleared_index != 2:
        assert win.vehicle_tab.get_data()["tractor_brand"] == TRACTOR_BRAND
    if cleared_index != 3:
        assert win.route_tab.get_data()["route"] == ROUTE
    if cleared_index != 4:
        assert len(win.cargo_tab.get_data()["vehicles"]) == 1
    if cleared_index != 5:
        assert win.crew_tab.get_data()["driver_full_name"] == DRIVER_NAME
    if cleared_index != 6:
        assert win.price_tab.get_data()["sum_wo_vat"] == FORM_AMOUNT


def _answer_recognition(window, data):
    """
    Отдаёт окну ответ модели так же, как это сделала бы задача распознавания.

    Задача передаётся аргументом слота: именно так её связывает
    _start_recognition (functools.partial), и так слот понимает, чей
    результат пришёл.
    """
    task = RecognitionTask(window.gigachat or FakeClient(), "текст", prompt="промпт")
    window.recognition_task = task

    window._on_recognition_finished(data, task=task)
    return task


def _document_text(doc) -> str:
    """Весь текст документа: абзацы и таблицы (включая вложенные)."""
    parts = [p.text for p in doc.paragraphs]

    def table_text(table):
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
                for nested in cell.tables:
                    table_text(nested)

    for table in doc.tables:
        table_text(table)

    return "\n".join(parts)


def _recognized_answer() -> dict:
    """
    Ответ модели в схеме промпта аренды (core/prompts/arenda_ts.py).

    Срок аренды, маршрут и обе стороны лежат в КОРНЕ ответа — ровно так их
    отдаёт промпт и поднимает сборщик данных (ui/windows/arenda_ts/data.py).
    """
    return {
        "customer": {"full_name": LESSEE_NAME, "short_name": LESSEE_SHORT,
                     "entity_type": "ООО", "edo": LESSEE_EDO},
        "lessee": {
            "entity_type": "ООО",
            "full_name": LESSEE_NAME,
            "short_name": LESSEE_SHORT,
            "inn": LESSEE_INN,
            "kpp": LESSEE_KPP,
            "ogrn": LESSEE_OGRN,
            "legal_address": LESSEE_ADDRESS,
            "actual_address": LESSEE_ADDRESS,
            "bank_account": LESSEE_ACCOUNT,
            "bank_name": LESSEE_BANK,
            "bik": LESSEE_BIK,
            "corr_account": LESSEE_CORR,
            "email": LESSEE_EMAIL,
            "edo": LESSEE_EDO,
            "director_position": "Генеральный директор",
            "director_name": LESSEE_DIRECTOR,
        },
        "lessor": {
            "full_name": LESSOR_NAME,
            "short_name": LESSOR_SHORT,
            "inn": LESSOR_INN,
            "ogrn": LESSOR_OGRN,
            "legal_address": LESSOR_ADDRESS,
            "actual_address": LESSOR_ADDRESS,
            "bank_account": LESSOR_ACCOUNT,
            "bank_name": LESSOR_BANK,
            "bik": LESSOR_BIK,
            "corr_account": LESSOR_CORR,
            "email": LESSOR_EMAIL,
            "edo": LESSOR_EDO,
            "director_position": "Директор",
            "director_name": LESSOR_DIRECTOR,
        },
        "route": ROUTE,
        "lease_start_date": LEASE_START,
        "lease_end_date": LEASE_END,
        "tractor": {"brand_model": "Volvo FH", "plate_number": "А001АА77",
                    "vehicle_type": "грузовой тягач седельный"},
        "trailer": {"brand_model": "Schmitz", "plate_number": "ВК999977"},
        "vehicles": [{"brand_model": "МОДЕЛЬ X", "vin": "XTC651150N0009001",
                      "loading_point": LOADING_ADDRESS,
                      "unloading_point": UNLOADING_ADDRESS}],
        "loadings": [{"address": LOADING_ADDRESS, "date": LOADING_DATE,
                      "time_from": LOADING_TIME_FROM, "time_to": LOADING_TIME_TO}],
        "unloadings": [{"address": UNLOADING_ADDRESS, "date": UNLOADING_DATE}],
        "driver": {
            "full_name": "Петров Пётр Петрович",
            "birth_date": "1985-05-05",
            "passport": "45 08 111222",
            "passport_issuer": "Отделом УФМС России по г. Твери",
            "passport_issue_date": "2021-04-04",
            "license": "69 25 333444",
            "license_issue_date": "2019-03-03",
            "address": "г. Тверь, ул. Новая, д. 7",
            "phone": "+7 (999) 000-11-22",
        },
        "contract": {
            "number": "ТЛ-2026-77",
            "date": "01.10.2026",
            "sum_wo_vat": 200000.0,
            "sum_vat": 44000.0,
            "sum_total": 244000.0,
            "vat_rate": "22%",
            "vat_rate_num": 22.0,
            "special_conditions": "Без дозагрузки",
        },
    }


# ─────────────────────────────────────────────────────────────
# 1. Окно и его вкладки
# ─────────────────────────────────────────────────────────────

def test_window_builds_with_seven_tabs(window):
    assert window.CONTRACT_TYPE == "arenda_ts"
    assert window.windowTitle() == "Разовая аренда"
    assert window.tabs.count() == 7
    assert window.tab_titles() == [title for title, _ in TAB_SPECS]
    assert window.side_nav.count() == 7


def test_window_builds_without_arguments(qt_app, quiet_messages):
    """Окно создаётся без аргументов — как его и создаёт менеджер окон."""
    win = ArendaTsWindow()
    try:
        assert win.parent() is None
        assert win.tabs.count() == 7
    finally:
        win.force_close()


def test_make_tab_returns_real_tab_classes(window):
    """Хук _make_tab отдаёт настоящие вкладки, а не заглушки «в разработке»."""
    for index, (title, tab_class) in enumerate(TAB_SPECS):
        tab = window.tabs.widget(index)
        assert isinstance(tab, tab_class), title

    for tab in window._tabs():
        assert hasattr(tab, "recognition_panel"), type(tab).__name__
        assert hasattr(tab, "get_data"), type(tab).__name__


def test_tabs_are_named_attributes(window):
    """Окно держит вкладки и по именам — на них опираются слоты."""
    assert isinstance(window.lessee_tab, LesseeTab)
    assert isinstance(window.lessor_tab, LessorTab)
    assert isinstance(window.vehicle_tab, VehicleTab)
    assert isinstance(window.route_tab, RouteTab)
    assert isinstance(window.cargo_tab, CargoTab)
    assert isinstance(window.crew_tab, CrewTab)
    assert isinstance(window.price_tab, PriceTab)


def test_tab_attribute_matches_position(window):
    """Имя вкладки соответствует её месту в TAB_CONFIGS."""
    named = {
        "lessee_tab": 0, "lessor_tab": 1, "vehicle_tab": 2, "route_tab": 3,
        "cargo_tab": 4, "crew_tab": 5, "price_tab": 6,
    }
    for attribute, index in named.items():
        assert getattr(window, attribute) is window.tabs.widget(index), attribute
        assert type(window.tabs.widget(index)).__name__.lower().startswith(
            attribute.split("_")[0]
        ), attribute


def test_unknown_tab_title_falls_back_to_placeholder(window):
    """Опечатка в TAB_CONFIGS не оставляет окно без страницы."""
    tab = window._make_tab("Такой вкладки нет", "contract.svg")

    assert tab is not None
    assert not hasattr(tab, "get_data")


def test_header_has_create_contract_button(window):
    """«Создать договор» есть и в шапке, и на вкладках (как в MainWindow)."""
    assert window.btn_create_contract.text() == "Создать договор"
    assert window.btn_create_contract.parent() is not None
    for tab in window._tabs():
        assert tab.btn_create_contract.text() == "Создать договор"


def test_header_button_stands_before_exit(window):
    """Кнопка шапки встаёт перед «Выход», а не после неё."""
    header = window.btn_exit.parentWidget().layout()
    assert header.indexOf(window.btn_create_contract) < header.indexOf(window.btn_exit)


def test_every_tab_has_buttons(window):
    """У всех семи вкладок есть панель действий, как у вкладок MainWindow."""
    for tab in window._tabs():
        assert tab.btn_create_contract.isEnabled() is True
        assert tab.btn_clear_form.isEnabled() is True
        # «Распознать данные» живёт на панели распознавания вкладки.
        assert tab.recognition_panel.btn_recognize.isEnabled() is True


def test_gigachat_is_not_created_in_init(window):
    """Клиент GigaChat — ленивый: простое открытие окна ключа не требует."""
    assert window.gigachat is None
    assert window.recognition_task is None


def test_type_is_registered_for_factory(window):
    """
    Окно прогревает реестр типов: фабрика знает «arenda_ts».

    Без этого GeneratorFactory взяла бы генератор по умолчанию, и договор
    аренды собирался бы по чужому шаблону.
    """
    spec = ContractTypeRegistry.find("arenda_ts")

    assert spec is not None
    assert spec.generator_class is ArendaTsGenerator
    assert spec.validator_class is ArendaTsValidator


# ─────────────────────────────────────────────────────────────
# 2. Сигналы вкладок подключены к слотам окна
# ─────────────────────────────────────────────────────────────
# Сигнал подключён к связанному методу окна, поэтому доказательство
# связи — вызов этого же пути: клик по кнопке вкладки доходит до
# _collect_data() и _confirm_validation(). Подменять сами слоты после
# создания окна бессмысленно: связь уже держит исходный метод.

def test_tab_create_signal_is_connected_to_window(window, monkeypatch):
    reports = []
    monkeypatch.setattr(
        ArendaTsWindow, "_confirm_validation",
        lambda self, report: reports.append(report) and False,
    )

    for tab in window._tabs():
        reports.clear()
        tab.btn_create_contract.click()
        assert len(reports) == 1, type(tab).__name__


def test_header_create_button_uses_the_same_slot_as_tabs(window, monkeypatch):
    """
    «Создать договор» в шапке и на вкладке ведут в один слот окна.

    Проверяется по исходнику: обе связи идут на _on_create_contract, а не
    на разные обработчики — иначе шапка и вкладка давали бы разный результат.
    """
    assert "_on_create_contract" in inspect.getsource(
        ArendaTsWindow._build_header_actions
    )
    assert "_on_create_contract" in inspect.getsource(ArendaTsWindow._bind_tabs)

    reports = []
    monkeypatch.setattr(
        ArendaTsWindow, "_confirm_validation",
        lambda self, report: reports.append(report) and False,
    )

    window.btn_create_contract.click()

    assert len(reports) == 1


#: Вкладки, у которых проверяется сигнал распознавания: (индекс, заголовок).
RECOGNIZE_SPECS = tuple(
    (index, title) for index, (title, _cls) in enumerate(TAB_SPECS)
)


@pytest.mark.parametrize(
    "index,title", RECOGNIZE_SPECS, ids=[spec[1] for spec in RECOGNIZE_SPECS]
)
def test_tab_recognize_signal_is_connected_to_window(
    window, monkeypatch, index, title
):
    """Кнопка «Распознать вкладку» каждой вкладки доходит до слота окна."""
    clients = []
    monkeypatch.setattr(
        ArendaTsWindow, "_ensure_gigachat",
        lambda self: clients.append(FakeClient()) or clients[-1],
    )

    tab = window.tabs.widget(index)
    tab.recognition_panel.text_edit.setPlainText("текст договора аренды")
    tab.recognition_panel.btn_recognize.click()

    assert len(clients) == 1, title
    assert window.recognition_task is not None
    assert window.recognition_task.text == "текст договора аренды"


#: Как проверить, что вкладка действительно очистилась: (индекс, проверка).
CLEAR_SPECS = (
    (0, "Арендатор", lambda tab: tab.get_data()["full_name"] == ""
        and tab.get_data()["carrier_type"] == "ООО"),
    (1, "Арендодатель", lambda tab: tab.get_data()["full_name"] == ""),
    (2, "ТС", lambda tab: tab.get_data()["tractor_brand"] == ""
        and tab.get_data()["contract_number"] == ""),
    (3, "Маршрут", lambda tab: tab.get_data()["route"] == ""
        and tab.get_data()["loadings"] == []),
    (4, "Груз", lambda tab: tab.get_data()["vehicles"] == []),
    (5, "Экипаж", lambda tab: tab.get_data()["driver_full_name"] == ""),
    (6, "Стоимость", lambda tab: tab.get_data()["sum_wo_vat"] == 0.0
        and tab.get_data()["special_conditions"] == ""),
)


@pytest.mark.parametrize("index,title,is_cleared", CLEAR_SPECS,
                         ids=[spec[1] for spec in CLEAR_SPECS])
def test_tab_clear_signal_clears_only_its_own_tab(
    window, index, title, is_cleared
):
    """Кнопка «Очистить форму» каждой вкладки чистит только свою вкладку."""
    _fill_all_tabs(window, cars=1)

    tab = window.tabs.widget(index)
    tab.btn_clear_form.click()

    assert is_cleared(tab), title
    _assert_other_tabs_intact(window, index)


def test_window_connect_does_not_use_lambda():
    """
    В connect — методы и partial, но не lambda: lambda, захватывающая окно
    или вкладку, создаёт цикл ссылок Python ↔ Qt и роняет процесс при
    завершении.
    """
    for method in (
        ArendaTsWindow.__init__,
        ArendaTsWindow._build_header_actions,
        ArendaTsWindow._bind_tabs,
        ArendaTsWindow._start_recognition,
    ):
        source = inspect.getsource(method)
        assert ".connect(lambda" not in source.replace(" ", ""), source.splitlines()[0]


def test_window_connects_only_bound_methods_and_partial():
    """
    Все connect окна ведут на связанные методы окна или на partial.

    Иначе говоря, среди связей нет замыканий на вкладку: вкладка живёт
    дольше сигнала, и цикл Python ↔ Qt не даёт сборщику мусора освободить
    ни окно, ни вкладки при выходе из программы.
    """
    source = inspect.getsource(arenda_window_module)
    targets = re.findall(r"\.connect\(([^)]*)", source)

    assert targets, "в модуле окна не нашлось ни одной связи сигналов"
    for target in targets:
        text = target.strip()
        assert text.startswith(("self._", "partial(")), text


def test_recognition_signals_are_bound_to_task_through_partial():
    """
    Сигналы задачи привязываются к самой задаче (partial), а не к lambda.

    Именованный аргумент task виден в подписи слота и отличает результат
    «своей» задачи от результата отменённой.
    """
    source = inspect.getsource(ArendaTsWindow._start_recognition)

    assert source.count("partial(") == 3
    assert source.count("task=task") == 3


def test_tabs_do_not_reference_the_window(window):
    """
    Вкладки не держат ссылок на окно — цикла Python ↔ Qt нет.

    Связь идёт только в одну сторону (окно → вкладка): обратная ссылка
    замкнула бы цикл, и ни окно, ни вкладки не освободились бы при выходе.
    """
    for tab in window._tabs():
        for name, value in vars(tab).items():
            assert value is not window, f"{type(tab).__name__}.{name}"


def test_closed_window_is_collected(qt_app, quiet_messages):
    """После force_close() окно освобождается сборщиком мусора."""
    win = ArendaTsWindow()
    reference = weakref.ref(win)

    win.force_close()
    qt_app.processEvents()
    del win
    gc.collect()

    assert reference() is None, "окно не освободилось — где-то остался цикл ссылок"


# ─────────────────────────────────────────────────────────────
# 3. Сбор данных
# ─────────────────────────────────────────────────────────────

def test_collect_data_returns_contract_data_from_all_tabs(window):
    _fill_all_tabs(window, cars=2)

    data = window._collect_data()

    assert isinstance(data, ContractData)
    assert data.contract["number"] == CONTRACT_NUMBER
    assert data.contract["date"] == CONTRACT_DATE
    assert data.contract["lease_start_date"] == LEASE_START
    assert data.contract["lease_end_date"] == LEASE_END
    assert data.contract["route"] == ROUTE
    assert data.contract["carrier_type"] == "ООО"
    assert data.contract["vat_rate"] == VAT_RATE_TEXT
    assert data.contract["vat_rate_num"] == 22.0
    assert data.contract["price_without_vat"] == FORM_AMOUNT
    assert data.contract["special_conditions"] == SPECIAL_CONDITIONS

    # Стороны: Арендатор — наша сторона (customer), Арендодатель — вторая.
    assert data.contract["lessee"]["full_name"] == LESSEE_NAME
    assert data.contract["lessee"]["inn"] == LESSEE_INN
    assert data.contract["lessee"]["bank_account"] == LESSEE_ACCOUNT
    assert data.contract["lessee"]["bank_name"] == LESSEE_BANK
    assert data.customer["full_name"] == LESSEE_NAME
    assert data.contract["lessor"]["full_name"] == LESSOR_NAME
    assert data.contract["lessor"]["inn"] == LESSOR_INN
    assert data.contract["lessor"]["bank_name"] == LESSOR_BANK
    assert data.carrier["full_name"] == LESSOR_NAME

    # Точки маршрута: адреса в приведённом виде, названия и время — в contract.
    assert [point["address"] for point in data.loadings] == [LOADING_ADDRESS]
    assert data.loadings[0]["time_window"] == "08:00-18:00"
    assert [point["name"] for point in data.contract["loadings"]] == [LOADING_NAME]
    assert [point["address"] for point in data.unloadings] == [UNLOADING_ADDRESS]
    assert [point["name"] for point in data.contract["unloadings"]] == [UNLOADING_NAME]

    assert len(data.vehicles) == 2
    assert data.vehicles[0]["vin"] == VIN_1

    assert data.tractor == {"brand_model": TRACTOR_BRAND,
                            "plate_number": TRACTOR_PLATE,
                            "vehicle_type": TRACTOR_TYPE}
    assert data.trailer == {"brand_model": TRAILER_BRAND,
                            "plate_number": TRAILER_PLATE}

    assert data.driver["full_name"] == DRIVER_NAME
    assert data.driver["passport_series"] == "18 22"
    assert data.driver["passport_number"] == "926830"
    assert data.driver["license_series"] == "99 36"


def test_collect_data_is_repeatable_and_does_not_change_ui(window):
    """_collect_data только читает: повторный вызов даёт тот же результат."""
    _fill_all_tabs(window, cars=2)

    first = window._collect_data()
    second = window._collect_data()

    assert first.to_generator_dict() == second.to_generator_dict()
    assert window.vehicle_tab.get_data()["contract_number"] == CONTRACT_NUMBER


def test_more_than_twelve_cars_are_capped_in_contract_data(window):
    """В бланк помещается 12 машин: лишние в ContractData не попадают."""
    _fill_all_tabs(window, cars=MAX_CARS + 3)
    assert len(window.cargo_tab.get_data()["vehicles"]) == MAX_CARS

    data = window._collect_data()

    assert len(data.vehicles) == MAX_CARS
    assert data.vehicles[-1]["vin"] == "XTC651150N0001012"


def test_empty_tabs_give_empty_contract_data(window):
    data = window._collect_data()

    assert data.vehicles == []
    assert data.loadings == []
    assert data.unloadings == []
    assert data.tractor == {}
    assert data.trailer == {}
    assert data.driver.get("full_name", "") == ""
    assert data.contract.get("number") is None
    assert data.contract.get("route") is None
    assert data.contract.get("price_without_vat") is None


def test_full_form_passes_validator_without_errors(window):
    """Заполненная форма проходит валидатор чисто — диалог не помешает."""
    _fill_all_tabs(window, cars=2)

    report = ArendaTsValidator().check(window._collect_data())

    assert report.errors == [], f"неожиданные ошибки: {report.errors}"
    assert report.warnings == [], f"неожиданные замечания: {report.warnings}"
    assert report.is_clean is True


# ─────────────────────────────────────────────────────────────
# 4. Очистка: только своя вкладка
# ─────────────────────────────────────────────────────────────

def test_clear_tab_clears_only_sender_tab(window):
    _fill_all_tabs(window, cars=1)

    window.crew_tab.btn_clear_form.click()

    assert window.crew_tab.get_data()["driver_full_name"] == ""
    assert window.lessee_tab.get_data()["full_name"] == LESSEE_NAME
    assert window.lessor_tab.get_data()["full_name"] == LESSOR_NAME
    assert window.vehicle_tab.get_data()["tractor_brand"] == TRACTOR_BRAND
    assert window.route_tab.get_data()["route"] == ROUTE
    assert window.price_tab.get_data()["sum_wo_vat"] == FORM_AMOUNT


def test_clear_slot_without_sender_is_safe(window):
    """Прямой вызов слота без отправителя ничего не чистит и не падает."""
    _fill_all_tabs(window, cars=1)

    window._on_clear_tab()

    assert window.lessee_tab.get_data()["full_name"] == LESSEE_NAME
    assert window.crew_tab.get_data()["driver_full_name"] == DRIVER_NAME


def test_clear_tab_does_not_touch_other_point_tables(window):
    """Очистка «Маршрута» не трогает таблицу машин вкладки «Груз»."""
    _fill_all_tabs(window, cars=1)

    window.route_tab.btn_clear_form.click()

    assert window.route_tab.get_data()["loadings"] == []
    assert window.route_tab.get_data()["unloadings"] == []
    assert window.cargo_tab.get_data()["vehicles"][0]["vin"] == VIN_1


# ─────────────────────────────────────────────────────────────
# 5. Закрытие окна
# ─────────────────────────────────────────────────────────────

def test_close_hides_window_without_destroying_it(qt_app, quiet_messages):
    win = ArendaTsWindow()
    win.show()
    _fill_all_tabs(win, cars=1)
    try:
        win.close()

        assert win.isVisible() is False
        assert win.vehicle_tab.get_data()["contract_number"] == CONTRACT_NUMBER
    finally:
        win.force_close()


def test_force_close_really_closes_window(qt_app, quiet_messages):
    win = ArendaTsWindow()
    win.show()
    win.force_close()

    assert win.isVisible() is False
    assert win._force_close is True


# ─────────────────────────────────────────────────────────────
# 6. «Создать договор»: валидатор → диалог → генератор → диалог успеха
# ─────────────────────────────────────────────────────────────

def test_empty_form_fails_validation_and_does_not_generate(
    window, monkeypatch, quiet_messages
):
    """Пустая форма: валидатор находит ошибки, генерация не запускается."""
    report = ArendaTsValidator().check(ContractData())

    assert report.has_errors is True
    assert "Не заполнен номер договора аренды" in report.errors
    assert "Не заполнено ФИО водителя (экипажа)" in report.errors
    assert "Добавьте хотя бы одну перевозимую машину" in report.errors
    assert "Не заполнена марка тягача" in report.errors
    assert "Не указан маршрут аренды" in report.errors
    assert "Стоимость без НДС должна быть больше нуля" in report.errors

    started = []
    monkeypatch.setattr(
        GeneratorFactory, "get_generator",
        classmethod(lambda cls, *args, **kwargs: started.append(args)),
    )
    # Диалог проверки данных: пользователь выбрал «Исправить».
    monkeypatch.setattr(QMessageBox, "exec_", lambda self: QMessageBox.No)

    window._on_create_contract()

    assert started == []
    assert window.statusBar().currentMessage() == "Договор не создан — исправьте данные"


def test_empty_form_validation_errors_cover_all_sections(
    window, monkeypatch, quiet_messages
):
    """
    Пустая форма называет все незаполненные разделы бланка.

    Ни даты договора, ни срока аренды в списке ошибок нет: вкладка «ТС»
    подставляет сегодняшнюю дату и срок по умолчанию (DEFAULT_LEASE_YEARS) —
    так же, как это делают вкладки MainWindow.
    """
    report = ArendaTsValidator().check(window._collect_data())

    for expected in (
        "Не заполнен номер договора аренды",
        "Не заполнено наименование Арендатора",
        "Не заполнен ИНН Арендатора",
        "Не заполнен КПП Арендатора",
        "Не заполнено наименование Арендодателя",
        "Не заполнен ИНН Арендодателя",
        "Не заполнена марка тягача",
        "Не заполнен госномер тягача",
        "Не заполнен тип ТС тягача",
        "Не заполнена марка прицепа",
        "Не заполнен госномер прицепа",
        "Добавьте хотя бы одну перевозимую машину",
        "Укажите хотя бы одну точку погрузки",
        "Укажите хотя бы одну точку выгрузки",
        "Не указан маршрут аренды",
        "Не заполнено ФИО водителя (экипажа)",
        "Не заполнены паспортные данные водителя",
        "Не заполнен адрес регистрации",
        "Не заполнен телефон водителя",
        "Стоимость без НДС должна быть больше нуля",
    ):
        assert expected in report.errors, expected

    assert "Не заполнена дата договора" not in report.errors
    assert "Не указана дата начала аренды" not in report.errors
    assert "Не указана дата окончания аренды" not in report.errors


def test_empty_form_has_default_lease_dates(window):
    """Свежая вкладка «ТС» уже отдаёт срок аренды — он не пустой."""
    data = window.vehicle_tab.get_data()

    assert data["lease_start_date"]
    assert data["lease_end_date"]
    assert data["lease_end_date"] > data["lease_start_date"]


def test_create_contract_can_be_cancelled_in_dialog(
    window, monkeypatch, quiet_messages
):
    """Данные с ошибкой, но пользователь нажал «Исправить» — генерации нет."""
    _fill_all_tabs(window, cars=1)
    window.vehicle_tab.trailer_plate.clear()      # госномер прицепа не заполнен

    report = ArendaTsValidator().check(window._collect_data())
    assert "Не заполнен госномер прицепа" in report.errors

    started = []
    monkeypatch.setattr(
        GeneratorFactory, "get_generator",
        classmethod(lambda cls, *args, **kwargs: started.append(args)),
    )
    monkeypatch.setattr(QMessageBox, "exec_", lambda self: QMessageBox.No)

    window._on_create_contract()

    assert started == []
    assert window.statusBar().currentMessage() == "Договор не создан — исправьте данные"


def test_create_contract_generates_docx_end_to_end(
    window, monkeypatch, quiet_messages, work_dir, templates_dir
):
    """
    Заполненная форма: генератор → DOCX в папке вывода → диалог успеха.

    Документ проверяется целиком: файл создан, плейсхолдеры шаблона
    заменены, данные всех семи вкладок дошли до бланка.
    """
    _fill_all_tabs(window, cars=2)

    generator = ArendaTsGenerator(templates_dir=str(templates_dir))
    output_dir = work_dir / "arenda_ts_window_output"
    output_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(
        GeneratorFactory, "get_generator",
        classmethod(lambda cls, *args, **kwargs: generator),
    )
    monkeypatch.setattr(
        generator, "default_output_dir", lambda: str(output_dir)
    )

    created = []
    monkeypatch.setattr(
        ArendaTsWindow, "_show_contract_created",
        lambda self, path: created.append(path),
    )
    # Данных достаточно, поэтому диалог проверки не показывается; ответ на
    # всякий случай — «Создать договор».
    monkeypatch.setattr(QMessageBox, "exec_", lambda: QMessageBox.Yes)

    window._on_create_contract()

    assert len(created) == 1
    path = Path(created[0])
    try:
        assert path.exists() and path.suffix == ".docx"
        assert CONTRACT_NUMBER in path.name
        assert path.parent == output_dir

        text = _document_text(Document(str(path)))

        # Плейсхолдеров шаблона не осталось — и «{{...}}», и «<<...>>».
        assert "{{" not in text, "в договоре остался плейсхолдер"
        assert "}}" not in text
        assert "<<" not in text, "в договоре остался плейсхолдер"

        # Данные вкладок дошли до бланка: обе стороны, тягач, экипаж, суммы.
        assert "Арендатор-Тест" in text
        assert "Арендодатель-Тест" in text
        assert TRACTOR_BRAND in text
        assert DRIVER_NAME in text
    finally:
        # За собой убираем: папка вывода теста не должна пухнуть от прогонов.
        try:
            path.unlink(missing_ok=True)
            output_dir.rmdir()
        except OSError:
            pass


def test_clean_data_skips_confirmation_dialog(
    window, monkeypatch, quiet_messages, templates_dir
):
    """
    Данные без замечаний: диалог проверки не показывается вовсе.

    Данные подставляются напрямую — так виден именно шаг проверки, а не то,
    что вкладки успели заполнить.
    """
    _fill_all_tabs(window, cars=1)
    data = window._collect_data()
    report = ArendaTsValidator().check(data)
    assert report.is_clean is True, report.format_text()

    generator = ArendaTsGenerator(templates_dir=str(templates_dir))
    monkeypatch.setattr(
        GeneratorFactory, "get_generator",
        classmethod(lambda cls, *args, **kwargs: generator),
    )
    # Данные подставляются напрямую: виден именно шаг проверки, а не то,
    # что вкладки успели заполнить.
    monkeypatch.setattr(ArendaTsWindow, "_collect_data", lambda self: data)

    generated = []
    monkeypatch.setattr(
        ArendaTsWindow, "_show_contract_created",
        lambda self, path: generated.append(path),
    )
    monkeypatch.setattr(
        generator, "generate",
        lambda payload: generated.append(payload) and "output/x.docx",
    )

    dialogs = []
    monkeypatch.setattr(
        QMessageBox, "exec_", lambda self: dialogs.append(self.windowTitle()) or 0
    )

    window._on_create_contract()

    assert dialogs == []
    assert generated[-1] is data


def test_generator_failure_is_reported(window, monkeypatch, quiet_messages):
    """Падение генератора не роняет окно: показываем ошибку."""
    _fill_all_tabs(window, cars=1)
    monkeypatch.setattr(
        GeneratorFactory, "get_generator",
        classmethod(lambda cls, *args, **kwargs: object()),
    )

    window._on_create_contract()

    assert any("Не удалось создать договор" in text
               for text in quiet_messages["critical"])


def _make_contract_file(work_dir, name: str):
    """
    Создаёт файл-заготовку договора и убирает его после теста.

    Диалогу достаточно существующего пути: содержимое DOCX не читается,
    а мусор в tests/_tmp после прогона не нужен.
    """
    path = work_dir / name
    path.write_bytes(b"PK\x03\x04")
    return path


def test_show_contract_created_offers_open_folder(
    window, monkeypatch, quiet_messages, work_dir
):
    """Диалог успеха: кнопка «Открыть папку» и «OK», как в MainWindow."""
    path = _make_contract_file(
        work_dir, "Договор_аренды_ТС_ТЛ-574_20260919.docx"
    )

    buttons = []
    origin_add_button = QMessageBox.addButton

    def spy_add_button(self, *args, **kwargs):
        button = origin_add_button(self, *args, **kwargs)
        buttons.append(button.text())
        return button

    monkeypatch.setattr(QMessageBox, "addButton", spy_add_button)
    # Диалог не показываем: интересна его начинка, а не модальное окно.
    monkeypatch.setattr(QMessageBox, "exec_", lambda self: 0)

    opened = []
    monkeypatch.setattr(
        ArendaTsWindow, "_open_folder", lambda self, folder: opened.append(folder)
    )
    try:
        window._show_contract_created(str(path))

        assert buttons == ["Открыть папку", "OK"]
        assert opened == []  # нажали «OK»: папку не открываем
    finally:
        path.unlink(missing_ok=True)


def test_show_contract_created_opens_folder_on_request(
    window, monkeypatch, quiet_messages, work_dir
):
    """Кнопка «Открыть папку» открывает папку готового договора."""
    path = _make_contract_file(
        work_dir, "Договор_аренды_ТС_ТЛ-575_20260919.docx"
    )

    opened = []
    monkeypatch.setattr(
        ArendaTsWindow, "_open_folder", lambda self, folder: opened.append(folder)
    )

    # Запоминаем кнопку «Открыть папку»: в диалоге она первая.
    open_button = None
    origin_add_button = QMessageBox.addButton

    def spy_add_button(self, *args, **kwargs):
        nonlocal open_button
        button = origin_add_button(self, *args, **kwargs)
        if button.text() == "Открыть папку":
            open_button = button
        return button

    monkeypatch.setattr(QMessageBox, "addButton", spy_add_button)

    def spy_exec(box):
        box._clicked = open_button   # нажата «Открыть папку»
        return 0

    monkeypatch.setattr(QMessageBox, "exec_", spy_exec)
    monkeypatch.setattr(
        QMessageBox, "clickedButton", lambda self: getattr(self, "_clicked", None)
    )

    try:
        window._show_contract_created(str(path))

        assert open_button is not None
        assert opened == [str(work_dir)]
    finally:
        path.unlink(missing_ok=True)


def test_open_folder_returns_false_for_missing_folder(window):
    """Несуществующая папка не открывается и не роняет окно."""
    assert window._open_folder("") is False
    assert window._open_folder(str(Path("такой-папки-нет-12345"))) is False


# ─────────────────────────────────────────────────────────────
# 7. Распознавание
# ─────────────────────────────────────────────────────────────

def test_recognize_request_starts_task_with_arenda_prompt(
    window, monkeypatch, qt_app
):
    """Промпт типа передаётся в recognize_text вторым аргументом."""
    client = FakeClient({"contract": {"number": CONTRACT_NUMBER}})
    monkeypatch.setattr(ArendaTsWindow, "_ensure_gigachat", lambda self: client)

    prompt = get_prompt("arenda_ts")
    assert prompt, "у типа «Разовая аренда» должен быть свой промпт"

    window._on_recognize_requested("текст договора аренды")
    task = window.recognition_task
    assert isinstance(task, RecognitionTask)
    assert task.prompt == prompt

    # Ждём пул: задача уже выполняется в потоке. Сигнал из потока приходит
    # в поток интерфейса, поэтому очередь событий доводим до конца —
    # иначе слот не успеет сбросить состояние задачи.
    window.thread_pool.waitForDone(5000)
    qt_app.processEvents()

    assert len(client.calls) == 1
    assert client.calls[0]["text"] == "текст договора аренды"
    assert client.calls[0]["prompt"] == prompt

    # Связь сигналов с partial(task=...) действительно сработала.
    assert window.recognition_task is None


def test_recognition_without_key_shows_message(window, monkeypatch, quiet_messages):
    """Нет ключа — показываем MISSING_KEY_MESSAGE, задачу не запускаем."""
    monkeypatch.setattr(ArendaTsWindow, "_ensure_gigachat", lambda self: None)

    window._on_recognize_requested("текст договора аренды")

    assert window.recognition_task is None
    assert any("set_key.py" in text for text in quiet_messages["critical"])


def test_empty_text_is_not_sent_to_recognition(window, monkeypatch, quiet_messages):
    """Пустой текст на распознавание не уходит."""
    called = []
    monkeypatch.setattr(ArendaTsWindow, "_ensure_gigachat",
                        lambda self: called.append("client"))

    window._on_recognize_requested("   ")

    assert called == []
    assert quiet_messages["warning"]


def test_recognition_fills_tabs(window, quiet_messages):
    """Ответ модели раскладывается по вкладкам: у каждой — свои поля."""
    _answer_recognition(window, _recognized_answer())

    # Арендатор: реквизиты блока lessee, счёт и банк — под именами вкладки.
    lessee = window.lessee_tab.get_data()
    assert lessee["full_name"] == LESSEE_NAME
    assert lessee["short_name"] == LESSEE_SHORT
    assert lessee["inn"] == LESSEE_INN
    assert lessee["kpp"] == LESSEE_KPP
    assert lessee["ogrn"] == LESSEE_OGRN
    assert lessee["address"] == LESSEE_ADDRESS
    assert lessee["actual_address"] == LESSEE_ADDRESS
    assert lessee["account"] == LESSEE_ACCOUNT
    assert lessee["bank"] == LESSEE_BANK
    assert lessee["bik"] == LESSEE_BIK
    assert lessee["corr_account"] == LESSEE_CORR
    assert lessee["email"] == LESSEE_EMAIL
    assert lessee["edo"] == LESSEE_EDO
    assert lessee["director_name"] == LESSEE_DIRECTOR
    assert lessee["carrier_type"] == "ООО"

    # Арендодатель: те же реквизиты второй стороны.
    lessor = window.lessor_tab.get_data()
    assert lessor["full_name"] == LESSOR_NAME
    assert lessor["inn"] == LESSOR_INN
    assert lessor["ogrn"] == LESSOR_OGRN
    assert lessor["address"] == LESSOR_ADDRESS
    assert lessor["account"] == LESSOR_ACCOUNT
    assert lessor["bank"] == LESSOR_BANK
    assert lessor["director_name"] == LESSOR_DIRECTOR

    # ТС: шапка договора и срок аренды из contract и корня ответа, тягач и
    # прицеп — из блоков схемы промпта (brand_model → tractor_brand).
    vehicle = window.vehicle_tab.get_data()
    assert vehicle["contract_number"] == "ТЛ-2026-77"
    assert vehicle["contract_date"] == "2026-10-01"
    assert vehicle["lease_start_date"] == LEASE_START
    assert vehicle["lease_end_date"] == LEASE_END
    assert vehicle["tractor_brand"] == "Volvo FH"
    assert vehicle["tractor_plate"] == "А001АА77"
    assert vehicle["tractor_type"] == "грузовой тягач седельный"
    assert vehicle["trailer_brand"] == "Schmitz"
    assert vehicle["trailer_plate"] == "ВК999977"

    # Маршрут: направление из корня ответа, точки — массивами.
    route = window.route_tab.get_data()
    assert route["route"] == ROUTE
    assert route["loadings"] == [{"name": "", "address": LOADING_ADDRESS,
                                  "date": LOADING_DATE,
                                  "time_from": LOADING_TIME_FROM,
                                  "time_to": LOADING_TIME_TO}]
    assert route["unloadings"] == [{"name": "", "address": UNLOADING_ADDRESS,
                                    "date": UNLOADING_DATE}]

    # Груз: таблица перерисована по ответу модели.
    assert window.cargo_tab.get_data()["vehicles"] == [
        {"brand_model": "МОДЕЛЬ X", "vin": "XTC651150N0009001",
         "loading_point": LOADING_ADDRESS, "unloading_point": UNLOADING_ADDRESS}
    ]

    # Экипаж: девять полей бланка из блока driver.
    crew = window.crew_tab.get_data()
    assert crew["driver_full_name"] == "Петров Пётр Петрович"
    assert crew["driver_birth_date"] == "1985-05-05"
    assert crew["driver_passport"] == "45 08 111222"
    assert crew["driver_passport_issuer"] == "Отделом УФМС России по г. Твери"
    assert crew["driver_passport_issue_date"] == "2021-04-04"
    assert crew["driver_license"] == "69 25 333444"
    assert crew["driver_license_issue_date"] == "2019-03-03"
    assert crew["driver_registration_address"] == "г. Тверь, ул. Новая, д. 7"
    assert crew["driver_phone"] == "+7 (999) 000-11-22"

    # Стоимость: суммы, ставка и особые условия — из блока contract.
    price = window.price_tab.get_data()
    assert price["sum_wo_vat"] == 200000.0
    assert price["vat_rate"] == "22%"
    assert price["vat_rate_num"] == 22.0
    assert price["special_conditions"] == "Без дозагрузки"

    assert window.recognition_task is None
    assert window.statusBar().currentMessage() == "Готово"


def test_recognition_ignores_unknown_sections(window, quiet_messages):
    """Разделы, которых у вкладок нет, раскладку не ломают."""
    _answer_recognition(window, {
        "contract": {"number": "ТЛ-1"},
        "shippers": [{"name": "ООО «Склад»", "address": "адрес погрузки"}],
        "consignees": [{"name": "ООО «Клиент»", "address": "адрес выгрузки"}],
        "cargo_count": 3,
        "чего-то_ещё": {"вложено": True},
    })

    assert window.vehicle_tab.get_data()["contract_number"] == "ТЛ-1"
    # Точки чужой схемы (shippers / consignees) тоже принимаются.
    assert window.route_tab.get_data()["loadings"] == [
        {"name": "ООО «Склад»", "address": "адрес погрузки",
         "date": "", "time_from": "", "time_to": ""}
    ]
    assert window.route_tab.get_data()["unloadings"] == [
        {"name": "ООО «Клиент»", "address": "адрес выгрузки", "date": ""}
    ]


def test_recognition_uses_sum_total_for_ip_without_vat(window, quiet_messages):
    """
    «НДС не облагается» (ИП без НДС): сумма документа лежит в sum_total.

    Нулевая sum_wo_vat вкладке не передаётся: иначе ноль занял бы первое
    место в списке приоритетов и спрятал за собой настоящую сумму.
    """
    _answer_recognition(window, {
        "lessee": {"entity_type": "ИП", "full_name": "ИП Смирнов С.С.",
                   "inn": "770123456789", "ogrn": "321770000123456",
                   "legal_address": "г. Москва, ул. ИП, д. 7",
                   "director_name": "Смирнов Сергей Сергеевич"},
        "contract": {"sum_wo_vat": 0.0, "sum_vat": 0.0, "sum_total": 269741.0,
                     "vat_rate": "0%", "vat_rate_num": 0.0},
    })

    lessee = window.lessee_tab.get_data()
    assert lessee["carrier_type"] == "ИП без НДС"
    assert lessee["full_name"] == "ИП Смирнов С.С."

    price = window.price_tab.get_data()
    assert price["sum_wo_vat"] == 269741.0
    assert price["vat_rate"] == "0%"
    assert price["vat_rate_num"] == 0.0


def test_recognition_derives_ip_with_vat_from_rate(window, quiet_messages):
    """«ИП» и ставка 22% — вариант ИП с НДС."""
    _answer_recognition(window, {
        "lessee": {"entity_type": "ИП", "full_name": "ИП Смирнов С.С."},
        "contract": {"sum_wo_vat": 100000.0, "vat_rate": "22%",
                     "vat_rate_num": 22.0},
    })

    assert window.lessee_tab.get_data()["carrier_type"] == "ИП с НДС"


def test_recognition_does_not_overwrite_manual_input(window, quiet_messages):
    """Пустые поля ответа не стирают то, что пользователь ввёл руками."""
    _fill_all_tabs(window, cars=1)

    _answer_recognition(window, {
        # Модель вернула разделы целиком, но с пустыми строками и нулями.
        "customer": {"full_name": "", "short_name": ""},
        "lessee": {"full_name": "", "short_name": "", "inn": "", "kpp": "",
                   "ogrn": "", "legal_address": "", "actual_address": "",
                   "bank_account": "", "bank_name": "", "bik": "",
                   "corr_account": "", "email": "", "edo": "",
                   "director_name": ""},
        "lessor": {"full_name": "", "inn": "", "ogrn": "", "legal_address": "",
                   "bank_account": "", "bank_name": "", "director_name": ""},
        "tractor": {"brand_model": "", "plate_number": "", "vehicle_type": ""},
        "trailer": {"brand_model": "", "plate_number": ""},
        "vehicles": [{"brand_model": "", "vin": ""}],
        "loadings": [{"address": "", "date": "", "time_from": "", "time_to": ""}],
        "unloadings": [{"address": "", "date": ""}],
        "route": "",
        "lease_start_date": "",
        "lease_end_date": "",
        "driver": {"full_name": "", "passport": "", "phone": ""},
        "contract": {"number": "", "date": "", "sum_wo_vat": 0.0, "sum_vat": 0.0,
                     "sum_total": 0.0, "vat_rate": "", "special_conditions": ""},
    })

    assert window.lessee_tab.get_data()["full_name"] == LESSEE_NAME
    assert window.lessee_tab.get_data()["inn"] == LESSEE_INN
    assert window.lessee_tab.get_data()["account"] == LESSEE_ACCOUNT
    assert window.lessor_tab.get_data()["full_name"] == LESSOR_NAME
    assert window.vehicle_tab.get_data()["contract_number"] == CONTRACT_NUMBER
    assert window.vehicle_tab.get_data()["tractor_brand"] == TRACTOR_BRAND
    assert window.vehicle_tab.get_data()["lease_start_date"] == LEASE_START
    assert window.route_tab.get_data()["route"] == ROUTE
    assert window.route_tab.get_data()["loadings"][0]["address"] == LOADING_ADDRESS
    assert window.cargo_tab.get_data()["vehicles"][0]["vin"] == VIN_1
    assert window.crew_tab.get_data()["driver_full_name"] == DRIVER_NAME
    assert window.price_tab.get_data()["sum_wo_vat"] == FORM_AMOUNT
    assert window.price_tab.get_data()["special_conditions"] == SPECIAL_CONDITIONS


def test_recognition_keeps_carrier_type_chosen_by_hand(window, quiet_messages):
    """
    Частичный ответ не меняет выбранный вручную вид Арендатора.

    Без entity_type вид не выводится: от него зависят и бланк, и состав
    сумм, поэтому подставлять «ООО» принудительно нельзя.
    """
    window.lessee_tab.carrier_type.setCurrentText("ИП без НДС")

    _answer_recognition(window, {"contract": {"sum_total": 100000.0,
                                              "vat_rate": "0%"}})

    assert window.lessee_tab.get_data()["carrier_type"] == "ИП без НДС"


def test_recognition_logs_have_no_personal_data(caplog, window, quiet_messages):
    """
    В лог распознавания не попадают данные договора.

    Пишутся только имена разделов, имена полей и количества: ФИО, адреса,
    VIN, названия организаций и суммы в логе быть не должно.
    """
    import logging

    with caplog.at_level(logging.DEBUG):
        _answer_recognition(window, {
            "lessee": {"full_name": "ООО «Секретная организация»",
                       "inn": "7712345678"},
            "lessor": {"full_name": "ООО «Второй секрет»"},
            "vehicles": [{"brand_model": "МОДЕЛЬ СЕКРЕТ",
                          "vin": "XTC651150N0009999"}],
            "route": "г. Секретоград — г. Тайна",
            "loadings": [{"address": "г. Тестоград, ул. Складская, д. 9"}],
            "driver": {"full_name": "Секретов Секрет Секретович"},
            "contract": {"number": "ТЛ-СЕКРЕТ", "sum_total": 987654.32},
        })

    messages = "\n".join(record.getMessage() for record in caplog.records)

    for fragment in (
        "Секретная организация",
        "Второй секрет",
        "7712345678",
        "Секретоград",
        "Тестоград",
        "XTC651150N0009999",
        "МОДЕЛЬ СЕКРЕТ",
        "Секретов",
        "ТЛ-СЕКРЕТ",
        "987654",
    ):
        assert fragment not in messages, f"в логе есть «{fragment}»"

    # А служебные сведения о разделах — есть.
    assert "разделы" in messages


def test_recognition_ui_action_log_has_no_personal_data(caplog, window,
                                                        quiet_messages):
    """Журнал действий окна тоже обходится без данных договора."""
    import logging

    with caplog.at_level(logging.INFO,
                         logger="ui.windows.arenda_ts.window"):
        _answer_recognition(window, _recognized_answer())

    messages = "\n".join(
        record.getMessage() for record in caplog.records
        if record.name == "ui.windows.arenda_ts.window"
    )

    for fragment in (LESSEE_NAME, LESSOR_NAME, DRIVER_NAME, LESSEE_INN, VIN_1):
        assert fragment not in messages, f"в логе есть «{fragment}»"
    assert "ответ разложен по вкладкам" in messages


def test_error_slot_without_sender_is_safe(window, quiet_messages):
    """Ошибка без отправителя: сообщение показываем, состояние сбрасываем."""
    window._on_recognition_error("Модель недоступна")

    assert any("Модель недоступна" in text for text in quiet_messages["critical"])
    assert window.recognition_task is None


def test_recognition_error_is_reported(window, quiet_messages):
    """Ошибка распознавания: QMessageBox.critical и сброс состояния."""
    task = RecognitionTask(FakeClient(), "текст", prompt="промпт")
    window.recognition_task = task

    window._on_recognition_error("Модель недоступна", task=task)

    assert any("Модель недоступна" in text for text in quiet_messages["critical"])
    assert window.recognition_task is None


def test_result_of_stale_task_is_ignored(window, quiet_messages):
    """Результат устаревшей (отменённой) задачи к вкладкам не применяется."""
    _fill_all_tabs(window, cars=1)
    stale = RecognitionTask(FakeClient(), "текст", prompt="промпт")

    window._on_recognition_finished(
        {"contract": {"number": "ЧУЖОЙ-НОМЕР"}}, task=stale
    )

    assert window.vehicle_tab.get_data()["contract_number"] == CONTRACT_NUMBER


def test_recognition_cancelled_task_does_not_fill_tabs(window, quiet_messages):
    """Отмена задачи освобождает интерфейс и не трогает вкладки."""
    task = RecognitionTask(FakeClient(), "текст", prompt="промпт")
    window.recognition_task = task

    window._on_recognition_cancelled(task=task)

    assert window.recognition_task is None
    assert window.statusBar().currentMessage() == "Готово"


def test_recognition_progress_is_shown_in_status_bar(window):
    """Ход распознавания виден в статус-баре."""
    window._on_recognition_progress(20, "Отправка в GigaChat")

    assert window.statusBar().currentMessage() == "Отправка в GigaChat (20%)"


# ─────────────────────────────────────────────────────────────
# 8. Задача распознавания (своя, а не из MainWindow)
# ─────────────────────────────────────────────────────────────

def test_recognition_task_passes_prompt_to_client():
    """Своя задача типа: промпт уходит в клиент вторым аргументом."""
    client = FakeClient({"contract": {"number": "ТЛ-1"}})
    task = RecognitionTask(client, "исходный текст", prompt="свой промпт")

    received = []
    task.signals.finished.connect(received.append)
    task.run()

    assert received == [{"contract": {"number": "ТЛ-1"}}]
    assert client.calls[0]["prompt"] == "свой промпт"


def test_recognition_task_reports_empty_answer():
    """None от клиента — сигнал ошибки, а не пустой результат."""
    class NoneClient(FakeClient):
        def recognize_text(self, text, prompt=None):
            return None

    task = RecognitionTask(NoneClient(), "текст", prompt="промпт")
    errors = []
    task.signals.error.connect(errors.append)
    task.run()

    assert errors == ["Модель вернула пустой ответ"]


def test_cancelled_recognition_task_does_not_apply_result():
    """Отменённая задача результат не отдаёт — интерфейс не тронут."""
    client = FakeClient({"contract": {"number": "ТЛ-1"}})
    task = RecognitionTask(client, "исходный текст", prompt="свой промпт")
    task.cancel()

    finished, cancelled = [], []
    task.signals.finished.connect(finished.append)
    task.signals.cancelled.connect(lambda: cancelled.append(True))  # noqa: B023
    task.run()

    assert finished == []
    assert cancelled == [True]
    assert task.is_cancelled is True


# ─────────────────────────────────────────────────────────────
# 9. Раскладка блоков ответа по вкладкам (прямые проверки карт ключей)
# ─────────────────────────────────────────────────────────────

def test_lessee_tab_data_merges_customer_and_lessee():
    """customer даёт наименования, lessee — реквизиты, и он важнее."""
    data = ArendaTsWindow._lessee_tab_data(
        {"full_name": "ООО «Из customer»", "short_name": "ООО «К»",
         "entity_type": "ООО", "edo": "AAA-111"},
        {"full_name": "ООО «Из lessee»", "inn": "7701234567",
         "bank_account": LESSEE_ACCOUNT, "bank_name": LESSEE_BANK,
         "legal_address": LESSEE_ADDRESS},
        None,
    )

    assert data["full_name"] == "ООО «Из lessee»"
    assert data["short_name"] == "ООО «К»"
    assert data["edo"] == "AAA-111"
    assert data["inn"] == "7701234567"
    # Счёт и банк переведены в имена полей вкладки.
    assert data["account"] == LESSEE_ACCOUNT
    assert data["bank"] == LESSEE_BANK
    assert "bank_account" not in data
    assert "bank_name" not in data
    assert data["legal_address"] == LESSEE_ADDRESS

    # Пустые блоки ничего не заполняют.
    assert ArendaTsWindow._lessee_tab_data({}, {}, None) == {}
    assert ArendaTsWindow._lessee_tab_data(
        {"full_name": ""}, {"inn": "  "}, None
    ) == {}


def test_lessee_tab_data_derives_carrier_type_from_entity_and_rate():
    """Вид Арендатора выводится из entity_type и ставки НДС."""
    assert ArendaTsWindow._lessee_tab_data(
        {}, {"entity_type": "ООО"}, {"vat_rate": "22%"}
    )["carrier_type"] == "ООО"

    assert ArendaTsWindow._lessee_tab_data(
        {}, {"entity_type": "ИП"}, {"vat_rate": "22%"}
    )["carrier_type"] == "ИП с НДС"

    assert ArendaTsWindow._lessee_tab_data(
        {}, {"entity_type": "ИП"}, {"vat_rate": "0%", "vat_rate_num": 0.0}
    )["carrier_type"] == "ИП без НДС"

    assert ArendaTsWindow._lessee_tab_data(
        {}, {"entity_type": "Индивидуальный предприниматель"},
        {"vat_rate": "Без НДС"}
    )["carrier_type"] == "ИП без НДС"


def test_lessee_tab_data_does_not_guess_without_entity_type():
    """Без entity_type вид не подставляется — вкладка оставляет свой."""
    data = ArendaTsWindow._lessee_tab_data({}, {"full_name": "ООО «Тест»"}, {})

    assert "carrier_type" not in data


def test_lessee_tab_data_keeps_explicit_carrier_type():
    """Явно пришедший carrier_type не переписывается выводом."""
    data = ArendaTsWindow._lessee_tab_data(
        {}, {"carrier_type": "ИП без НДС", "entity_type": "ООО"},
        {"vat_rate": "22%"},
    )

    assert data["carrier_type"] == "ИП без НДС"


def test_lessor_tab_data_maps_bank_keys():
    """Блок lessor: счёт и банк переводятся в имена полей вкладки."""
    data = ArendaTsWindow._lessor_tab_data({
        "full_name": LESSOR_NAME,
        "inn": LESSOR_INN,
        "legal_address": LESSOR_ADDRESS,
        "bank_account": LESSOR_ACCOUNT,
        "bank_name": LESSOR_BANK,
    })

    assert data == {
        "full_name": LESSOR_NAME,
        "inn": LESSOR_INN,
        "legal_address": LESSOR_ADDRESS,
        "account": LESSOR_ACCOUNT,
        "bank": LESSOR_BANK,
    }

    # КПП у Арендодателя нет ни в бланке, ни в раскладке окна.
    assert "kpp" not in ArendaTsWindow._lessor_tab_data({"inn": LESSOR_INN})


def test_party_keys_keeps_explicit_tab_name():
    """Явное имя вкладки (account) важнее имени блока (bank_account)."""
    data = ArendaTsWindow._party_keys({
        "account": "40702810000000000009",
        "bank_account": LESSEE_ACCOUNT,
        "bank": "АО «Свой банк»",
        "bank_name": LESSEE_BANK,
    })

    assert data["account"] == "40702810000000000009"
    assert data["bank"] == "АО «Свой банк»"
    assert "bank_account" not in data
    assert "bank_name" not in data


def test_vehicle_tab_data_maps_prompt_schema():
    """
    Прямая проверка раскладки блоков tractor / trailer (схема промпта).

    Отдельно от окна: так видно саму карту ключей, а не её косвенное
    действие через вкладку.
    """
    data = ArendaTsWindow._vehicle_tab_data(
        None,
        {"brand_model": TRACTOR_BRAND, "plate_number": TRACTOR_PLATE,
         "vehicle_type": TRACTOR_TYPE},
        {"brand_model": TRAILER_BRAND, "plate_number": TRAILER_PLATE},
        None,
    )

    assert data == {
        "tractor_brand": TRACTOR_BRAND,
        "tractor_plate": TRACTOR_PLATE,
        "tractor_type": TRACTOR_TYPE,
        "trailer_brand": TRAILER_BRAND,
        "trailer_plate": TRAILER_PLATE,
    }

    # Пустые значения схема обещает как "" — они не должны ничего затирать.
    assert ArendaTsWindow._vehicle_tab_data(
        None, {"brand_model": "", "plate_number": ""}, {}, None
    ) == {}


def test_vehicle_tab_data_reads_nested_vehicle_block():
    """Тягач и прицеп принимаются и вложенными в блок vehicle."""
    data = ArendaTsWindow._vehicle_tab_data(
        {"tractor": {"brand_model": TRACTOR_BRAND, "plate_number": TRACTOR_PLATE,
                     "vehicle_type": TRACTOR_TYPE},
         "trailer": {"brand_model": TRAILER_BRAND, "plate_number": TRAILER_PLATE,
                     "color": "Серый"}},
        None, None, None,
    )

    assert data == {
        "tractor_brand": TRACTOR_BRAND,
        "tractor_plate": TRACTOR_PLATE,
        "tractor_type": TRACTOR_TYPE,
        "trailer_brand": TRAILER_BRAND,
        "trailer_plate": TRAILER_PLATE,
    }


def test_vehicle_tab_data_tab_keys_win_over_prompt_keys():
    """
    Ключи вкладки (tractor_brand) важнее ключей схемы промпта (brand_model).

    Ответ может прийти в обоих видах: справочник машин отдаёт имена полей
    вкладки, а схема промпта аренды — краткие имена. Явное значение должно
    побеждать, в каком бы порядке ключи ни пришли.
    """
    forward = ArendaTsWindow._vehicle_tab_data(
        None,
        {"tractor_brand": "Явная марка", "brand_model": "Из промпта",
         "plate_number": TRACTOR_PLATE},
        None, None,
    )
    backward = ArendaTsWindow._vehicle_tab_data(
        None,
        {"brand_model": "Из промпта", "tractor_brand": "Явная марка",
         "plate_number": TRACTOR_PLATE},
        None, None,
    )

    assert forward["tractor_brand"] == "Явная марка"
    assert backward["tractor_brand"] == "Явная марка"


def test_vehicle_tab_data_reads_header_and_lease_dates():
    """Номер и дата договора — из contract, срок аренды — из корня ответа."""
    data = ArendaTsWindow._vehicle_tab_data(
        {"contract_number": "ИЗ-ВКЛАДКИ"},
        None, None,
        {"number": "ТЛ-77", "date": "01.10.2026",
         "lease_start_date": "2026-10-02", "lease_end_date": "2026-10-09"},
        lease_start_date="2026-11-01",
        lease_end_date="2026-11-30",
    )

    assert data["contract_number"] == "ТЛ-77"
    assert data["contract_date"] == "01.10.2026"
    # Корень ответа важнее блока contract: именно так их отдаёт промпт.
    assert data["lease_start_date"] == "2026-11-01"
    assert data["lease_end_date"] == "2026-11-30"


def test_vehicle_tab_data_skips_unknown_keys():
    """Полей, которых нет в бланке (цвет), вкладка не получает."""
    data = ArendaTsWindow._vehicle_tab_data(
        None,
        {"color": "Белый", "year": "2020", "brand_model": TRACTOR_BRAND},
        None, None,
    )

    assert data == {"tractor_brand": TRACTOR_BRAND}


def test_route_tab_data_maps_points_and_route():
    """Маршрут — строкой, точки — массивами loadings / unloadings."""
    data = ArendaTsWindow._route_tab_data(
        {
            "route": ROUTE,
            "loadings": [{"address": LOADING_ADDRESS, "date": LOADING_DATE,
                          "time_from": LOADING_TIME_FROM,
                          "time_to": LOADING_TIME_TO}],
            "unloadings": [{"address": UNLOADING_ADDRESS, "date": UNLOADING_DATE}],
        },
        None,
    )

    assert data["route"] == ROUTE
    assert data["loadings"][0]["address"] == LOADING_ADDRESS
    assert data["loadings"][0]["time_from"] == LOADING_TIME_FROM
    assert data["unloadings"][0]["address"] == UNLOADING_ADDRESS

    # Пустые точки и пустой маршрут — пустой словарь (вкладка не тронута).
    assert ArendaTsWindow._route_tab_data({}, {}) == {}
    assert ArendaTsWindow._route_tab_data(
        {"loadings": [{"address": ""}], "route": "   "}, None
    ) == {}


def test_route_tab_data_accepts_foreign_point_names():
    """Точки чужой схемы (shippers / consignees) тоже принимаются."""
    data = ArendaTsWindow._route_tab_data({
        "shippers": [{"name": "ООО «Склад»", "address": "адрес погрузки"}],
        "consignees": [{"name": "ООО «Клиент»", "address": "адрес выгрузки"}],
    })

    assert data["loadings"] == [{"name": "ООО «Склад»",
                                 "address": "адрес погрузки"}]
    assert data["unloadings"] == [{"name": "ООО «Клиент»",
                                   "address": "адрес выгрузки"}]


def test_route_tab_data_reads_points_from_contract():
    """Точки и маршрут принимаются и внутри блока contract."""
    data = ArendaTsWindow._route_tab_data(
        {"route": "г. Москва — г. Калуга"},
        {"loadings": [{"address": LOADING_ADDRESS}],
         "unloadings": [{"address": UNLOADING_ADDRESS}]},
    )

    assert data["route"] == ROUTE
    assert data["loadings"] == [{"address": LOADING_ADDRESS}]
    assert data["unloadings"] == [{"address": UNLOADING_ADDRESS}]


def test_cargo_tab_data_drops_empty_rows():
    """Пустые строки ответа таблицу груза не очищают."""
    assert ArendaTsWindow._cargo_tab_data([{"brand_model": "", "vin": ""}]) == {}
    assert ArendaTsWindow._cargo_tab_data(None) == {}
    assert ArendaTsWindow._cargo_tab_data([]) == {}
    assert ArendaTsWindow._cargo_tab_data("мусор") == {}
    assert ArendaTsWindow._cargo_tab_data(
        [{"brand_model": "МОДЕЛЬ 1", "vin": ""}]
    ) == {"vehicles": [{"brand_model": "МОДЕЛЬ 1", "vin": ""}]}


def test_crew_tab_data_maps_driver_block():
    """Блок driver: краткие имена схемы → имена полей вкладки «Экипаж»."""
    data = ArendaTsWindow._crew_tab_data({
        "full_name": DRIVER_NAME,
        "birth_date": DRIVER_BIRTH,
        "passport": DRIVER_PASSPORT,
        "passport_issuer": DRIVER_PASSPORT_ISSUER,
        "passport_issue_date": DRIVER_PASSPORT_DATE,
        "license": DRIVER_LICENSE,
        "license_issue_date": DRIVER_LICENSE_DATE,
        "address": DRIVER_ADDRESS,
        "phone": DRIVER_PHONE,
    })

    assert data == {
        "driver_full_name": DRIVER_NAME,
        "driver_birth_date": DRIVER_BIRTH,
        "driver_passport": DRIVER_PASSPORT,
        "driver_passport_issuer": DRIVER_PASSPORT_ISSUER,
        "driver_passport_issue_date": DRIVER_PASSPORT_DATE,
        "driver_license": DRIVER_LICENSE,
        "driver_license_issue_date": DRIVER_LICENSE_DATE,
        "driver_registration_address": DRIVER_ADDRESS,
        "driver_phone": DRIVER_PHONE,
    }


def test_crew_tab_data_skips_empty_and_unknown_keys():
    """Пустой блок и посторонние ключи вкладку не трогают."""
    assert ArendaTsWindow._crew_tab_data({}) == {}
    assert ArendaTsWindow._crew_tab_data(None) == {}
    assert ArendaTsWindow._crew_tab_data({
        "full_name": "", "phone": "   ", "что-то_своё": "значение",
    }) == {}


def test_crew_tab_data_keeps_separate_document_parts():
    """Серия и номер, пришедшие раздельно, доходят до вкладки."""
    data = ArendaTsWindow._crew_tab_data({
        "passport_series": "18 22", "passport_number": "926830",
        "license_series": "99 36", "license_number": "123456",
    })

    assert data == {
        "driver_passport_series": "18 22",
        "driver_passport_number": "926830",
        "driver_license_series": "99 36",
        "driver_license_number": "123456",
    }


def test_price_tab_data_maps_sums_and_rate():
    """Суммы, ставка и особые условия уходят во вкладку «Стоимость»."""
    data = ArendaTsWindow._price_tab_data({
        "sum_wo_vat": FORM_AMOUNT, "sum_vat": VAT_AMOUNT,
        "sum_total": TOTAL_AMOUNT, "vat_rate": VAT_RATE_TEXT,
        "vat_rate_num": 22.0, "special_conditions": SPECIAL_CONDITIONS,
    })

    assert data == {
        "sum_wo_vat": FORM_AMOUNT,
        "sum_vat": VAT_AMOUNT,
        "sum_total": TOTAL_AMOUNT,
        "vat_rate": VAT_RATE_TEXT,
        "vat_rate_num": 22.0,
        "special_conditions": SPECIAL_CONDITIONS,
    }


def test_price_tab_data_skips_zero_and_empty():
    """
    Ноль и пустота — «суммы не было»: они не занимают место в словаре.

    У ИП без НДС единственная сумма лежит в sum_total: нулевая sum_wo_vat
    не должна встать первой в списке приоритетов вкладки.
    """
    assert ArendaTsWindow._price_tab_data({}) == {}
    assert ArendaTsWindow._price_tab_data(None) == {}
    assert ArendaTsWindow._price_tab_data({
        "sum_wo_vat": 0.0, "sum_vat": 0.0, "sum_total": 0.0, "vat_rate": "",
    }) == {}

    data = ArendaTsWindow._price_tab_data({
        "sum_wo_vat": 0.0, "sum_vat": 0.0, "sum_total": TOTAL_AMOUNT,
        "vat_rate": "0%",
    })

    assert data == {"sum_total": TOTAL_AMOUNT, "vat_rate": "0%"}


def test_price_tab_data_reads_amount_from_text():
    """Сумма строкой с разделителями тысяч и запятой распознаётся числом."""
    assert ArendaTsWindow._price_tab_data({
        "sum_wo_vat": "221 099,18",
    }) == {"sum_wo_vat": FORM_AMOUNT}

    assert ArendaTsWindow._price_tab_data({"sum_wo_vat": "не число"}) == {}


def test_vat_rate_is_zero_handles_text_and_number():
    """Ставка «0%», «Без НДС», 0 и 0.0 — нулевая; 22 — нет."""
    assert ArendaTsWindow._vat_rate_is_zero({"vat_rate": "0%"}) is True
    assert ArendaTsWindow._vat_rate_is_zero({"vat_rate": "Без НДС"}) is True
    assert ArendaTsWindow._vat_rate_is_zero(
        {"vat_rate": "НДС не облагается"}
    ) is True
    assert ArendaTsWindow._vat_rate_is_zero({"vat_rate_num": 0.0}) is True
    assert ArendaTsWindow._vat_rate_is_zero({"vat_rate_num": 0}) is True
    assert ArendaTsWindow._vat_rate_is_zero({"vat_rate_num": 22}) is False
    assert ArendaTsWindow._vat_rate_is_zero({"vat_rate": "22%"}) is False
    assert ArendaTsWindow._vat_rate_is_zero({}) is False
    assert ArendaTsWindow._vat_rate_is_zero(None) is False
