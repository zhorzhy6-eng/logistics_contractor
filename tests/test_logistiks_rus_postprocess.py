#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты шагов постобработки заявки «Логистикс Рус» (ЭТАП 3.1.C.A.4).

Каждый шаг проверяется ОТДЕЛЬНО на программно собранном Document() — без
шаблонов templates/*.docx и без полного цикла генерации:

  * RemoveEmptyVehicleRowsStep — таблица на 12 машин: пустые строки данных
    удаляются, заполненные хотя бы одним значением остаются, чужие таблицы
    (другая шапка, шапки нет) не трогаются;
  * RemoveEmptyShipperConsigneeBlocksStep — раздел 1 (ОДИН грузоотправитель
    и до 10 строк «Адрес погрузки №N:») и раздел 2 (до 10 пар «название +
    адрес»): незаполненные строки адресов погрузки удаляются поодиночке,
    метка «Грузоотправитель:» не удаляется никогда, пара грузополучателя
    без обоих значений удаляется, частично заполненная — сохраняется.

Интеграционные проверки (тот же результат в готовом документе после
generate()) остаются в tests/test_logistiks_rus_generator.py.

Все данные синтетические, реальных ПДн нет.
"""

import gc

import pytest
from docx import Document

from core.contracts.logistiks_rus.generator import (
    MAX_CARS,
    MAX_POINTS,
    LogistiksRusGenerator,
)
from core.contracts.logistiks_rus.postprocess import (
    RemoveEmptyShipperConsigneeBlocksStep,
    RemoveEmptyVehicleRowsStep,
)

#: Заголовки таблицы автомобилей — как в бланке («Марка, модель» с запятой).
CARGO_HEADERS = ("№", "Марка, модель", "VIN-номер")

#: Шапка таблицы Формики: слэш вместо запятой — для этого бланка чужая.
FOREIGN_CARGO_HEADERS = ("№", "Марка/Модель", "VIN-номер")

#: Метки строк точек: (начало строки названия, начало строки адреса).
#: У грузополучателя название — с номером блока, у адреса погрузки —
#: с номером адреса (в разделе 1 грузоотправитель ОДИН).
POINT_PREFIXES = {
    "shipper": ("Грузоотправитель:", "Адрес погрузки №"),
    "consignee": ("Грузополучатель №", "Адрес выгрузки:"),
}

#: Заголовки разделов точек — остаются в документе при любом числе точек.
POINT_SECTION_TITLES = {"shipper": "1. ПОГРУЗКА", "consignee": "2. ВЫГРУЗКА"}

#: Имя единственного грузоотправителя в собранном документе (раздел 1).
SHIPPER_NAME = "ООО «Грузоотправитель 1»"

#: Абзац-«хвост» после блоков точек: проверяем, что он не задет.
POINT_SECTION_TAIL = "ОСОБЫЕ УСЛОВИЯ РЕЙСА"


@pytest.fixture(autouse=True)
def _release_documents():
    """Word-файлы на Windows освобождаем до удаления временных копий."""
    yield
    gc.collect()


@pytest.fixture
def generator(tmp_path) -> LogistiksRusGenerator:
    """Генератор без реальных шаблонов: шагам они не нужны."""
    return LogistiksRusGenerator(templates_dir=str(tmp_path))


# ─────────────────────────────────────────────────────────────
# Сборка документа: таблица автомобилей
# ─────────────────────────────────────────────────────────────

def _brand(number: int) -> str:
    """Синтетическая марка модели для строки таблицы."""
    return f"МОДЕЛЬ {number}"


def _vin(number: int) -> str:
    """Синтетический VIN для строки таблицы (не настоящий)."""
    return f"TESTVIN{number:08d}"


def _cargo_table_doc(filled: int, rows: int = MAX_CARS, headers=CARGO_HEADERS):
    """
    Документ с таблицей автомобилей бланка: шапка + `rows` строк данных.

    Первые `filled` строк заполнены (как после docxtpl), у остальных
    заполнен только номер — ровно так выглядит бланк до постобработки.
    """
    doc = Document()
    table = doc.add_table(rows=rows + 1, cols=3)
    for cell, title in zip(table.rows[0].cells, headers):
        cell.text = title

    for number in range(1, rows + 1):
        cells = table.rows[number].cells
        cells[0].text = str(number)
        if number <= filled:
            cells[1].text = _brand(number)
            cells[2].text = _vin(number)

    return doc


def _row_values(table):
    """Значения всех строк таблицы — списком списков."""
    return [[cell.text for cell in row.cells] for row in table.rows]


# ─────────────────────────────────────────────────────────────
# RemoveEmptyVehicleRowsStep
# ─────────────────────────────────────────────────────────────

def test_remove_empty_vehicle_rows_step_removes_all_rows_without_cars(generator):
    """0 машин: все 12 строк данных удалены, шапка остаётся на месте."""
    doc = _cargo_table_doc(filled=0)
    assert len(doc.tables[0].rows) == MAX_CARS + 1

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    table = doc.tables[0]
    assert len(table.rows) == 1
    assert [cell.text for cell in table.rows[0].cells] == list(CARGO_HEADERS)


@pytest.mark.parametrize("cars", [1, 2, 3, MAX_CARS])
def test_remove_empty_vehicle_rows_step_keeps_one_row_per_car(generator, cars):
    """1, 2, 3 и 12 машин: сколько машин, столько строк и остаётся."""
    doc = _cargo_table_doc(filled=cars)

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    table = doc.tables[0]
    assert len(table.rows) == cars + 1
    for number in range(1, cars + 1):
        assert _row_values(table)[number] == [
            str(number), _brand(number), _vin(number),
        ]


def test_remove_empty_vehicle_rows_step_keeps_row_with_only_brand(generator):
    """Марка есть, VIN пуст — строка сохраняется: данные не теряем."""
    doc = _cargo_table_doc(filled=0, rows=1)
    doc.tables[0].rows[1].cells[1].text = "МОДЕЛЬ БЕЗ VIN"

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 2
    assert _row_values(doc.tables[0])[1][1] == "МОДЕЛЬ БЕЗ VIN"


def test_remove_empty_vehicle_rows_step_keeps_row_with_only_vin(generator):
    """VIN есть, марка пуста — строка тоже сохраняется."""
    doc = _cargo_table_doc(filled=0, rows=1)
    doc.tables[0].rows[1].cells[2].text = "TESTVIN00000001"

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 2
    assert _row_values(doc.tables[0])[1][2] == "TESTVIN00000001"


def test_remove_empty_vehicle_rows_step_treats_whitespace_as_empty(generator):
    """Пробелы вместо значений — та же пустая строка."""
    doc = _cargo_table_doc(filled=2)
    table = doc.tables[0]
    table.rows[1].cells[1].text = "   "
    table.rows[1].cells[2].text = "  "

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(table.rows) == 2
    assert _row_values(table)[1] == ["2", _brand(2), _vin(2)]


def test_remove_empty_vehicle_rows_step_ignores_foreign_tables(generator):
    """Чужие таблицы (подписи, реквизиты) не трогаются: шапка не та."""
    doc = Document()
    signatures = doc.add_table(rows=1, cols=2)
    signatures.rows[0].cells[0].text = "Заказчик:"
    signatures.rows[0].cells[1].text = "Экспедитор:"

    requisites = doc.add_table(rows=2, cols=3)
    for cell, title in zip(requisites.rows[0].cells, ("ИНН", "КПП", "ОГРН")):
        cell.text = title

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 1
    assert len(doc.tables[1].rows) == 2


def test_remove_empty_vehicle_rows_step_ignores_slash_header(generator):
    """Шапка «Марка/Модель» (слэш) — таблица другого типа, не наша."""
    doc = _cargo_table_doc(filled=1, headers=FOREIGN_CARGO_HEADERS)

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == MAX_CARS + 1


def test_remove_empty_vehicle_rows_step_ignores_table_without_header(generator):
    """Таблица без шапки бланка не наша — даже если строки пустые."""
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    for cell, value in zip(table.rows[0].cells, ("1", _brand(1), _vin(1))):
        cell.text = value

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 3


def test_remove_empty_vehicle_rows_step_keeps_header_only_table(generator):
    """Таблица из одной шапки: удалять нечего, шапка остаётся."""
    doc = _cargo_table_doc(filled=0, rows=0)

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 1


# ─────────────────────────────────────────────────────────────
# RemoveEmptyShipperConsigneeBlocksStep
# ─────────────────────────────────────────────────────────────

def _point_block(kind: str, number: int, name: str, address: str):
    """Строки блока точки: название и адрес.

    Раздел 1 (shipper): грузоотправитель ОДИН — имя без номера, а адрес
    погрузки нумерован («Адрес погрузки №N»). Раздел 2 (consignee): пара
    «Грузополучатель №N» + «Адрес выгрузки».
    """
    if kind == "shipper":
        return f"Грузоотправитель: {name}", f"Адрес погрузки №{number}: {address}"
    return f"Грузополучатель №{number}: {name}", f"Адрес выгрузки: {address}"


def _points_doc(kind: str, filled: int, total: int = MAX_POINTS):
    """
    Документ с `total` строками точек, из которых заполнены первые `filled`.

    Для грузополучателей незаполненный блок — это строки «Грузополучатель
    №N:» и «Адрес выгрузки:» без значений: ровно так выглядит бланк после
    docxtpl. Для адресов погрузки пустая строка — «Адрес погрузки №N:»;
    метка «Грузоотправитель:» в разделе 1 стоит одна и печатается всегда.
    """
    doc = Document()
    doc.add_paragraph(POINT_SECTION_TITLES[kind])

    if kind == "shipper":
        doc.add_paragraph(f"Грузоотправитель: {SHIPPER_NAME}")

    for number in range(1, total + 1):
        is_filled = number <= filled
        name = f"ООО «Точка {number}»" if is_filled else ""
        address = f"Адрес точки {number}" if is_filled else ""
        if kind == "shipper":
            doc.add_paragraph(_point_block(kind, number, name, address)[1])
        else:
            for text in _point_block(kind, number, name, address):
                doc.add_paragraph(text)

    doc.add_paragraph(POINT_SECTION_TAIL)
    return doc


def _texts(doc):
    """Тексты абзацев верхнего уровня без крайних пробелов."""
    return [p.text.strip() for p in doc.paragraphs]


def _count(doc, prefix: str) -> int:
    """Сколько абзацев начинается с указанной метки."""
    return sum(1 for text in _texts(doc) if text.startswith(prefix))


@pytest.mark.parametrize("kind", ["shipper", "consignee"])
@pytest.mark.parametrize("filled", [1, 2, 5, MAX_POINTS])
def test_remove_empty_point_blocks_keeps_only_filled_blocks(generator, kind, filled):
    """Сколько строк заполнено, столько и остаётся (обоих видов)."""
    doc = _points_doc(kind, filled=filled)

    RemoveEmptyShipperConsigneeBlocksStep(generator).apply(doc, None)

    name_prefix, address_prefix = POINT_PREFIXES[kind]
    assert _count(doc, address_prefix) == filled
    if kind == "consignee":
        assert _count(doc, name_prefix) == filled
    else:
        # Грузоотправитель один и печатается всегда — от числа адресов
        # он не зависит.
        assert _count(doc, name_prefix) == 1
        assert _texts(doc)[1] == f"Грузоотправитель: {SHIPPER_NAME}"

    # Шапка раздела и хвост документа остаются на месте.
    assert _texts(doc)[0] == POINT_SECTION_TITLES[kind]
    assert _texts(doc)[-1] == POINT_SECTION_TAIL


def test_remove_empty_point_blocks_removes_all_ten_when_empty(generator):
    """0 заполненных: все 10 адресов погрузки удалены, раздел и метка остаются."""
    doc = _points_doc("shipper", filled=0)

    RemoveEmptyShipperConsigneeBlocksStep(generator).apply(doc, None)

    assert _texts(doc) == [
        POINT_SECTION_TITLES["shipper"],
        f"Грузоотправитель: {SHIPPER_NAME}",
        POINT_SECTION_TAIL,
    ]


def test_remove_empty_point_blocks_removes_eight_of_ten_shippers(generator):
    """2 заполненных адреса из 10: удаляются 8 пустых, метка остаётся."""
    doc = _points_doc("shipper", filled=2)

    RemoveEmptyShipperConsigneeBlocksStep(generator).apply(doc, None)

    assert _texts(doc) == [
        "1. ПОГРУЗКА",
        f"Грузоотправитель: {SHIPPER_NAME}",
        "Адрес погрузки №1: Адрес точки 1",
        "Адрес погрузки №2: Адрес точки 2",
        "ОСОБЫЕ УСЛОВИЯ РЕЙСА",
    ]


def test_shipper_label_is_never_removed(generator):
    """Метка «Грузоотправитель:» не удаляется никогда, даже без адресов."""
    doc = Document()
    doc.add_paragraph("Грузоотправитель:")
    doc.add_paragraph("Адрес погрузки №1: ")
    doc.add_paragraph("Адрес погрузки №2:")
    doc.add_paragraph("Адрес погрузки: старый формат без номера")

    RemoveEmptyShipperConsigneeBlocksStep(generator).apply(doc, None)

    assert _texts(doc) == [
        "Грузоотправитель:",
        "Адрес погрузки: старый формат без номера",
    ]


def test_remove_empty_point_blocks_removes_nothing_when_all_filled(generator):
    """Все 10 блоков заполнены: документ не меняется."""
    doc = _points_doc("consignee", filled=MAX_POINTS)
    before = _texts(doc)

    RemoveEmptyShipperConsigneeBlocksStep(generator).apply(doc, None)

    assert _texts(doc) == before
    assert len(doc.paragraphs) == MAX_POINTS * 2 + 2


@pytest.mark.parametrize(
    "name, address",
    [("ООО «Только имя»", ""), ("", "Только адрес")],
)
def test_remove_empty_point_blocks_keeps_partially_filled_consignee(
    generator, name, address
):
    """Блок грузополучателя с одним заполненным полем сохраняется целиком."""
    doc = Document()
    for text in _point_block("consignee", 1, name, address):
        doc.add_paragraph(text)

    RemoveEmptyShipperConsigneeBlocksStep(generator).apply(doc, None)

    assert _texts(doc) == [
        text.strip() for text in _point_block("consignee", 1, name, address)
    ]


def test_remove_empty_shipper_address_keeps_filled_one(generator):
    """Адрес погрузки с заполненным значением сохраняется."""
    doc = Document()
    doc.add_paragraph("Грузоотправитель: ООО «Первый»")
    doc.add_paragraph("Адрес погрузки №1: Адрес первый")

    RemoveEmptyShipperConsigneeBlocksStep(generator).apply(doc, None)

    assert _texts(doc) == [
        "Грузоотправитель: ООО «Первый»",
        "Адрес погрузки №1: Адрес первый",
    ]


def test_shipper_address_with_only_label_is_removed_but_name_stays(generator):
    """Строка «Адрес погрузки №N:» без значения уходит, метка имени — нет."""
    doc = Document()
    doc.add_paragraph("Грузоотправитель: ООО «Первый»")
    doc.add_paragraph("Адрес погрузки №1: ")
    doc.add_paragraph("Адрес погрузки №2: Адрес второй")

    RemoveEmptyShipperConsigneeBlocksStep(generator).apply(doc, None)

    assert _texts(doc) == [
        "Грузоотправитель: ООО «Первый»",
        "Адрес погрузки №2: Адрес второй",
    ]


def test_remove_empty_point_blocks_handles_both_kinds_in_one_document(generator):
    """Адреса погрузки и грузополучатели разбираются независимо."""
    doc = Document()
    doc.add_paragraph(f"Грузоотправитель: {SHIPPER_NAME}")
    doc.add_paragraph("Адрес погрузки №1: Адрес первый")
    doc.add_paragraph("Адрес погрузки №2:")
    for number in (1, 2):
        for text in _point_block("consignee", number, "", ""):
            doc.add_paragraph(text)

    RemoveEmptyShipperConsigneeBlocksStep(generator).apply(doc, None)

    assert _texts(doc) == [
        f"Грузоотправитель: {SHIPPER_NAME}",
        "Адрес погрузки №1: Адрес первый",
    ]


def test_remove_empty_point_blocks_ignores_foreign_paragraphs(generator):
    """
    Чужие абзацы не удаляются.

    Метка грузополучателя — только с номером блока («Грузополучатель №1:»),
    метка адреса погрузки — только с номером адреса («Адрес погрузки №1:»),
    строки дат и других разделов под метки не подходят.
    """
    doc = Document()
    doc.add_paragraph("Дата / время погрузки: 26.09.2026 г. Время с 08:00 по 20:00")
    doc.add_paragraph("Плановая дата / время завершения выгрузки: 01.10.2026 г.")
    doc.add_paragraph("Грузополучатель: ООО «Без номера»")
    doc.add_paragraph("Адрес погрузки: старый формат")
    doc.add_paragraph("Общее количество: 3 шт.")

    RemoveEmptyShipperConsigneeBlocksStep(generator).apply(doc, None)

    assert len(doc.paragraphs) == 5


def test_remove_empty_point_blocks_does_not_swallow_next_paragraph(generator):
    """Одиночная пустая метка удаляется одна: следующий абзац — не её адрес."""
    doc = Document()
    doc.add_paragraph("Грузополучатель №1:")
    doc.add_paragraph("Дата / время погрузки: 26.09.2026 г.")

    RemoveEmptyShipperConsigneeBlocksStep(generator).apply(doc, None)

    assert _texts(doc) == ["Дата / время погрузки: 26.09.2026 г."]


# ─────────────────────────────────────────────────────────────
# Изоляция шагов
# ─────────────────────────────────────────────────────────────

def _mixed_doc():
    """Документ сразу с таблицей автомобилей и строками точек."""
    doc = _cargo_table_doc(filled=1)
    doc.add_paragraph(f"Грузоотправитель: {SHIPPER_NAME}")
    doc.add_paragraph("Адрес погрузки №1: Адрес первый")
    doc.add_paragraph("Адрес погрузки №2:")
    return doc


def test_steps_touch_only_their_own_part_of_document(generator):
    """Шаги изолированы: строки таблицы и строки точек друг друга не трогают."""
    assert RemoveEmptyVehicleRowsStep.name != RemoveEmptyShipperConsigneeBlocksStep.name

    doc = _mixed_doc()
    points_before = _texts(doc)

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)
    assert len(doc.tables[0].rows) == 2  # пустые строки груза убраны
    assert _texts(doc) == points_before  # строки точек не тронуты

    RemoveEmptyShipperConsigneeBlocksStep(generator).apply(doc, None)
    assert _texts(doc) == [
        f"Грузоотправитель: {SHIPPER_NAME}",
        "Адрес погрузки №1: Адрес первый",
    ]
    assert len(doc.tables[0].rows) == 2  # таблица не тронута вторым шагом


def test_steps_on_empty_document_are_noops(generator):
    """Пустой документ: оба шага ничего не делают и не падают."""
    doc = Document()

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)
    RemoveEmptyShipperConsigneeBlocksStep(generator).apply(doc, None)

    assert not doc.tables
    assert not doc.paragraphs
