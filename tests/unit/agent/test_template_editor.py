#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты безопасной правки шаблонов (template_editor.py).

Проверяется главный закон правки:

  1. предложение НЕ меняет файл;
  2. абзац с плейсхолдером не правится никогда (safe: False);
  3. перед записью создаётся .bak;
  4. после записи шаблон проверяется, и при поломке откатывается;
  5. применяется только отмеченное и только безопасное.

Шаблоны собираются на месте (python-docx / openpyxl): тест не зависит от
боевых шаблонов и проверяет именно механику правки.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest

from services.agent.tools import template_analyzer as analyzer
from services.agent.tools import template_editor as te


# ─────────────────────────────────────────────────────────────
# Помощники
# ─────────────────────────────────────────────────────────────

def make_docx(path: Path, paragraphs, runs=None) -> Path:
    """
    Создаёт DOCX: абзацы обычные или по частям (runs).

    :param paragraphs: тексты абзацев;
    :param runs: {номер абзаца: [куски]}: абзац собирается из нескольких
        runs — так проверяется случай «текст разбит оформлением».
    """
    from docx import Document

    document = Document()
    for index, text in enumerate(paragraphs):
        paragraph = document.add_paragraph()
        parts = (runs or {}).get(index)
        if parts:
            for part in parts:
                paragraph.add_run(part)
        else:
            paragraph.add_run(text)
    document.save(str(path))
    return path


def make_xlsx(path: Path, rows) -> Path:
    """Создаёт XLSX с заданными строками ячеек."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(list(row))
    workbook.save(str(path))
    return path


def docx_texts(path: Path):
    """Тексты абзацев DOCX (для проверки, что именно изменилось)."""
    from docx import Document

    return [paragraph.text for paragraph in Document(str(path)).paragraphs]


def file_hash(path: Path) -> str:
    """Контрольная сумма файла."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ─────────────────────────────────────────────────────────────
# 1. Предложение правки
# ─────────────────────────────────────────────────────────────

def test_proposal_finds_safe_change(tmp_path):
    """Обычный абзац с «НДС 20%» — безопасная правка."""
    path = make_docx(tmp_path / "a.docx", ["Стоимость с НДС 20%."])
    plan = te.propose_template_fix(path, "НДС 20%")
    assert plan["ok"] is True
    assert len(plan["changes"]) == 1
    change = plan["changes"][0]
    assert change["safe"] is True
    assert change["old"] == "НДС 20%"
    assert change["new"] == "НДС 22%"
    assert change["has_placeholder"] is False


def test_proposal_does_not_change_file(tmp_path):
    """Предложение не пишет в шаблон ни байта."""
    path = make_docx(tmp_path / "b.docx", ["НДС 20%", "{{contract_number}}"])
    before = file_hash(path)
    te.propose_template_fix(path, "НДС 20%")
    assert file_hash(path) == before


def test_proposal_marks_placeholder_paragraph_unsafe(tmp_path):
    """Абзац с плейсхолдером — только руками, с причиной."""
    path = make_docx(tmp_path / "c.docx", [
        "НДС 20% в том числе {{contract_number}}",
    ])
    plan = te.propose_template_fix(path, "НДС 20%")
    change = plan["changes"][0]
    assert change["safe"] is False
    assert change["has_placeholder"] is True
    assert "плейсхолдер" in change["reason"]


def test_proposal_marks_split_runs_unsafe(tmp_path):
    """Текст, разбитый на несколько runs, не правится автоматически."""
    path = make_docx(tmp_path / "d.docx", [""],
                     runs={0: ["НДС ", "20%"]})
    plan = te.propose_template_fix(path, "НДС 20%")
    change = plan["changes"][0]
    assert change["safe"] is False
    assert "фрагментов оформления" in change["reason"]


def test_proposal_without_known_replacement_is_unsafe(tmp_path):
    """Замену не знаем — предлагать нечего: safe False, причина названа."""
    path = make_docx(tmp_path / "e.docx", ["НДС 20%"])
    plan = te.propose_template_fix(path, "НДС 20%", old="НДС 20%", new="")
    change = plan["changes"][0]
    assert change["safe"] is False
    assert "замена" in change["reason"]


def test_proposal_accepts_explicit_old_and_new(tmp_path):
    """Явные old/new важнее справочника известных замен."""
    path = make_docx(tmp_path / "f.docx", ["Оплата за 45 дней"])
    plan = te.propose_template_fix(path, "срок оплаты", old="45 дней",
                                   new="30 дней")
    assert plan["changes"][0]["new"] == "30 дней"
    assert plan["changes"][0]["safe"] is True


def test_proposal_not_found_says_so(tmp_path):
    """Текста в шаблоне нет — предупреждение, а не пустой список молча."""
    path = make_docx(tmp_path / "g.docx", ["Чистый текст"])
    plan = te.propose_template_fix(path, "НДС 20%")
    assert plan["ok"] is False
    assert any("не найдено" in w for w in plan["warnings"])


def test_proposal_for_missing_file(tmp_path):
    """Нет шаблона — понятный отказ."""
    plan = te.propose_template_fix(tmp_path / "нет.docx", "НДС 20%")
    assert plan["ok"] is False


def test_proposal_finds_xlsx_cells(tmp_path):
    """В XLSX правка предлагается по ячейке."""
    path = make_xlsx(tmp_path / "h.xlsx", [["НДС 20%"], ["{{vin}}"]])
    plan = te.propose_template_fix(path, "НДС 20%")
    change = plan["changes"][0]
    assert change["sheet"]
    assert change["cell"] == "A1"
    assert change["safe"] is True


def test_proposal_xlsx_placeholder_cell_is_unsafe(tmp_path):
    """Ячейка с плейсхолдером — руками."""
    path = make_xlsx(tmp_path / "i.xlsx", [["НДС 20% {{vin}}"]])
    plan = te.propose_template_fix(path, "НДС 20%")
    assert plan["changes"][0]["safe"] is False
    assert "плейсхолдер" in plan["changes"][0]["reason"]


def test_format_proposal_names_manual_places(tmp_path):
    """Текст предложения: сколько мест и что придётся делать руками."""
    path = make_docx(tmp_path / "j.docx", [
        "НДС 20%",
        "НДС 20% {{contract_number}}",
    ])
    plan = te.propose_template_fix(path, "НДС 20%")
    text = te.format_proposal(plan)
    assert "Найдено мест: 2" in text
    assert "автоматически: 1" in text
    assert "руками" in text


# ─────────────────────────────────────────────────────────────
# 2. Применение
# ─────────────────────────────────────────────────────────────

def test_apply_changes_safe_paragraph_only(tmp_path):
    """Правка меняет безопасный абзац и НЕ трогает абзац с плейсхолдером."""
    path = make_docx(tmp_path / "k.docx", [
        "Стоимость с НДС 20%.",
        "НДС 20% указан в {{contract_number}}",
    ])
    plan = te.propose_template_fix(path, "НДС 20%")
    result = te.apply_template_changes(path, plan["changes"])
    assert result["ok"] is True
    assert result["applied"] == 1
    texts = docx_texts(path)
    assert texts[0] == "Стоимость с НДС 22%."
    assert "{{contract_number}}" in texts[1]
    assert "НДС 20%" in texts[1]


def test_apply_makes_backup(tmp_path):
    """Перед правкой создаётся .bak с исходным содержимым."""
    path = make_docx(tmp_path / "l.docx", ["НДС 20%"])
    before = file_hash(path)
    plan = te.propose_template_fix(path, "НДС 20%")
    result = te.apply_template_changes(path, plan["changes"])
    backup = Path(result["backup"])
    assert backup.is_file()
    assert file_hash(backup) == before
    assert backup.name.endswith(te.BACKUP_SUFFIX)


def test_backup_is_not_overwritten_by_second_edit(tmp_path):
    """Вторая правка не затирает первую копию: она — состояние ДО всех правок."""
    path = make_docx(tmp_path / "m.docx", ["НДС 20%"])
    original = file_hash(path)
    plan = te.propose_template_fix(path, "НДС 20%")
    te.apply_template_changes(path, plan["changes"])
    # Вторая правка: файл уже изменён, .bak должен остаться прежним.
    te.make_backup(path)
    assert file_hash(te.backup_path(path)) == original


def test_apply_keeps_placeholders_intact(tmp_path):
    """После правки все плейсхолдеры на месте (главная проверка)."""
    path = make_docx(tmp_path / "n.docx", [
        "Договор {{contract_number}} от {{contract_date}}",
        "НДС 20% в стоимости",
    ])
    before = analyzer.analyze_template(path)
    plan = te.propose_template_fix(path, "НДС 20%")
    result = te.apply_template_changes(path, plan["changes"])
    after = analyzer.analyze_template(path)
    assert result["ok"] is True
    assert analyzer.regressions(before, after) == []
    assert after["placeholder_count"] == before["placeholder_count"]


def test_apply_case_insensitive(tmp_path):
    """Замена находит текст в другом регистре, но не портит остальное."""
    path = make_docx(tmp_path / "o.docx", ["стоимость с ндс 20% и всё"])
    plan = te.propose_template_fix(path, "НДС 20%")
    te.apply_template_changes(path, plan["changes"])
    assert docx_texts(path)[0] == "стоимость с НДС 22% и всё"


def test_apply_selected_only(tmp_path):
    """Применяется только отмеченное оператором."""
    path = make_docx(tmp_path / "p.docx", ["НДС 20% раз", "НДС 20% два"])
    plan = te.propose_template_fix(path, "НДС 20%")
    result = te.apply_template_changes(path, plan["changes"], selected=[1])
    assert result["applied"] == 1
    texts = docx_texts(path)
    assert texts[0] == "НДС 20% раз"
    assert texts[1] == "НДС 22% два"


def test_apply_without_selection_refuses(tmp_path):
    """Ничего не отмечено — файл не трогаем."""
    path = make_docx(tmp_path / "q.docx", ["НДС 20%"])
    before = file_hash(path)
    plan = te.propose_template_fix(path, "НДС 20%")
    result = te.apply_template_changes(path, plan["changes"], selected=[])
    assert result["ok"] is False
    assert "не выбрано" in result["problems"][0]
    assert file_hash(path) == before


def test_apply_missing_file(tmp_path):
    """Нет файла — отказ без исключения."""
    result = te.apply_template_changes(tmp_path / "нет.docx", [{"safe": True}])
    assert result["ok"] is False


def test_apply_rolls_back_when_validation_fails(tmp_path, monkeypatch):
    """Проверка после записи не прошла — файл восстановлен из .bak."""
    path = make_docx(tmp_path / "r.docx", ["НДС 20%"])
    before = file_hash(path)
    plan = te.propose_template_fix(path, "НДС 20%")

    monkeypatch.setattr(analyzer, "regressions",
                        lambda before_, after_: ["пропал плейсхолдер"])
    result = te.apply_template_changes(path, plan["changes"])
    assert result["ok"] is False
    assert result["rolled_back"] is True
    assert result["problems"] == ["пропал плейсхолдер"]
    assert file_hash(path) == before, "шаблон должен вернуться как было"


def test_apply_xlsx_replaces_cell(tmp_path):
    """Правка XLSX меняет ячейку и не портит соседние."""
    path = make_xlsx(tmp_path / "s.xlsx", [["НДС 20%", "{{vin}}"]])
    plan = te.propose_template_fix(path, "НДС 20%")
    result = te.apply_template_changes(path, plan["changes"])
    assert result["ok"] is True
    from openpyxl import load_workbook

    workbook = load_workbook(str(path))
    sheet = workbook.active
    assert sheet["A1"].value == "НДС 22%"
    assert sheet["B1"].value == "{{vin}}"
    workbook.close()


def test_selected_changes_never_takes_unsafe():
    """Небезопасная правка не применяется, даже если её «отметили»."""
    changes = [
        {"safe": True, "old": "a"},
        {"safe": False, "old": "b"},
    ]
    chosen = te.selected_changes(changes, selected=[0, 1])
    assert len(chosen) == 1
    assert chosen[0]["old"] == "a"


def test_restore_backup(tmp_path):
    """Откат из .bak возвращает исходное содержимое."""
    path = make_docx(tmp_path / "t.docx", ["НДС 20%"])
    before = file_hash(path)
    plan = te.propose_template_fix(path, "НДС 20%")
    te.apply_template_changes(path, plan["changes"])
    assert file_hash(path) != before
    assert te.restore_backup(path) is True
    assert file_hash(path) == before


def test_apply_plan_wrapper(tmp_path):
    """apply_plan работает по словарю предложения (как диалог)."""
    path = make_docx(tmp_path / "u.docx", ["НДС 20%"])
    plan = te.propose_template_fix(path, "НДС 20%")
    result = te.apply_plan(plan)
    assert result["ok"] is True
    assert result["applied"] == 1


def test_apply_skips_stale_change(tmp_path):
    """Файл изменился после предложения — правка пропущена, а не применена вслепую."""
    path = make_docx(tmp_path / "v.docx", ["НДС 20%"])
    plan = te.propose_template_fix(path, "НДС 20%")
    # Пока оператор думал, текст уже поправили руками.
    from docx import Document

    document = Document(str(path))
    document.paragraphs[0].text = "НДС 22%"
    document.save(str(path))
    result = te.apply_template_changes(path, plan["changes"])
    assert result["ok"] is True
    assert result["applied"] == 0
    assert result["skipped"]

