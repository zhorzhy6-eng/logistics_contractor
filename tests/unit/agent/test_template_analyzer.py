#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты анализатора шаблонов (template_analyzer.py).

Проверяется то, ради чего анализатор существует: шаблон можно ПРОЧИТАТЬ
до правки — увидеть плейсхолдеры, парность блоков, стили и устаревший
текст. Отдельно проверяется главное обещание модуля: он НЕ меняет файл
(байты шаблона до и после анализа совпадают).

Шаблоны для тестов собираются на месте (python-docx / openpyxl): так тест
не зависит от содержимого боевых шаблонов и проверяет именно разбор.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest

from services.agent.tools import template_analyzer as ta

#: Корень проекта — по маркеру (pytest.ini), а не по числу parents.
PROJECT_ROOT = next(
    _path for _path in Path(__file__).resolve().parents
    if (_path / "pytest.ini").is_file()
)


# ─────────────────────────────────────────────────────────────
# Сборка шаблонов для тестов
# ─────────────────────────────────────────────────────────────

def make_docx(path: Path, paragraphs, table_rows=None) -> Path:
    """Создаёт DOCX с заданными абзацами и (необязательно) таблицей."""
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
    document.save(str(path))
    return path


def make_xlsx(path: Path, rows) -> Path:
    """Создаёт XLSX с заданными строками ячеек."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(list(row))
    workbook.save(str(path))
    return path


def file_hash(path: Path) -> str:
    """Контрольная сумма файла: доказательство, что анализатор не писал."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ─────────────────────────────────────────────────────────────
# 1. Плейсхолдеры
# ─────────────────────────────────────────────────────────────

def test_placeholders_are_found_with_counts(tmp_path):
    """Плейсхолдеры находятся, повтор считаетcя, порядок — как в документе."""
    path = make_docx(tmp_path / "a.docx", [
        "Договор № {{contract_number}} от {{contract_date}}",
        "Сторона: {{contract_number}}",
    ])
    result = ta.analyze_template(path)
    names = [item["name"] for item in result["placeholders"]]
    assert names == ["contract_number", "contract_date"]
    counts = {item["name"]: item["count"] for item in result["placeholders"]}
    assert counts["contract_number"] == 2
    assert result["placeholder_count"] == 2
    assert result["placeholders"][0]["raw"] == "{{contract_number}}"


def test_placeholders_in_table_cells_are_found(tmp_path):
    """Плейсхолдеры в таблице тоже видны: там их большинство."""
    path = make_docx(tmp_path / "b.docx", ["Таблица:"],
                     table_rows=[["{{vin}}", "{{brand_model}}"]])
    result = ta.analyze_template(path)
    names = {item["name"] for item in result["placeholders"]}
    assert {"vin", "brand_model"} <= names
    assert result["tables_count"] == 1


def test_document_without_placeholders_warns(tmp_path):
    """Шаблон без плейсхолдеров — предупреждение: возможно, это не тот файл."""
    path = make_docx(tmp_path / "c.docx", ["Просто текст без подстановок"])
    result = ta.analyze_template(path)
    assert result["placeholder_count"] == 0
    assert any("Плейсхолдеров не найдено" in w for w in result["warnings"])


# ─────────────────────────────────────────────────────────────
# 2. Условные блоки
# ─────────────────────────────────────────────────────────────

def test_closed_if_block_is_recognised(tmp_path):
    """Закрытый if помечен closed, предупреждений нет."""
    path = make_docx(tmp_path / "d.docx", [
        "{%p if is_vat_free %}",
        "Без НДС",
        "{%p endif %}",
        "{{contract_number}}",
    ])
    result = ta.analyze_template(path)
    kinds = [block["kind"] for block in result["conditional_blocks"]]
    assert kinds == ["if", "endif"]
    assert result["conditional_blocks"][0]["closed"] is True
    assert result["unclosed_tags"] == []


def test_unclosed_if_is_a_warning(tmp_path):
    """Незакрытый if — документ не соберётся, это предупреждение."""
    path = make_docx(tmp_path / "e.docx", [
        "{%p if is_vat_free %}",
        "Без НДС",
        "{{contract_number}}",
    ])
    result = ta.analyze_template(path)
    assert len(result["unclosed_tags"]) == 1
    assert any("Не закрыт блок" in w for w in result["warnings"])


def test_for_loop_pairing(tmp_path):
    """Цикл for/endfor парный и не считается ошибкой."""
    path = make_docx(tmp_path / "f.docx", [
        "{% for vehicle in vehicles %}",
        "{{vehicle.vin}}",
        "{% endfor %}",
    ])
    result = ta.analyze_template(path)
    assert result["unclosed_tags"] == []
    assert [block["kind"] for block in result["conditional_blocks"]] == ["for", "endfor"]


def test_extra_endif_is_a_warning(tmp_path):
    """Закрывающий тег без открытия — тоже поломка."""
    path = make_docx(tmp_path / "g.docx", ["{%p endif %}", "{{contract_number}}"])
    result = ta.analyze_template(path)
    assert len(result["extra_closing_tags"]) == 1
    assert any("без открытия" in w for w in result["warnings"])


# ─────────────────────────────────────────────────────────────
# 3. Устаревшие тексты
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("phrase", [
    "Стоимость указана с НДС 20%.",
    "В том числе НДС 20 %.",
    "20% НДС включены в цену.",
])
def test_vat_20_is_detected(tmp_path, phrase):
    """«НДС 20%» в трёх написаниях — устаревшая ставка."""
    path = make_docx(tmp_path / "h.docx", [phrase, "{{contract_number}}"])
    result = ta.analyze_template(path)
    assert result["has_vat_20"] is True
    assert result["outdated"][0]["kind"] == "vat_20"
    assert any("22%" in w for w in result["warnings"])


def test_current_vat_is_not_outdated(tmp_path):
    """Ставка 22% — действующая, предупреждения нет."""
    path = make_docx(tmp_path / "i.docx", ["Стоимость с НДС 22%.", "{{contract_number}}"])
    result = ta.analyze_template(path)
    assert result["has_vat_20"] is False
    assert result["outdated"] == []


def test_vat_18_is_outdated(tmp_path):
    """Ставка 18% — тоже устаревшая."""
    path = make_docx(tmp_path / "j.docx", ["НДС 18%", "{{contract_number}}"])
    result = ta.analyze_template(path)
    assert result["outdated"][0]["kind"] == "vat_18"


# ─────────────────────────────────────────────────────────────
# 4. Файлы, формат, паспорт
# ─────────────────────────────────────────────────────────────

def test_missing_file_is_reported_not_raised(tmp_path):
    """Нет файла — результат с ошибкой, а не исключение."""
    result = ta.analyze_template(tmp_path / "нет.docx")
    assert result["ok"] is False
    assert any("файла нет" in w for w in result["warnings"])


def test_unsupported_format_is_reported(tmp_path):
    """TXT не шаблон — сказано, какие форматы поддерживаются."""
    path = tmp_path / "x.txt"
    path.write_text("{{contract_number}}", encoding="utf-8")
    result = ta.analyze_template(path)
    assert result["ok"] is False
    assert any("не поддерживается" in w for w in result["warnings"])


def test_broken_docx_is_reported(tmp_path):
    """Битый DOCX — понятное предупреждение, окно не падает."""
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a docx at all")
    result = ta.analyze_template(path)
    assert result["ok"] is False
    assert any("не открывается" in w for w in result["warnings"])


def test_xlsx_template_is_analyzed(tmp_path):
    """XLSX анализируется: подписи и подстановки в ячейках."""
    path = make_xlsx(tmp_path / "k.xlsx", [
        ["Марка", "VIN"],
        ["{{brand_model}}", "{{vin}}"],
    ])
    result = ta.analyze_template(path)
    assert result["ok"] is True
    assert result["format"] == "xlsx"
    names = {item["name"] for item in result["placeholders"]}
    assert names == {"brand_model", "vin"}


def test_styles_and_counts_are_collected(tmp_path):
    """Стили абзацев и число абзацев собираются для проверки правки."""
    path = make_docx(tmp_path / "l.docx", ["Первый", "Второй", "{{contract_number}}"])
    result = ta.analyze_template(path)
    assert result["paragraphs_count"] >= 3
    assert "Normal" in result["styles"]
    assert result["styles_count"] >= 1


def test_analysis_does_not_change_the_file(tmp_path):
    """Главное обещание: анализ не пишет в шаблон ни байта."""
    path = make_docx(tmp_path / "m.docx", ["{{contract_number}}"])
    before = file_hash(path)
    ta.analyze_template(path)
    assert file_hash(path) == before


def test_real_project_template_is_whole():
    """Боевой шаблон перевозки: открывается, плейсхолдеры есть, блоки парные."""
    path = PROJECT_ROOT / "templates" / "shablon_ooo.docx"
    if not path.is_file():
        pytest.skip("шаблона нет в этой сборке")
    result = ta.analyze_template(path)
    assert result["ok"] is True
    assert result["placeholder_count"] > 50
    assert result["unclosed_tags"] == []
    assert result["extra_closing_tags"] == []
    assert ta.is_template_safe(path) is True


# ─────────────────────────────────────────────────────────────
# 5. Сравнение «до / после»
# ─────────────────────────────────────────────────────────────

def _analysis(placeholders, styles=None, unclosed=None):
    """Результат анализа минимальным набором полей (для regressions)."""
    return {
        "placeholders": [{"name": name, "count": count}
                          for name, count in placeholders],
        "styles": list(styles or ["Normal"]),
        "unclosed_tags": list(unclosed or []),
        "ok": True,
    }


def test_regressions_empty_when_nothing_broke():
    """Тот же набор плейсхолдеров и стилей — проблем нет."""
    before = _analysis([("contract_number", 1)])
    after = _analysis([("contract_number", 1)])
    assert ta.regressions(before, after) == []


def test_regressions_lost_placeholder():
    """Пропавший плейсхолдер — поломка (поле договора останется пустым)."""
    problems = ta.regressions(_analysis([("contract_number", 1)]),
                              _analysis([]))
    assert any("Пропали плейсхолдеры" in p for p in problems)


def test_regressions_changed_count():
    """Задвоенный плейсхолдер — тоже поломка."""
    problems = ta.regressions(_analysis([("vin", 1)]), _analysis([("vin", 2)]))
    assert any("Изменилось число вхождений" in p for p in problems)


def test_regressions_new_placeholder():
    """Новый плейсхолдер без данных в генераторе — ошибка сборки."""
    problems = ta.regressions(_analysis([("vin", 1)]),
                              _analysis([("vin", 1), ("new_field", 1)]))
    assert any("новые плейсхолдеры" in p for p in problems)


def test_regressions_unclosed_block():
    """Незакрытый блок после правки."""
    problems = ta.regressions(_analysis([("vin", 1)]),
                              _analysis([("vin", 1)], unclosed=[{"kind": "if"}]))
    assert any("незакрытые блоки" in p for p in problems)


def test_regressions_lost_style():
    """Пропавший стиль абзаца — переоформление текста."""
    problems = ta.regressions(_analysis([("vin", 1)], ["Normal", "Заголовок 1"]),
                              _analysis([("vin", 1)], ["Normal"]))
    assert any("стили абзацев" in p for p in problems)


def test_is_template_safe_false_without_placeholders(tmp_path):
    """Шаблон без подстановок не считается целым."""
    path = make_docx(tmp_path / "n.docx", ["Текст"])
    assert ta.is_template_safe(path) is False


# ─────────────────────────────────────────────────────────────
# 6. Регресс по всем шаблонам проекта
# ─────────────────────────────────────────────────────────────

def test_all_project_templates_open_and_are_whole():
    """
    Каждый шаблон проекта открывается, блоки парные, плейсхолдеры целы.

    Регресс на правку шаблонов: если файл случайно пересохранили или
    сломали тег, тест покраснеет ДО того, как договор не соберётся у
    оператора. Проверяются и публичные, и приватные бланки (если папка
    приватной части подключена).
    """
    templates = ta.list_templates()
    assert templates, "ни одного шаблона не найдено — анализатор сломан"
    broken = []
    for item in templates:
        result = ta.analyze_template(item["path"])
        if not result.get("ok"):
            broken.append(f"{item['name']}: не открывается")
            continue
        if result.get("unclosed_tags"):
            broken.append(f"{item['name']}: незакрытый блок")
        if result.get("extra_closing_tags"):
            broken.append(f"{item['name']}: закрывающий тег без открытия")
    assert broken == [], "шаблоны сломаны: " + "; ".join(broken)


def test_shablon_docx_templates_have_placeholders():
    """
    DOCX-бланк, который начинается с shablon_, обязан быть заполняемым.

    Имя shablon_ означает рабочий бланк договора: если плейсхолдеров в нём
    не осталось, docxtpl выдаст пустой документ. XLSX здесь не проверяются:
    заявку (shablon_havaly.xlsx) генератор заполняет ПО КООРДИНАТАМ ячеек
    (core/contracts/zayavka/generator.py), плейсхолдеров в ней и не должно
    быть.
    """
    empty = [
        item["name"] for item in ta.list_templates()
        if item["name"].casefold().startswith("shablon_")
        and item["suffix"] == ".docx"
        and ta.analyze_template(item["path"])["placeholder_count"] == 0
    ]
    assert empty == [], f"в шаблонах нет плейсхолдеров: {empty}"
