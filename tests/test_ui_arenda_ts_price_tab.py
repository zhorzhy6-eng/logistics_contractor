#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты вкладки «Стоимость» окна «Разовая аренда».

Вкладка считает НДС по ЕДИНОМУ правилу ядра — «НДС В ТОМ ЧИСЛЕ»
(core/vat.py::compute_vat), тому же, что в Логистиксе, Формике и договоре
перевозки:

    sum_total  = введённое число (ИТОГ договора — то, что видит заказчик);
    sum_wo_vat = sum_total / (1 + ставка/100);
    sum_vat    = sum_total − sum_wo_vat.

Проверяется:

  * список ставок — шесть пунктов ядра, по умолчанию «22%»;
  * вкладка считает тем же ядром (база + НДС = итог, итог не сдвигается);
  * единые ключи get_data() (vat_rate / price_with_vat / price_without_vat /
    vat_amount) и исторические имена сумм аренды;
  * загрузка данных: итог из price_with_vat / sum_total, база из
    price_without_vat / sum_wo_vat (старые записи — через прежнюю формулу);
  * срок оплаты — поле «Срок оплаты, банковских дней», по умолчанию 30,
    ключ payment_days доходит до contract;
  * стык с генератором: суммы в готовом договоре.

Qt — в offscreen-режиме. Данные синтетические, реальных ПДн нет.
"""

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from PyQt5.QtWidgets import QApplication, QComboBox, QDoubleSpinBox, QSpinBox  # noqa: E402

from core.contract_data import ContractData  # noqa: E402
from core.vat import VAT_RATES, compute_vat  # noqa: E402
from ui.windows.arenda_ts.data import collect_arenda_ts_data  # noqa: E402
from ui.windows.arenda_ts.tabs import price_tab as price_tab_module  # noqa: E402
from ui.windows.arenda_ts.tabs.price_tab import (  # noqa: E402
    DEFAULT_VAT_RATE,
    PriceTab,
    base_from_total,
    rate_to_factor,
    total_from_base,
)

# ─────────────────────────────────────────────────────────────
# Константы тестовых данных
# ─────────────────────────────────────────────────────────────

#: Итог 230 000,00 при ставке 22%: 188 524,59 + 41 475,41.
TOTAL_230K = 230000.00
BASE_230K = 188524.59
VAT_230K = 41475.41

#: Итог из задания на сведение НДС: 250 000,00 при 22% и 5%.
TOTAL_250K = 250000.00
BASE_250K_22 = 204918.03
VAT_250K_22 = 45081.97
BASE_250K_5 = 238095.24
VAT_250K_5 = 11904.76

#: Суммы «из документа» для проверки загрузки.
BASE_SUM = 221099.18
VAT_SUM = 48641.82
TOTAL_SUM = 269741.00

#: Ставки.
RATE_22 = 22.0
ZERO_RATE = 0.0

#: Срок оплаты из задания и значение по умолчанию.
PAYMENT_DAYS_CUSTOM = 45
PAYMENT_DAYS_DEFAULT = 30


# ─────────────────────────────────────────────────────────────
# Фикстуры
# ─────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def qt_app():
    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture
def tab(qt_app):
    """Свежая вкладка «Стоимость» (ставка по умолчанию — 22%)."""
    widget = PriceTab()
    yield widget
    widget.deleteLater()


# ─────────────────────────────────────────────────────────────
# Вспомогательное
# ─────────────────────────────────────────────────────────────

def _flatten(text: str) -> str:
    """Текст одной строкой: любые пробелы (в том числе неразрывные) — по одному."""
    import re

    return re.sub(r"\s+", " ", text.replace("\u00a0", " ")).strip()


def _document_text(doc) -> str:
    """Весь текст документа: абзацы и таблицы, пробелы нормализованы."""
    parts = [_flatten(p.text) for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(_flatten(cell.text))
    return "\n".join(parts)


# ─────────────────────────────────────────────────────────────
# Математика НДС — правила ядра, которыми считает вкладка
# ─────────────────────────────────────────────────────────────

def test_rate_to_factor():
    """22% → 1.22, 0% → 1.0 (множитель ядра)."""
    assert rate_to_factor(22.0) == pytest.approx(1.22)
    assert rate_to_factor(ZERO_RATE) == pytest.approx(1.0)
    assert rate_to_factor("Без НДС") == pytest.approx(1.0)


def test_total_from_base_adds_vat():
    """Итог = база × (1 + ставка/100), до копеек (обратный ход, старые записи)."""
    assert total_from_base(BASE_230K, RATE_22) == pytest.approx(TOTAL_230K)
    assert total_from_base(1000.0, ZERO_RATE) == pytest.approx(1000.0)


def test_base_from_total_extracts_vat():
    """База = итог / (1 + ставка/100), до копеек."""
    assert base_from_total(TOTAL_230K, RATE_22) == pytest.approx(BASE_230K)
    assert base_from_total(1000.0, ZERO_RATE) == pytest.approx(1000.0)


def test_task_example_250k_with_22_percent():
    """Пример задания: 250 000 + 22% → база 204 918,03, НДС 45 081,97."""
    vat = compute_vat(TOTAL_250K, "22%")

    assert vat["sum_wo_nds"] == BASE_250K_22
    assert vat["sum_nds"] == VAT_250K_22
    assert vat["sum_total"] == TOTAL_250K


def test_task_example_250k_with_5_percent():
    """Пример задания: 250 000 + 5% → база 238 095,24, НДС 11 904,76."""
    vat = compute_vat(TOTAL_250K, "5%")

    assert vat["sum_wo_nds"] == BASE_250K_5
    assert vat["sum_nds"] == VAT_250K_5
    assert vat["sum_total"] == TOTAL_250K


@pytest.mark.parametrize("rate", ["Без НДС", "0%", "5%", "7%", "10%", "22%"])
def test_base_and_total_round_trip(rate):
    """Пересчёт «туда и обратно» не сдвигает сумму больше чем на копейку."""
    for total in (TOTAL_230K, TOTAL_SUM, 999.99, 100.0, 1.0):
        base = base_from_total(total, rate)
        assert abs(total_from_base(base, rate) - total) <= 0.01, (rate, total)


@pytest.mark.parametrize("rate", ["Без НДС", "0%", "5%", "7%", "10%", "22%"])
def test_vat_and_base_sum_to_total(rate):
    """НДС и база в сумме дают ровно итог: ни одной потерянной копейки."""
    vat = compute_vat(TOTAL_250K, rate)

    assert round(vat["sum_wo_nds"] + vat["sum_nds"], 2) == vat["sum_total"]


# ─────────────────────────────────────────────────────────────
# Вкладка: состав полей и ставки
# ─────────────────────────────────────────────────────────────

def test_vat_rates_are_the_core_six():
    """Список ставок — шесть пунктов ядра, локального списка у вкладки нет."""
    assert price_tab_module.VAT_RATES is VAT_RATES
    assert [tab_rate for tab_rate in VAT_RATES] == [
        "Без НДС", "0%", "5%", "7%", "10%", "22%",
    ]


def test_tab_rate_list_and_default(tab):
    """Выпадающий список ставок — те же шесть пунктов; по умолчанию «22%»."""
    assert isinstance(tab.vat_rate, QComboBox)
    assert [tab.vat_rate.itemText(i) for i in range(tab.vat_rate.count())] \
        == list(VAT_RATES)
    assert tab.vat_rate.currentText() == "22%"
    assert DEFAULT_VAT_RATE == "22%"


def test_tab_has_single_editable_amount_field(tab):
    """Одно редактируемое поле суммы; НДС и база — расчётные."""
    assert isinstance(tab.sum_total, QDoubleSpinBox)
    assert tab.sum_total.isReadOnly() is False
    assert tab.sum_vat.isReadOnly() is True
    assert tab.sum_wo_vat.isReadOnly() is True


def test_tab_has_no_amount_mode_switch(tab):
    """Переключателя «Считать от» больше нет: правило НДС одно."""
    assert not hasattr(tab, "amount_mode")
    assert tab.amount_mode_text() == "С НДС"
    assert tab.calculates_from_total() is True


def test_tab_has_payment_days_field(tab):
    """Поле срока оплаты — целое число банковских дней."""
    assert isinstance(tab.payment_days, QSpinBox)
    assert not isinstance(tab.payment_days, QDoubleSpinBox)
    assert tab.payment_days.value() == PAYMENT_DAYS_DEFAULT
    assert tab.payment_days.minimum() == price_tab_module.MIN_PAYMENT_DAYS
    assert tab.payment_days.maximum() == price_tab_module.MAX_PAYMENT_DAYS


# ─────────────────────────────────────────────────────────────
# Расчёт: «НДС в том числе»
# ─────────────────────────────────────────────────────────────

def test_entered_amount_is_the_total(tab):
    """Введённое число — ИТОГ; база и НДС выводятся из него."""
    tab.sum_total.setValue(TOTAL_250K)
    data = tab.get_data()

    assert data["price_with_vat"] == pytest.approx(TOTAL_250K)
    assert data["price_without_vat"] == pytest.approx(BASE_250K_22)
    assert data["vat_amount"] == pytest.approx(VAT_250K_22)


def test_amount_is_never_inflated_by_the_rate(tab):
    """Итог в договоре равен введённому числу: налог не прибавляется сверху."""
    tab.sum_total.setValue(1000.0)
    data = tab.get_data()

    assert data["price_with_vat"] == pytest.approx(1000.0)
    assert data["price_without_vat"] == pytest.approx(819.67)
    assert data["vat_amount"] == pytest.approx(180.33)


@pytest.mark.parametrize("rate,total,base,vat", [
    ("22%", TOTAL_250K, BASE_250K_22, VAT_250K_22),
    ("5%", TOTAL_250K, BASE_250K_5, VAT_250K_5),
    ("10%", TOTAL_250K, 227272.73, 22727.27),
    ("7%", TOTAL_250K, 233644.86, 16355.14),
    ("0%", TOTAL_250K, TOTAL_250K, 0.0),
    ("Без НДС", TOTAL_250K, TOTAL_250K, 0.0),
])
def test_rates_table(tab, rate, total, base, vat):
    """Таблица ставок на итоге 250 000,00: база, налог и итог."""
    tab.vat_rate.setCurrentText(rate)
    tab.sum_total.setValue(total)
    data = tab.get_data()

    assert data["price_without_vat"] == pytest.approx(base)
    assert data["vat_amount"] == pytest.approx(vat)
    assert data["price_with_vat"] == pytest.approx(total)


@pytest.mark.parametrize("rate", list(VAT_RATES))
def test_base_and_vat_always_sum_to_total(tab, rate):
    """При любой ставке база + НДС = итог, а итог не уходит от введённого."""
    tab.vat_rate.setCurrentText(rate)
    tab.sum_total.setValue(TOTAL_250K)
    data = tab.get_data()

    assert round(data["price_without_vat"] + data["vat_amount"], 2) \
        == pytest.approx(data["price_with_vat"])


def test_words_follow_the_total(tab):
    """Сумма прописью — от итога: он и печатается в договоре."""
    tab.sum_total.setValue(TOTAL_230K)

    assert tab.sum_total_words.text().startswith("Двести тридцать тысяч")


def test_recalculation_on_rate_change(tab):
    """Смена ставки пересчитывает базу и налог, не трогая введённый итог."""
    tab.sum_total.setValue(TOTAL_250K)
    tab.vat_rate.setCurrentText("5%")

    assert tab.input_amount() == pytest.approx(TOTAL_250K)
    assert tab.sum_wo_vat.value() == pytest.approx(BASE_250K_5)
    assert tab.sum_vat.value() == pytest.approx(VAT_250K_5)


# ─────────────────────────────────────────────────────────────
# Единые ключи и загрузка данных
# ─────────────────────────────────────────────────────────────

def test_get_data_has_unified_keys(tab):
    """Единые ключи типа отдаются наравне с историческими именами аренды."""
    tab.sum_total.setValue(TOTAL_250K)
    data = tab.get_data()

    assert {"vat_rate", "price_with_vat", "price_without_vat", "vat_amount"} \
        <= set(data)
    assert data["sum_wo_vat"] == data["price_without_vat"]
    assert data["sum_vat"] == data["vat_amount"]
    assert data["sum_total"] == data["price_with_vat"]
    assert data["vat_rate"] == "22%"
    assert data["vat_rate_num"] == pytest.approx(22.0)


@pytest.mark.parametrize("key", ["price_with_vat", "sum_total"])
def test_fill_data_reads_the_total(tab, key):
    """Итог в поле — из price_with_vat или исторического sum_total."""
    tab.fill_data({key: TOTAL_250K, "vat_rate": "22%"})

    assert tab.input_amount() == pytest.approx(TOTAL_250K)
    assert tab.get_data()["price_without_vat"] == pytest.approx(BASE_250K_22)


def test_fill_data_reads_recognized_triple(tab):
    """Распознанный блок: три согласованные суммы дают прежние числа."""
    tab.fill_data({
        "sum_wo_vat": BASE_SUM, "sum_vat": VAT_SUM, "sum_total": TOTAL_SUM,
        "vat_rate": "22%",
    })

    assert tab.input_amount() == pytest.approx(TOTAL_SUM)
    assert tab.get_data()["price_without_vat"] == pytest.approx(BASE_SUM)
    assert tab.get_data()["vat_amount"] == pytest.approx(VAT_SUM)


def test_fill_data_restores_total_from_base(tab):
    """Старая запись хранит только базу: итог восстанавливается прежней формулой."""
    tab.fill_data({"price_without_vat": BASE_230K, "vat_rate": "22%"})

    assert tab.input_amount() == pytest.approx(TOTAL_230K)
    assert tab.get_data()["price_with_vat"] == pytest.approx(TOTAL_230K)


def test_fill_data_reads_vat_free_document(tab):
    """Документ «НДС не облагается»: единственная сумма — она же итог."""
    tab.fill_data({
        "sum_wo_vat": 0.0, "sum_vat": 0.0, "sum_total": TOTAL_SUM,
        "vat_rate": "0%",
    })

    data = tab.get_data()
    assert data["price_with_vat"] == pytest.approx(TOTAL_SUM)
    assert data["price_without_vat"] == pytest.approx(TOTAL_SUM)
    assert data["vat_amount"] == 0.0


def test_fill_data_reads_rate_number_and_free(tab):
    """Ставка приходит и числом (22), и словом («Без НДС»)."""
    tab.fill_data({"price_with_vat": TOTAL_250K, "vat_rate_num": 22})
    assert tab.get_data()["vat_rate"] == "22%"

    tab.fill_data({"price_with_vat": TOTAL_250K, "vat_rate": "Без НДС"})
    assert tab.get_data()["vat_rate"] == "Без НДС"
    assert tab.get_data()["vat_amount"] == 0.0


def test_fill_data_ignores_empty_and_zero(tab):
    """Пустое и нулевое значение не сбрасывают введённую сумму."""
    tab.fill_data({"price_with_vat": TOTAL_250K, "vat_rate": "22%"})

    tab.fill_data({"price_with_vat": 0.0, "sum_total": None})
    tab.fill_data({})

    assert tab.input_amount() == pytest.approx(TOTAL_250K)


# ─────────────────────────────────────────────────────────────
# Срок оплаты
# ─────────────────────────────────────────────────────────────

def test_payment_days_default_is_thirty(tab):
    """Значение по умолчанию — 30 банковских дней (как в бланке)."""
    assert tab.get_data()["payment_days"] == PAYMENT_DAYS_DEFAULT
    assert PAYMENT_DAYS_DEFAULT == price_tab_module.DEFAULT_PAYMENT_DAYS


def test_payment_days_accepts_custom_value(tab):
    tab.payment_days.setValue(PAYMENT_DAYS_CUSTOM)

    assert tab.get_data()["payment_days"] == PAYMENT_DAYS_CUSTOM


def test_payment_days_is_integer(tab):
    tab.payment_days.setValue(PAYMENT_DAYS_CUSTOM)

    value = tab.get_data()["payment_days"]
    assert isinstance(value, int)
    assert not isinstance(value, bool)


def test_payment_days_zero_means_not_set(tab):
    """Ноль — «срок не задан»: поле это принимает, о нём скажет валидатор."""
    tab.payment_days.setValue(0)

    assert tab.get_data()["payment_days"] == 0


def test_payment_days_rejects_negative(tab):
    """Отрицательного срока не бывает: поле не уходит в минус."""
    tab.payment_days.setValue(-5)

    assert tab.get_data()["payment_days"] == 0


def test_payment_days_has_upper_bound(tab):
    """Верхняя граница не даёт ввести «срок на десятилетие»."""
    tab.payment_days.setValue(9999)

    assert tab.get_data()["payment_days"] == price_tab_module.MAX_PAYMENT_DAYS


def test_payment_days_fill_and_read_back(tab):
    tab.fill_data({"price_with_vat": TOTAL_SUM, "payment_days": PAYMENT_DAYS_CUSTOM})

    assert tab.get_data()["payment_days"] == PAYMENT_DAYS_CUSTOM


def test_payment_days_ignores_empty_and_garbage(tab):
    """Пустое и нечисловое значение не сбрасывают введённый срок."""
    tab.payment_days.setValue(PAYMENT_DAYS_CUSTOM)

    tab.fill_data({"payment_days": ""})
    tab.fill_data({"payment_days": None})
    tab.fill_data({"payment_days": "мусор"})

    assert tab.get_data()["payment_days"] == PAYMENT_DAYS_CUSTOM


def test_payment_days_does_not_depend_on_amount(tab):
    """Срок оплаты не зависит ни от суммы, ни от ставки."""
    tab.payment_days.setValue(PAYMENT_DAYS_CUSTOM)
    tab.sum_total.setValue(TOTAL_230K)
    tab.vat_rate.setCurrentText("5%")

    assert tab.get_data()["payment_days"] == PAYMENT_DAYS_CUSTOM


def test_clear_returns_to_defaults(tab):
    tab.payment_days.setValue(PAYMENT_DAYS_CUSTOM)
    tab.sum_total.setValue(TOTAL_230K)

    tab.clear()

    assert tab.get_data()["payment_days"] == PAYMENT_DAYS_DEFAULT
    assert tab.input_amount() == 0.0
    assert tab.vat_rate.currentText() == DEFAULT_VAT_RATE


# ─────────────────────────────────────────────────────────────
# Вкладка → данные договора
# ─────────────────────────────────────────────────────────────

def _contract_of(tab: PriceTab) -> dict:
    """ContractData из одной вкладки «Стоимость» (остальные разделы пусты)."""
    data = collect_arenda_ts_data({"price": tab})
    assert isinstance(data, ContractData)
    return data.contract


def test_payment_days_reaches_contract(qt_app):
    """Ключ payment_days собирается в contract."""
    tab = PriceTab()
    tab.sum_total.setValue(TOTAL_230K)
    tab.payment_days.setValue(PAYMENT_DAYS_CUSTOM)

    contract = _contract_of(tab)

    assert contract["payment_days"] == PAYMENT_DAYS_CUSTOM


@pytest.mark.parametrize("rate,total,base,vat", [
    ("22%", TOTAL_250K, BASE_250K_22, VAT_250K_22),
    ("5%", TOTAL_250K, BASE_250K_5, VAT_250K_5),
    ("0%", TOTAL_250K, TOTAL_250K, 0.0),
])
def test_contract_sums_come_from_the_core_rule(qt_app, rate, total, base, vat):
    """В contract уходят base / НДС / итог, посчитанные ядром."""
    tab = PriceTab()
    tab.vat_rate.setCurrentText(rate)
    tab.sum_total.setValue(total)

    contract = _contract_of(tab)

    assert contract["price_without_vat"] == pytest.approx(base)
    assert contract["price_with_vat"] == pytest.approx(total)
    assert contract["vat_amount"] == pytest.approx(vat)
    assert contract["sum_wo_vat"] == pytest.approx(base)
    assert contract["sum_vat"] == pytest.approx(vat)
    assert contract["sum_total"] == pytest.approx(total)


def test_contract_totals_are_consistent(qt_app):
    """База + НДС = итог в собранном contract: копейка не теряется."""
    tab = PriceTab()
    tab.vat_rate.setCurrentText("22%")
    tab.sum_total.setValue(TOTAL_250K)

    contract = _contract_of(tab)

    assert round(contract["price_without_vat"] + contract["vat_amount"], 2) \
        == contract["price_with_vat"]


def test_zero_payment_days_is_not_written_to_contract(qt_app):
    """Ноль в поле — «срок не задан»: в contract ключ не попадает."""
    tab = PriceTab()
    tab.sum_total.setValue(TOTAL_230K)
    tab.payment_days.setValue(0)

    contract = _contract_of(tab)

    assert "payment_days" not in contract


def test_unparsable_payment_days_is_not_written():
    """Мусор вместо срока в contract не попадает: бланк печатает число."""
    contract = collect_arenda_ts_data({
        "price": {"price_with_vat": TOTAL_SUM, "payment_days": "тридцать"},
    }).contract

    assert "payment_days" not in contract


def test_ip_without_vat_contract_has_single_sum(qt_app):
    """ИП без НДС: в contract одна сумма, ставка нулевая, ключей НДС нет."""
    tab = PriceTab()
    tab.sum_total.setValue(TOTAL_250K)

    contract = collect_arenda_ts_data({
        "lessee": {"carrier_type": "ИП без НДС", "full_name": "ИП Тестов"},
        "price": tab,
    }).contract

    assert contract["price_without_vat"] == pytest.approx(TOTAL_250K)
    assert contract["price_with_vat"] == pytest.approx(TOTAL_250K)
    assert contract["vat_rate_num"] == 0.0
    assert "sum_wo_vat" not in contract
    assert "sum_vat" not in contract


# ─────────────────────────────────────────────────────────────
# Стык с генератором: суммы и срок оплаты в готовом договоре
# ─────────────────────────────────────────────────────────────

def _filled_tabs(qt_app, amount: float = TOTAL_230K,
                 rate: str = "22%") -> dict:
    """
    Семь вкладок окна аренды с минимально достаточными данными.

    Вкладка «Стоимость» заполняется через интерфейс (проверяется именно она),
    остальные разделы — словарями: у каждой вкладки свой набор тестов.

    :param amount: число, которое пользователь ввёл в поле суммы, — ИТОГ.
    """
    from PyQt5.QtCore import QDate

    price = PriceTab()
    price.vat_rate.setCurrentText(rate)
    price.sum_total.setValue(amount)
    price.payment_days.setValue(PAYMENT_DAYS_CUSTOM)

    return {
        "lessee": {
            "carrier_type": "ООО",
            "full_name": "ООО «Арендатор-Тест»",
            "short_name": "ООО «АТ»",
            "inn": "7701234567",
            "kpp": "770101001",
            "ogrn": "1027700132195",
            "address": "г. Москва, ул. Арендаторская, д. 1",
            "account": "40702810000000000001",
            "bank": "ПАО Сбербанк",
            "bik": "044525225",
            "corr_account": "30101810400000000225",
            "director_position": "Генеральный директор",
            "director_name": "Петров Пётр Петрович",
        },
        "lessor": {
            "full_name": "ООО «Арендодатель-Тест»",
            "short_name": "ООО «АД»",
            "inn": "7707654321",
            "ogrn": "1027700261234",
            "address": "г. Москва, ул. Арендодательская, д. 2",
            "account": "40702810000000000002",
            "bank": "АО «Банк Второй»",
            "bik": "044525226",
            "corr_account": "30101810400000000226",
            "director_position": "Директор",
            "director_name": "Сидоров Сидор Сидорович",
        },
        "vehicle": {
            "contract_number": "01/2026",
            "contract_date": QDate(2026, 9, 23).toString("yyyy-MM-dd"),
            "lease_start_date": QDate(2026, 9, 28).toString("yyyy-MM-dd"),
            "lease_end_date": QDate(2026, 10, 5).toString("yyyy-MM-dd"),
            "tractor_brand": "DAF XF",
            "tractor_plate": "М342СА761",
            "tractor_type": "Седельный тягач",
            "trailer_brand": "KRONE SD",
            "trailer_plate": "ВК123478",
        },
        "route": {
            "route": "г. Москва — г. Казань",
            "loadings": [{"name": "Склад", "address": "г. Москва, ул. Складская, д. 1",
                          "date": "2026-09-28", "time_from": "08:00",
                          "time_to": "18:00"}],
            "unloadings": [{"name": "Приёмка", "address": "г. Казань, ул. Приёмная, д. 3",
                            "date": "2026-10-02"}],
        },
        "cargo": {
            "vehicles": [{"brand_model": "JETOUR T2", "vin": "EC3TEUMB0T0002608",
                          "loading_point": "Москва", "unloading_point": "Казань"}],
        },
        "crew": {
            "driver_full_name": "Иванов Иван Иванович",
            "driver_birth_date": "1980-01-01",
            "driver_passport": "18 22 926830",
            "driver_passport_issuer": "Отделом УФМС России по г. Москве",
            "driver_passport_issue_date": "2023-01-30",
            "driver_license": "99 36 123456",
            "driver_license_issue_date": "2020-01-01",
            "driver_registration_address": "г. Москва, ул. Тестовая, д. 1",
            "driver_phone": "+7 (999) 123-45-67",
        },
        "price": price,
    }


def _generated_document(qt_app, templates_dir, work_dir,
                        folder: str = "", amount: float = TOTAL_230K,
                        rate: str = "22%"):
    """
    Готовый договор из настоящей вкладки «Стоимость» и вкладок-словарей.

    Имя выходного файла зависит только от номера и даты договора, поэтому
    каждому варианту теста даётся своя папка — иначе второй договор затрёт
    первый, и оба теста прочитали бы один и тот же документ.
    """
    from core.contracts.arenda_ts.generator import ArendaTsGenerator

    if folder:
        work_dir = work_dir / folder
        work_dir.mkdir(parents=True, exist_ok=True)

    tabs = _filled_tabs(qt_app, amount, rate)
    contract_data = collect_arenda_ts_data(tabs)

    generator = ArendaTsGenerator(templates_dir=str(templates_dir))
    path = generator.generate(contract_data, output_dir=str(work_dir))
    return Document(path)


def test_vat_calculation_reaches_the_document(qt_app, templates_dir, work_dir):
    """
    Введённый ИТОГ доходит до бланка тремя согласованными числами.

    Ввод «230 000 с НДС» → в п. 4.1 бланка «188 524,59 без НДС», «НДС 22% —
    41 475,41» и «Итого с НДС: 230 000,00», каждое число ещё и прописью.
    """
    doc = _generated_document(qt_app, templates_dir, work_dir, folder="from_vat")
    text = _document_text(doc)

    assert ("– 188 524,59 руб. (Сто восемьдесят восемь тысяч пятьсот двадцать "
            "четыре рубля пятьдесят девять копеек) — стоимость без НДС;") in text
    assert ("– НДС 22% — 41 475,41 руб. (Сорок одна тысяча четыреста "
            "семьдесят пять рублей сорок одна копейка);") in text
    assert "Итого с НДС: 230 000,00 руб." in text
    assert "{{" not in text


def test_same_input_gives_the_same_document(qt_app, templates_dir, work_dir):
    """Один и тот же итог и ставка дают один и тот же договор."""
    first = _document_text(_generated_document(
        qt_app, templates_dir, work_dir, folder="same_1"))
    second = _document_text(_generated_document(
        qt_app, templates_dir, work_dir, folder="same_2"))

    assert "Итого с НДС: 230 000,00 руб." in first
    assert first == second


def test_five_percent_document(qt_app, templates_dir, work_dir):
    """Итог 250 000,00 при 5%: база 238 095,24 и налог 11 904,76 в бланке."""
    doc = _generated_document(qt_app, templates_dir, work_dir,
                              folder="rate_5", amount=TOTAL_250K, rate="5%")
    text = _document_text(doc)

    assert "238 095,24 руб." in text
    assert "11 904,76 руб." in text
    assert "НДС 5%" in text
    assert "Итого с НДС: 250 000,00 руб." in text


def test_generator_replacements_carry_the_payment_days(qt_app, templates_dir):
    """Ключ payment_days доходит до сборщика данных договора."""
    from core.contracts.arenda_ts.generator import ArendaTsGenerator

    tabs = _filled_tabs(qt_app)
    contract_data = collect_arenda_ts_data(tabs)

    assert contract_data.contract["payment_days"] == PAYMENT_DAYS_CUSTOM

    generator = ArendaTsGenerator(templates_dir=str(templates_dir))
    replacements = generator.build_replacements(contract_data)

    # Плейсхолдера {payment_days} в бланке аренды пока нет (см. STATE.md,
    # FIX-1) — проверяем, что значение собрано и готово к подстановке.
    assert contract_data.contract["payment_days"] == 45
    assert replacements["sum_total"] == "230\u00a0000,00"
    assert replacements["sum_wo_vat"] == "188\u00a0524,59"
    assert replacements["sum_vat"] == "41\u00a0475,41"
