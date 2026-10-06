#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты окна типа «Логистикс Рус» на реальных вкладках (ЭТАП 3.1.C.B.3).

Проверяют то, что появилось на этом шаге: окно собирается на шести
настоящих вкладках (а не на заглушках «в разработке»), сигналы вкладок
подключены к слотам окна, данные всех вкладок уходят в один ContractData,
«Очистить форму» чистит только свою вкладку, «Создать договор» проходит
путь валидатор → диалог → генератор → диалог успеха, а распознавание
раскладывает ответ модели по вкладкам, не стирая ручной ввод.

Отдельная группа тестов — сама раскладка блоков ответа по вкладкам
(_customer_tab_data, _route_tab_data, _vehicle_tab_data, _price_tab_data):
промпт Логистикс Рус отдаёт блоки customer / shippers / consignees /
vehicles / tractor / trailer / driver / contract, а вкладки ждут СВОИ
имена полей. Проверять это через окно неудобно, поэтому карта ключей
проверяется напрямую — без поднятия интерфейса.

Qt — в offscreen-режиме. Сеть и системное хранилище ключей не трогаются:
клиент GigaChat подменяется заглушкой, а промпт типа проверяется по
аргументам, с которыми окно позвало recognize_text.

Все данные синтетические, реальных ПДн нет.
"""

import gc
import inspect
import os
import re
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from PyQt5.QtWidgets import QApplication, QMessageBox  # noqa: E402

import ui.windows.logistiks_rus.window as logistiks_window_module  # noqa: E402
from core.contract_data import ContractData  # noqa: E402
from core.contracts.factory import GeneratorFactory  # noqa: E402
from core.contracts.logistiks_rus.generator import LogistiksRusGenerator  # noqa: E402
from core.contracts.logistiks_rus.validator import LogistiksRusValidator  # noqa: E402
from core.contracts.registry import ContractTypeRegistry  # noqa: E402
from core.prompts import get_prompt  # noqa: E402
from ui.windows.logistiks_rus import LogistiksRusWindow  # noqa: E402
from ui.windows.logistiks_rus.tabs import (  # noqa: E402
    CargoTab, CustomerTab, DriverTab, PriceTab, RouteTab, VehicleTab,
)
from ui.windows.logistiks_rus.window import RecognitionTask  # noqa: E402

#: Сколько машин помещается в таблицу бланка (см. CargoTab.MAX_CARS).
MAX_CARS = 12

#: Ожидаемый состав вкладок: (заголовок, класс) — порядок как в окне.
TAB_SPECS = (
    ("Заказчик", CustomerTab),
    ("Груз", CargoTab),
    ("Маршрут", RouteTab),
    ("Водитель", DriverTab),
    ("ТС", VehicleTab),
    ("Стоимость", PriceTab),
)

#: Заказчик этой заявки: подставлен вкладкой по умолчанию.
CUSTOMER_NAME = "ООО «ДжейСиСиТиЭс Интернейшнл Логистикс Рус»"

#: Номер заявки, которым заполняются вкладки в тестах.
CONTRACT_NUMBER = "ЛР-2026-1"

#: Сумма без НДС из формы (её читает сборщик данных).
FORM_AMOUNT = 221099.18


# ─────────────────────────────────────────────────────────────
# Фикстуры и помощники
# ─────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def qt_app():
    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture(autouse=True)
def _release_documents():
    """Word-файлы на Windows освобождаем до удаления временных копий."""
    yield
    gc.collect()


@pytest.fixture
def quiet_messages(monkeypatch):
    """
    Глушит диалоги QMessageBox и запоминает их тексты.

    :return: словарь со списками текстов information / warning / critical.
    """
    seen = {"information": [], "warning": [], "critical": []}

    def recorder(kind):
        def _record(parent, title, text, *args, **kwargs):
            seen[kind].append(text)
            return QMessageBox.Ok
        return staticmethod(_record)

    for kind in seen:
        monkeypatch.setattr(QMessageBox, kind, recorder(kind))
    return seen


@pytest.fixture
def window(qt_app, quiet_messages):
    """
    Свежее окно Логистикс Рус; по завершении теста закрывается по-настоящему.

    Перед закрытием ждём пул распознавания: поток, переживший окно, роняет
    процесс при разрушении QThreadPool (access violation на выходе).
    """
    win = LogistiksRusWindow()
    yield win
    win.thread_pool.waitForDone(5000)
    qt_app.processEvents()
    win.force_close()


class FakeClient:
    """Заглушка клиента GigaChat: в сеть не ходит, помнит аргументы вызова."""

    def __init__(self, answer=None):
        self.answer = answer if answer is not None else {}
        self.calls = []

    def recognize_text(self, text, prompt=None):
        self.calls.append({"text": text, "prompt": prompt})
        return self.answer


def _fill_all_tabs(win, cars=1):
    """Заполняет все шесть вкладок данными, которых хватает валидатору."""
    win.customer_tab.fill_data({
        "number": CONTRACT_NUMBER,
        "date": "2026-09-24",
        "name": CUSTOMER_NAME,
    })
    win.cargo_tab.fill_data({"vehicles": [
        {"brand_model": f"МОДЕЛЬ {number}",
         "vin": f"XTC651150N0001{number:03d}"}
        for number in range(1, cars + 1)
    ]})
    win.route_tab.fill_data({
        "route": "Москва - Казань",
        "shippers": [{"name": "ООО «Склад 1»",
                      "address": "г. Москва, ул. Складская, д. 1"}],
        "consignees": [{"name": "ООО «Клиент 1»",
                        "address": "г. Казань, ул. Заводская, д. 2"}],
        "loading_date": "2026-09-26",
        "loading_time_from": "08:00",
        "loading_time_to": "20:00",
        "unloading_date": "2026-10-01",
        "unloading_time_from": "08:00",
        "unloading_time_to": "20:00",
    })
    win.driver_tab.fill_data({"full_name": "Иванов Иван Иванович"})
    win.vehicle_tab.fill_data({
        "tractor_brand": "DAF XF 95.430",
        "tractor_plate": "М342СА761",
        "trailer_brand": "KRONE SD",
        "trailer_plate": "ВК123478",
    })
    win.price_tab.fill_data({
        "amount_without_vat": FORM_AMOUNT,
        "vat_rate": "22%",
        "special_conditions": "Погрузка круглосуточно, простой не более 24 часов.",
    })


def _assert_other_tabs_intact(win, cleared_index: int) -> None:
    """Соседние вкладки после «Очистить форму» остались заполненными."""
    if cleared_index != 0:
        assert win.customer_tab.get_data()["number"] == CONTRACT_NUMBER
    if cleared_index != 1:
        assert len(win.cargo_tab.get_data()["vehicles"]) == 1
    if cleared_index != 2:
        assert win.route_tab.get_data()["route"] == "Москва - Казань"
    if cleared_index != 3:
        assert win.driver_tab.get_data()["full_name"] == "Иванов Иван Иванович"
    if cleared_index != 4:
        assert win.vehicle_tab.get_data()["tractor_brand"] == "DAF XF 95.430"
    if cleared_index != 5:
        assert win.price_tab.get_data()["amount_without_vat"] == FORM_AMOUNT


def _answer_recognition(window, data):
    """
    Отдаёт окну ответ модели так же, как это сделала бы задача распознавания.

    Задача передаётся аргументом слота: именно так её связывает
    _start_recognition (functools.partial), и так слот понимает, чей
    результат пришёл.
    """
    task = RecognitionTask(window.gigachat or FakeClient(), "текст", prompt="промпт")
    window.recognition_task = task

    window._on_recognition_finished(data, task=task)
    return task


def _document_text(doc) -> str:
    """Весь текст документа: абзацы и таблицы (включая вложенные)."""
    parts = [p.text for p in doc.paragraphs]

    def table_text(table):
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
                for nested in cell.tables:
                    table_text(nested)

    for table in doc.tables:
        table_text(table)

    return "\n".join(parts)


def _valid_contract_data() -> ContractData:
    """
    Полный ContractData заявки — без вкладок.

    Нужен там, где проверяется сам путь «проверка → генерация»: данные
    подставляются напрямую, и видно именно шаг проверки.
    """
    return ContractData(
        driver={"full_name": "Иванов Иван Иванович"},
        customer={"full_name": CUSTOMER_NAME, "short_name": CUSTOMER_NAME},
        vehicles=[{"brand_model": "МОДЕЛЬ 1", "vin": "XTC651150N0001001"}],
        tractor={"brand_model": "DAF XF 95.430", "plate_number": "М342СА761"},
        trailer={"brand_model": "KRONE SD", "plate_number": "ВК123478"},
        contract={
            "number": CONTRACT_NUMBER,
            "date": "2026-09-24",
            "carrier_type": "ООО",
            "price_without_vat": FORM_AMOUNT,
            "vat_rate": "22%",
            "vat_rate_num": 22,
            "special_conditions": "Погрузка круглосуточно.",
            "loading_date": "2026-09-26",
            "loading_time_from": "08:00",
            "loading_time_to": "20:00",
            "unloading_date": "2026-10-01",
            "unloading_time_from": "08:00",
            "unloading_time_to": "20:00",
            "loadings": [{"name": "ООО «Склад 1»",
                          "address": "г. Москва, ул. Складская, д. 1",
                          "date": "2026-09-26", "time_window": "08:00-20:00"}],
            "unloadings": [{"name": "ООО «Клиент 1»",
                            "address": "г. Казань, ул. Заводская, д. 2",
                            "date": "2026-10-01", "time_window": "08:00-20:00"}],
        },
    )


# ─────────────────────────────────────────────────────────────
# 1. Окно и его вкладки
# ─────────────────────────────────────────────────────────────

def test_window_builds_with_six_tabs(window):
    assert window.CONTRACT_TYPE == "logistiks_rus"
    assert window.windowTitle() == "Логистикс Рус"
    assert window.tabs.count() == 6
    assert window.tab_titles() == [title for title, _ in TAB_SPECS]
    assert window.side_nav.count() == 6


def test_make_tab_returns_real_tab_classes(window):
    """Хук _make_tab отдаёт настоящие вкладки, а не заглушки «в разработке»."""
    for index, (title, tab_class) in enumerate(TAB_SPECS):
        tab = window.tabs.widget(index)
        assert isinstance(tab, tab_class), title

    for tab in window._tabs():
        assert hasattr(tab, "recognition_panel"), type(tab).__name__
        assert hasattr(tab, "get_data"), type(tab).__name__


def test_tabs_are_named_attributes(window):
    """Окно держит вкладки и по именам — на них опираются слоты."""
    assert isinstance(window.customer_tab, CustomerTab)
    assert isinstance(window.cargo_tab, CargoTab)
    assert isinstance(window.route_tab, RouteTab)
    assert isinstance(window.driver_tab, DriverTab)
    assert isinstance(window.vehicle_tab, VehicleTab)
    assert isinstance(window.price_tab, PriceTab)


def test_unknown_tab_title_falls_back_to_placeholder(window):
    """Опечатка в TAB_CONFIGS не оставляет окно без страницы."""
    tab = window._make_tab("Такой вкладки нет", "contract.svg")

    assert tab is not None
    assert not hasattr(tab, "get_data")


def test_header_has_create_contract_button(window):
    """«Создать договор» есть и в шапке, и на вкладках (как в MainWindow)."""
    assert window.btn_create_contract.text() == "Создать договор"
    assert window.btn_create_contract.parent() is not None
    for tab in window._tabs():
        assert tab.btn_create_contract.text() == "Создать договор"


def test_gigachat_is_not_created_in_init(window):
    """Клиент GigaChat — ленивый: простое открытие окна ключа не требует."""
    assert window.gigachat is None
    assert window.recognition_task is None


def test_type_is_registered_for_factory(window):
    """
    Окно прогревает реестр типов: фабрика знает «logistiks_rus».

    Без этого GeneratorFactory взяла бы генератор по умолчанию, и заявка
    Логистикс Рус собиралась бы по чужому шаблону.
    """
    spec = ContractTypeRegistry.find("logistiks_rus")

    assert spec is not None
    assert spec.generator_class is LogistiksRusGenerator
    assert spec.validator_class is LogistiksRusValidator


# ─────────────────────────────────────────────────────────────
# 2. Сигналы вкладок подключены к слотам окна
# ─────────────────────────────────────────────────────────────
# Сигнал подключён к связанному методу окна, поэтому доказательство
# связи — вызов этого же пути: клик по кнопке вкладки доходит до
# _collect_data() и _confirm_validation(). Подменять сами слоты после
# создания окна бессмысленно: связь уже держит исходный метод.

def test_tab_create_signal_is_connected_to_window(window, monkeypatch):
    reports = []
    monkeypatch.setattr(
        LogistiksRusWindow, "_confirm_validation",
        lambda self, report: reports.append(report) and False,
    )

    for tab in window._tabs():
        reports.clear()
        tab.btn_create_contract.click()
        assert len(reports) == 1, type(tab).__name__


def test_header_create_button_uses_the_same_slot_as_tabs(window, monkeypatch):
    """
    «Создать договор» в шапке и на вкладке ведут в один слот окна.

    Проверяется по исходнику: обе связи идут на _on_create_contract, а не
    на разные обработчики — иначе шапка и вкладка давали бы разный результат.
    """
    assert "_on_create_contract" in inspect.getsource(
        LogistiksRusWindow._build_header_actions
    )
    assert "_on_create_contract" in inspect.getsource(LogistiksRusWindow._bind_tabs)

    reports = []
    monkeypatch.setattr(
        LogistiksRusWindow, "_confirm_validation",
        lambda self, report: reports.append(report) and False,
    )

    window.btn_create_contract.click()

    assert len(reports) == 1


def test_tab_recognize_signal_is_connected_to_window(window, monkeypatch):
    """Кнопка «Распознать вкладку» доходит до слота окна."""
    clients = []
    monkeypatch.setattr(
        LogistiksRusWindow, "_ensure_gigachat",
        lambda self: clients.append(FakeClient()) or clients[-1],
    )

    window.route_tab.recognition_panel.text_edit.setPlainText("текст заявки")
    window.route_tab.recognition_panel.btn_recognize.click()

    assert len(clients) == 1
    assert window.recognition_task is not None
    assert window.recognition_task.text == "текст заявки"


#: Как проверить, что вкладка действительно очистилась: (индекс, проверка).
CLEAR_SPECS = (
    (0, "Заказчик", lambda tab: tab.get_data()["number"] == ""),
    (1, "Груз", lambda tab: tab.get_data()["vehicles"] == []),
    (2, "Маршрут", lambda tab: tab.get_data()["route"] == ""
        and tab.get_data()["loading_addresses"] == []),
    (3, "Водитель", lambda tab: tab.get_data()["full_name"] == ""),
    (4, "ТС", lambda tab: not any(tab.get_data().values())),
    (5, "Стоимость", lambda tab: tab.get_data()["amount_without_vat"] == 0.0
        and tab.get_data()["special_conditions"] == ""),
)


@pytest.mark.parametrize("index,title,is_cleared", CLEAR_SPECS,
                         ids=[spec[1] for spec in CLEAR_SPECS])
def test_tab_clear_signal_is_connected_to_window(window, index, title, is_cleared):
    """Кнопка «Очистить форму» каждой вкладки доходит до слота окна."""
    _fill_all_tabs(window, cars=1)

    tab = window.tabs.widget(index)
    tab.btn_clear_form.click()

    assert is_cleared(tab), title
    _assert_other_tabs_intact(window, index)


def test_window_connect_does_not_use_lambda():
    """
    В connect — методы и partial, но не lambda: lambda, захватывающая окно
    или вкладку, создаёт цикл ссылок Python ↔ Qt и роняет процесс при
    завершении.
    """
    for method in (
        LogistiksRusWindow.__init__,
        LogistiksRusWindow._build_header_actions,
        LogistiksRusWindow._bind_tabs,
        LogistiksRusWindow._start_recognition,
    ):
        source = inspect.getsource(method)
        assert ".connect(lambda" not in source.replace(" ", ""), source.splitlines()[0]


def test_window_connects_only_bound_methods_and_partial():
    """
    Все connect окна ведут на связанные методы окна или на partial.

    Иначе говоря, среди связей нет замыканий на вкладку: вкладка живёт
    дольше сигнала, и цикл Python ↔ Qt не даёт сборщику мусора освободить
    ни окно, ни вкладки при выходе из программы.
    """
    source = inspect.getsource(logistiks_window_module)
    targets = re.findall(r"\.connect\(([^)]*)", source)

    assert targets, "в модуле окна не нашлось ни одной связи сигналов"
    for target in targets:
        text = target.strip()
        assert text.startswith(("self._", "partial(")), text


def test_recognition_signals_are_bound_to_task_through_partial():
    """
    Сигналы задачи привязываются к самой задаче (partial), а не к lambda.

    Именованный аргумент task виден в подписи слота и отличает результат
    «своей» задачи от результата отменённой.
    """
    source = inspect.getsource(LogistiksRusWindow._start_recognition)

    assert source.count("partial(") == 3
    assert source.count("task=task") == 3


# ─────────────────────────────────────────────────────────────
# 3. Сбор данных
# ─────────────────────────────────────────────────────────────

def test_collect_data_returns_contract_data_from_all_tabs(window):
    _fill_all_tabs(window, cars=2)

    data = window._collect_data()

    assert isinstance(data, ContractData)
    assert data.contract["number"] == CONTRACT_NUMBER
    assert data.contract["date"] == "2026-09-24"
    assert data.contract["route"] == "Москва - Казань"
    assert data.contract["vat_rate"] == "22%"
    assert data.contract["vat_rate_num"] == 22.0
    assert data.contract["carrier_type"] == "ООО"
    assert data.contract["price_without_vat"] == FORM_AMOUNT
    assert data.contract["special_conditions"].startswith("Погрузка")

    # Точки маршрута: адреса в приведённом виде, названия — в contract.
    assert [point["address"] for point in data.loadings] == [
        "г. Москва, ул. Складская, д. 1"
    ]
    assert data.loadings[0]["time_window"] == "08:00-20:00"
    assert [point["name"] for point in data.contract["loadings"]] == [
        "ООО «Склад 1»"
    ]
    assert [point["address"] for point in data.unloadings] == [
        "г. Казань, ул. Заводская, д. 2"
    ]
    assert [point["name"] for point in data.contract["unloadings"]] == [
        "ООО «Клиент 1»"
    ]

    assert len(data.vehicles) == 2
    assert data.vehicles[0]["vin"] == "XTC651150N0001001"

    assert data.tractor == {"brand_model": "DAF XF 95.430",
                            "plate_number": "М342СА761"}
    assert data.trailer == {"brand_model": "KRONE SD",
                            "plate_number": "ВК123478"}

    assert data.driver == {"full_name": "Иванов Иван Иванович"}
    assert data.customer["full_name"] == CUSTOMER_NAME
    assert data.customer["short_name"] == CUSTOMER_NAME

    # Экспедитор в этой заявке фиксирован шаблоном — окно его не выдумывает.
    assert data.carrier == {}


def test_collect_data_is_repeatable_and_does_not_change_ui(window):
    """_collect_data только читает: повторный вызов даёт тот же результат."""
    _fill_all_tabs(window, cars=2)

    first = window._collect_data()
    second = window._collect_data()

    assert first.to_generator_dict() == second.to_generator_dict()
    assert window.customer_tab.get_data()["number"] == CONTRACT_NUMBER


def test_more_than_twelve_cars_are_capped_in_contract_data(window):
    """В бланк помещается 12 машин: лишние в ContractData не попадают."""
    _fill_all_tabs(window, cars=MAX_CARS + 3)
    assert len(window.cargo_tab.get_data()["vehicles"]) == MAX_CARS

    data = window._collect_data()

    assert len(data.vehicles) == MAX_CARS
    assert data.vehicles[-1]["vin"] == "XTC651150N0001012"


def test_empty_tabs_give_empty_contract_data(window):
    data = window._collect_data()

    assert data.vehicles == []
    assert data.loadings == []
    assert data.unloadings == []
    assert data.tractor == {}
    assert data.trailer == {}
    assert data.driver.get("full_name", "") == ""
    assert data.contract.get("number") is None
    assert data.contract.get("route") is None
    assert data.contract.get("price_without_vat") is None
    # Заказчик этой заявки подставлен вкладкой по умолчанию — он не пуст.
    assert data.customer["full_name"] == CUSTOMER_NAME


def test_full_form_passes_validator_without_errors(window):
    """Заполненная форма проходит валидатор чисто — диалог не помешает."""
    _fill_all_tabs(window, cars=2)

    report = LogistiksRusValidator().check(window._collect_data())

    assert report.errors == [], f"неожиданные ошибки: {report.errors}"
    assert report.warnings == [], f"неожиданные замечания: {report.warnings}"


# ─────────────────────────────────────────────────────────────
# 4. Очистка: только своя вкладка
# ─────────────────────────────────────────────────────────────

def test_clear_tab_clears_only_sender_tab(window):
    _fill_all_tabs(window, cars=1)

    window.driver_tab.btn_clear_form.click()

    assert window.driver_tab.get_data()["full_name"] == ""
    assert window.customer_tab.get_data()["number"] == CONTRACT_NUMBER
    assert window.route_tab.get_data()["route"] == "Москва - Казань"
    assert window.price_tab.get_data()["amount_without_vat"] == FORM_AMOUNT
    assert window.vehicle_tab.get_data()["tractor_brand"] == "DAF XF 95.430"


def test_clear_slot_without_sender_is_safe(window):
    """Прямой вызов слота без отправителя ничего не чистит и не падает."""
    _fill_all_tabs(window, cars=1)

    window._on_clear_tab()

    assert window.customer_tab.get_data()["number"] == CONTRACT_NUMBER
    assert window.driver_tab.get_data()["full_name"] == "Иванов Иван Иванович"


# ─────────────────────────────────────────────────────────────
# 5. Закрытие окна
# ─────────────────────────────────────────────────────────────

def test_close_hides_window_without_destroying_it(qt_app):
    win = LogistiksRusWindow()
    win.show()
    _fill_all_tabs(win, cars=1)
    try:
        win.close()

        assert win.isVisible() is False
        assert win.customer_tab.get_data()["number"] == CONTRACT_NUMBER
    finally:
        win.force_close()


def test_force_close_really_closes_window(qt_app):
    win = LogistiksRusWindow()
    win.show()
    win.force_close()

    assert win.isVisible() is False
    assert win._force_close is True


# ─────────────────────────────────────────────────────────────
# 6. «Создать договор»: валидатор → диалог → генератор → диалог успеха
# ─────────────────────────────────────────────────────────────

def test_empty_form_fails_validation_and_does_not_generate(
    window, monkeypatch, quiet_messages
):
    """Пустая форма: валидатор находит ошибки, генерация не запускается."""
    report = LogistiksRusValidator().check(ContractData())

    assert report.has_errors is True
    assert "Не заполнен номер заявки" in report.errors
    assert "Не заполнено ФИО водителя" in report.errors
    assert "Добавьте хотя бы одну перевозимую машину" in report.errors
    assert "Не заполнена марка тягача" in report.errors

    started = []
    monkeypatch.setattr(
        GeneratorFactory, "get_generator",
        classmethod(lambda cls, *args, **kwargs: started.append(args)),
    )
    # Диалог проверки данных: пользователь выбрал «Исправить».
    monkeypatch.setattr(QMessageBox, "exec_", lambda self: QMessageBox.No)

    window._on_create_contract()

    assert started == []
    assert window.statusBar().currentMessage() == "Заявка не создана — исправьте данные"


def test_create_contract_can_be_cancelled_in_dialog(
    window, monkeypatch, quiet_messages
):
    """Данные с ошибкой, но пользователь нажал «Исправить» — генерации нет."""
    _fill_all_tabs(window, cars=1)
    window.vehicle_tab.fill_data({"trailer_plate": ""})
    window.vehicle_tab.trailer_plate.clear()      # госномер прицепа не заполнен

    report = LogistiksRusValidator().check(window._collect_data())
    assert "Не заполнен госномер прицепа" in report.errors

    started = []
    monkeypatch.setattr(
        GeneratorFactory, "get_generator",
        classmethod(lambda cls, *args, **kwargs: started.append(args)),
    )
    monkeypatch.setattr(QMessageBox, "exec_", lambda self: QMessageBox.No)

    window._on_create_contract()

    assert started == []
    assert window.statusBar().currentMessage() == "Заявка не создана — исправьте данные"


def test_create_contract_generates_docx_end_to_end(
    window, monkeypatch, quiet_messages, work_dir, templates_dir
):
    """
    Заполненная форма: генератор → DOCX в папке вывода → диалог успеха.

    Документ проверяется целиком: файл создан, плейсхолдеры шаблона
    заменены, данные вкладок дошли до бланка.
    """
    _fill_all_tabs(window, cars=2)

    generator = LogistiksRusGenerator(templates_dir=str(templates_dir))
    output_dir = work_dir / "logistiks_rus_window_output"
    output_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(
        GeneratorFactory, "get_generator",
        classmethod(lambda cls, *args, **kwargs: generator),
    )
    monkeypatch.setattr(
        generator, "default_output_dir", lambda: str(output_dir)
    )

    created = []
    monkeypatch.setattr(
        LogistiksRusWindow, "_show_contract_created",
        lambda self, path: created.append(path),
    )
    # Данных достаточно, поэтому диалог проверки не показывается; ответ на
    # всякий случай — «Создать договор».
    monkeypatch.setattr(QMessageBox, "exec_", lambda self: QMessageBox.Yes)

    window._on_create_contract()

    assert len(created) == 1
    path = Path(created[0])
    try:
        assert path.exists() and path.suffix == ".docx"
        assert CONTRACT_NUMBER in path.name
        assert path.parent == output_dir

        text = _document_text(Document(str(path)))

        # Плейсхолдеров шаблона не осталось — и «{{...}}», и одиночных скобок.
        assert "{{" not in text, "в заявке остался плейсхолдер"
        assert "}}" not in text
        assert "<<PERSON" not in text

        # Данные вкладок дошли до бланка: заказчик, водитель, точки маршрута.
        assert "ДжейСиСиТиЭс" in text
        assert "Иванов Иван Иванович" in text
        assert "Складская" in text
    finally:
        # За собой убираем: папка вывода теста не должна пухнуть от прогонов.
        try:
            path.unlink(missing_ok=True)
            output_dir.rmdir()
        except OSError:
            pass


def test_clean_data_skips_confirmation_dialog(
    window, monkeypatch, quiet_messages, templates_dir
):
    """
    Данные без замечаний: диалог проверки не показывается вовсе.

    Ответ модели подставляется напрямую — так виден именно шаг проверки,
    а не то, что вкладки успели заполнить.
    """
    data = _valid_contract_data()
    report = LogistiksRusValidator().check(data)
    assert report.is_clean is True, report.format_text()

    generator = LogistiksRusGenerator(templates_dir=str(templates_dir))
    monkeypatch.setattr(
        GeneratorFactory, "get_generator",
        classmethod(lambda cls, *args, **kwargs: generator),
    )
    monkeypatch.setattr(LogistiksRusWindow, "_collect_data", lambda self: data)

    generated = []
    monkeypatch.setattr(
        LogistiksRusWindow, "_show_contract_created",
        lambda self, path: generated.append(path),
    )
    monkeypatch.setattr(
        generator, "generate",
        lambda payload: generated.append(payload) and "output/x.docx",
    )

    dialogs = []
    monkeypatch.setattr(
        QMessageBox, "exec_", lambda self: dialogs.append(self.windowTitle()) or 0
    )

    window._on_create_contract()

    assert dialogs == []
    assert generated == [data]


def _make_contract_file(work_dir, name: str):
    """
    Создаёт файл-заготовку заявки и убирает его после теста.

    Диалогу достаточно существующего пути: содержимое DOCX не читается,
    а мусор в tests/_tmp после прогона не нужен.
    """
    path = work_dir / name
    path.write_bytes(b"PK\x03\x04")
    return path


def test_show_contract_created_offers_open_folder(
    window, monkeypatch, quiet_messages, work_dir
):
    """Диалог успеха: кнопка «Открыть папку» и «OK», как в MainWindow."""
    path = _make_contract_file(
        work_dir, "Заявка_Логистикс_Рус_ЛР-2026-1_20260924.docx"
    )

    buttons = []
    origin_add_button = QMessageBox.addButton

    def spy_add_button(self, *args, **kwargs):
        button = origin_add_button(self, *args, **kwargs)
        buttons.append(button.text())
        return button

    monkeypatch.setattr(QMessageBox, "addButton", spy_add_button)
    # Диалог не показываем: интересна его начинка, а не модальное окно.
    monkeypatch.setattr(QMessageBox, "exec_", lambda self: 0)

    opened = []
    monkeypatch.setattr(
        LogistiksRusWindow, "_open_folder", lambda self, folder: opened.append(folder)
    )
    try:
        window._show_contract_created(str(path))

        assert buttons == ["Открыть папку", "OK"]
        assert opened == []  # нажали «OK»: папку не открываем
    finally:
        path.unlink(missing_ok=True)


def test_show_contract_created_opens_folder_on_request(
    window, monkeypatch, quiet_messages, work_dir
):
    """Кнопка «Открыть папку» открывает папку готовой заявки."""
    path = _make_contract_file(
        work_dir, "Заявка_Логистикс_Рус_ЛР-2026-2_20260924.docx"
    )

    opened = []
    monkeypatch.setattr(
        LogistiksRusWindow, "_open_folder", lambda self, folder: opened.append(folder)
    )

    # Запоминаем кнопку «Открыть папку»: в диалоге она первая.
    open_button = None
    origin_add_button = QMessageBox.addButton

    def spy_add_button(self, *args, **kwargs):
        nonlocal open_button
        button = origin_add_button(self, *args, **kwargs)
        if button.text() == "Открыть папку":
            open_button = button
        return button

    monkeypatch.setattr(QMessageBox, "addButton", spy_add_button)

    def spy_exec(box):
        box._clicked = open_button   # нажата «Открыть папку»
        return 0

    monkeypatch.setattr(QMessageBox, "exec_", spy_exec)
    monkeypatch.setattr(
        QMessageBox, "clickedButton", lambda self: getattr(self, "_clicked", None)
    )

    try:
        window._show_contract_created(str(path))

        assert open_button is not None
        assert opened == [str(work_dir)]
    finally:
        path.unlink(missing_ok=True)


# ─────────────────────────────────────────────────────────────
# 7. Распознавание
# ─────────────────────────────────────────────────────────────

def test_recognize_request_starts_task_with_logistiks_prompt(
    window, monkeypatch, qt_app
):
    """Промпт типа передаётся в recognize_text вторым аргументом."""
    client = FakeClient({"contract": {"number": CONTRACT_NUMBER}})
    monkeypatch.setattr(LogistiksRusWindow, "_ensure_gigachat", lambda self: client)

    prompt = get_prompt("logistiks_rus")
    assert prompt, "у типа «Логистикс Рус» должен быть свой промпт"

    window._on_recognize_requested("текст заявки")
    task = window.recognition_task
    assert isinstance(task, RecognitionTask)
    assert task.prompt == prompt

    # Ждём пул: задача уже выполняется в потоке. Сигнал из потока приходит
    # в поток интерфейса, поэтому очередь событий доводим до конца —
    # иначе слот не успеет сбросить состояние задачи.
    window.thread_pool.waitForDone(5000)
    qt_app.processEvents()

    assert len(client.calls) == 1
    assert client.calls[0]["text"] == "текст заявки"
    assert client.calls[0]["prompt"] == prompt

    # Связь сигналов с partial(task=...) действительно сработала.
    assert window.recognition_task is None


def test_recognition_without_key_shows_message(window, monkeypatch, quiet_messages):
    """Нет ключа — показываем MISSING_KEY_MESSAGE, задачу не запускаем."""
    monkeypatch.setattr(LogistiksRusWindow, "_ensure_gigachat", lambda self: None)

    window._on_recognize_requested("текст заявки")

    assert window.recognition_task is None
    assert any("set_key.py" in text for text in quiet_messages["critical"])


def test_empty_text_is_not_sent_to_recognition(window, monkeypatch, quiet_messages):
    """Пустой текст на распознавание не уходит."""
    called = []
    monkeypatch.setattr(LogistiksRusWindow, "_ensure_gigachat",
                        lambda self: called.append("client"))

    window._on_recognize_requested("   ")

    assert called == []
    assert quiet_messages["warning"]


def test_recognition_fills_tabs(window, quiet_messages):
    """Ответ модели раскладывается по вкладкам: у каждой — свои поля."""
    _fill_all_tabs(window, cars=1)

    _answer_recognition(window, {
        "customer": {"full_name": "ООО «Новый заказчик»",
                     "short_name": "ООО «Новый заказчик»"},
        "shipper_name": "ООО «Склад 2»",
        "loading_addresses": ["г. Тверь, ул. Новая, д. 3"],
        "consignees": [{"name": "ООО «Клиент 2»",
                        "address": "г. Сочи, ул. Морская, д. 4"}],
        "vehicles": [{"brand_model": "МОДЕЛЬ X", "vin": "XTC651150N0009001"}],
        "tractor": {"brand_model": "Volvo FH", "plate_number": "А001АА77"},
        "trailer": {"brand_model": "Schmitz", "plate_number": "ВК999977"},
        "driver": {"full_name": "Петров Пётр Петрович"},
        "contract": {
            "number": "ЛР-2026-77", "date": "01.10.2026", "cargo_count": 1,
            "loading_date": "02.10.2026",
            "loading_time_from": "09:00", "loading_time_to": "18:00",
            "unloading_date": "05.10.2026",
            "unloading_time_from": "10:00", "unloading_time_to": "17:00",
            "sum_wo_vat": 200000.0, "sum_vat": 44000.0, "sum_total": 244000.0,
            "vat_rate": "22%", "special_conditions": "Без дозагрузки",
        },
    })

    # Заказчик: наименование из customer, номер и дата — из contract.
    customer = window.customer_tab.get_data()
    assert customer["name"] == "ООО «Новый заказчик»"
    assert customer["number"] == "ЛР-2026-77"
    assert customer["date"] == "2026-10-01"

    # Груз: таблица перерисована по ответу модели.
    assert window.cargo_tab.get_data()["vehicles"] == [
        {"brand_model": "МОДЕЛЬ X", "vin": "XTC651150N0009001"}
    ]

    # Маршрут: грузоотправитель и адреса погрузки — своими полями,
    # план — из блока contract.
    route = window.route_tab.get_data()
    assert route["shipper_name"] == "ООО «Склад 2»"
    assert route["loading_addresses"] == ["г. Тверь, ул. Новая, д. 3"]
    assert route["consignees"] == [{"name": "ООО «Клиент 2»",
                                    "address": "г. Сочи, ул. Морская, д. 4"}]
    assert route["loading_date"] == "2026-10-02"
    assert route["loading_time_from"] == "09:00"
    assert route["unloading_date"] == "2026-10-05"
    assert route["unloading_time_to"] == "17:00"

    # Водитель: у этого типа печатается только ФИО.
    assert window.driver_tab.get_data()["full_name"] == "Петров Пётр Петрович"

    # Автовоз: ключи промпта (brand_model / plate_number) раскладываются
    # в поля вкладки с префиксами tractor_ / trailer_.
    vehicle = window.vehicle_tab.get_data()
    assert vehicle["tractor_brand"] == "Volvo FH"
    assert vehicle["tractor_plate"] == "А001АА77"
    assert vehicle["trailer_brand"] == "Schmitz"
    assert vehicle["trailer_plate"] == "ВК999977"

    # Стоимость: сумма без НДС берётся из sum_wo_vat, ставка и условия — из
    # того же блока.
    price = window.price_tab.get_data()
    assert price["amount_without_vat"] == 200000.0
    assert price["vat_rate"] == "22%"
    assert price["vat_rate_num"] == 22.0
    assert price["special_conditions"] == "Без дозагрузки"

    assert window.recognition_task is None
    assert window.statusBar().currentMessage() == "Готово"


def test_recognition_uses_sum_total_for_ip_document(window, quiet_messages):
    """
    «Без НДС» (ИП): сумма документа лежит в sum_total.

    В ИП-бланке одна сумма, поэтому её место занимает sum_total, а ставка
    «0%» уходит во вкладку строкой.
    """
    _answer_recognition(window, {
        "contract": {"sum_total": 269741.0, "vat_rate": "0%",
                     "sum_wo_vat": 0.0},
    })

    price = window.price_tab.get_data()
    assert price["amount_without_vat"] == 269741.0
    assert price["vat_rate"] == "0%"
    assert price["vat_rate_num"] == 0.0


def test_recognition_does_not_overwrite_manual_input(window, quiet_messages):
    """Пустые поля ответа не стирают то, что пользователь ввёл руками."""
    _fill_all_tabs(window, cars=1)

    _answer_recognition(window, {
        # Модель вернула разделы целиком, но с пустыми строками и нулями.
        "customer": {"full_name": "", "short_name": ""},
        "shipper_name": "",
        "loading_addresses": ["", "   "],
        "consignees": [{"name": "", "address": ""}],
        "vehicles": [{"brand_model": "", "vin": ""}],
        "tractor": {"brand_model": "", "plate_number": ""},
        "trailer": {"brand_model": "", "plate_number": ""},
        "driver": {"full_name": ""},
        "contract": {"number": "", "date": "", "loading_date": "",
                     "sum_wo_vat": 0.0, "sum_vat": 0.0, "sum_total": 0.0,
                     "vat_rate": "", "special_conditions": ""},
    })

    assert window.customer_tab.get_data()["number"] == CONTRACT_NUMBER
    assert window.customer_tab.get_data()["name"] == CUSTOMER_NAME
    assert window.cargo_tab.get_data()["vehicles"] == [
        {"brand_model": "МОДЕЛЬ 1", "vin": "XTC651150N0001001"}
    ]
    assert window.route_tab.get_data()["shipper_name"] == "ООО «Склад 1»"
    assert window.route_tab.get_data()["loading_addresses"] == [
        "г. Москва, ул. Складская, д. 1"
    ]
    assert window.route_tab.get_data()["route"] == "Москва - Казань"
    assert window.driver_tab.get_data()["full_name"] == "Иванов Иван Иванович"
    assert window.vehicle_tab.get_data()["tractor_brand"] == "DAF XF 95.430"
    assert window.price_tab.get_data()["amount_without_vat"] == FORM_AMOUNT
    assert window.price_tab.get_data()["special_conditions"].startswith("Погрузка")


def test_recognition_keeps_carrier_type_chosen_by_hand(window, quiet_messages):
    """
    Частичный ответ не меняет выбранный вручную тип экспедитора.

    В схеме промпта carrier_type нет (блок «carrier» запрещён), поэтому
    окно не подставляет «ООО» принудительно: от типа зависят и бланк,
    и расчёт сумм.
    """
    window.price_tab.carrier_type.setCurrentText("ИП")

    _answer_recognition(window, {"contract": {"sum_total": 100000.0,
                                              "vat_rate": "0%"}})

    assert window.price_tab.get_data()["carrier_type"] == "ИП"


def test_recognition_logs_have_no_personal_data(caplog, window, quiet_messages):
    """
    В лог распознавания не попадают данные заявки.

    Пишутся только имена разделов, имена полей и количества: ФИО, адреса,
    VIN и названия организаций в логе быть не должно.
    """
    import logging

    with caplog.at_level(logging.DEBUG):
        _answer_recognition(window, {
            "customer": {"full_name": "ООО «Секретная организация»"},
            "shippers": [{"name": "ООО «Склад 9»",
                          "address": "г. Тестоград, ул. Складская, д. 9"}],
            "vehicles": [{"brand_model": "МОДЕЛЬ СЕКРЕТ", "vin": "XTC651150N0009999"}],
            "driver": {"full_name": "Секретов Секрет Секретович"},
        })

    messages = "\n".join(record.getMessage() for record in caplog.records)

    for fragment in (
        "Секретная организация",
        "Склад 9",
        "Тестоград",
        "XTC651150N0009999",
        "МОДЕЛЬ СЕКРЕТ",
        "Секретов",
    ):
        assert fragment not in messages, f"в логе есть «{fragment}»"

    # А служебные сведения о разделах — есть.
    assert "разделы" in messages


def test_error_slot_without_sender_is_safe(window, quiet_messages):
    """Ошибка без отправителя: сообщение показываем, состояние сбрасываем."""
    window._on_recognition_error("Модель недоступна")

    assert any("Модель недоступна" in text for text in quiet_messages["critical"])
    assert window.recognition_task is None


def test_recognition_error_is_reported(window, quiet_messages):
    """Ошибка распознавания: QMessageBox.critical и сброс состояния."""
    task = RecognitionTask(FakeClient(), "текст", prompt="промпт")
    window.recognition_task = task

    window._on_recognition_error("Модель недоступна", task=task)

    assert any("Модель недоступна" in text for text in quiet_messages["critical"])
    assert window.recognition_task is None


def test_result_of_stale_task_is_ignored(window, quiet_messages):
    """Результат устаревшей (отменённой) задачи к вкладкам не применяется."""
    _fill_all_tabs(window, cars=1)
    stale = RecognitionTask(FakeClient(), "текст", prompt="промпт")

    window._on_recognition_finished(
        {"contract": {"number": "ЧУЖОЙ-НОМЕР"}}, task=stale
    )

    assert window.customer_tab.get_data()["number"] == CONTRACT_NUMBER


# ─────────────────────────────────────────────────────────────
# 8. Задача распознавания (своя, а не из MainWindow)
# ─────────────────────────────────────────────────────────────

def test_recognition_task_passes_prompt_to_client():
    """Своя задача типа: промпт уходит в клиент вторым аргументом."""
    client = FakeClient({"contract": {"number": "ЛР-1"}})
    task = RecognitionTask(client, "исходный текст", prompt="свой промпт")

    received = []
    task.signals.finished.connect(received.append)
    task.run()

    assert received == [{"contract": {"number": "ЛР-1"}}]
    assert client.calls[0]["prompt"] == "свой промпт"


def test_recognition_task_reports_empty_answer():
    """None от клиента — сигнал ошибки, а не пустой результат."""
    class NoneClient(FakeClient):
        def recognize_text(self, text, prompt=None):
            return None

    task = RecognitionTask(NoneClient(), "текст", prompt="промпт")
    errors = []
    task.signals.error.connect(errors.append)
    task.run()

    assert errors == ["Модель вернула пустой ответ"]


def test_cancelled_recognition_task_does_not_apply_result():
    """Отменённая задача результат не отдаёт — интерфейс не тронут."""
    client = FakeClient({"contract": {"number": "ЛР-1"}})
    task = RecognitionTask(client, "исходный текст", prompt="свой промпт")
    task.cancel()

    finished, cancelled = [], []
    task.signals.finished.connect(finished.append)
    task.signals.cancelled.connect(lambda: cancelled.append(True))  # noqa: B023
    task.run()

    assert finished == []
    assert cancelled == [True]


# ─────────────────────────────────────────────────────────────
# 9. Раскладка блоков ответа по вкладкам (прямые проверки карт ключей)
# ─────────────────────────────────────────────────────────────

def test_customer_tab_data_takes_header_from_contract():
    """customer даёт наименование, contract — номер и дату заявки."""
    data = LogistiksRusWindow._customer_tab_data(
        {"full_name": "ООО «Заказчик»"},
        {"number": "ЛР-1", "date": "2026-10-01", "sum_total": 5.0},
    )

    assert data == {"name": "ООО «Заказчик»", "number": "ЛР-1",
                    "date": "2026-10-01"}

    # Пустые значения (и блоки целиком) ничего не заполняют.
    assert LogistiksRusWindow._customer_tab_data({}, {}) == {}
    assert LogistiksRusWindow._customer_tab_data(
        {"full_name": ""}, {"number": "", "date": ""}
    ) == {}

    # Пустой customer не мешает заполнить шапку заявки.
    assert LogistiksRusWindow._customer_tab_data({}, {"number": "ЛР-2"}) == {
        "number": "ЛР-2"
    }


def test_cargo_tab_data_drops_empty_rows():
    """Пустые строки ответа таблицу груза не очищают."""
    assert LogistiksRusWindow._cargo_tab_data([{"brand_model": "", "vin": ""}]) == {}
    assert LogistiksRusWindow._cargo_tab_data(None) == {}
    assert LogistiksRusWindow._cargo_tab_data([]) == {}
    assert LogistiksRusWindow._cargo_tab_data(
        [{"brand_model": "МОДЕЛЬ 1", "vin": ""}]
    ) == {"vehicles": [{"brand_model": "МОДЕЛЬ 1", "vin": ""}]}


def test_route_tab_data_maps_points_and_plan():
    """Грузоотправитель и адреса погрузки — из своего блока, план — из contract."""
    data = LogistiksRusWindow._route_tab_data({
        "shipper_name": "ООО «Склад»",
        "loading_addresses": ["адрес погрузки", "  ", "адрес погрузки 2"],
        "consignees": [{"name": "ООО «Приёмка»", "address": "адрес выгрузки"}],
        "contract": {"loading_date": "02.10.2026", "loading_time_from": "09:00",
                     "sum_wo_vat": 100.0, "number": "ЛР-1"},
    })

    assert data == {
        "loading_date": "02.10.2026",
        "loading_time_from": "09:00",
        "shipper_name": "ООО «Склад»",
        "loading_addresses": ["адрес погрузки", "адрес погрузки 2"],
        "consignees": [{"name": "ООО «Приёмка»", "address": "адрес выгрузки"}],
    }

    # Пустые значения и пустой план — пустой словарь (вкладка не тронута).
    assert LogistiksRusWindow._route_tab_data({}) == {}
    assert LogistiksRusWindow._route_tab_data(
        {"loading_addresses": ["", "  "], "contract": {"loading_date": ""}}
    ) == {}


def test_route_tab_data_accepts_old_shippers_array():
    """Старый формат ответа (массив shippers) раскладывается как fallback."""
    data = LogistiksRusWindow._route_tab_data({
        "shippers": [
            {"name": "ООО «Склад Север»", "address": "адрес 1"},
            {"name": "ООО «Склад Юг»", "address": "адрес 2"},
        ],
        "consignees": [{"name": "ООО «Приёмка»", "address": "адрес выгрузки"}],
    })

    assert data["shipper_name"] == "ООО «Склад Север»"
    assert data["loading_addresses"] == ["адрес 1", "адрес 2"]
    assert data["consignees"] == [
        {"name": "ООО «Приёмка»", "address": "адрес выгрузки"},
    ]


def test_route_tab_data_new_format_wins_over_old():
    """Если пришли оба формата, новый важнее — старый не подмешивается."""
    data = LogistiksRusWindow._route_tab_data({
        "shipper_name": "ООО «Новый»",
        "loading_addresses": ["новый адрес"],
        "shippers": [{"name": "ООО «Старый»", "address": "старый адрес"}],
    })

    assert data["shipper_name"] == "ООО «Новый»"
    assert data["loading_addresses"] == ["новый адрес"]


def test_driver_tab_data_requires_full_name():
    """У водителя этого типа печатается только ФИО."""
    assert LogistiksRusWindow._driver_tab_data(
        {"full_name": "Иванов Иван Иванович", "phone": "+7 (999) 000-00-00"}
    ) == {"full_name": "Иванов Иван Иванович"}
    assert LogistiksRusWindow._driver_tab_data({"full_name": "   "}) == {}
    assert LogistiksRusWindow._driver_tab_data({}) == {}


def test_vehicle_tab_data_maps_prompt_schema():
    """
    Прямая проверка раскладки блоков tractor / trailer (схема промпта).

    Отдельно от окна: так видно саму карту ключей, а не её косвенное
    действие через вкладку.
    """
    data = LogistiksRusWindow._vehicle_tab_data({
        "tractor": {"brand_model": "DAF XF 95.430", "plate_number": "М342СА761"},
        "trailer": {"brand_model": "KRONE SD", "plate_number": "ВК123478"},
    })

    assert data == {
        "tractor_brand": "DAF XF 95.430",
        "tractor_plate": "М342СА761",
        "trailer_brand": "KRONE SD",
        "trailer_plate": "ВК123478",
    }

    # Пустые значения схема обещает как "" — они не должны ничего затирать.
    assert LogistiksRusWindow._vehicle_tab_data({
        "tractor": {"brand_model": "", "plate_number": ""},
        "trailer": {},
    }) == {}

    # Полей, которых нет в бланке этого типа (цвет, год), вкладка не получает.
    assert LogistiksRusWindow._vehicle_tab_data({
        "tractor": {"brand_model": "", "color": "Белый", "year": "2020"},
    }) == {}


def test_vehicle_tab_data_tab_keys_win_over_prompt_keys():
    """
    Ключи вкладки (tractor_brand) важнее ключей схемы промпта (brand_model).

    Ответ может прийти в обоих видах: справочник машин и повторное
    распознавание отдают имена полей вкладки, а схема промпта Логистикс
    Рус — краткие имена. Явное значение должно побеждать.
    """
    data = LogistiksRusWindow._vehicle_tab_data({
        "tractor": {"tractor_brand": "Явная марка", "brand_model": "Из промпта",
                    "plate_number": "М342СА761"},
        "trailer": {"brand_model": "KRONE SD"},
    })

    assert data == {
        "tractor_brand": "Явная марка",
        "tractor_plate": "М342СА761",
        "trailer_brand": "KRONE SD",
    }


def test_price_tab_data_uses_sum_wo_vat_then_sum_total():
    """Сумма: сначала «Стоимость услуг» (sum_wo_vat), затем итог (sum_total)."""
    assert LogistiksRusWindow._price_tab_data({
        "sum_wo_vat": 221099.18, "sum_vat": 48641.82, "sum_total": 269741.0,
        "vat_rate": "22%",
    }) == {
        "amount_without_vat": 221099.18,
        "vat_rate": "22%",
    }

    # sum_wo_vat пуста (или нулевая) — берётся итог документа.
    assert LogistiksRusWindow._price_tab_data({
        "sum_wo_vat": 0.0, "sum_total": 269741.0, "vat_rate": "0%",
        "vat_rate_num": 0.0,
    }) == {
        "amount_without_vat": 269741.0,
        "vat_rate": "0%",
        "vat_rate_num": 0.0,
    }

    # Запасное имя на случай, если модель ответила ключами ContractData.
    assert LogistiksRusWindow._price_tab_data({
        "price_without_vat": 100000.0,
    }) == {"amount_without_vat": 100000.0}

    # Сумма строкой с разделителями тысяч и запятой.
    assert LogistiksRusWindow._price_tab_data({
        "sum_wo_vat": "221 099,18",
    }) == {"amount_without_vat": 221099.18}


def test_price_tab_data_skips_empty_and_zero():
    """Ноль и пустота — «суммы не было»: вкладка не заполняется."""
    assert LogistiksRusWindow._price_tab_data({}) == {}
    assert LogistiksRusWindow._price_tab_data(None) == {}
    assert LogistiksRusWindow._price_tab_data({
        "sum_wo_vat": 0.0, "sum_vat": 0.0, "sum_total": 0.0, "vat_rate": "",
    }) == {}
    assert LogistiksRusWindow._price_tab_data({
        "special_conditions": "",
    }) == {}


def test_price_tab_data_keeps_explicit_carrier_type():
    """
    Тип экспедитора из ответа уходит во вкладку как есть.

    В схеме промпта его нет, но если модель вернула carrier_type (или его
    положил другой слой), окно не подменяет значение дефолтом.
    """
    assert LogistiksRusWindow._price_tab_data({
        "carrier_type": "ИП", "sum_total": 100000.0, "vat_rate": "0%",
    }) == {
        "carrier_type": "ИП",
        "amount_without_vat": 100000.0,
        "vat_rate": "0%",
    }

    # Без carrier_type ключ не подставляется: у вкладки свой дефолт (ООО),
    # и частичный ответ не должен стирать выбранный вручную тип.
    assert "carrier_type" not in LogistiksRusWindow._price_tab_data({
        "sum_total": 100000.0,
    })
