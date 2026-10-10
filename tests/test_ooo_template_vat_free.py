#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Бланк ООО: условный блок «Без НДС» в п. 4.1 (шаг «Ставки НДС в UI»).

У ООО на УСН с доходом до 20 млн ₽ налога нет вовсе (п. 1 ст. 145 НК).
До этого шага п. 4.1 бланка всегда печатал ставку и сумму НДС, а п. 4.2 —
«не является плательщиком НДС»: в одном договоре выходило противоречие.

Что проверяется:

  * в `templates/shablon_ooo.docx` стоит условный блок `{%p if is_vat_free %}`
    с ветками `if` / `else` и текстом «НДС не облагается»;
  * ветка `else` — ДОСЛОВНО прежний п. 4.1 (три абзаца со ставкой);
  * при `is_vat_free=True` в документе «НДС не облагается», строки про
    ставку нет; при `is_vat_free=False` — наоборот;
  * формулировка совпадает с бланком ИП без НДС;
  * инструмент правки идемпотентен и меняет ровно одну запись пакета;
  * остальные бланки перевозки этот шаг не тронул.

Бланки только ЧИТАЮТСЯ: правка выполняется на копии в рабочей папке тестов.
Данные синтетические, ПДн нет.
"""

import re
import shutil
import zipfile
from pathlib import Path

import pytest

from docx import Document

from core.contract_generator import ContractGenerator
from tools.fix_ooo_template_vat_free import (
    ELSE_TAG, ENDIF_TAG, FLAG, IF_TAG, PARAGRAPH_RE, TOTAL_FREE_TEXT,
    VAT_FREE_TEXT, check_state, fix_document_xml,
)
from tools.make_golden import strip_embedded_fonts

PROJECT_ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = PROJECT_ROOT / "templates"

OOO_TEMPLATE = TEMPLATES / "shablon_ooo.docx"
IP_WITHOUT_VAT_TEMPLATE = TEMPLATES / "shablon_ip_without_vat.docx"

#: Прежний текст п. 4.1 — ветка `else` (печатается, когда налог есть).
OLD_CLAUSE_FRAGMENTS = (
    "{{sum_wo_nds}} руб. ({{sum_wo_nds_words}}) — стоимость услуг без НДС;",
    "НДС по ставке, действующей на дату оказания услуг",
    "(в настоящее время {{vat_rate}})",
    "{{sum_nds}} руб. ({{sum_nds_words}}).",
    "Итоговая сумма Договора (стоимость услуг с НДС): {{sum_total}} руб.",
)

#: Формулировка «налога нет» — как в бланке ИП без НДС.
VAT_FREE_SENTENCE = "НДС не облагается (упрощённая система налогообложения)."


def _document_xml(path: Path) -> str:
    with zipfile.ZipFile(str(path)) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def _plain(xml: str) -> str:
    return re.sub(r"<[^>]+>", "", xml)


def _paragraph_texts(xml: str):
    """Тексты абзацев документа (без разметки) — по одному на абзац."""
    return [_plain(match.group(0)) for match in PARAGRAPH_RE.finditer(xml)]


def _clause_branches(xml: str):
    """
    Ветки условного блока п. 4.1: (ветка `if`, ветка `else`).

    Границы ищутся ПО АБЗАЦАМ: теги стоят отдельными абзацами, а якорь —
    абзац прежнего п. 4.1. Искать текст в сыром XML нельзя: плейсхолдер
    разбит на несколько run'ов («{{» + «sum_wo_nds» + «}}»), и строкой
    «{{sum_wo_nds}}» он там не встречается.
    """
    matches = list(PARAGRAPH_RE.finditer(xml))
    texts = _paragraph_texts(xml)

    anchor_at = next(
        (i for i, text in enumerate(texts) if "— стоимость услуг без НДС;" in text),
        None,
    )
    assert anchor_at is not None, "в бланке нет прежнего текста п. 4.1"

    if_at = next(
        (i for i in range(anchor_at - 1, -1, -1) if texts[i].strip() == IF_TAG),
        None,
    )
    assert if_at is not None, "у п. 4.1 нет ветки if"

    else_at = next(
        (i for i in range(if_at + 1, anchor_at) if texts[i].strip() == ELSE_TAG),
        None,
    )
    assert else_at is not None, "у п. 4.1 нет ветки else"

    endif_at = next(
        (i for i in range(anchor_at, len(texts)) if texts[i].strip() == ENDIF_TAG),
        None,
    )
    assert endif_at is not None, "у п. 4.1 нет закрывающего тега"

    return (
        xml[matches[if_at].end():matches[else_at].start()],
        xml[matches[else_at].end():matches[endif_at].start()],
    )


# ─────────────────────────────────────────────────────────────
# Бланк: структура условного блока
# ─────────────────────────────────────────────────────────────

def test_template_has_conditional_block():
    """В бланке ООО стоят все три тега, и тег `if` — по флагу is_vat_free."""
    xml = _document_xml(OOO_TEMPLATE)

    assert IF_TAG in xml
    assert ELSE_TAG in xml
    assert ENDIF_TAG in xml
    assert FLAG in xml


def test_if_branch_prints_vat_free_sentence():
    """Ветка `if` — итог и «НДС не облагается», без ставки и суммы налога."""
    branch_if, _branch_else = _clause_branches(_document_xml(OOO_TEMPLATE))
    text = _plain(branch_if)

    assert VAT_FREE_TEXT in text
    assert "sum_total" in text
    assert "sum_total_words" in text
    assert "vat_rate" not in text
    assert "sum_nds" not in text


def test_else_branch_keeps_the_old_clause():
    """Ветка `else` — дословно прежний п. 4.1 (иначе изменились бы договоры)."""
    _branch_if, branch_else = _clause_branches(_document_xml(OOO_TEMPLATE))

    for fragment in OLD_CLAUSE_FRAGMENTS:
        assert fragment in _plain(branch_else), fragment


def test_else_branch_keeps_three_paragraphs_in_order():
    """В ветке `else` те же ТРИ абзаца и в том же порядке, что и раньше."""
    _branch_if, branch_else = _clause_branches(_document_xml(OOO_TEMPLATE))
    texts = [text.strip() for text in _paragraph_texts(branch_else) if text.strip()]

    assert len(texts) == 3, texts
    assert "без НДС;" in texts[0]
    assert texts[1].startswith("НДС по ставке")
    assert texts[2].startswith("Итоговая сумма Договора")


def test_wording_matches_ip_without_vat_template():
    """Тот же случай в бланке ИП без НДС написан теми же словами."""
    ooo = _plain(_document_xml(OOO_TEMPLATE))
    ip = _plain(_document_xml(IP_WITHOUT_VAT_TEMPLATE))

    assert VAT_FREE_SENTENCE in ip
    assert VAT_FREE_SENTENCE in ooo


def test_other_templates_are_untouched():
    """Бланки ИП этот шаг не трогал: условного блока is_vat_free в них нет."""
    for name in ("shablon_ip_with_vat.docx", "shablon_ip_without_vat.docx"):
        xml = _document_xml(TEMPLATES / name)
        assert IF_TAG not in xml, name


# ─────────────────────────────────────────────────────────────
# Инструмент правки
# ─────────────────────────────────────────────────────────────

def test_tool_state_reports_the_block():
    """`--check` видит блок и текст «НДС не облагается»."""
    state = check_state(OOO_TEMPLATE)

    assert "блок есть" in state
    assert "НДС не облагается" in state
    assert state.count("else") >= 1


def test_tool_is_idempotent():
    """Повторный прогон ничего не меняет: правка ровно одна."""
    xml = _document_xml(OOO_TEMPLATE)

    fixed, changes = fix_document_xml(xml)

    assert changes == 0
    assert fixed == xml


def _xml_without_block(xml: str) -> str:
    """
    Снимает условный блок — тем же разбором абзацев, что и инструмент.

    Нужно, чтобы проверить саму правку: инструмент должен собрать блок
    обратно БАЙТ В БАЙТ, а не «примерно так же».
    """
    matches = list(PARAGRAPH_RE.finditer(xml))
    if_at = next(i for i, m in enumerate(matches) if IF_TAG in _plain(m.group(0)))
    endif_at = next(
        i for i in range(if_at + 1, len(matches))
        if _plain(matches[i].group(0)).strip() == ENDIF_TAG
    )

    # Внутри блока: тег if, два новых абзаца, тег else, три прежних абзаца.
    first_old = matches[if_at + 4]
    last_old = matches[endif_at - 1]

    return (
        xml[:matches[if_at].start()]
        + xml[first_old.start():last_old.end()]
        + xml[matches[endif_at].end():]
    )


def test_tool_restores_the_block_byte_for_byte():
    """Инструмент собирает блок побайтово так же, как он стоит в бланке."""
    xml = _document_xml(OOO_TEMPLATE)
    without_block = _xml_without_block(xml)

    assert IF_TAG not in without_block
    assert "— стоимость услуг без НДС;" in _plain(without_block), (
        "прежний п. 4.1 должен остаться"
    )

    fixed, changes = fix_document_xml(without_block)

    assert changes == 1
    assert fixed == xml


def test_tool_reports_missing_clause():
    """Без п. 4.1 инструмент честно падает, а не портит бланк молча."""
    with pytest.raises(ValueError):
        fix_document_xml("<w:document><w:body></w:body></w:document>")


def test_tool_keeps_other_parts_untouched(tmp_path):
    """Правка меняет ровно `word/document.xml`: шрифты и части — побайтово."""
    target = tmp_path / "shablon_ooo.docx"
    shutil.copy2(OOO_TEMPLATE, target)

    with zipfile.ZipFile(str(target)) as archive:
        parts_before = {info.filename: archive.read(info.filename)
                        for info in archive.infolist()}

    # Возвращаем бланк в состояние «до правки» и правим инструментом.
    xml = parts_before["word/document.xml"].decode("utf-8")
    without_block = _xml_without_block(xml)
    fixed, changes = fix_document_xml(without_block)
    assert changes == 1

    from tools.fix_ooo_template_vat_free import _rewrite_document_part

    _rewrite_document_part(target, fixed.encode("utf-8"))

    with zipfile.ZipFile(str(target)) as archive:
        parts_after = {info.filename: archive.read(info.filename)
                       for info in archive.infolist()}

    assert set(parts_after) == set(parts_before)
    assert parts_after["word/document.xml"] == parts_before["word/document.xml"]
    for name in parts_before:
        if name != "word/document.xml":
            assert parts_after[name] == parts_before[name], name

    fonts = [name for name in parts_after if name.startswith("word/fonts/")]
    assert len(fonts) == 3, fonts


# ─────────────────────────────────────────────────────────────
# Договор: обе ветки на живом рендере
# ─────────────────────────────────────────────────────────────

#: Синтетический перевозчик-ООО (ПДн нет).
CARRIER = {
    "full_name": "ООО «Ромашка»", "short_name": "ООО «Ромашка»",
    "inn": "7701234567", "kpp": "770101001", "ogrn": "1027700132195",
    "legal_address": "г. Москва, ул. Тестовая, д. 1",
    "bank_account": "40702810000000000001", "bik": "044525225",
    "correspondent_account": "30101810400000000225", "bank_name": "ПАО Сбербанк",
    "director_name": "Петров Пётр Петрович",
    "director_position": "Генеральный директор", "entity_type": "ООО",
}
CUSTOMER = {
    "full_name": "ООО «Заказчик»", "short_name": "ООО «Заказчик»",
    "inn": "7701234567", "kpp": "770101001",
}
CONTRACT = {
    "number": "НДС-БЛАНК-1", "date": "2026-10-10", "route": "Мурманск - Пятигорск",
    "price_without_vat": 180300.0, "payment_days": 10,
}


@pytest.fixture(scope="module")
def light_templates(tmp_path_factory) -> Path:
    """Облегчённые копии бланков (без встроенных шрифтов) — как в golden."""
    work = tmp_path_factory.mktemp("vat_free_templates")
    for name in ("shablon_ooo.docx", "shablon_ip_with_vat.docx",
                 "shablon_ip_without_vat.docx"):
        strip_embedded_fonts(TEMPLATES / name, work / name)
    return work


def _render(light_templates: Path, out_path: Path, **contract) -> str:
    generator = ContractGenerator(templates_dir=str(light_templates))
    generator.generate_docx(
        {
            "driver": {"full_name": "Иванов Иван Иванович"},
            "carrier": dict(CARRIER),
            "customer": dict(CUSTOMER),
            "vehicles": [{"brand_model": "JETOUR T2", "vin": "EC3TEUMB0T0002608",
                          "vehicle_type": "Легковой автомобиль"}],
            "contract": dict(CONTRACT, **contract),
            "loadings": [{"address": "183052, г.Мурманск, пр.Кольский, д.53",
                          "date": "2026-10-11", "time_window": "09:00-18:00"}],
            "unloadings": [{"address": "г. Пятигорск, шоссе 17",
                            "date": "2026-10-14", "time_window": ""}],
        },
        str(out_path),
    )
    return _document_text(out_path)


def _document_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_contract_without_vat_prints_the_free_sentence(light_templates, work_file):
    """Ветвь `if`: «НДС не облагается», ставки и суммы налога в п. 4.1 нет."""
    text = _render(
        light_templates, work_file("vat_free_ooo.docx"),
        carrier_type="ООО (без НДС)", vat_rate="Без НДС", vat_rate_num=0,
    )

    assert VAT_FREE_SENTENCE in text
    assert "НДС по ставке, действующей на дату оказания услуг" not in text
    assert "Итоговая сумма Договора (стоимость услуг с НДС)" not in text
    # Итог печатается один раз — он же стоимость без НДС.
    assert "180300.00 руб. (Сто восемьдесят тысяч триста рублей 00 копеек)." in text
    assert "не является плательщиком НДС" in text


def test_contract_with_vat_keeps_the_rate_line(light_templates, work_file):
    """Ветвь `else`: ставка и сумма налога на месте, «не облагается» нет."""
    text = _render(
        light_templates, work_file("vat_ooo.docx"),
        carrier_type="ООО (с НДС)", vat_rate="22%", vat_rate_num=22,
    )

    assert "НДС по ставке, действующей на дату оказания услуг" in text
    assert "(в настоящее время 22%)" in text
    assert "39666.00 руб." in text
    assert "Итоговая сумма Договора (стоимость услуг с НДС): 219966.00 руб." in text
    assert VAT_FREE_SENTENCE not in text


@pytest.mark.parametrize("rate, nds, total", [
    ("5%", 9015.00, 189315.00),
    ("7%", 12621.00, 192921.00),
    ("10%", 18030.00, 198330.00),
    ("0%", 0.00, 180300.00),
])
def test_rate_line_follows_the_rate(light_templates, work_file, rate, nds, total):
    """Ставка из формы печатается как есть, налог и итог считаются по ней."""
    text = _render(
        light_templates, work_file(f"vat_{rate.strip('%')}.docx"),
        carrier_type="ООО (с НДС)", vat_rate=rate,
        vat_rate_num=float(rate.strip("%")),
    )

    assert f"(в настоящее время {rate})" in text
    assert f"{nds:.2f} руб." in text
    assert f"Итоговая сумма Договора (стоимость услуг с НДС): {total:.2f} руб." in text
    if rate == "0%":
        # Ставка 0 % — это НЕ «не облагается»: ставка есть, налог нулевой.
        assert VAT_FREE_SENTENCE not in text


def test_no_leftover_jinja_tags(light_templates, work_file):
    """В готовом договоре не остаётся ни тегов, ни плейсхолдеров шаблона."""
    for name, contract in (
        ("tags_vat.docx", {"carrier_type": "ООО (с НДС)", "vat_rate": "22%",
                           "vat_rate_num": 22}),
        ("tags_free.docx", {"carrier_type": "ООО (без НДС)",
                            "vat_rate": "Без НДС", "vat_rate_num": 0}),
    ):
        text = _render(light_templates, work_file(name), **contract)

        for marker in ("{{", "}}", "{%", "%}"):
            assert marker not in text, (name, marker)
