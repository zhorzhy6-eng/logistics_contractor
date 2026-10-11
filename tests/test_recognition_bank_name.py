#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты поля «Наименование банка» (шаг «Четыре бага», баг D).

ЧТО ЗА БАГ
---------
Наименование банка искалось правилом с ОБЯЗАТЕЛЬНЫМ двоеточием:

    ^банк\\s+получателя\\s*[:|]\\s*(.+)$

А на самом бланке реквизитов двоеточия нет вовсе — подпись «Банк получателя»
стоит отдельной ячейкой, значение в соседней. В стресс-тесте (1000
синтетических документов) из-за этого поле терялось в 197 случаях: строки
«Банк получателя ПАО «Открытие»» и «Банк получателя. ДО «Газпромбанк»»
(точка вместо двоеточия после OCR) не подходили под правило.

КАК ПОЧИНЕНО
------------
Двоеточие (и точка, и вертикальная черта) в правиле стали НЕобязательными,
а метка «Банк получателя» осталась обязательной; добавлены запасные метки
«Наименование банка» и «Банк:» — так банк подписан в карточке контрагента и
в выписках.

ЧТО ПРОВЕРЯЕТСЯ
---------------
1. Значение извлекается из всех форм метки.
2. В значение НЕ попадают БИК, к/с и р/с — ни со своей строки, ни с соседней.
3. БИК, расчётный и корреспондентский счёт разбираются как раньше.
4. Значение доходит до вкладки «Перевозчик» (её get_data) и до бланка
   договора-заявки на перевозку (плейсхолдер carrier_bank / has_carrier_bank).

Данные синтетические, реальных документов и ПДн здесь нет.
"""

import pytest

from core.document_import_service import extract_local_fields
from ui.tabs.carrier_tab import CarrierTab

#: Реквизиты, которые обязаны разбираться наравне с названием банка.
BANK_ACCOUNT = "40702810000000000002"
BIK = "044525225"
CORR_ACCOUNT = "30101810400000000225"

#: Название банка — так, как оно напечатано в бланке.
BANK_NAME = "ПАО Сбербанк"


def _carrier(text: str) -> dict:
    """Раздел «carrier» разбора (первая запись)."""
    fields = extract_local_fields(text)
    carrier = fields.get("carrier")
    assert carrier, f"раздел «carrier» не найден: {list(fields)}"
    return carrier[0]


def _requisites(bank_line: str) -> str:
    """
    Синтетический бланк реквизитов: получатель, ИНН/КПП, банк, БИК, счета.

    `bank_line` — строка с названием банка: ею и проверяются формы метки.
    """
    return "\n".join((
        "РЕКВИЗИТЫ",
        "Получатель ООО «Пример-Логистик»",
        "ИНН/КПП 7707654321 770701001",
        bank_line,
        f"БИК {BIK}",
        f"К/с {CORR_ACCOUNT}",
        f"номер счёта получателя {BANK_ACCOUNT}",
        "Телефон +7 (495) 000-11-22",
    ))


#: (имя кейса, строка с названием банка, ожидаемое значение).
BANK_LABEL_CASES = (
    ("bank_of_recipient_no_colon", "Банк получателя ПАО Сбербанк", "ПАО Сбербанк"),
    ("bank_of_recipient_colon", "Банк получателя: ПАО Сбербанк", "ПАО Сбербанк"),
    ("bank_of_recipient_dot", "Банк получателя. АО «Газпромбанк»",
     "АО «Газпромбанк»"),
    ("bank_of_recipient_pipe", "Банк получателя | ПАО «Открытие»",
     "ПАО «Открытие»"),
    ("bank_of_recipient_lowercase", "банк получателя ПАО Сбербанк", "ПАО Сбербанк"),
    ("bank_name_label", "Наименование банка: АО Банк", "АО Банк"),
    ("bank_name_label_no_colon", "Наименование банка АО Банк", "АО Банк"),
    ("short_label", "Банк: ПАО Сбербанк", "ПАО Сбербанк"),
    ("quoted_name", "Банк получателя Банк «Пример»", "Банк «Пример»"),
    ("name_with_spaces", "Банк получателя  АО  «Тинькофф  Банк»",
     "АО  «Тинькофф  Банк»"),
)


@pytest.mark.parametrize(
    "line",
    [case[1] for case in BANK_LABEL_CASES],
    ids=[case[0] for case in BANK_LABEL_CASES],
)
def test_bank_name_is_read_from_every_label_form(line):
    """Название банка извлекается из всех форм подписи — с двоеточием и без."""
    expected = dict((case[1], case[2]) for case in BANK_LABEL_CASES)[line]

    assert _carrier(_requisites(line))["bank_name"] == expected


def test_bank_name_does_not_swallow_bik_and_accounts():
    """В название банка НЕ попадают БИК, к/с и р/с — ни со своей строки."""
    carrier = _carrier(_requisites("Банк получателя ПАО Сбербанк"))

    assert carrier["bank_name"] == BANK_NAME
    for value in (BIK, CORR_ACCOUNT, BANK_ACCOUNT):
        assert value not in carrier["bank_name"]


def test_bank_name_stops_before_the_next_line():
    """Значение берётся со СВОЕЙ строки: соседние строки реквизитов не липнут."""
    carrier = _carrier(_requisites("Банк получателя: ПАО Сбербанк"))

    assert carrier["bank_name"] == BANK_NAME
    assert "БИК" not in carrier["bank_name"]
    assert "Телефон" not in carrier["bank_name"]


def test_bank_name_is_not_taken_from_bik_line():
    """БИК без названия банка не становится названием: поле остаётся пустым."""
    text = "\n".join((
        "РЕКВИЗИТЫ",
        "Получатель ООО «Пример-Логистик»",
        f"БИК {BIK}",
        f"номер счёта получателя {BANK_ACCOUNT}",
    ))
    carrier = _carrier(text)

    assert carrier.get("bank_name", "") == ""
    assert carrier["bik"] == BIK


# ─────────────────────────────────────────────────────────────
# Регресс: остальные реквизиты разбираются как раньше
# ─────────────────────────────────────────────────────────────

def test_bik_account_and_corr_account_are_still_read():
    """БИК, расчётный и корреспондентский счёт — на месте и без лишних цифр."""
    carrier = _carrier(_requisites("Банк получателя ПАО Сбербанк"))

    assert carrier["bik"] == BIK
    assert carrier["bank_account"] == BANK_ACCOUNT
    assert carrier["correspondent_account"] == CORR_ACCOUNT


def test_organisation_requisites_are_untouched():
    """ИНН, КПП и получатель не затронуты правкой названия банка."""
    carrier = _carrier(_requisites("Банк получателя ПАО Сбербанк"))

    assert carrier["inn"] == "7707654321"
    assert carrier["kpp"] == "770701001"
    assert carrier["full_name"] == "ООО «Пример-Логистик»"
    assert carrier["phone"] == "+7 (495) 000-11-22"


def test_inn_of_bank_is_still_not_the_carrier_inn():
    """ИНН банка получателя по-прежнему НЕ становится ИНН перевозчика."""
    text = "\n".join((
        "ИП Примеров Тест Тестович",
        "Реквизиты:",
        "Получатель: Примеров Тест Тестович",
        f"Номер счёта Получателя: {BANK_ACCOUNT}",
        "Банк Получателя: Банк Пример",
        "ИНН Банка Получателя: 1234567890",
        f"БИК Банка Получателя: {BIK}",
        f"К/С Банка Получателя: {CORR_ACCOUNT}",
    ))
    carrier = _carrier(text)

    assert carrier["bank_name"] == "Банк Пример"
    assert "inn" not in carrier


# ─────────────────────────────────────────────────────────────
# Значение доходит до вкладки и до бланка
# ─────────────────────────────────────────────────────────────

def _carrier_form_values(carrier: dict) -> dict:
    """Данные разбора — в значения полей вкладки «Перевозчик» (как в окне)."""
    return {
        "full_name": carrier.get("full_name", ""),
        "inn": carrier.get("inn", ""),
        "kpp": carrier.get("kpp", ""),
        "bank_name": carrier.get("bank_name", ""),
        "bank_account": carrier.get("bank_account", ""),
        "bik": carrier.get("bik", ""),
        "correspondent_account": carrier.get("correspondent_account", ""),
    }


@pytest.fixture(scope="module")
def qt_app():
    from PyQt5.QtWidgets import QApplication

    app = QApplication.instance() or QApplication([])
    yield app


def test_bank_name_reaches_the_carrier_tab(qt_app):
    """Распознанное название банка доезжает до вкладки «Перевозчик»."""
    carrier = _carrier(_requisites("Банк получателя ПАО Сбербанк"))

    tab = CarrierTab()
    try:
        tab.fill_data(_carrier_form_values(carrier))
        data = tab.get_data()
    finally:
        tab.deleteLater()

    assert data["bank_name"] == BANK_NAME
    assert data["bik"] == BIK
    assert data["bank_account"] == BANK_ACCOUNT
    assert data["correspondent_account"] == CORR_ACCOUNT


@pytest.fixture
def templates_dir():
    from pathlib import Path

    return Path(__file__).resolve().parent.parent / "templates"


def test_bank_name_reaches_the_contract_blank(templates_dir, tmp_path, qt_app):
    """
    Сквозная проверка: распознанный банк печатается в бланке перевозки.

    Конвейер тот же, что в приложении: разбор → поля вкладки «Перевозчик» →
    ContractData → карта замен генератора → плейсхолдер carrier_bank бланка.
    """
    from core.contract_data import ContractData
    from core.contracts.perevozka.generator import ContractGenerator

    carrier = _carrier(_requisites("Банк получателя ПАО Сбербанк"))
    tab = CarrierTab()
    try:
        tab.fill_data(_carrier_form_values(carrier))
        carrier_data = tab.get_data()
    finally:
        tab.deleteLater()

    data = ContractData.coerce({
        "carrier": carrier_data,
        "customer": {"full_name": "ООО «Заказчик»"},
        "contract": {"number": "ПР-1", "date": "2026-10-11"},
    })

    generator = ContractGenerator(templates_dir=str(templates_dir))
    replacements = generator._build_replacements_map(data)

    assert replacements["carrier_bank"] == BANK_NAME
    assert replacements["has_carrier_bank"] is True
    assert replacements["carrier_bik"] == BIK

    from docx import Document

    output = tmp_path / "perevozka_bank.docx"
    generator.generate_docx(data, str(output))
    text = _document_text(Document(str(output)))

    assert BANK_NAME in text, "название банка не попало в бланк"
    assert "{{" not in text and "}}" not in text


def _document_text(doc) -> str:
    """Весь текст документа: абзацы и таблицы (включая вложенные)."""
    parts = [p.text for p in doc.paragraphs]

    def table_text(table):
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
                for nested in cell.tables:
                    table_text(nested)

    for table in doc.tables:
        table_text(table)
    return "\n".join(parts)
