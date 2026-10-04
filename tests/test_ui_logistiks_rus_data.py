#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты сборки данных «Логистикс Рус» из вкладок (ЭТАП 3.1.C.B.1).

Проверяется ui/windows/logistiks_rus/data.py::collect_logistiks_rus_data:
раскладка полей вкладок по ContractData, три маппинга этого шага (точки
с name, shippers/consignees → loadings/unloadings, суммы sum_* →
price_without_vat) и устойчивость сборки (нет вкладки, нет get_data(),
get_data() упал, вернул не словарь).

Qt не нужен: вкладки подменяются простыми объектами-заглушками со своим
get_data(). Отдельные тесты проверяют, что собранных данных достаточно
валидатору и генератору типа, а модуль сборки не тянет PyQt5.

Все данные синтетические, реальных ПДн нет.
"""

import logging
import re
from pathlib import Path

import pytest

from core.contract_data import ContractData
from ui.windows.logistiks_rus.data import (
    DEFAULT_CARRIER_TYPE,
    DEFAULT_VAT_RATE_NUM,
    MAX_CARS,
    MAX_POINTS,
    build,
    collect_logistiks_rus_data,
)

# ─────────────────────────────────────────────────────────────
# Константы тестовых данных
# ─────────────────────────────────────────────────────────────

CUSTOMER_NAME = "ООО «ДжейСиСиТиЭс Интернейшнл Логистикс Рус»"

#: VIN из conftest (17 символов, ISO 3779) — на нём данных хватает валидатору.
VIN_1 = "EC3TEUMB0T0002608"
VIN_2 = "XTC651150N0001001"

#: Суммы ООО-варианта: 221 099,18 + 22% = 48 641,82 → 269 741,00.
PRICE_WITHOUT_VAT = 221099.18
PRICE_WITH_VAT = 269741.00
VAT_RATE_NUM = 22.0

#: Сумма ИП-варианта: одна строка «Без НДС», ставка не применяется.
IP_PRICE = 135833.00

#: Те же суммы в форме ответа распознавания (core/prompts/logistiks_rus.py).
RECOGNIZED_SUM_WO_VAT = 221099.18
RECOGNIZED_SUM_VAT = 48641.82
RECOGNIZED_SUM_TOTAL = 269741.00

#: Незаменённый плейсхолдер бланка («{{car_12_brand}}») в готовом документе.
PLACEHOLDER_RE = re.compile(r"\{\{[^{}]*\}\}")


class StubTab:
    """Вкладка-заглушка: отдаёт заранее заданный словарь."""

    def __init__(self, data):
        self._data = data
        self.calls = 0

    def get_data(self):
        self.calls += 1
        return self._data


class BrokenTab:
    """Вкладка, чей get_data() падает."""

    def get_data(self):
        raise RuntimeError("вкладка сломана")


class NoGetDataTab:
    """Вкладка без get_data() (например, ещё не написанная)."""


class NotADictTab:
    """Вкладка, вернувшая не словарь."""

    def get_data(self):
        return ["не", "словарь"]


# ─────────────────────────────────────────────────────────────
# Пустой вход и устойчивость
# ─────────────────────────────────────────────────────────────

def test_empty_input_gives_empty_contract_data():
    cd = collect_logistiks_rus_data({})

    assert isinstance(cd, ContractData)
    assert cd.driver == {}
    assert cd.customer == {}
    assert cd.carrier == {}
    assert cd.vehicles == []
    assert cd.tractor == {}
    assert cd.trailer == {}
    assert cd.loadings == []
    assert cd.unloadings == []
    # Массивы точек остаются в contract всегда (генератор и валидатор читают
    # именно их), но пустой вход не добавляет ни одного значения.
    assert cd.contract == {"loadings": [], "unloadings": []}


def test_empty_input_keeps_city_empty():
    """Город в сборе не задаём: ContractData.resolved_city() подставит «Москва»."""
    cd = collect_logistiks_rus_data({})

    assert cd.city == ""
    assert cd.resolved_city() == "Москва"


def test_partial_tabs_do_not_break_collection():
    cd = collect_logistiks_rus_data({"customer": {
        "number": "ЛР-2026-17", "date": "2026-09-24", "name": CUSTOMER_NAME,
    }})

    assert cd.contract["number"] == "ЛР-2026-17"
    assert cd.customer["full_name"] == CUSTOMER_NAME
    assert cd.vehicles == []
    assert cd.driver == {}


def test_tab_object_with_get_data_is_used():
    tab = StubTab({"number": "  ЛР-2026-17  ", "date": " 2026-09-24 "})

    cd = collect_logistiks_rus_data({"customer": tab})

    assert tab.calls == 1
    assert cd.contract["number"] == "ЛР-2026-17"
    assert cd.contract["date"] == "2026-09-24"


def test_failing_get_data_does_not_break_collection():
    tabs = {
        "customer": BrokenTab(),
        "route": {"route": "Москва - Казань"},
    }

    cd = collect_logistiks_rus_data(tabs)

    assert cd.contract.get("number") is None
    assert cd.contract["route"] == "Москва - Казань"


def test_tab_without_get_data_is_skipped():
    cd = collect_logistiks_rus_data({"driver": NoGetDataTab()})

    assert cd.driver == {}


def test_tab_returning_not_a_dict_is_skipped():
    cd = collect_logistiks_rus_data({"cargo": NotADictTab()})

    assert cd.vehicles == []


def test_tabs_not_a_mapping_gives_empty_result():
    """Даже совсем неверный вход не роняет сборку."""
    cd = collect_logistiks_rus_data(None)

    assert isinstance(cd, ContractData)
    assert cd.contract == {"loadings": [], "unloadings": []}
    assert cd.vehicles == []


def test_data_module_does_not_import_qt():
    """Сборка данных не зависит от интерфейса — PyQt5 в модуле не нужен."""
    import ui.windows.logistiks_rus.data as data_module

    assert not hasattr(data_module, "QWidget")
    assert "PyQt5" not in getattr(data_module, "__dict__", {})


# ─────────────────────────────────────────────────────────────
# Заказчик
# ─────────────────────────────────────────────────────────────

def test_customer_name_goes_to_full_and_short_name():
    """Заказчик фиксирован: вкладка отдаёт одно имя, оно же и краткое."""
    cd = collect_logistiks_rus_data({"customer": {
        "number": "ЛР-2026-17", "date": "2026-09-24", "name": CUSTOMER_NAME,
    }})

    assert cd.customer["full_name"] == CUSTOMER_NAME
    assert cd.customer["short_name"] == cd.customer["full_name"]


def test_empty_customer_name_is_skipped():
    cd = collect_logistiks_rus_data({"customer": {"name": "   "}})

    assert "full_name" not in cd.customer
    assert "short_name" not in cd.customer


def test_empty_number_and_date_are_skipped():
    cd = collect_logistiks_rus_data({"customer": {
        "number": "   ", "date": "", "name": CUSTOMER_NAME,
    }})

    assert "number" not in cd.contract
    assert "date" not in cd.contract
    assert cd.customer["full_name"] == CUSTOMER_NAME


# ─────────────────────────────────────────────────────────────
# Груз: перевозимые автомобили
# ─────────────────────────────────────────────────────────────

def test_vehicles_are_collected():
    cd = collect_logistiks_rus_data({"cargo": {"vehicles": [
        {"brand_model": "JETOUR T2", "vin": VIN_1},
    ]}})

    assert cd.vehicles == [{"brand_model": "JETOUR T2", "vin": VIN_1}]


def test_empty_vehicle_rows_are_dropped():
    cd = collect_logistiks_rus_data({"cargo": {"vehicles": [
        {"brand_model": "", "vin": ""},
        {"brand_model": "  ", "vin": "  "},
        {"brand_model": "JETOUR T2", "vin": VIN_1},
    ]}})

    assert len(cd.vehicles) == 1
    assert cd.vehicles[0]["brand_model"] == "JETOUR T2"


def test_vehicle_with_only_brand_or_only_vin_is_kept():
    """Заполнена хотя бы одна колонка — строку не выбрасываем."""
    cd = collect_logistiks_rus_data({"cargo": {"vehicles": [
        {"brand_model": "JETOUR T2", "vin": ""},
        {"brand_model": "", "vin": VIN_1},
    ]}})

    assert len(cd.vehicles) == 2
    assert cd.vehicles[0]["brand_model"] == "JETOUR T2"
    assert cd.vehicles[1]["vin"] == VIN_1


def test_vehicles_are_limited_to_max_cars():
    rows = [{"brand_model": f"МОДЕЛЬ {i}", "vin": VIN_1} for i in range(1, 16)]

    cd = collect_logistiks_rus_data({"cargo": {"vehicles": rows}})

    assert MAX_CARS == 12
    assert len(cd.vehicles) == MAX_CARS
    assert cd.vehicles[-1]["brand_model"] == f"МОДЕЛЬ {MAX_CARS}"


def test_vehicles_not_a_list_is_ignored():
    cd = collect_logistiks_rus_data({"cargo": {"vehicles": "мусор"}})

    assert cd.vehicles == []


# ─────────────────────────────────────────────────────────────
# Маршрут: план погрузки и выгрузки
# ─────────────────────────────────────────────────────────────

@pytest.fixture
def route_tab() -> dict:
    """Вкладка «Маршрут»: два грузоотправителя и один грузополучатель."""
    return {
        "route": "Москва - Казань",
        "shippers": [
            {"name": "ООО «Склад Север»", "address": "г. Москва, ул. Складская, д. 1"},
            {"name": "ООО «Склад Юг»", "address": "г. Москва, ул. Южная, д. 2"},
        ],
        "consignees": [
            {"name": "ООО «Приёмка»", "address": "г. Казань, ул. Приёмная, д. 3"},
        ],
        "loading_date": "2026-09-26",
        "loading_time_from": "08:00",
        "loading_time_to": "20:00",
        "unloading_date": "2026-10-01",
        "unloading_time_from": "09:00",
        "unloading_time_to": "18:00",
    }


def test_route_fields_are_collected(route_tab):
    cd = collect_logistiks_rus_data({"route": route_tab})

    assert cd.contract["route"] == "Москва - Казань"
    assert cd.contract["loading_date"] == "2026-09-26"
    assert cd.contract["loading_time_from"] == "08:00"
    assert cd.contract["loading_time_to"] == "20:00"
    assert cd.contract["unloading_date"] == "2026-10-01"
    assert cd.contract["unloading_time_from"] == "09:00"
    assert cd.contract["unloading_time_to"] == "18:00"


def test_shippers_and_consignees_go_to_contract_points(route_tab):
    """Маппинг 2: shippers / consignees → contract.loadings / unloadings."""
    cd = collect_logistiks_rus_data({"route": route_tab})

    assert [p["name"] for p in cd.contract["loadings"]] == [
        "ООО «Склад Север»", "ООО «Склад Юг»",
    ]
    assert [p["address"] for p in cd.contract["loadings"]] == [
        "г. Москва, ул. Складская, д. 1", "г. Москва, ул. Южная, д. 2",
    ]
    assert [p["name"] for p in cd.contract["unloadings"]] == ["ООО «Приёмка»"]
    assert [p["address"] for p in cd.contract["unloadings"]] == [
        "г. Казань, ул. Приёмная, д. 3",
    ]


def test_points_in_contract_keep_name_date_and_time_window(route_tab):
    """Маппинг 1: точки лежат в contract полным набором — с name."""
    cd = collect_logistiks_rus_data({"route": route_tab})

    assert cd.contract["loadings"][0] == {
        "name": "ООО «Склад Север»",
        "address": "г. Москва, ул. Складская, д. 1",
        "date": "2026-09-26",
        "time_window": "08:00-20:00",
    }
    assert cd.contract["unloadings"][0] == {
        "name": "ООО «Приёмка»",
        "address": "г. Казань, ул. Приёмная, д. 3",
        "date": "2026-10-01",
        "time_window": "09:00-18:00",
    }


def test_top_level_points_stay_without_name(route_tab):
    """
    Верхнеуровневые точки — в приведённом виде, без name.

    ContractData хранит точки как {address, date, time_window}: название
    в них теряется (core.contract_data._as_point_list) — поэтому названия
    и кладутся отдельно, в contract["loadings"] / ["unloadings"].
    """
    cd = collect_logistiks_rus_data({"route": route_tab})

    assert cd.loadings == [
        {"address": "г. Москва, ул. Складская, д. 1",
         "date": "2026-09-26", "time_window": "08:00-20:00"},
        {"address": "г. Москва, ул. Южная, д. 2",
         "date": "2026-09-26", "time_window": "08:00-20:00"},
    ]
    assert cd.unloadings == [
        {"address": "г. Казань, ул. Приёмная, д. 3",
         "date": "2026-10-01", "time_window": "09:00-18:00"},
    ]
    assert "name" not in cd.loadings[0]


def test_points_without_name_or_address_are_dropped():
    cd = collect_logistiks_rus_data({"route": {
        "shippers": [
            {"name": "", "address": ""},
            {"name": "  ", "address": "  "},
            {"name": "ООО «Склад Север»", "address": ""},
        ],
    }})

    assert len(cd.contract["loadings"]) == 1
    assert cd.contract["loadings"][0]["name"] == "ООО «Склад Север»"
    assert cd.contract["loadings"][0]["address"] == ""


def test_points_are_limited_to_max_points():
    route = {
        "shippers": [{"name": f"ООО «Склад {i}»", "address": f"адрес {i}"}
                     for i in range(1, 14)],
        "consignees": [{"name": f"ООО «Приёмка {i}»", "address": f"адрес {i}"}
                       for i in range(1, 13)],
    }

    cd = collect_logistiks_rus_data({"route": route})

    assert MAX_POINTS == 10
    assert len(cd.contract["loadings"]) == MAX_POINTS
    assert len(cd.contract["unloadings"]) == MAX_POINTS
    assert cd.contract["loadings"][-1]["name"] == f"ООО «Склад {MAX_POINTS}»"
    assert cd.contract["unloadings"][-1]["name"] == f"ООО «Приёмка {MAX_POINTS}»"


def test_route_without_shippers_has_no_loadings():
    cd = collect_logistiks_rus_data({"route": {
        "consignees": [{"name": "ООО «Приёмка»", "address": "г. Казань"}],
    }})

    assert cd.contract["loadings"] == []
    assert cd.loadings == []
    assert [p["name"] for p in cd.contract["unloadings"]] == ["ООО «Приёмка»"]


def test_route_points_not_a_list_are_ignored():
    cd = collect_logistiks_rus_data({"route": {"shippers": "мусор"}})

    assert cd.contract["loadings"] == []


def test_time_window_with_one_boundary():
    cd = collect_logistiks_rus_data({"route": {
        "shippers": [{"name": "ООО «Склад Север»", "address": "г. Москва"}],
        "loading_time_from": "08:00",
    }})

    assert cd.contract["loadings"][0]["time_window"] == "08:00"
    assert cd.contract["loadings"][0]["date"] == ""


# ─────────────────────────────────────────────────────────────
# Водитель (только ФИО)
# ─────────────────────────────────────────────────────────────

def test_driver_keeps_only_full_name():
    cd = collect_logistiks_rus_data({"driver": {
        "full_name": "  Иванов Иван Иванович  ",
        "passport_series": "18 22",
        "phone": "+7 (999) 123-45-67",
    }})

    assert cd.driver == {"full_name": "Иванов Иван Иванович"}


def test_driver_without_full_name_is_empty():
    cd = collect_logistiks_rus_data({"driver": {"full_name": "  "}})

    assert cd.driver == {}


# ─────────────────────────────────────────────────────────────
# ТС: тягач и прицеп
# ─────────────────────────────────────────────────────────────

def test_tractor_and_trailer_are_collected():
    cd = collect_logistiks_rus_data({"vehicle": {
        "tractor_brand": "DAF XF 95.430",
        "tractor_plate": "М342СА761",
        "trailer_brand": "KRONE SD",
        "trailer_plate": "ВК123478",
    }})

    assert cd.tractor == {"brand_model": "DAF XF 95.430", "plate_number": "М342СА761"}
    assert cd.trailer == {"brand_model": "KRONE SD", "plate_number": "ВК123478"}


def test_empty_tractor_and_trailer_fields_are_skipped():
    cd = collect_logistiks_rus_data({"vehicle": {
        "tractor_brand": "DAF XF 95.430",
        "tractor_plate": "   ",
        "trailer_brand": "",
        "trailer_plate": "",
    }})

    assert cd.tractor == {"brand_model": "DAF XF 95.430"}
    assert cd.trailer == {}


# ─────────────────────────────────────────────────────────────
# Стоимость: ООО
# ─────────────────────────────────────────────────────────────

@pytest.fixture
def ooo_price_tab() -> dict:
    """Вкладка «Стоимость» экспедитора-ООО."""
    return {
        "carrier_type": "ООО",
        "amount_without_vat": PRICE_WITHOUT_VAT,
        "amount_with_vat": PRICE_WITH_VAT,
        "vat_rate": "22%",
        "vat_rate_num": 22,
        "special_conditions": "  Простой не более 24 часов  ",
    }


def test_ooo_price_fields_are_collected(ooo_price_tab):
    cd = collect_logistiks_rus_data({"price": ooo_price_tab})

    assert cd.contract["carrier_type"] == "ООО"
    assert cd.contract["price_without_vat"] == PRICE_WITHOUT_VAT
    assert cd.contract["price_with_vat"] == PRICE_WITH_VAT
    assert cd.contract["vat_rate_num"] == VAT_RATE_NUM
    assert cd.contract["vat_rate"] == "22%"
    assert cd.contract["special_conditions"] == "Простой не более 24 часов"


def test_ooo_vat_rate_is_parsed_from_text():
    """Ставка может прийти только строкой «22%» — число разбирается из неё."""
    cd = collect_logistiks_rus_data({"price": {
        "carrier_type": "ООО",
        "amount_without_vat": PRICE_WITHOUT_VAT,
        "vat_rate": "22%",
    }})

    assert cd.contract["vat_rate_num"] == VAT_RATE_NUM
    assert cd.contract["vat_rate"] == "22%"


def test_ooo_default_carrier_type_and_vat_rate():
    """Пустой тип экспедитора — это ООО (вариант бланка по умолчанию)."""
    cd = collect_logistiks_rus_data({"price": {"amount_without_vat": PRICE_WITHOUT_VAT}})

    assert cd.contract["carrier_type"] == DEFAULT_CARRIER_TYPE
    assert cd.contract["vat_rate_num"] == DEFAULT_VAT_RATE_NUM
    assert cd.contract["vat_rate"] == "22%"


def test_amounts_as_text_are_parsed():
    cd = collect_logistiks_rus_data({"price": {
        "carrier_type": "ООО",
        "amount_without_vat": "221 099,18",
        "amount_with_vat": "269\u00a0741,00",
    }})

    assert cd.contract["price_without_vat"] == PRICE_WITHOUT_VAT
    assert cd.contract["price_with_vat"] == PRICE_WITH_VAT


def test_empty_price_fields_are_skipped():
    cd = collect_logistiks_rus_data({"price": {
        "carrier_type": "ООО",
        "amount_without_vat": None,
        "amount_with_vat": "   ",
        "special_conditions": "",
    }})

    assert "price_without_vat" not in cd.contract
    assert "price_with_vat" not in cd.contract
    assert "special_conditions" not in cd.contract
    # Тип экспедитора и ставка НДС заполняются всегда: по ним выбирается
    # вариант бланка и считается НДС.
    assert cd.contract["carrier_type"] == "ООО"
    assert cd.contract["vat_rate_num"] == DEFAULT_VAT_RATE_NUM


# ─────────────────────────────────────────────────────────────
# Стоимость: ИП
# ─────────────────────────────────────────────────────────────

def test_ip_price_has_zero_vat_rate():
    """У ИП ставка НДС не применяется, даже если в форме стоит «22%»."""
    cd = collect_logistiks_rus_data({"price": {
        "carrier_type": "ИП",
        "amount_without_vat": IP_PRICE,
        "vat_rate": "22%",
        "vat_rate_num": 22,
    }})

    assert cd.contract["carrier_type"] == "ИП"
    assert cd.contract["price_without_vat"] == IP_PRICE
    assert cd.contract["vat_rate_num"] == 0.0
    assert cd.contract["vat_rate"] == "0%"


def test_ip_without_vat_rate_fields_still_gets_zero():
    cd = collect_logistiks_rus_data({"price": {
        "carrier_type": "ИП Хейгетян Е.В.",
        "amount_without_vat": IP_PRICE,
    }})

    assert cd.contract["carrier_type"] == "ИП"
    assert cd.contract["vat_rate_num"] == 0.0
    assert cd.contract["vat_rate"] == "0%"


# ─────────────────────────────────────────────────────────────
# Стоимость: распознанные суммы (sum_*) → price_without_vat
# ─────────────────────────────────────────────────────────────

def test_ooo_recognized_sum_wo_vat_is_mapped():
    """Маппинг 3 (ООО): sum_wo_vat промпта → price_without_vat генератора."""
    cd = collect_logistiks_rus_data({"price": {
        "carrier_type": "ООО",
        "sum_wo_vat": RECOGNIZED_SUM_WO_VAT,
        "sum_vat": RECOGNIZED_SUM_VAT,
        "sum_total": RECOGNIZED_SUM_TOTAL,
        "vat_rate": "22%",
    }})

    assert cd.contract["price_without_vat"] == RECOGNIZED_SUM_WO_VAT
    assert cd.contract["vat_rate_num"] == VAT_RATE_NUM


def test_ooo_falls_back_to_recognized_sum_total():
    """У ООО в документе бывает только итог — тогда сумма берётся из него."""
    cd = collect_logistiks_rus_data({"price": {
        "carrier_type": "ООО",
        "sum_wo_vat": 0.0,
        "sum_total": RECOGNIZED_SUM_TOTAL,
        "vat_rate": "22%",
    }})

    assert cd.contract["price_without_vat"] == RECOGNIZED_SUM_TOTAL


def test_ip_recognized_sum_total_is_mapped():
    """Маппинг 3 (ИП): единственная сумма «Без НДС» лежит в sum_total."""
    cd = collect_logistiks_rus_data({"price": {
        "carrier_type": "ИП",
        "sum_wo_vat": 0.0,
        "sum_vat": 0.0,
        "sum_total": IP_PRICE,
        "vat_rate": "0%",
    }})

    assert cd.contract["price_without_vat"] == IP_PRICE
    assert cd.contract["vat_rate_num"] == 0.0


def test_form_amount_wins_over_recognized_sum():
    """Введённое в форме значение не перебивается распознанным."""
    cd = collect_logistiks_rus_data({"price": {
        "carrier_type": "ООО",
        "amount_without_vat": PRICE_WITHOUT_VAT,
        "sum_total": RECOGNIZED_SUM_TOTAL,
    }})

    assert cd.contract["price_without_vat"] == PRICE_WITHOUT_VAT


def test_cargo_count_is_not_invented():
    """
    cargo_count в сборке не выдумывается: его кладёт только распознавание.

    Число машин бланк считает сам ({{cargo_count}} печатается из списка), а
    выдуманное значение спорило бы с таблицей и давало замечание валидатора.
    """
    cd = collect_logistiks_rus_data({
        "cargo": {"vehicles": [{"brand_model": "JETOUR T2", "vin": VIN_1}]},
    })

    assert "cargo_count" not in cd.contract


# ─────────────────────────────────────────────────────────────
# Полный сценарий: все шесть вкладок вместе
# ─────────────────────────────────────────────────────────────

@pytest.fixture
def full_tabs() -> dict:
    """Все шесть вкладок окна «Логистикс Рус», заполненные как в жизни."""
    return {
        "customer": StubTab({
            "number": "ЛР-2026-17",
            "date": "2026-09-24",
            "name": CUSTOMER_NAME,
        }),
        "cargo": StubTab({"vehicles": [
            {"brand_model": "МОДЕЛЬ 1", "vin": VIN_1},
            {"brand_model": "МОДЕЛЬ 2", "vin": VIN_2},
            {"brand_model": "", "vin": ""},
        ]}),
        "route": StubTab({
            "route": "Москва - Казань",
            "shippers": [
                {"name": "ООО «Склад Север»",
                 "address": "г. Москва, ул. Складская, д. 1"},
                {"name": "ООО «Склад Юг»", "address": "г. Москва, ул. Южная, д. 2"},
            ],
            "consignees": [
                {"name": "ООО «Приёмка»", "address": "г. Казань, ул. Приёмная, д. 3"},
            ],
            "loading_date": "2026-09-26",
            "loading_time_from": "08:00",
            "loading_time_to": "20:00",
            "unloading_date": "2026-10-01",
            "unloading_time_from": "09:00",
            "unloading_time_to": "18:00",
        }),
        "driver": StubTab({"full_name": "Иванов Иван Иванович"}),
        "vehicle": StubTab({
            "tractor_brand": "DAF XF 95.430",
            "tractor_plate": "М342СА761",
            "trailer_brand": "KRONE SD",
            "trailer_plate": "ВК123478",
        }),
        "price": StubTab({
            "carrier_type": "ООО",
            "amount_without_vat": PRICE_WITHOUT_VAT,
            "amount_with_vat": PRICE_WITH_VAT,
            "vat_rate": "22%",
            "vat_rate_num": 22,
            "special_conditions": "Простой не более 24 часов",
        }),
    }


def test_full_scenario_collects_everything(full_tabs):
    cd = collect_logistiks_rus_data(full_tabs)

    assert cd.contract["number"] == "ЛР-2026-17"
    assert cd.contract["date"] == "2026-09-24"
    assert cd.contract["route"] == "Москва - Казань"
    assert cd.contract["loading_date"] == "2026-09-26"
    assert cd.contract["unloading_date"] == "2026-10-01"
    assert cd.contract["price_without_vat"] == PRICE_WITHOUT_VAT
    assert cd.contract["vat_rate_num"] == VAT_RATE_NUM
    assert cd.contract["special_conditions"] == "Простой не более 24 часов"

    assert cd.customer["full_name"] == CUSTOMER_NAME
    assert cd.customer["short_name"] == CUSTOMER_NAME
    assert len(cd.vehicles) == 2
    assert cd.driver == {"full_name": "Иванов Иван Иванович"}
    assert cd.tractor["plate_number"] == "М342СА761"
    assert cd.trailer["plate_number"] == "ВК123478"
    assert len(cd.contract["loadings"]) == 2
    assert len(cd.contract["unloadings"]) == 1
    assert len(cd.loadings) == 2


def test_full_scenario_keeps_carrier_empty(full_tabs):
    """Экспедитор фиксирован шаблоном заявки — своих значений не выдумываем."""
    cd = collect_logistiks_rus_data(full_tabs)

    assert cd.carrier == {}


def test_full_scenario_summary_is_safe_for_logs(full_tabs):
    """Сводка для логов не падает и не содержит самих данных."""
    cd = collect_logistiks_rus_data(full_tabs)
    summary = cd.summary()

    assert "vehicles=2" in summary
    assert "loadings=2" in summary
    assert CUSTOMER_NAME not in summary
    assert "Иванов" not in summary


def test_logs_do_not_contain_personal_data(full_tabs, caplog):
    """
    В INFO сборки нет персональных данных: только имена полей и количества.

    Адреса, ФИО, VIN, названия организаций и марок машин в лог не попадают.
    """
    with caplog.at_level(logging.INFO, logger="ui.windows.logistiks_rus.data"):
        cd = collect_logistiks_rus_data(full_tabs)
        assert cd.contract["number"] == "ЛР-2026-17"

    messages = "\n".join(
        record.getMessage() for record in caplog.records
        if record.name == "ui.windows.logistiks_rus.data"
    )
    assert messages, "сборка ничего не записала в свой лог"

    for fragment in (
        "ДжейСиСиТиЭс",     # заказчик
        "Иванов",           # ФИО водителя
        "Склад Север",      # название грузоотправителя
        "Складская",        # адрес точки
        VIN_1,              # VIN
        "МОДЕЛЬ",           # марка машины
        "DAF",              # тягач
        "KRONE",            # прицеп
    ):
        assert fragment not in messages, f"в логе есть «{fragment}»"

    # Имена полей и количества — можно.
    assert "машин=2" in messages
    assert "грузоотправителей=2" in messages
    assert "грузополучателей=1" in messages


def test_full_scenario_passes_logistiks_rus_validator(full_tabs):
    """Собранных данных достаточно, чтобы валидатор типа не нашёл ошибок."""
    from core.contracts.logistiks_rus.validator import LogistiksRusValidator

    report = LogistiksRusValidator().check(collect_logistiks_rus_data(full_tabs))

    assert report.errors == [], report.errors
    assert report.warnings == [], report.warnings


def test_full_scenario_generator_gets_expected_placeholders(full_tabs, templates_dir):
    """Генератор читает из собранных данных ровно то, что ждёт."""
    from core.contracts.logistiks_rus.generator import LogistiksRusGenerator

    generator = LogistiksRusGenerator(templates_dir=str(templates_dir))
    replacements = generator.build_replacements(collect_logistiks_rus_data(full_tabs))

    assert replacements["contract_number"] == "ЛР-2026-17"
    assert replacements["customer_name"] == CUSTOMER_NAME
    assert replacements["car_1_vin"] == VIN_1
    assert replacements["car_2_vin"] == VIN_2
    assert replacements["car_3_vin"] == ""
    assert replacements["tractor_plate"] == "М342СА761"
    assert replacements["trailer_plate"] == "ВК123478"
    assert replacements["driver_name"] == "Иванов Иван Иванович"
    assert replacements["loading_date"] == "26.09.2026"
    assert replacements["loading_time_from"] == "08:00"
    assert replacements["unloading_date"] == "01.10.2026"
    assert replacements["cargo_count"] == "2"

    # Названия грузоотправителей и грузополучателей доходят до бланка:
    # это и есть маппинг 1 и 2 (contract["loadings"] / ["unloadings"]).
    assert replacements["shipper_1_name"] == "ООО «Склад Север»"
    assert replacements["shipper_2_name"] == "ООО «Склад Юг»"
    assert replacements["shipper_1_address"] == "г. Москва, ул. Складская, д. 1"
    assert replacements["consignee_1_name"] == "ООО «Приёмка»"
    assert replacements["consignee_1_address"] == "г. Казань, ул. Приёмная, д. 3"

    # Стоимость: сумма без НДС и ставка — из формы, НДС и итог считает генератор.
    assert replacements["sum_wo_vat"] == "221\u00a0099,18"
    assert replacements["sum_vat"] == "48\u00a0641,82"
    assert replacements["sum_total"] == "269\u00a0741,00"
    assert replacements["vat_rate"] == "22%"


def test_full_scenario_with_ip_price_builds_single_sum(full_tabs, templates_dir):
    """ИП-вариант: собранных данных хватает на одну сумму «Без НДС»."""
    from core.contracts.logistiks_rus.generator import LogistiksRusGenerator

    tabs = dict(full_tabs)
    tabs["price"] = StubTab({
        "carrier_type": "ИП",
        "amount_without_vat": IP_PRICE,
    })

    cd = collect_logistiks_rus_data(tabs)
    assert cd.contract["vat_rate_num"] == 0.0
    assert cd.contract["vat_rate"] == "0%"

    generator = LogistiksRusGenerator(templates_dir=str(templates_dir))
    replacements = generator.build_replacements(cd)

    assert replacements["sum_total"] == "135\u00a0833,00"
    for absent in ("sum_wo_vat", "sum_vat", "vat_rate"):
        assert absent not in replacements, f"в ИП-вариант положен {absent!r}"


def test_full_scenario_renders_docx_with_point_names(full_tabs, templates_dir, work_dir):
    """
    Собранные данные доходят до ГОТОВОГО документа — вместе с названиями точек.

    Это проверка стыка «сборщик вкладок → генератор → бланк»: если бы точки
    с name не попали в contract["loadings"] / ["unloadings"], в документе
    остались бы одни адреса, а блоки грузоотправителей/грузополучателей
    выглядели бы незаполненными.
    """
    from docx import Document

    from core.contracts.logistiks_rus.generator import LogistiksRusGenerator

    output_dir = work_dir / "ui_logistiks_rus_data_full"
    output_dir.mkdir(parents=True, exist_ok=True)

    generator = LogistiksRusGenerator(templates_dir=str(templates_dir))
    path = Path(generator.generate(collect_logistiks_rus_data(full_tabs),
                                   output_dir=str(output_dir)))
    try:
        assert path.exists(), f"файл не создан: {path}"

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]

        def table_text(table):
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
                    for nested in cell.tables:
                        table_text(nested)

        for table in doc.tables:
            table_text(table)
        text = "\n".join(parts)

        assert PLACEHOLDER_RE.search(text) is None, "в документе остался плейсхолдер"
        assert "ЗАЯВКА № ЛР-2026-17" in text
        assert "Грузоотправитель: ООО «Склад Север»" in text
        assert "Адрес погрузки: г. Москва, ул. Складская, д. 1" in text
        assert "Грузоотправитель: ООО «Склад Юг»" in text
        assert "Грузополучатель №1: ООО «Приёмка»" in text
        assert "Адрес выгрузки: г. Казань, ул. Приёмная, д. 3" in text
        assert "Тягач: DAF XF 95.430 гос. №: М342СА761" in text
        assert "Водитель: Иванов Иван Иванович" in text
        assert "Итого: 269\u00a0741,00 руб." in text
    finally:
        path.unlink(missing_ok=True)


# ─────────────────────────────────────────────────────────────
# Точка входа build()
# ─────────────────────────────────────────────────────────────

def test_build_uses_same_layout_as_collect(full_tabs):
    """build() — та же сборка, только вкладки переданы по именам."""
    cd = build(
        full_tabs["customer"],
        full_tabs["cargo"],
        full_tabs["route"],
        full_tabs["driver"],
        full_tabs["vehicle"],
        full_tabs["price"],
    )

    assert cd.contract["number"] == "ЛР-2026-17"
    assert cd.contract["price_without_vat"] == PRICE_WITHOUT_VAT
    assert cd.customer["full_name"] == CUSTOMER_NAME
    assert cd.driver == {"full_name": "Иванов Иван Иванович"}
    assert [p["name"] for p in cd.contract["loadings"]] == [
        "ООО «Склад Север»", "ООО «Склад Юг»",
    ]
    assert len(cd.vehicles) == 2


def test_build_accepts_none_tabs():
    """Вкладок может ещё не быть — сборка не падает."""
    cd = build(None, None, None, None, None, None)

    assert isinstance(cd, ContractData)
    assert cd.contract == {"loadings": [], "unloadings": []}
    assert cd.driver == {}
    assert cd.vehicles == []
