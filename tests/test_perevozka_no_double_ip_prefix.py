#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Двойная приставка «Индивидуальный предприниматель» (ШАГ «Фикс сохранения +
динамические стороны в шаблонах», часть B.3).

Что было
--------
В справочнике ИП записан полным наименованием: «Индивидуальный предприниматель
Добросоцкий А.Н.». Бланк печатает вид стороны отдельным плейсхолдером
(`{{*_legal_form}}`), а в ФИО приставка оставалась — в договоре выходило
«Индивидуальный предприниматель Индивидуальный предприниматель …».

Что стало
---------
Для ИП из full_name снимается приставка («Индивидуальный предприниматель »
или «ИП ») — и у перевозчика, и у заказчика. У ООО наименование не трогается
вовсе.

Данные синтетические, ПДн нет.
"""

import re

import pytest
from docx import Document

from core.contract_generator import ContractGenerator

IP_PREFIX = "Индивидуальный предприниматель"
DOUBLE = f"{IP_PREFIX} {IP_PREFIX}"

OOO_CARRIER = {
    "full_name": "ООО «Перевозчик»", "short_name": "ООО «Перевозчик»",
    "inn": "7701234567", "kpp": "770101001", "ogrn": "1027700132195",
    "legal_address": "г. Москва", "actual_address": "",
    "bank_account": "40702810000000000002", "bik": "044525225",
    "correspondent_account": "30101810400000000225", "bank_name": "ПАО Сбербанк",
    "director_name": "Сидоров Сидор Сидорович",
    "director_position": "Генеральный директор",
    "phone": "", "email": "", "entity_type": "ООО",
}

OOO_CUSTOMER = {
    "full_name": "Общество с ограниченной ответственностью «Заказчик»",
    "short_name": "ООО «Заказчик»",
    "inn": "7707654321", "kpp": "770701001", "ogrn": "1027700132195",
    "legal_address": "г. Москва", "actual_address": "",
    "bank_account": "40702810000000000001", "bik": "044525225",
    "correspondent_account": "30101810400000000225", "bank_name": "ПАО Сбербанк",
    "director_name": "Петров Пётр Петрович",
    "director_position": "Генеральный директор",
    "phone": "", "email": "",
}


@pytest.fixture
def generator(templates_dir) -> ContractGenerator:
    return ContractGenerator(templates_dir=str(templates_dir))


def _ip_organization(full_name: str) -> dict:
    """ИП с наименованием из справочника (с приставкой или без)."""
    return {
        "full_name": full_name, "short_name": "", "inn": "770123456789",
        "kpp": "", "ogrn": "315770000000012", "legal_address": "г. Волгоград",
        "actual_address": "", "bank_account": "", "bik": "",
        "correspondent_account": "", "bank_name": "",
        "director_name": "", "director_position": "",
        "phone": "", "email": "", "entity_type": "ИП",
    }


def _payload(carrier, customer, carrier_type) -> dict:
    return {
        "driver": {"full_name": "Иванов Иван Иванович"},
        "carrier": carrier, "customer": customer,
        "vehicles": [{"vin": "EC3TEUMB0T0002608", "brand_model": "JETOUR T2",
                      "vehicle_type": "Легковой автомобиль"}],
        "tractor": {"brand_model": "Foton", "plate_number": "O844XY196"},
        "trailer": {"brand_model": "YANGMINDA", "plate_number": "71ABF18"},
        "contract": {
            "number": "PREFIX-1", "date": "2026-10-09",
            "carrier_type": carrier_type,
            "vat_rate": "22%" if carrier_type.startswith("ООО") else "0%",
            "vat_rate_num": 22 if carrier_type.startswith("ООО") else 0,
            "price_without_vat": 180300.0, "payment_days": 10,
        },
        "loadings": [{"address": "г. Воронеж", "date": "2026-10-10",
                      "time_window": ""}],
        "unloadings": [{"address": "г. Москва", "date": "2026-10-13",
                        "time_window": ""}],
    }


def _render(generator, carrier, customer, carrier_type, output) -> Document:
    generator.generate_docx(
        _payload(carrier, customer, carrier_type), str(output)
    )
    assert output.exists(), f"файл не создан: {output}"
    return Document(str(output))


def _flatten(text: str) -> str:
    """Текст одной строкой: пробелы (в том числе неразрывные) по одному."""
    return re.sub(r"\s+", " ", text.replace("\u00a0", " ")).strip()


def _document_text(doc) -> str:
    parts = [_flatten(p.text) for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(_flatten(cell.text))
    return "\n".join(parts)


def _clause_text(doc, label: str) -> str:
    """Абзац, идущий сразу за строкой-меткой («1.2. Перевозчик:»)."""
    texts = [p.text for p in doc.paragraphs]
    for index, text in enumerate(texts):
        if text.strip().startswith(label):
            for following in texts[index + 1:]:
                if following.strip():
                    return _flatten(following)
    raise AssertionError(f"абзац {label!r} не найден")


# ─────────────────────────────────────────────────────────────
# Перевозчик-ИП
# ─────────────────────────────────────────────────────────────

def test_ip_carrier_full_name_cleaned(generator, work_file):
    """
    ФИО ИП-перевозчика печатается ОДИН раз: приставку даёт legal_form.

    В справочнике запись лежит как «Индивидуальный предприниматель Смирнов
    Алексей Николаевич» — в договоре это же и должно быть, а не удвоение.
    """
    carrier = _ip_organization(
        f"{IP_PREFIX} Смирнов Алексей Николаевич"
    )
    doc = _render(generator, carrier, OOO_CUSTOMER, "ИП без НДС",
                  work_file("prefix_carrier.docx"))

    text = _document_text(doc)
    assert DOUBLE not in text
    clause = _clause_text(doc, "1.2. Перевозчик:")
    assert clause.startswith(
        "Индивидуальный предприниматель Смирнов Алексей Николаевич, ИНН"
    )


def test_short_ip_prefix_cleaned(generator, work_file):
    """Короткая приставка «ИП » снимается так же, как полная."""
    carrier = _ip_organization("ИП Смирнов Алексей Николаевич")
    doc = _render(generator, carrier, OOO_CUSTOMER, "ИП без НДС",
                  work_file("prefix_carrier_short.docx"))

    clause = _clause_text(doc, "1.2. Перевозчик:")
    assert clause.startswith(
        "Индивидуальный предприниматель Смирнов Алексей Николаевич, ИНН"
    )
    assert "ИП ИП" not in _document_text(doc)


def test_ip_carrier_name_is_not_doubled_in_requisites(generator, work_file):
    """П. 9 печатает наименование перевозчика тоже один раз."""
    carrier = _ip_organization(
        f"{IP_PREFIX} Смирнов Алексей Николаевич"
    )
    doc = _render(generator, carrier, OOO_CUSTOMER, "ИП без НДС",
                  work_file("prefix_carrier_req.docx"))

    text = _document_text(doc)
    assert DOUBLE not in text
    assert "Индивидуальный предприниматель Смирнов Алексей Николаевич" in text


# ─────────────────────────────────────────────────────────────
# Заказчик-ИП и ООО
# ─────────────────────────────────────────────────────────────

def test_ip_client_full_name_cleaned(generator, work_file):
    """У заказчика-ИП приставка тоже снимается — правило одно для сторон."""
    customer = _ip_organization(
        f"{IP_PREFIX} Смирнова Елена Владимировна"
    )
    doc = _render(generator, OOO_CARRIER, customer, "ООО (с НДС)",
                  work_file("prefix_client.docx"))

    text = _document_text(doc)
    assert DOUBLE not in text
    clause = _clause_text(doc, "1.1. Заказчик:")
    assert clause.startswith(
        "Индивидуальный предприниматель Смирнова Елена Владимировна, ИНН"
    )


def test_ooo_full_name_not_touched(generator, work_file):
    """У ООО наименование не чистится: приставки ИП в нём и быть не может."""
    customer = dict(
        OOO_CUSTOMER,
        full_name="Общество с ограниченной ответственностью «Ромашка»",
        short_name="ООО «Ромашка»",
    )
    doc = _render(generator, OOO_CARRIER, customer, "ООО (с НДС)",
                  work_file("prefix_ooo.docx"))

    clause = _clause_text(doc, "1.1. Заказчик:")
    assert clause.startswith(
        "Общество с ограниченной ответственностью «Ромашка» (ООО «Ромашка»), "
        "ИНН 7707654321"
    )
    assert "Индивидуальный предприниматель" not in clause
