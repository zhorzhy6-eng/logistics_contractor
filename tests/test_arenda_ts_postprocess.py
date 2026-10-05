#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты шагов постобработки договора аренды ТС с экипажем (ЭТАП 3.1.D.A.4).

Каждый шаг проверяется ОТДЕЛЬНО на программно собранном Document() — без
шаблонов templates/*.docx и без полного цикла генерации:

  * RemoveEmptyVehicleRowsStep — таблица машин п. 3.1 рассчитана на 12 строк:
    строка удаляется, только если пусты ВСЕ четыре колонки данных (марка,
    VIN и обе точки маршрута), строка с любым заполненным значением
    остаётся, чужие таблицы (другая шапка, шапки нет) не трогаются;
  * RemoveEmptyLoadingUnloadingBlocksStep — по 10 строк точек погрузки и
    выгрузки: строка точки удаляется по ПУСТОМУ АДРЕСУ (дата и время не
    важны), строка с адресом остаётся; если у раздела не осталось ни одной
    точки, удаляется и его заголовок.

Интеграционные проверки (тот же результат в готовом документе после
generate()) остаются в tests/test_arenda_ts_generator.py.

Все данные синтетические, реальных ПДн нет.
"""

import gc
import re

import pytest
from docx import Document

from core.contracts.arenda_ts.generator import (
    MAX_CARS,
    MAX_POINTS,
    ArendaTsGenerator,
)
from core.contracts.arenda_ts.postprocess import (
    RemoveEmptyLoadingUnloadingBlocksStep,
    RemoveEmptyVehicleRowsStep,
)

#: Заголовки таблицы машин п. 3.1 — как в бланке («Марка, модель» с запятой).
CAR_HEADERS = ("№", "Марка, модель", "VIN-номер", "Точка погрузки",
               "Точка выгрузки")

#: Вариант шапки со слэшем: генератор принимает его наравне с запятой
#: (CAR_TABLE_BRAND_HEADERS — «как в Формике»), поэтому такая таблица СВОЯ.
SLASH_CAR_HEADERS = ("№", "Марка/Модель", "VIN-номер", "Точка погрузки",
                     "Точка выгрузки")

#: Заголовки разделов точек маршрута.
SECTION_TITLES = {
    "loading": "3.2. Согласованные точки погрузки:",
    "unloading": "3.3. Согласованные точки выгрузки:",
}

#: Нумерованные строки точек: «3.2.N. …» (погрузка), «3.3.N. …» (выгрузка).
POINT_LINE_RES = {
    "loading": re.compile(r"^3\.2\.\d+\."),
    "unloading": re.compile(r"^3\.3\.\d+\."),
}

#: Абзац-«хвост» после разделов точек: проверяем, что он не задет.
ROUTE_TAIL = "3.4. Согласованный маршрут: Москва — Калуга."


@pytest.fixture(autouse=True)
def _release_documents():
    """Word-файлы на Windows освобождаем до удаления временных копий."""
    yield
    gc.collect()


@pytest.fixture
def generator(tmp_path) -> ArendaTsGenerator:
    """Генератор без реальных шаблонов: шагам они не нужны."""
    return ArendaTsGenerator(templates_dir=str(tmp_path))


# ─────────────────────────────────────────────────────────────
# Сборка документа: таблица машин п. 3.1
# ─────────────────────────────────────────────────────────────

def _brand(number: int) -> str:
    """Синтетическая марка модели для строки таблицы."""
    return f"МОДЕЛЬ {number}"


def _vin(number: int) -> str:
    """Синтетический VIN для строки таблицы (не настоящий)."""
    return f"TESTVIN{number:08d}"


def _car_table_doc(filled: int, rows: int = MAX_CARS, headers=CAR_HEADERS):
    """
    Документ с таблицей машин бланка: шапка + `rows` строк данных.

    Первые `filled` строк заполнены (как после docxtpl), у остальных
    заполнен только номер — ровно так выглядит бланк до постобработки.
    """
    doc = Document()
    table = doc.add_table(rows=rows + 1, cols=5)
    for cell, title in zip(table.rows[0].cells, headers):
        cell.text = title

    for number in range(1, rows + 1):
        cells = table.rows[number].cells
        cells[0].text = str(number)
        if number <= filled:
            cells[1].text = _brand(number)
            cells[2].text = _vin(number)
            cells[3].text = f"Точка погрузки {number}"
            cells[4].text = f"Точка выгрузки {number}"

    return doc


def _row_values(table):
    """Значения всех строк таблицы — списком списков."""
    return [[cell.text for cell in row.cells] for row in table.rows]


# ─────────────────────────────────────────────────────────────
# RemoveEmptyVehicleRowsStep
# ─────────────────────────────────────────────────────────────

def test_remove_empty_vehicle_rows_step_removes_all_rows_without_cars(generator):
    """0 машин: все 12 строк данных удалены, шапка остаётся на месте."""
    doc = _car_table_doc(filled=0)
    assert len(doc.tables[0].rows) == MAX_CARS + 1

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    table = doc.tables[0]
    assert len(table.rows) == 1
    assert [cell.text for cell in table.rows[0].cells] == list(CAR_HEADERS)


@pytest.mark.parametrize("cars", [1, 2, 3, MAX_CARS])
def test_remove_empty_vehicle_rows_step_keeps_one_row_per_car(generator, cars):
    """1, 2, 3 и 12 машин: сколько машин, столько строк и остаётся."""
    doc = _car_table_doc(filled=cars)

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    table = doc.tables[0]
    assert len(table.rows) == cars + 1
    for number in range(1, cars + 1):
        assert _row_values(table)[number] == [
            str(number), _brand(number), _vin(number),
            f"Точка погрузки {number}", f"Точка выгрузки {number}",
        ]


def test_remove_empty_vehicle_rows_step_keeps_row_with_only_vin(generator):
    """Марка пуста, VIN заполнен — строка сохраняется: данные не теряем."""
    doc = _car_table_doc(filled=0, rows=1)
    doc.tables[0].rows[1].cells[2].text = _vin(1)

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 2
    assert _row_values(doc.tables[0])[1][2] == _vin(1)


def test_remove_empty_vehicle_rows_step_keeps_row_with_only_brand(generator):
    """Марка заполнена, VIN пуст — строка тоже сохраняется."""
    doc = _car_table_doc(filled=0, rows=1)
    doc.tables[0].rows[1].cells[1].text = _brand(1)

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 2
    assert _row_values(doc.tables[0])[1][1] == _brand(1)


def test_remove_empty_vehicle_rows_step_keeps_row_with_only_loading_point(generator):
    """Заполнена только точка погрузки (ни марки, ни VIN) — строка остаётся."""
    doc = _car_table_doc(filled=0, rows=1)
    doc.tables[0].rows[1].cells[3].text = "Точка погрузки 1"

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 2
    assert _row_values(doc.tables[0])[1][3] == "Точка погрузки 1"


def test_remove_empty_vehicle_rows_step_keeps_row_with_only_unloading_point(generator):
    """Симметрично: одна точка выгрузки — тоже данные, строку не удаляем."""
    doc = _car_table_doc(filled=0, rows=1)
    doc.tables[0].rows[1].cells[4].text = "Точка выгрузки 1"

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 2
    assert _row_values(doc.tables[0])[1][4] == "Точка выгрузки 1"


def test_remove_empty_vehicle_rows_step_treats_whitespace_as_empty(generator):
    """Пробелы вместо значений — та же пустая строка."""
    doc = _car_table_doc(filled=2)
    table = doc.tables[0]
    table.rows[1].cells[1].text = "   "
    table.rows[1].cells[2].text = "  "
    table.rows[1].cells[3].text = " "
    table.rows[1].cells[4].text = "\t"

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(table.rows) == 2
    assert _row_values(table)[1] == [
        "2", _brand(2), _vin(2), "Точка погрузки 2", "Точка выгрузки 2",
    ]


def test_remove_empty_vehicle_rows_step_ignores_foreign_tables(generator):
    """Чужие таблицы (подписи, реквизиты) не трогаются: шапка не та."""
    doc = Document()
    signatures = doc.add_table(rows=2, cols=2)
    signatures.rows[0].cells[0].text = "АРЕНДАТОР:"
    signatures.rows[0].cells[1].text = "АРЕНДОДАТЕЛЬ:"

    requisites = doc.add_table(rows=2, cols=3)
    for cell, title in zip(requisites.rows[0].cells, ("ИНН", "КПП", "ОГРН")):
        cell.text = title

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 2
    assert len(doc.tables[1].rows) == 2


def test_remove_empty_vehicle_rows_step_ignores_table_without_header(generator):
    """Таблица без шапки бланка не наша — даже если строки пустые."""
    doc = Document()
    table = doc.add_table(rows=3, cols=5)
    for cell, value in zip(table.rows[0].cells, ("1", _brand(1), _vin(1), "", "")):
        cell.text = value

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 3


def test_remove_empty_vehicle_rows_step_accepts_slash_brand_header(generator):
    """
    «Марка/Модель» (слэш) — для ЭТОГО типа шапка своя, а не чужая.

    Генератор аренды принимает оба написания марки наравне
    (CAR_TABLE_BRAND_HEADERS, «как в Формике»), поэтому таблица со слэшем
    обрабатывается: пустые строки данных из неё удаляются.
    """
    doc = _car_table_doc(filled=1, headers=SLASH_CAR_HEADERS)

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 2
    assert _row_values(doc.tables[0])[0] == list(SLASH_CAR_HEADERS)


def test_remove_empty_vehicle_rows_step_keeps_header_only_table(generator):
    """Таблица из одной шапки: удалять нечего, шапка остаётся."""
    doc = _car_table_doc(filled=0, rows=0)

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 1


# ─────────────────────────────────────────────────────────────
# RemoveEmptyLoadingUnloadingBlocksStep
# ─────────────────────────────────────────────────────────────

def _loading_line(number: int, address: str = "", date: str = "21.09.2026 г.") -> str:
    """Строка точки погрузки бланка: «3.2.N. Точка погрузки № N — <адрес>. …»."""
    return (f"3.2.{number}. Точка погрузки № {number} — {address}. Плановая дата "
            f"и время подачи ТС: {date}, с 08:00 до 18:00.")


def _unloading_line(number: int, address: str = "", date: str = "27.09.2026 г.") -> str:
    """Строка точки выгрузки бланка: «3.3.N. Точка выгрузки № N — <адрес>. …»."""
    return (f"3.3.{number}. Точка выгрузки № {number} — {address}. Плановая дата "
            f"завершения: {date}.")


def _point_line(kind: str, number: int, address: str = "", date: str = None) -> str:
    """
    Строка точки нужного вида с подставленными адресом и датой.

    date=None — дата по умолчанию для этого вида (заполненная точка бланка);
    date="" — пустая дата (в бланке плейсхолдер не заполнен).
    """
    line = _loading_line if kind == "loading" else _unloading_line
    if date is None:
        return line(number, address=address)
    return line(number, address=address, date=date)


def _points_doc(kind: str, filled: int, total: int = MAX_POINTS) -> Document:
    """
    Документ с `total` строками точек одного вида, заполнены первые `filled`.

    Незаполненная строка — «3.2.N. Точка погрузки № N — . Плановая дата…»
    (ни адреса, ни даты): ровно так выглядит бланк после docxtpl, когда
    точек в договоре меньше, чем строк в бланке.
    """
    doc = Document()
    doc.add_paragraph(SECTION_TITLES[kind])

    for number in range(1, total + 1):
        if number <= filled:
            doc.add_paragraph(
                _point_line(kind, number, address=f"Адрес точки {number}")
            )
        else:
            doc.add_paragraph(_point_line(kind, number, date=""))

    doc.add_paragraph(ROUTE_TAIL)
    return doc


def _texts(doc):
    """Тексты абзацев верхнего уровня без крайних пробелов."""
    return [p.text.strip() for p in doc.paragraphs]


def _count_point_lines(doc, kind: str) -> int:
    """Сколько нумерованных строк точек указанного вида осталось."""
    return sum(1 for text in _texts(doc) if POINT_LINE_RES[kind].match(text))


def _count_section_titles(doc, kind: str) -> int:
    """Сколько заголовков раздела указанного вида осталось (0 или 1)."""
    return sum(1 for text in _texts(doc) if text == SECTION_TITLES[kind])


@pytest.mark.parametrize("kind", ["loading", "unloading"])
@pytest.mark.parametrize("filled", [1, 2, 5, MAX_POINTS])
def test_remove_empty_point_blocks_keeps_only_filled_points(generator, kind, filled):
    """Сколько точек заполнено, столько строк и остаётся (обоих видов)."""
    doc = _points_doc(kind, filled=filled)

    RemoveEmptyLoadingUnloadingBlocksStep(generator).apply(doc, None)

    assert _count_point_lines(doc, kind) == filled
    # Заголовок раздела и хвост документа остаются на месте.
    assert _count_section_titles(doc, kind) == 1
    assert _texts(doc)[0] == SECTION_TITLES[kind]
    assert _texts(doc)[-1] == ROUTE_TAIL
    assert len(doc.paragraphs) == filled + 2


@pytest.mark.parametrize("kind", ["loading", "unloading"])
def test_remove_empty_point_blocks_removes_section_when_no_points(generator, kind):
    """0 заполненных точек: все 10 строк удалены вместе с заголовком раздела."""
    doc = _points_doc(kind, filled=0)

    RemoveEmptyLoadingUnloadingBlocksStep(generator).apply(doc, None)

    assert _count_point_lines(doc, kind) == 0
    assert _count_section_titles(doc, kind) == 0
    assert _texts(doc) == [ROUTE_TAIL]


def test_remove_empty_point_blocks_removes_both_sections_when_empty(generator):
    """0 погрузок и 0 выгрузок: оба раздела (с заголовками) удалены целиком."""
    doc = Document()
    for kind in ("loading", "unloading"):
        doc.add_paragraph(SECTION_TITLES[kind])
        for number in range(1, MAX_POINTS + 1):
            doc.add_paragraph(_point_line(kind, number))
    doc.add_paragraph(ROUTE_TAIL)

    RemoveEmptyLoadingUnloadingBlocksStep(generator).apply(doc, None)

    assert _texts(doc) == [ROUTE_TAIL]
    assert len(doc.paragraphs) == 1


def test_remove_empty_point_blocks_keeps_section_title_for_one_point(generator):
    """1 заполненная точка из 10: 9 строк удалены, заголовок раздела остаётся."""
    doc = _points_doc("loading", filled=1)

    RemoveEmptyLoadingUnloadingBlocksStep(generator).apply(doc, None)

    assert _texts(doc) == [
        SECTION_TITLES["loading"],
        _point_line("loading", 1, address="Адрес точки 1"),
        ROUTE_TAIL,
    ]
    assert len(doc.paragraphs) == 3


def test_remove_empty_point_blocks_removes_nothing_when_all_filled(generator):
    """Все 10 точек заполнены: документ не меняется."""
    doc = _points_doc("loading", filled=MAX_POINTS)
    before = _texts(doc)

    RemoveEmptyLoadingUnloadingBlocksStep(generator).apply(doc, None)

    assert _texts(doc) == before
    assert len(doc.paragraphs) == MAX_POINTS + 2


@pytest.mark.parametrize("kind", ["loading", "unloading"])
def test_remove_empty_point_blocks_keeps_point_with_address_without_date(
    generator, kind
):
    """Точка с адресом, но без даты и времени — это данные: строку не удаляем."""
    doc = Document()
    doc.add_paragraph(SECTION_TITLES[kind])
    doc.add_paragraph(_point_line(kind, 1, address="Адрес точки 1", date=""))
    doc.add_paragraph(ROUTE_TAIL)

    RemoveEmptyLoadingUnloadingBlocksStep(generator).apply(doc, None)

    assert _count_point_lines(doc, kind) == 1
    assert _count_section_titles(doc, kind) == 1
    assert len(doc.paragraphs) == 3


def test_remove_empty_point_blocks_removes_point_with_date_without_address(generator):
    """
    Точка с датой, но БЕЗ адреса — незаполненная: её удаляем.

    Признак незаполненной точки в этом типе — пустой адрес (в бланке точки
    нет, а плейсхолдер даты мог остаться заполненным). Дата сама по себе
    точку маршрута не описывает, поэтому строка уходит.
    """
    doc = Document()
    doc.add_paragraph(SECTION_TITLES["loading"])
    doc.add_paragraph(_loading_line(1, address="", date="21.09.2026 г."))
    doc.add_paragraph(ROUTE_TAIL)

    RemoveEmptyLoadingUnloadingBlocksStep(generator).apply(doc, None)

    assert _count_point_lines(doc, "loading") == 0
    assert _count_section_titles(doc, "loading") == 0
    assert _texts(doc) == [ROUTE_TAIL]


def test_remove_empty_point_blocks_removes_only_empty_section(generator):
    """0 погрузок, но есть выгрузки: уходит только раздел погрузки."""
    doc = Document()
    doc.add_paragraph(SECTION_TITLES["loading"])
    for number in range(1, MAX_POINTS + 1):
        doc.add_paragraph(_loading_line(number))
    doc.add_paragraph(SECTION_TITLES["unloading"])
    for number in range(1, 4):
        doc.add_paragraph(_unloading_line(number, address=f"Адрес точки {number}"))
    doc.add_paragraph(ROUTE_TAIL)

    RemoveEmptyLoadingUnloadingBlocksStep(generator).apply(doc, None)

    assert _count_point_lines(doc, "loading") == 0
    assert _count_section_titles(doc, "loading") == 0
    # Раздел выгрузок не тронут: заголовок и все три точки на месте.
    assert _count_section_titles(doc, "unloading") == 1
    assert _count_point_lines(doc, "unloading") == 3
    assert _texts(doc)[0] == SECTION_TITLES["unloading"]
    assert _texts(doc)[-1] == ROUTE_TAIL


def test_remove_empty_point_blocks_ignores_foreign_paragraphs(generator):
    """Чужие абзацы (маршрут, экипаж, реквизиты) не удаляются."""
    doc = Document()
    doc.add_paragraph(ROUTE_TAIL)
    doc.add_paragraph("3.5. Член экипажа Арендодателя (водитель):")
    doc.add_paragraph("Плановая дата и время подачи ТС: 21.09.2026 г.")
    doc.add_paragraph("Общее количество: 3 шт.")

    RemoveEmptyLoadingUnloadingBlocksStep(generator).apply(doc, None)

    assert len(doc.paragraphs) == 4


# ─────────────────────────────────────────────────────────────
# Изоляция шагов и идемпотентность
# ─────────────────────────────────────────────────────────────

def _mixed_doc():
    """Документ сразу с таблицей машин и с незаполненными точками маршрута."""
    doc = _car_table_doc(filled=1)
    doc.add_paragraph(SECTION_TITLES["loading"])
    doc.add_paragraph(_loading_line(1, address="Адрес точки 1"))
    doc.add_paragraph(_loading_line(2))
    return doc


def test_steps_touch_only_their_own_part_of_document(generator):
    """Шаги изолированы: строки таблицы и точки маршрута друг друга не трогают."""
    assert (RemoveEmptyVehicleRowsStep.name
            != RemoveEmptyLoadingUnloadingBlocksStep.name)

    doc = _mixed_doc()
    points_before = _texts(doc)

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)
    assert len(doc.tables[0].rows) == 2  # пустые строки машин убраны
    assert _texts(doc) == points_before  # точки маршрута не тронуты

    RemoveEmptyLoadingUnloadingBlocksStep(generator).apply(doc, None)
    assert _texts(doc) == [
        SECTION_TITLES["loading"],
        _loading_line(1, address="Адрес точки 1"),
    ]
    assert len(doc.tables[0].rows) == 2  # таблица не тронута вторым шагом


def test_steps_on_empty_document_are_noops(generator):
    """Пустой документ: оба шага ничего не делают и не падают."""
    doc = Document()

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)
    RemoveEmptyLoadingUnloadingBlocksStep(generator).apply(doc, None)

    assert not doc.tables
    assert not doc.paragraphs


def test_vehicle_rows_step_is_idempotent(generator):
    """Повторный вызов шага на обработанном документе ничего не меняет."""
    doc = _car_table_doc(filled=3)
    step = RemoveEmptyVehicleRowsStep(generator)

    step.apply(doc, None)
    after_first = _row_values(doc.tables[0])
    step.apply(doc, None)

    assert _row_values(doc.tables[0]) == after_first


def test_point_blocks_step_is_idempotent(generator):
    """Повторный вызов шага точек: удалять уже нечего, документ тот же."""
    doc = _points_doc("loading", filled=1)
    step = RemoveEmptyLoadingUnloadingBlocksStep(generator)

    step.apply(doc, None)
    after_first = _texts(doc)
    step.apply(doc, None)

    assert _texts(doc) == after_first


def test_point_blocks_step_is_idempotent_after_section_removed(generator):
    """Раздел удалён целиком — повторный вызов не ищет удалённый заголовок."""
    doc = _points_doc("loading", filled=0)
    step = RemoveEmptyLoadingUnloadingBlocksStep(generator)

    step.apply(doc, None)
    step.apply(doc, None)

    assert _texts(doc) == [ROUTE_TAIL]


def test_steps_are_idempotent_together(generator):
    """Полный конвейер на смешанном документе: второй проход — no-op."""
    doc = _mixed_doc()

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)
    RemoveEmptyLoadingUnloadingBlocksStep(generator).apply(doc, None)
    rows_after_first = _row_values(doc.tables[0])
    texts_after_first = _texts(doc)

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)
    RemoveEmptyLoadingUnloadingBlocksStep(generator).apply(doc, None)

    assert _row_values(doc.tables[0]) == rows_after_first
    assert _texts(doc) == texts_after_first
