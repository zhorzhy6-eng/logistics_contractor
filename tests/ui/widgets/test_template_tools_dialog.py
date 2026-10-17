#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты окна «Шаблоны» (ui/template_tools_dialog.py).

Окно — единственный путь к безопасной правке шаблона из интерфейса:
список бланков, анализ выбранного, предложение правки через диалог
«было / стало». Проверяется, что окно:

  * показывает список шаблонов и анализ выбранного;
  * включает «Предложить правку» только когда есть устаревший текст;
  * само в файлы не пишет (применение — в диалоге подтверждения);
  * не показывает панель без кнопок (обе кнопки на месте).

Qt поднимается в режиме offscreen (tests/ui/conftest.py).
"""

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest

from services.agent.tools import template_editor
from ui.template_tools_dialog import (
    TemplateToolsDialog,
    analysis_text,
)


def make_docx(path, paragraphs):
    """Создаёт DOCX с абзацами."""
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))
    return path


@pytest.fixture
def templates(tmp_path):
    """Два шаблона: с устаревшей ставкой и без."""
    outdated = make_docx(tmp_path / "old.docx", [
        "Стоимость с НДС 20%.",
        "{{contract_number}}",
    ])
    clean = make_docx(tmp_path / "new.docx", [
        "Стоимость с НДС 22%.",
        "{{contract_number}}",
    ])
    return [outdated, clean]


# ─────────────────────────────────────────────────────────────
# 1. Текст анализа
# ─────────────────────────────────────────────────────────────

def test_analysis_text_reports_counts():
    """В тексте анализа — файл, плейсхолдеры, абзацы и предупреждения."""
    result = {
        "path": "templates/shablon.docx",
        "format": "docx",
        "placeholders": [{"name": "contract_number", "count": 2}],
        "placeholder_count": 1,
        "paragraphs_count": 10,
        "tables_count": 1,
        "styles_count": 2,
        "conditional_blocks": [],
        "unclosed_tags": [],
        "outdated": [{"kind": "vat_20", "message": "Устаревшая ставка НДС 20%"}],
        "warnings": ["проверьте шаблон"],
    }
    text = analysis_text(result)
    assert "shablon.docx" in text
    assert "Плейсхолдеров: 1" in text
    assert "{{contract_number}}" in text
    assert "Устаревший текст" in text
    assert "проверьте шаблон" in text


def test_analysis_text_without_outdated():
    """Чистый шаблон: блока устаревшего текста в анализе нет."""
    text = analysis_text({"path": "x.docx", "format": "docx",
                          "placeholders": [], "placeholder_count": 0})
    assert "Устаревший текст" not in text


# ─────────────────────────────────────────────────────────────
# 2. Окно
# ────────────────────────────────────────────────────────────────────

def test_dialog_lists_templates(qt_app, templates):
    """Список содержит оба шаблона, выбран первый."""
    dialog = TemplateToolsDialog(templates=templates)
    try:
        assert dialog.list.count() == 2
        assert dialog.current_template() == templates[0]
    finally:
        dialog.close()
        dialog.deleteLater()


def test_dialog_shows_analysis_of_selected(qt_app, templates):
    """При выборе шаблона показывается его анализ."""
    dialog = TemplateToolsDialog(templates=templates)
    try:
        assert "Плейсхолдеров: 1" in dialog.details.toPlainText()
        dialog.list.setCurrentRow(1)
        assert "НДС 20%" not in dialog.details.toPlainText()
    finally:
        dialog.close()
        dialog.deleteLater()


def test_propose_button_follows_outdated_text(qt_app, templates):
    """Кнопка правки включена только у шаблона с устаревшим текстом."""
    dialog = TemplateToolsDialog(templates=templates)
    try:
        assert dialog.fix_issues() == ["НДС 20%"]
        assert dialog.propose_button.isEnabled() is True
        dialog.list.setCurrentRow(1)
        assert dialog.fix_issues() == []
        assert dialog.propose_button.isEnabled() is False
    finally:
        dialog.close()
        dialog.deleteLater()


def test_propose_fix_builds_plan_without_writing(qt_app, templates, monkeypatch):
    """Предложение строится, но файл не меняется (диалог не открываем)."""
    before = template_editor.propose_template_fix(templates[0], "НДС 20%")
    assert before["ok"] is True

    opened = []

    def fake_open_diff(plan):
        opened.append(plan)

    dialog = TemplateToolsDialog(templates=templates)
    monkeypatch.setattr(dialog, "_open_diff", fake_open_diff)
    try:
        plan = dialog.propose_fix()
        assert plan is not None and plan["ok"] is True
        assert opened and opened[0] is plan
        assert len(opened[0]["changes"]) == 1
    finally:
        dialog.close()
        dialog.deleteLater()


def test_propose_fix_does_nothing_without_issues(qt_app, templates):
    """У чистого шаблона предлагать нечего: план не строится."""
    dialog = TemplateToolsDialog(templates=[templates[1]])
    try:
        assert dialog.propose_fix() is None
        assert "нечего" in dialog.status.text()
    finally:
        dialog.close()
        dialog.deleteLater()


# ─────────────────────────────────────────────────────────────
# 3. Кнопки панели Максимилиана
# ────────────────────────────────────────────────────────────────────

def test_panel_has_legal_and_templates_buttons(qt_app):
    """В шапке панели есть «⚖ Юридическая проверка» и «📄 Шаблоны»."""
    from ui.agent_widget import AgentPanel

    panel = AgentPanel()
    try:
        assert "Юридическая проверка" in panel.btn_legal.toolTip()
        assert "Шаблоны" in panel.btn_templates.toolTip()
    finally:
        panel.close()
        panel.deleteLater()

