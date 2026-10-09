#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Скобки с сокращённым наименованием (ШАГ «Фикс сохранения + динамические
стороны в шаблонах», часть C.7).

Что было
--------
Бланк печатал «{{carrier_full_name}} ({{carrier_name}})» — скобка стояла в
самом шаблоне. В рабочей базе `short_name` часто повторяет полное
наименование (у ИП его нет вовсе), поэтому в договоре выходило
«ООО «Ромашка» (ООО «Ромашка»)» и «Индивидуальный предприниматель … ()».

Что стало
---------
Скобку собирает генератор в ключе `*_name_in_parens`: она появляется только
тогда, когда сокращённое наименование ЕСТЬ и ОТЛИЧАЕТСЯ от полного.

Данные синтетические, ПДн нет.
"""

import re

import pytest
from docx import Document

from core.contract_generator import ContractGenerator

OOO_CARRIER = {
    "full_name": "Общество с ограниченной ответственностью «Перевозчик»",
    "short_name": "ООО «Перевозчик»",
    "inn": "7701234567", "kpp": "770101001", "ogrn": "1027700132195",
    "legal_address": "г. Москва", "actual_address": "",
    "bank_account": "40702810000000000002", "bik": "044525225",
    "correspondent_account": "30101810400000000225", "bank_name": "ПАО Сбербанк",
    "director_name": "Сидоров Сидор Сидорович",
    "director_position": "Генеральный директор",
    "phone": "", "email": "", "entity_type": "ООО",
}

OOO_CUSTOMER = {
    "full_name": "Общество с ограниченной ответственностью «Заказчик»",
    "short_name": "ООО «Заказчик»",
    "inn": "7707654321", "kpp": "770701001", "ogrn": "1027700132195",
    "legal_address": "г. Москва", "actual_address": "",
    "bank_account": "40702810000000000001", "bik": "044525225",
    "correspondent_account": "30101810400000000225", "bank_name": "ПАО Сбербанк",
    "director_name": "Петров Пётр Петрович",
    "director_position": "Генеральный директор",
    "phone": "", "email": "",
}

IP_CUSTOMER = {
    "full_name": "Индивидуальный предприниматель Смирнова Елена Владимировна",
    "short_name": "ИП Смирнова Е.В.",
    "inn": "770123456789", "kpp": "", "ogrn": "315770000000012",
    "legal_address": "г. Волгоград", "actual_address": "",
    "bank_account": "", "bik": "", "correspondent_account": "",
    "bank_name": "", "director_name": "",
    "director_position": "Генеральный директор",
    "phone": "", "email": "",
}


@pytest.fixture
def generator(templates_dir) -> ContractGenerator:
    return ContractGenerator(templates_dir=str(templates_dir))


def _payload(carrier, customer, carrier_type="ООО (с НДС)") -> dict:
    return {
        "driver": {"full_name": "Иванов Иван Иванович"},
        "carrier": carrier, "customer": customer,
        "vehicles": [{"vin": "EC3TEUMB0T0002608", "brand_model": "JETOUR T2",
                      "vehicle_type": "Легковой автомобиль"}],
        "tractor": {"brand_model": "Foton", "plate_number": "O844XY196"},
        "trailer": {"brand_model": "YANGMINDA", "plate_number": "71ABF18"},
        "contract": {
            "number": "PARENS-1", "date": "2026-10-09",
            "carrier_type": carrier_type,
            "vat_rate": "22%" if carrier_type.startswith("ООО") else "0%",
            "vat_rate_num": 22 if carrier_type.startswith("ООО") else 0,
            "price_without_vat": 180300.0, "payment_days": 10,
        },
        "loadings": [{"address": "г. Воронеж", "date": "2026-10-10",
                      "time_window": ""}],
        "unloadings": [{"address": "г. Москва", "date": "2026-10-13",
                        "time_window": ""}],
    }


def _render(generator, output, carrier=None, customer=None,
            carrier_type="ООО (с НДС)") -> Document:
    generator.generate_docx(
        _payload(carrier or OOO_CARRIER, customer or OOO_CUSTOMER, carrier_type),
        str(output),
    )
    assert output.exists(), f"файл не создан: {output}"
    return Document(str(output))


def _flatten(text: str) -> str:
    """Текст одной строкой: пробелы (в том числе неразрывные) по одному."""
    return re.sub(r"\s+", " ", text.replace("\u00a0", " ")).strip()


def _document_text(doc) -> str:
    parts = [_flatten(p.text) for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(_flatten(cell.text))
    return "\n".join(parts)


def _clause_text(doc, label: str) -> str:
    """Абзац, идущий сразу за строкой-меткой («1.2. Перевозчик:»)."""
    texts = [p.text for p in doc.paragraphs]
    for index, text in enumerate(texts):
        if text.strip().startswith(label):
            for following in texts[index + 1:]:
                if following.strip():
                    return _flatten(following)
    raise AssertionError(f"абзац {label!r} не найден")


# ─────────────────────────────────────────────────────────────
# Перевозчик
# ─────────────────────────────────────────────────────────────

def test_ooo_carrier_parens_printed_when_short_differs(generator, work_file):
    """Сокращённое наименование отличается — скобка печатается (как раньше)."""
    doc = _render(generator, work_file("parens_carrier_diff.docx"))

    clause = _clause_text(doc, "1.2. Перевозчик:")
    assert clause.startswith(
        "Общество с ограниченной ответственностью «Перевозчик» "
        "(ООО «Перевозчик»), ИНН 7701234567"
    )


def test_ooo_carrier_no_parens_when_short_equals_full(generator, work_file):
    """Сокращённое совпадает с полным — скобки нет (было «(ООО «Ромашка»)»)."""
    carrier = dict(
        OOO_CARRIER,
        full_name="ООО «Перевозчик»", short_name="ООО «Перевозчик»",
    )
    doc = _render(generator, work_file("parens_carrier_same.docx"),
                  carrier=carrier)

    clause = _clause_text(doc, "1.2. Перевозчик:")
    assert clause.startswith("ООО «Перевозчик», ИНН 7701234567")
    assert "(" not in clause


def test_carrier_no_parens_when_short_empty(generator, work_file):
    """Сокращённого наименования нет — пустой скобки быть не должно."""
    carrier = dict(OOO_CARRIER, short_name="")
    doc = _render(generator, work_file("parens_carrier_empty.docx"),
                  carrier=carrier)

    clause = _clause_text(doc, "1.2. Перевозчик:")
    assert "(" not in clause
    assert "()" not in _document_text(doc)


def test_ip_carrier_no_empty_parens(generator, work_file):
    """У ИП ветка бланка скобок не печатает вовсе."""
    carrier = dict(OOO_CARRIER, entity_type="ИП",
                   full_name="Индивидуальный предприниматель Смирнов А.Н.",
                   short_name="")
    doc = _render(generator, work_file("parens_carrier_ip.docx"),
                  carrier=carrier, carrier_type="ИП без НДС")

    clause = _clause_text(doc, "1.2. Перевозчик:")
    assert "(" not in clause
    assert "()" not in _document_text(doc)


# ─────────────────────────────────────────────────────────────
# Заказчик
# ─────────────────────────────────────────────────────────────

def test_ooo_client_parens_printed_when_short_differs(generator, work_file):
    """У заказчика скобка работает по тому же правилу."""
    doc = _render(generator, work_file("parens_client_diff.docx"))

    clause = _clause_text(doc, "1.1. Заказчик:")
    assert clause.startswith(
        "Общество с ограниченной ответственностью «Заказчик» "
        "(ООО «Заказчик»), ИНН 7707654321"
    )
    assert "именуемое в дальнейшем «Заказчик»" in clause


def test_ooo_client_no_parens_when_short_equals_full(generator, work_file):
    """Короткое имя заказчика повторяет полное — скобки нет."""
    customer = dict(
        OOO_CUSTOMER,
        full_name="ООО «Заказчик»", short_name="ООО «Заказчик»",
    )
    doc = _render(generator, work_file("parens_client_same.docx"),
                  customer=customer)

    clause = _clause_text(doc, "1.1. Заказчик:")
    assert clause.startswith(
        "ООО «Заказчик», ИНН 7707654321, КПП 770701001, ОГРН 1027700132195, "
        "именуемое"
    )
    assert "(" not in clause


def test_ip_client_short_name_not_printed_in_parens(generator, work_file):
    """У ИП-заказчика «(ИП Смирнова Е.В.)» в п. 1.1 не печатается."""
    doc = _render(generator, work_file("parens_client_ip.docx"),
                  customer=IP_CUSTOMER)

    clause = _clause_text(doc, "1.1. Заказчик:")
    assert clause.startswith(
        "Индивидуальный предприниматель Смирнова Елена Владимировна, "
        "ИНН 770123456789"
    )
    assert "(" not in clause


# ─────────────────────────────────────────────────────────────
# Ключи карты замен
# ─────────────────────────────────────────────────────────────

def test_paren_keys_and_old_name_keys_are_kept(generator):
    """Новые ключи скобок есть, старые `*_name` не удалены."""
    replacements = generator._build_replacements_map(
        _payload(OOO_CARRIER, OOO_CUSTOMER)
    )

    assert replacements["carrier_name_in_parens"] == " (ООО «Перевозчик»)"
    assert replacements["client_name_in_parens"] == " (ООО «Заказчик»)"
    # Старые ключи остаются: их читает внешний код и legacy-блоки.
    assert replacements["carrier_name"] == "ООО «Перевозчик»"
    assert replacements["client_name"] == "ООО «Заказчик»"
