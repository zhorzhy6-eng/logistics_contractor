#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сводные (e2e) тесты Excel-заявки Хавалов — полный цикл (ЭТАП 3.1.E.A.4).

ЧЕМ ЭТОТ ФАЙЛ ОТЛИЧАЕТСЯ ОТ tests/test_havaly_e2e.py
----------------------------------------------------
В A.3 проверялись СТЫКИ: промпт ↔ генератор, валидатор ↔ генератор,
генератор ↔ файл, фабрика ↔ генератор, присланный файл ↔ готовый файл.
Здесь проверяются СЦЕНАРИИ работы — то, как всем этим пользуется человек:

  1. полный цикл «импорт → правка → сохранение → повторный импорт»:
     данные → файл → данные → правка одного поля → файл → данные;
  2. раскладка файла с ровно 10 машинами совпадает с эталонным бланком
     (обрезать нечего: в бланке ровно 10 строк данных);
  3. одна машина: пустые строки таблицы обрезаны, нижний блок сторон
     поднялся, но сам остался на месте — относительно конца таблицы;
  4. ноль машин: файл всё равно создаётся и читается, а валидатор
     говорит об этом замечанием, а не ошибкой;
  5. присланный файл с данными: берём образец заказчика, программно
     заполняем машины, читаем обратно;
  6. пять типов договоров в одном прогоне: генератор Хавалов не ломает
     остальные четыре, реестр отдаёт zayavka_excel, промпт непустой.

Сверка с эталоном (пункты 2 и 3) идёт по геометрии листа: номера строк,
объединённые диапазоны, размеры, подписи. Что именно написано в ячейках —
не сравнивается: данные у готового файла свои, а ПДн синтетические.

Файлы тестов создаются в tests/_tmp/_havaly_e2e_full/<префикс прогона>/
и удаляются в finally. Префикс уникален для прогона: фиксированные имена
в tests/_tmp уже один раз сыграли злую шутку (грабля A.3).

Логи всей цепочки — без ПДн: только имена полей, количества и длины.
"""

import hashlib
import logging
import re
import shutil
import sys
import uuid
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core.contracts.arenda_ts.generator import ArendaTsGenerator  # noqa: E402
from core.contracts.arenda_ts.validator import ArendaTsValidator  # noqa: E402
from core.contracts.contract_types import ContractType  # noqa: E402
from core.contracts.factory import GeneratorFactory  # noqa: E402
from core.contracts.formika.generator import FormikaGenerator  # noqa: E402
from core.contracts.formika.validator import FormikaValidator  # noqa: E402
from core.contracts.logistiks_rus.generator import (  # noqa: E402
    LogistiksRusGenerator,
)
from core.contracts.logistiks_rus.validator import (  # noqa: E402
    LogistiksRusValidator,
)
from core.contracts.perevozka.generator import PerevozkaGenerator  # noqa: E402
from core.contracts.perevozka.validator import PerevozkaValidator  # noqa: E402
from core.contracts.registry import ContractTypeRegistry  # noqa: E402
from core.contracts.zayavka.generator import (  # noqa: E402
    CARRIER_NAME,
    CUSTOMER_NAME,
    DATE_LABEL,
    FILE_PREFIX,
    FIRST_HEADER,
    HEADER_OF,
    MAX_VEHICLES,
    PARTY_LABELS,
    SHEET_NAME,
    VEHICLE_KEYS,
    ZayavkaExcelGenerator,
)
from core.contracts.zayavka.postprocess import TrimVehicleRowsStep  # noqa: E402
from core.contracts.zayavka.validator import (  # noqa: E402
    ZayavkaExcelValidator,
)
from core.prompts import get_prompt  # noqa: E402
from core.prompts.havaly import PROMPT  # noqa: E402

# ─────────────────────────────────────────────────────────────
# Константы прогона
# ─────────────────────────────────────────────────────────────

#: Образец заказчика: та же форма, другое оформление, шесть размеченных
#: строк данных. Источник — только чтение.
SAMPLE_NAME = "Хавалы_образец.xlsx"

#: Папка тестовых форм прогона и его метка (уникальна: см. граблю A.3).
STAGE_DIR = Path(__file__).resolve().parent / "_tmp" / "_havaly_e2e_full"
_RUN_ID = uuid.uuid4().hex[:8]

#: Дата заявки: от неё зависит имя готового файла.
DATE = "05.10.2026"
DATE_ISO = "2026-10-05"

#: Плейсхолдеры распознавания: в файл и обратно идут как есть.
PLACEHOLDER_FIO = "<<PERSON_1>>"
PLACEHOLDER_PHONE = "<<PHONE_1>>"

#: Отпечатки синтетических ПДн: их не должно быть в логах. Значения —
#: из make_zayavka/make_vehicles ниже и из правок сценария 1.
PII_FRAGMENTS = (
    "Тестов", "Тестович", "9901", "926830", "1822",
    "EC3TEUMB0T0002601", "Промышленная", "Складской", "Тестовая",
    "О844ХУ196", "71ABF18", "+7 (999)", "123456.78", "000-11-22",
)

#: Пять типов договоров в одном прогоне: ключ → классы генератора и
#: валидатора. Ядро — zayavka_excel (Excel), остальные четыре — DOCX.
FIVE_TYPES = {
    "perevozka": (PerevozkaGenerator, PerevozkaValidator),
    "formika": (FormikaGenerator, FormikaValidator),
    "logistiks_rus": (LogistiksRusGenerator, LogistiksRusValidator),
    "arenda_ts": (ArendaTsGenerator, ArendaTsValidator),
    "zayavka_excel": (ZayavkaExcelGenerator, ZayavkaExcelValidator),
}

#: Частичные классы типов из реестра (проверка «генератор — в своём модуле»).
TYPE_CLASS_MODULES = {
    "perevozka": "core.contracts.perevozka",
    "formika": "core.contracts.formika",
    "logistiks_rus": "core.contracts.logistiks_rus",
    "arenda_ts": "core.contracts.arenda_ts",
    "zayavka_excel": "core.contracts.zayavka",
}


# ─────────────────────────────────────────────────────────────
# Данные
# ─────────────────────────────────────────────────────────────

def make_zayavka(**overrides) -> dict:
    """Общие сведения заявки: синтетические, все поля схемы промпта."""
    data = {
        "date": DATE,
        "lot_number": "LOT-42",
        "loading_city": "Калуга",
        "loading_point": "ул. Промышленная, 12",
        "unloading_city": "Москва",
        "unloading_point": "Складской проезд, 5",
        "carrier_name": CARRIER_NAME,
        "customer_name": CUSTOMER_NAME,
        "tractor_brand": "Foton Auman",
        "tractor_color": "Белый",
        "tractor_plate": "О844ХУ196",
        "trailer_brand": "YANGMINDA",
        "trailer_plate": "71ABF18",
        "driver_last_name": "Тестов",
        "driver_first_name": "Тест",
        "driver_middle_name": "Тестович",
        "driver_license_number": "9901 123456",
        "driver_license_issue_date": "12.03.2019",
        "driver_passport_series": "1822",
        "driver_passport_number": "926830",
        "driver_passport_issuer": "Отделом УФМС России по г. Москве",
        "driver_passport_issue_date": "30.01.2023",
        "driver_citizenship": "РФ",
        "driver_birth_date": "01.01.1980",
        "driver_registration": "г. Москва, ул. Тестовая, д. 1",
        "driver_phone": "+7 (999) 123-45-67",
        "loading_plan_date": "07.10.2026",
        "loading_plan_time": "09:00",
        "price_with_vat": 123456.78,
        "vat_rate": "22%",
    }
    data.update(overrides)
    return data


def make_vehicles(count: int) -> list:
    """Машины: VIN, марка, модель, дилер, код дилера — по одной на строку."""
    return [
        {
            "vin": f"EC3TEUMB0T00026{index:02d}",
            "brand": f"МАРКА {index}",
            "model": f"МОДЕЛЬ {index}",
            "dealer": f"Дилер {index}",
            "dealer_code": f"D-{index:03d}",
        }
        for index in range(1, count + 1)
    ]


def make_data(vehicles=None, **overrides) -> dict:
    """Данные заявки в схеме промпта: {"zayavka": {...}, "vehicles": [...]}."""
    return {
        "zayavka": make_zayavka(**overrides),
        "vehicles": vehicles if vehicles is not None else make_vehicles(2),
    }


def expected_zayavka(data: dict) -> dict:
    """
    Что генератор вернёт при чтении записанных данных.

    Два отличия от исходных данных, оба известные (см. docstring
    генератора): ``vat_rate`` в бланке негде напечатать — читается пустым,
    а стороны в бланке напечатаны и подставляются константами, а не
    читаются из ячеек.
    """
    expected = dict(data["zayavka"])
    expected["vat_rate"] = ""
    expected["customer_name"] = CUSTOMER_NAME
    expected["carrier_name"] = CARRIER_NAME
    return expected


# ─────────────────────────────────────────────────────────────
# Файлы прогона
# ─────────────────────────────────────────────────────────────

@pytest.fixture(scope="session", autouse=True)
def clean_stage_dir():
    """
    Убирает папку форм прогона целиком — в том числе при падении теста.

    Имена файлов внутри уникальны (метка прогона в имени папки), поэтому
    остатки прошлых прогонов новой проверке не мешают.
    """
    try:
        yield
    finally:
        shutil.rmtree(STAGE_DIR, ignore_errors=True)


@pytest.fixture
def stage() -> Path:
    """Папка форм этого теста: STAGE_DIR/<метка прогона>_<имя теста>."""
    return stage_dir()


def stage_dir(name: str = "") -> Path:
    """Папка форм прогона (с необязательным уточнением в имени)."""
    path = STAGE_DIR / f"{_RUN_ID}_{name}" if name else STAGE_DIR / _RUN_ID
    path.mkdir(parents=True, exist_ok=True)
    return path


@pytest.fixture
def generator() -> ZayavkaExcelGenerator:
    return ZayavkaExcelGenerator()


@pytest.fixture
def validator() -> ZayavkaExcelValidator:
    return ZayavkaExcelValidator()


@pytest.fixture
def saved(generator, tmp_output):
    """Готовые файлы теста: удаляются после проверки."""
    created: list = []

    def _save(data, method="generate", source=None, out_dir=None):
        target = str(out_dir or tmp_output)
        if method == "generate":
            path = generator.generate(data, target)
        else:
            path = generator.fill_from_template(
                source or generator.template_path(), data, target
            )
        created.append(Path(path))
        return Path(path)

    yield _save

    for path in created:
        try:
            path.unlink(missing_ok=True)
        except OSError:
            pass


@pytest.fixture
def tmp_output() -> Path:
    """
    Папка вывода теста — внутри папки форм прогона.

    Отдельная от output/ проекта: тесты не должны оставлять файлы рядом
    с рабочими договорами.
    """
    path = stage_dir() / "output"
    path.mkdir(parents=True, exist_ok=True)
    return path


def read_back(generator, path) -> dict:
    """Данные из готового файла — тем же генератором, что его записал."""
    return generator.read_template(str(path))


def sha256(path) -> str:
    """Отпечаток файла: по нему видно, что записан именно он, а не эталон."""
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


# ─────────────────────────────────────────────────────────────
# Разбор листа: где что стоит
# ─────────────────────────────────────────────────────────────

def header_row_of(sheet) -> int:
    """Номер строки шапки — по заголовку первой колонки, как в генераторе."""
    for row in sheet.iter_rows():
        for cell in row:
            if str(cell.value or "").strip() == FIRST_HEADER:
                return cell.row
    raise AssertionError(f"в листе нет строки шапки «{FIRST_HEADER}»")


def label_row_of(sheet, *labels: str) -> int:
    """
    Номер строки, в которой стоят ВСЕ подписи.

    Подписи ищутся по тексту ячейки — как их ищет сам генератор: в бланке
    подпись даты стоит с двоеточием («Дата заявки:»), и требовать точного
    совпадения строк нельзя. Номера строк в бланке не заданы жёстко: после
    обрезки пустых строк низ бланка переезжает наверх.
    """
    wanted = [label.strip().lower() for label in labels]
    for row in sheet.iter_rows():
        found = [str(cell.value or "").strip().lower() for cell in row]
        if all(
            any(label in value for value in found)
            for label in wanted
        ):
            return row[0].row
    raise AssertionError(f"не найдена строка с подписями {wanted}")


def layout_of(path) -> dict:
    """
    Геометрия листа: строки, объединения, размеры.

    Содержимое ячеек сюда не входит — только то, по чему видно, что
    обрезка строк не сломала раскладку бланка.
    """
    sheet = load_workbook(Path(path))[SHEET_NAME]
    return {
        "max_row": sheet.max_row,
        "max_column": sheet.max_column,
        "header_row": header_row_of(sheet),
        "party_row": label_row_of(sheet, *PARTY_LABELS),
        "signature_row": label_row_of(sheet, "ФИО, подпись, печать"),
        "merged": sorted(str(item) for item in sheet.merged_cells.ranges),
        "column_widths": {
            key: item.width for key, item in sheet.column_dimensions.items()
        },
        "row_heights": {
            key: item.height for key, item in sheet.row_dimensions.items()
        },
    }


def last_table_row_of(path) -> int:
    """Последняя строка таблицы: три строки выше подписи «Заказчик»."""
    sheet = load_workbook(Path(path))[SHEET_NAME]
    return label_row_of(sheet, *PARTY_LABELS) - 3


# ─────────────────────────────────────────────────────────────
# Сценарий 1: полный цикл «импорт → правка → сохранение → импорт»
# ─────────────────────────────────────────────────────────────

def test_full_cycle_generate_read_edit_save_read(generator, saved, tmp_output,
                                                 stage):
    """
    Полный круг дважды: данные → файл → данные, правка, снова файл → данные.

    Так выглядит рабочая правка: заявку сгенерировали, в одном поле
    ошиблись, поле поправили — остальное должно остаться нетронутым.

    Правленый файл кладётся в ДРУГУЮ папку, и это не прихоть теста: имя
    готового файла зависит только от даты заявки, поэтому правка «в тот же
    день и в ту же папку» целилась бы в исходный файл. Генератор такое
    запрещает (``test_same_date_edit_is_refused``) — исходный файл должен
    остаться целым.
    """
    data = make_data(vehicles=make_vehicles(3))

    first = saved(data)
    assert first.name == f"{FILE_PREFIX}_{DATE_ISO}.xlsx"
    assert first.parent == tmp_output

    back = read_back(generator, first)
    assert back["vehicles"] == data["vehicles"]
    assert back["zayavka"] == expected_zayavka(data)

    edited = {
        "zayavka": {**data["zayavka"], "driver_phone": "+7 (999) 000-11-22"},
        "vehicles": data["vehicles"],
    }
    edited_dir = stage / "edited"
    second = saved(edited, method="fill", source=first, out_dir=edited_dir)

    assert second != first
    after = read_back(generator, second)
    assert after["zayavka"]["driver_phone"] == "+7 (999) 000-11-22"
    assert after["vehicles"] == data["vehicles"]

    expected = expected_zayavka(edited)
    assert [key for key in expected if after["zayavka"][key] != expected[key]] \
        == []
    assert [key for key in expected if expected[key] != expected_zayavka(data)[key]] \
        == ["driver_phone"]

    # Первый файл правка не задела.
    assert read_back(generator, first)["zayavka"]["driver_phone"] == \
        data["zayavka"]["driver_phone"]


def test_same_date_edit_is_refused(generator, saved):
    """
    Правка «в тот же день и в ту же папку» отклоняется — файл цел.

    Ловушка имени: оно зависит только от даты заявки, поэтому повторная
    запись в ту же папку затёрла бы исходный файл. Генератор говорит об
    этом прямо, а не портит документ.
    """
    from core.contracts.zayavka.generator import ZayavkaTemplateError

    data = make_data(vehicles=make_vehicles(1))
    first = saved(data)
    before = sha256(first)

    edited = {"zayavka": {**data["zayavka"], "lot_number": "LOT-99"},
              "vehicles": data["vehicles"]}
    with pytest.raises(ZayavkaTemplateError, match="совпадает с исходным"):
        generator.fill_from_template(str(first), edited, str(first.parent))

    assert sha256(first) == before
    assert read_back(generator, first)["zayavka"]["lot_number"] == "LOT-42"


def test_edit_changes_only_the_edited_field(generator, saved, stage):
    """Правка одной машины не задевает соседние строки и общие сведения."""
    data = make_data(vehicles=make_vehicles(3))
    first = saved(data)

    edited_vehicles = [dict(item) for item in data["vehicles"]]
    edited_vehicles[1]["dealer_code"] = "D-999"
    second = saved(
        {"zayavka": dict(data["zayavka"]), "vehicles": edited_vehicles},
        method="fill",
        source=first,
        out_dir=stage / "edited_one",
    )

    after = read_back(generator, second)
    assert [item["dealer_code"] for item in after["vehicles"]] == \
        ["D-001", "D-999", "D-003"]
    assert after["vehicles"][0] == data["vehicles"][0]
    assert after["vehicles"][2] == data["vehicles"][2]
    assert after["zayavka"] == expected_zayavka(data)


def test_round_trip_keeps_prompt_placeholders(generator, saved, stage):
    """
    Обезличенные данные проходят два уровня круга без изменений.

    Распознавание отдаёт плейсхолдеры (<<PERSON_1>>, <<PHONE_1>>):
    генератор их не трогает, а повторный импорт возвращает ровно те же
    токены — иначе второй круг «съел» бы обезличивание.
    """
    data = make_data(
        vehicles=make_vehicles(1),
        driver_last_name=PLACEHOLDER_FIO,
        driver_first_name="",
        driver_middle_name="",
        driver_phone=PLACEHOLDER_PHONE,
    )
    first = saved(data)
    second = saved(data, method="fill", source=first,
                   out_dir=stage / "placeholders")
    after = read_back(generator, second)

    assert after["zayavka"]["driver_last_name"] == PLACEHOLDER_FIO
    assert after["zayavka"]["driver_phone"] == PLACEHOLDER_PHONE
    assert after["vehicles"] == data["vehicles"]


def test_second_generation_does_not_touch_the_first_file(generator, saved,
                                                         stage):
    """Правка пишет НОВЫЙ файл: исходный готовый файл остаётся целым."""
    data = make_data(vehicles=make_vehicles(2))
    first = saved(data)
    before = sha256(first)

    saved(
        {"zayavka": {**data["zayavka"], "lot_number": "LOT-77"},
         "vehicles": data["vehicles"]},
        method="fill",
        source=first,
        out_dir=stage / "second_generation",
    )

    assert sha256(first) == before
    assert read_back(generator, first)["zayavka"]["lot_number"] == "LOT-42"


# ─────────────────────────────────────────────────────────────
# Сценарий 2: ровно 10 машин — раскладка как в эталоне
# ─────────────────────────────────────────────────────────────

def test_ten_vehicles_keep_the_template_layout(generator, saved):
    """
    Десять машин — предельный случай: обрезать нечего.

    В эталонном бланке ровно 10 размеченных строк данных, поэтому готовый
    файл обязан совпасть с бланком по геометрии: последняя строка таблицы,
    положение блока сторон и места под подписи, объединения, размеры.
    """
    data = make_data(vehicles=make_vehicles(MAX_VEHICLES))
    path = saved(data)

    assert layout_of(path) == layout_of(generator.template_path())


def test_ten_vehicles_are_read_back_in_order(generator, saved):
    """Порядок машин при чтении тот же, что при записи: строка к строке."""
    vehicles = make_vehicles(MAX_VEHICLES)
    vehicles[0]["vin"] = "EC3TEUMB0T0002600"
    vehicles[-1]["vin"] = "EC3TEUMB0T0002699"
    data = make_data(vehicles=vehicles)
    path = saved(data)

    back = read_back(generator, path)

    assert len(back["vehicles"]) == MAX_VEHICLES
    assert back["vehicles"] == vehicles
    assert [item["vin"] for item in back["vehicles"]] == \
        [item["vin"] for item in vehicles]
    assert back["vehicles"][-1]["dealer_code"] == f"D-{MAX_VEHICLES:03d}"
    # Поля машины — ровно те, что объявлены схемой промпта.
    for vehicle in back["vehicles"]:
        assert set(vehicle) == set(VEHICLE_KEYS)


def test_ten_vehicles_change_the_file_itself(generator, saved):
    """Файл с данными — не копия бланка: отпечатки различаются."""
    path = saved(make_data(vehicles=make_vehicles(MAX_VEHICLES)))

    assert sha256(path) != sha256(generator.template_path())


# ─────────────────────────────────────────────────────────────
# Сценарий 3: одна машина
# ─────────────────────────────────────────────────────────────

def test_one_vehicle_is_trimmed_and_read_back(generator, saved):
    """Одна машина: лишние строки таблицы обрезаны, читается ровно одна."""
    data = make_data(vehicles=make_vehicles(1))
    path = saved(data)

    back = read_back(generator, path)

    assert back["vehicles"] == data["vehicles"]
    assert back["zayavka"] == expected_zayavka(data)
    # В бланке десять строк данных (7..16), после обрезки остаётся одна.
    sheet = load_workbook(path)[SHEET_NAME]
    assert last_table_row_of(path) == header_row_of(sheet) + 1


def test_trimming_removes_the_rows_and_moves_the_party_block(generator, saved):
    """
    Нижний блок сторон переехал наверх вместе с таблицей — и не потерялся.

    Здесь проверяется и работа шага постобработки, и то, что он не уносит
    с собой подписи и места под подписи: без них заявку не подписать.
    """
    path = saved(make_data(vehicles=make_vehicles(1)))
    sheet = load_workbook(path)[SHEET_NAME]
    template = load_workbook(generator.template_path())[SHEET_NAME]

    template_party = label_row_of(template, *PARTY_LABELS)
    party = label_row_of(sheet, *PARTY_LABELS)
    signature = label_row_of(sheet, "ФИО, подпись, печать")

    assert sheet.max_row < template.max_row
    assert party < template_party
    assert signature < label_row_of(template, "ФИО, подпись, печать")
    # Блок стоит там же относительно конца таблицы, что и в бланке:
    # между последней строкой таблицы и подписью — те же PARTY_TAIL_ROWS.
    assert party - last_table_row_of(path) == \
        template_party - last_table_row_of(generator.template_path())
    # Заголовок над таблицей остался объединённым: обрезка его не задела.
    assert "B2:C2" in {str(item) for item in sheet.merged_cells.ranges}


def test_trim_step_is_what_removes_the_rows(generator, saved):
    """Строки убирает именно шаг конвейера — он объявлен и назван."""
    steps = generator.postprocess_steps()

    assert [step.name for step in steps] == [TrimVehicleRowsStep.name]
    assert isinstance(steps[0], TrimVehicleRowsStep)


def test_one_vehicle_file_differs_from_template(generator, saved):
    """
    Отпечаток готового файла не совпадает с бланком.

    Ожидаемо: в файле есть данные (дата, машина, водитель, ставка), а в
    бланке — только пустые строки. Если отпечатки совпали, значит данные
    не записались.
    """
    path = saved(make_data(vehicles=make_vehicles(1)))

    assert sha256(path) != sha256(generator.template_path())


def test_one_vehicle_file_keeps_the_rate_and_the_parties(generator, saved):
    """Ставка, дата и стороны на месте: файл читается как документ заявки."""
    data = make_data(vehicles=make_vehicles(1))
    path = saved(data)
    back = read_back(generator, path)

    assert back["zayavka"]["price_with_vat"] == pytest.approx(
        data["zayavka"]["price_with_vat"]
    )
    assert back["zayavka"]["date"] == DATE
    assert back["zayavka"]["customer_name"] == CUSTOMER_NAME
    assert back["zayavka"]["carrier_name"] == CARRIER_NAME


# ─────────────────────────────────────────────────────────────
# Сценарий 4: ноль машин
# ─────────────────────────────────────────────────────────────

def test_zero_vehicles_file_is_created_and_readable(generator, saved,
                                                    tmp_output):
    """Заявка без машин всё равно сохраняется и читается."""
    data = make_data(vehicles=[])
    path = saved(data)

    assert path.exists()
    assert path.name == f"{FILE_PREFIX}_{DATE_ISO}.xlsx"
    assert path.parent == tmp_output

    back = read_back(generator, path)
    assert back["vehicles"] == []


def test_zero_vehicles_keep_the_request_date(generator, saved):
    """
    Дата заявки — единственное, что попадает в файл без машин.

    Общие сведения бланк раскладывает ПО СТРОКАМ таблицы: нет машин —
    писать их некуда. Дата лежит отдельной ячейкой над таблицей и
    остаётся на месте (см. «Известные расхождения» в docs/STATE.md).
    """
    path = saved(make_data(vehicles=[]))
    back = read_back(generator, path)

    assert back["zayavka"]["date"] == DATE

    sheet = load_workbook(path)[SHEET_NAME]
    label = label_row_of(sheet, DATE_LABEL)
    value = sheet.cell(row=label, column=2).value
    assert value.strftime("%d.%m.%Y") == DATE


def test_zero_vehicles_are_warned_about_not_blocked(validator, generator,
                                                    saved):
    """Валидатор про ноль машин предупреждает — и не мешает сохранению."""
    data = make_data(vehicles=[])
    report = validator.check(data)

    assert report.errors == []
    assert any("нет ни одной машины" in item for item in report.warnings)

    path = saved(data)
    assert read_back(generator, path)["vehicles"] == []


def test_zero_vehicles_file_passes_the_file_checks(validator, generator,
                                                   saved):
    """Структурные проверки файла без машин ошибок не дают: форма цела."""
    path = saved(make_data(vehicles=[]))
    report = validator.check(str(path))

    assert report.errors == []
    # Блок сторон в файле есть и стороны в нём те самые — замечаний нет.
    assert not any("другая организация" in item for item in report.warnings)
    assert not any("нет блока сторон" in item for item in report.warnings)


def test_zero_vehicles_trim_all_data_rows(generator, saved):
    """Все размеченные строки таблицы удалены: пустого «хвоста» не осталось."""
    path = saved(make_data(vehicles=[]))
    template = generator.template_path()

    assert last_table_row_of(path) < last_table_row_of(template)
    assert read_back(generator, path)["vehicles"] == []


# ─────────────────────────────────────────────────────────────
# Сценарий 5: присланный файл с данными
# ─────────────────────────────────────────────────────────────

def fill_sample(sample: Path, data: dict, target: Path) -> Path:
    """
    Заполняет КОПИЮ образца заказчика, как это сделал бы человек.

    Колонки берутся из шапки файла по тексту заголовка (как в генераторе),
    значения ставятся в первые строки под шапкой. Так проверяется чтение
    ЧУЖОГО файла с данными, а не наша раскладка: сравнивать результат с
    generate() нельзя — оформление и число строк у образца свои.
    """
    book = load_workbook(str(sample))
    sheet = book[SHEET_NAME]

    header = header_row_of(sheet)
    positions = {}
    for row in sheet.iter_rows(min_row=header, max_row=header):
        for cell in row:
            title = str(cell.value or "").strip()
            if title and title not in positions:
                positions[title] = cell.column

    ranges = list(sheet.merged_cells.ranges)

    def anchor(row: int, column: int):
        """Левый верхний угол объединения, накрывающего координату."""
        for merged in ranges:
            if merged.min_row <= row <= merged.max_row and \
                    merged.min_col <= column <= merged.max_col:
                return merged.min_row, merged.min_col
        return row, column

    def write(row: int, column: int, value) -> None:
        # Внутри объединённой ячейки писать можно только в её левый
        # верхний угол (openpyxl: остальные — MergedCell, только чтение).
        # Так же ведёт себя _anchor_cell в самом генераторе.
        row, column = anchor(row, column)
        sheet.cell(row=row, column=column, value=value)

    write(header - 1, 1 + 1, data["zayavka"]["date"])
    for offset, vehicle in enumerate(data["vehicles"]):
        merged = {**data["zayavka"], **vehicle}
        for key, title in HEADER_OF.items():
            column = positions.get(title)
            if column is None:
                continue
            write(header + 1 + offset, column, merged.get(key))

    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(str(target))
    return target


def test_incoming_filled_sample_is_read(generator, templates_dir, stage):
    """
    Присланный файл с данными читается: три машины и дата заявки.

    С генерацией результат не сравнивается: у образца своя раскладка
    (шесть размеченных строк, объединённая ставка), и совпадение с нашим
    файлом здесь ничего не доказывало бы.
    """
    sample = templates_dir / SAMPLE_NAME
    if not sample.exists():
        pytest.skip(f"нет образца {SAMPLE_NAME}")

    data = make_data(vehicles=make_vehicles(3))
    incoming = fill_sample(sample, data, stage / "incoming_filled.xlsx")

    back = generator.read_template(str(incoming))

    assert len(back["vehicles"]) == 3
    assert back["vehicles"] == data["vehicles"]
    assert back["zayavka"]["date"] == DATE
    assert back["zayavka"]["lot_number"] == data["zayavka"]["lot_number"]


def test_incoming_sample_is_not_modified(generator, templates_dir, stage):
    """
    Образец заказчика — источник структуры, а не место записи.

    Заполняется КОПИЯ: чтение и запись идут по ней, а сам образец в
    templates/ остаётся ровно таким, каким пришёл от заказчика.
    """
    sample = templates_dir / SAMPLE_NAME
    if not sample.exists():
        pytest.skip(f"нет образца {SAMPLE_NAME}")

    before = sha256(sample)
    data = make_data(vehicles=make_vehicles(3))
    incoming = fill_sample(sample, data, stage / "incoming_copy.xlsx")

    generator.read_template(str(incoming))

    assert sha256(sample) == before
    # Заполнилась именно копия: данные в ней есть. Побайтового совпадения
    # копии с образцом ждать нельзя — openpyxl пересобирает архив книги.
    back = generator.read_template(str(incoming))
    assert back["vehicles"] == data["vehicles"]
    assert len(back["vehicles"]) != len(
        generator.read_template(str(sample))["vehicles"]
    )


def test_incoming_filled_sample_can_be_filled_again(generator, templates_dir,
                                                    stage, saved):
    """
    Присланный заполненный файл годится как бланк для следующей заявки.

    Шесть размеченных строк образца, три машины: лишние строки обрезаются,
    блок сторон поднимается — и данные читаются обратно.
    """
    sample = templates_dir / SAMPLE_NAME
    if not sample.exists():
        pytest.skip(f"нет образца {SAMPLE_NAME}")

    incoming = fill_sample(sample, make_data(vehicles=make_vehicles(3)),
                           stage / "incoming_source.xlsx")
    data = make_data(vehicles=make_vehicles(2))
    target = saved(data, method="fill", source=incoming)

    back = read_back(generator, target)
    assert back["vehicles"] == data["vehicles"]
    assert back["zayavka"] == expected_zayavka(data)
    assert label_row_of(load_workbook(target)[SHEET_NAME], *PARTY_LABELS) < \
        label_row_of(load_workbook(incoming)[SHEET_NAME], *PARTY_LABELS)


# ─────────────────────────────────────────────────────────────
# Сценарий 6: пять типов договоров в одном прогоне
# ─────────────────────────────────────────────────────────────

@pytest.fixture
def five_generators():
    """Генераторы всех пяти типов, полученные через фабрику."""
    ContractTypeRegistry.load_builtin()
    if "zayavka_excel" not in ContractTypeRegistry.known_types():
        import importlib

        importlib.reload(importlib.import_module("core.contracts.zayavka"))
    assert not ContractTypeRegistry.failures(), ContractTypeRegistry.failures()
    return {
        key: GeneratorFactory.get_generator(key, strict=True)
        for key in FIVE_TYPES
    }


def test_registry_returns_all_five_types(five_generators):
    """Реестр знает все пять типов, и классы у них — те, что ожидались."""
    assert set(five_generators) == set(FIVE_TYPES)
    assert set(ContractTypeRegistry.known_types()) >= set(FIVE_TYPES)

    for key, (generator_class, validator_class) in FIVE_TYPES.items():
        spec = ContractTypeRegistry.get(key, strict=True)
        assert spec.generator_class is generator_class, key
        assert spec.validator_class is validator_class, key
        assert spec.title, key


def test_excel_generator_does_not_break_the_other_types(five_generators):
    """
    Генератор Хавалов не подменяет остальные четыре типа.

    Проверяются именно ОБЩИЕ атрибуты, которыми пользуются все типы:
    ключ типа, свой класс, шаблоны в своей папке, папка вывода. Если
    Excel-генератор сломает этот контракт, пострадают и DOCX-типы.

    ``TITLE`` здесь НЕ проверяется: у DOCX-типов его нет — название для
    заголовка окна берётся из реестра (``ui/contract_picker.py``), а
    ``TITLE`` есть только у генератора Хавалов. Проверка «у всех есть
    TITLE» ловила бы не поломку, а разницу в контрактах.
    """
    for key, (generator_class, _) in FIVE_TYPES.items():
        generator = five_generators[key]

        assert generator.CONTRACT_TYPE == key
        assert isinstance(generator, generator_class)
        assert generator_class.__module__.startswith(TYPE_CLASS_MODULES[key])
        assert ContractTypeRegistry.get(key, strict=True).title
        assert generator.TEMPLATE_NAMES
        assert generator.templates
        assert callable(generator.default_output_dir)

        for template_key, filename in generator.TEMPLATE_NAMES.items():
            path = Path(generator.templates[template_key])
            assert path.name == filename, key
            assert path.exists(), f"{key}: нет шаблона {path.name}"
            assert path.parent == Path(generator.templates_dir)
            assert Path(generator.default_output_dir()).name == "output"


def test_factory_validators_of_all_five_types(five_generators):
    """Фабрика отдаёт валидатор каждого типа, и это класс своего типа."""
    for key, (_, validator_class) in FIVE_TYPES.items():
        built = GeneratorFactory.get_validator(key, strict=True)

        assert isinstance(built, validator_class), key
        assert built.CONTRACT_TYPE == key


def test_excel_generator_alone_is_not_docx_based(five_generators):
    """Из пяти типов от DOCX-базы не наследуется только Excel-заявка."""
    from core.contracts.base_generator import BaseContractGenerator

    for key, (generator_class, _) in FIVE_TYPES.items():
        is_docx = issubclass(generator_class, BaseContractGenerator)
        assert is_docx == (key != "zayavka_excel"), key


def test_havaly_type_is_not_a_stub_anymore(five_generators):
    """
    zayavka_excel — рабочий тип, а не заглушка шага 8.

    Признаки заглушки: пустой TITLE, дочерний класс реестра без своих
    шаблонов, «(заглушка)» в названии. Ни одного из них быть не должно.
    """
    generator = five_generators["zayavka_excel"]
    spec = ContractTypeRegistry.get("zayavka_excel", strict=True)

    assert isinstance(generator, ZayavkaExcelGenerator)
    assert spec.generator_class is ZayavkaExcelGenerator
    assert spec.validator_class is ZayavkaExcelValidator
    assert "заглушка" not in spec.title.lower()
    assert spec.title == ZayavkaExcelGenerator.TITLE
    assert spec.module == ZayavkaExcelGenerator.__module__
    assert generator.template_path().endswith(".xlsx")


def test_prompt_of_every_type_is_available(five_generators):
    """
    Промпты пяти типов в одном прогоне: у четырёх свой, у перевозки — нет.

    ``perevozka`` сознательно осталась без своего промпта: для неё
    ``get_prompt`` возвращает None, и распознавание берёт дефолтный промпт
    GigaChat (core/prompts/perevozka.py). Это НЕ регрессия — так и задумано.
    А вот у Хавалов промпт обязан быть: он единственный источник имён
    полей, и без него распознавание вернуло бы чужую схему.
    """
    prompts = {key: get_prompt(key) for key in FIVE_TYPES}

    assert prompts["perevozka"] is None
    for key, prompt in prompts.items():
        if key == "perevozka":
            continue
        assert prompt, f"тип {key!r}: промпт пуст"
        assert prompt.strip(), f"тип {key!r}: промпт из пробелов"

    assert prompts["zayavka_excel"] == PROMPT
    assert "СХЕМА ОТВЕТА" in prompts["zayavka_excel"]


def test_every_type_can_still_work_after_the_excel_one(generator, saved,
                                                       five_generators):
    """
    Прогон Хавалов не мешает остальным типам: они работают после него.

    Генератор Хавалов — единственный, кто пишет .xlsx, и он не наследуется
    от DOCX-базы. Если бы он тянул за собой чужое состояние (общие
    словари шаблонов, общий реестр), это вылезло бы здесь.
    """
    path = saved(make_data(vehicles=make_vehicles(2)))
    assert read_back(generator, path)["vehicles"] == make_vehicles(2)

    docx_keys = [key for key in FIVE_TYPES if key != "zayavka_excel"]
    for key in docx_keys:
        spec = ContractTypeRegistry.get(key, strict=True)
        rebuilt = GeneratorFactory.get_generator(key, strict=True)

        assert spec.generator_class is FIVE_TYPES[key][0]
        assert isinstance(rebuilt, spec.generator_class)
        assert rebuilt.templates == five_generators[key].templates
        assert Path(rebuilt.default_output_dir()) == \
            Path(five_generators[key].default_output_dir())

    # Экземпляр Excel-типа из фабрики тоже цел: шаблон тот же.
    assert Path(
        five_generators["zayavka_excel"].template_path()
    ) == Path(generator.template_path())


def test_contract_type_enum_lists_all_five(five_generators):
    """Пять ключей есть и в перечислении ContractType, а не только в реестре."""
    values = {member.value for member in ContractType}

    assert set(FIVE_TYPES) <= values
    assert ContractType.ZAYAVKA_EXCEL.value == "zayavka_excel"


# ─────────────────────────────────────────────────────────────
# Логи без ПДн
# ─────────────────────────────────────────────────────────────

def log_messages(caplog) -> str:
    """Сообщения генератора, валидатора и реестра — одной строкой."""
    names = (
        "core.contract_generator",
        "core.contracts.zayavka.validator",
        "core.contracts.factory",
        "core.contracts.registry",
    )
    return "\n".join(
        record.getMessage() for record in caplog.records
        if record.name in names
    )


def test_full_cycle_logs_have_no_personal_data(generator, stage, caplog):
    """
    Полный круг на десяти машинах: валидация → файл → чтение → файл.

    Файл кладётся рядом с формами прогона (фикстура stage), а не в общую
    папку вывода: тест про логи, и уборка за ним — общая, в finally.
    """
    data = make_data(vehicles=make_vehicles(MAX_VEHICLES))
    output = stage
    # Второй файл — в подпапку: имя зависит только от даты заявки, и
    # запись в ту же папку целилась бы в первый файл (генератор такое
    # запрещает, см. test_same_date_edit_is_refused).
    refill = stage / "refill"

    with caplog.at_level(logging.DEBUG):
        report = ZayavkaExcelValidator().check(data)
        first = Path(generator.generate(data, str(output)))
        generator.read_template(str(first))
        second = Path(generator.fill_from_template(str(first), data, str(refill)))
        generator.read_template(str(second))

    messages = log_messages(caplog)
    assert report.errors == []
    assert messages
    for fragment in PII_FRAGMENTS:
        assert fragment not in messages, f"в логе есть «{fragment}»"


def test_logs_report_work_done(generator, saved, caplog):
    """В логе видно, ЧТО сделано: сколько машин и сколько строк удалено."""
    with caplog.at_level(logging.INFO, logger="core.contract_generator"):
        path = saved(make_data(vehicles=make_vehicles(3)))
        read_back(generator, path)

    messages = log_messages(caplog)
    assert "машин=3" in messages
    assert "Заявка Хавалов" in messages
    assert "Прочитан бланк" in messages


def test_logs_have_no_cell_values(generator, saved, caplog):
    """
    Значения ячеек в лог не пишутся вообще.

    В логе о ячейке есть координата, имя поля и длина значения — по длине
    видно, что данные доехали, а сами данные остаются в файле.
    """
    with caplog.at_level(logging.DEBUG, logger="core.contract_generator"):
        saved(make_data(vehicles=make_vehicles(1)))

    messages = log_messages(caplog)
    assert "длина значения=" in messages
    # Ни одного значения из данных: ни VIN, ни ФИО, ни адреса.
    assert not re.search(re.escape("EC3TEUMB0T0002601"), messages)
    assert not re.search(re.escape(make_zayavka()["driver_phone"]), messages)
