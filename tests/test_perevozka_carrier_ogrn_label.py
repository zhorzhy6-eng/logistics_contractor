#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ОГРН / ОГРНИП у перевозчика (ШАГ «Фикс сохранения + динамические стороны
в шаблонах», часть C.3).

Что было
--------
В бланках ИП стояло захардкоженное «ОГРНИП {{carrier_ogrn}}», а в бланке ООО —
«ОГРН {{carrier_ogrn}}»: метка не зависела от данных, и договор с ООО в бланке
ИП печатал «ОГРНИП» у тринадцатизначного номера. Заодно в п. 1.2 стояло
жёсткое «действующего» (у ИП — «действующий») и всегда печаталось «КПП ,»
даже когда КПП нет.

Что стало
---------
П. 1.2 стал условным блоком по ключу `is_carrier_ip`, метка берётся из
`carrier_ogrn_label` (ОГРН / ОГРНИП), причастие — из `carrier_acting`,
а строка КПП печатается только при заполненном КПП.

Данные синтетические, ПДн нет.
"""

import re

import pytest
from docx import Document

from core.contract_generator import ContractGenerator

IP_CARRIER = {
    "full_name": "Индивидуальный предприниматель Смирнов Алексей Николаевич",
    "short_name": "", "inn": "770123456789", "kpp": "770701001",
    "ogrn": "315770000000012", "legal_address": "г. Волгоград",
    "actual_address": "", "bank_account": "", "bik": "",
    "correspondent_account": "", "bank_name": "",
    "director_name": "", "director_position": "",
    "phone": "", "email": "", "entity_type": "ИП",
}

OOO_CARRIER = {
    "full_name": "ООО «Перевозчик»", "short_name": "ООО «Перевозчик»",
    "inn": "7701234567", "kpp": "770101001", "ogrn": "1027700132195",
    "legal_address": "г. Москва", "actual_address": "",
    "bank_account": "40702810000000000002", "bik": "044525225",
    "correspondent_account": "30101810400000000225", "bank_name": "ПАО Сбербанк",
    "director_name": "Сидоров Сидор Сидорович",
    "director_position": "Генеральный директор",
    "phone": "", "email": "", "entity_type": "ООО",
}

CUSTOMER = {
    "full_name": "ООО «Заказчик»", "short_name": "ООО «Заказчик»",
    "inn": "7707654321", "kpp": "770701001", "ogrn": "1027700132195",
    "legal_address": "г. Москва", "actual_address": "",
    "bank_account": "40702810000000000001", "bik": "044525225",
    "correspondent_account": "30101810400000000225", "bank_name": "ПАО Сбербанк",
    "director_name": "Петров Пётр Петрович",
    "director_position": "Генеральный директор",
    "phone": "", "email": "",
}


@pytest.fixture
def generator(templates_dir) -> ContractGenerator:
    return ContractGenerator(templates_dir=str(templates_dir))


def _payload(carrier, carrier_type) -> dict:
    return {
        "driver": {"full_name": "Иванов Иван Иванович"},
        "carrier": carrier, "customer": CUSTOMER,
        "vehicles": [{"vin": "EC3TEUMB0T0002608", "brand_model": "JETOUR T2",
                      "vehicle_type": "Легковой автомобиль"}],
        "tractor": {"brand_model": "Foton", "plate_number": "O844XY196"},
        "trailer": {"brand_model": "YANGMINDA", "plate_number": "71ABF18"},
        "contract": {
            "number": "OGRN-1", "date": "2026-10-09",
            "carrier_type": carrier_type,
            "vat_rate": "22%" if carrier_type.startswith("ООО") else "0%",
            "vat_rate_num": 22 if carrier_type.startswith("ООО") else 0,
            "price_without_vat": 180300.0, "payment_days": 10,
        },
        "loadings": [{"address": "г. Воронеж", "date": "2026-10-10",
                      "time_window": ""}],
        "unloadings": [{"address": "г. Москва", "date": "2026-10-13",
                        "time_window": ""}],
    }


def _render(generator, carrier, carrier_type, output) -> Document:
    generator.generate_docx(_payload(carrier, carrier_type), str(output))
    assert output.exists(), f"файл не создан: {output}"
    return Document(str(output))


def _flatten(text: str) -> str:
    """Текст одной строкой: пробелы (в том числе неразрывные) по одному."""
    return re.sub(r"\s+", " ", text.replace("\u00a0", " ")).strip()


def _document_text(doc) -> str:
    parts = [_flatten(p.text) for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(_flatten(cell.text))
    return "\n".join(parts)


def _clause_text(doc, label: str) -> str:
    """Абзац, идущий сразу за строкой-меткой («1.2. Перевозчик:»)."""
    texts = [p.text for p in doc.paragraphs]
    for index, text in enumerate(texts):
        if text.strip().startswith(label):
            for following in texts[index + 1:]:
                if following.strip():
                    return _flatten(following)
    raise AssertionError(f"абзац {label!r} не найден")


# ─────────────────────────────────────────────────────────────
# П. 1.2: метка ОГРН / ОГРНИП
# ─────────────────────────────────────────────────────────────

def test_ooo_carrier_ogrn_label_is_ogrn(generator, work_file):
    """Перевозчик-ООО: в п. 1.2 «ОГРН» и тринадцать цифр."""
    doc = _render(generator, OOO_CARRIER, "ООО (с НДС)",
                  work_file("ogrn_ooo.docx"))

    clause = _clause_text(doc, "1.2. Перевозчик:")
    assert "ОГРН 1027700132195" in clause
    assert "ОГРНИП" not in clause
    assert re.search(r"ОГРН \d{13}\b", clause), "метка ОГРН без тринадцати цифр"


def test_ip_carrier_ogrn_label_is_ogrnip(generator, work_file):
    """Перевозчик-ИП: в п. 1.2 «ОГРНИП» и пятнадцать цифр."""
    doc = _render(generator, IP_CARRIER, "ИП без НДС",
                  work_file("ogrn_ip.docx"))

    clause = _clause_text(doc, "1.2. Перевозчик:")
    assert "ОГРНИП 315770000000012" in clause
    assert re.search(r"ОГРНИП \d{15}\b", clause), "метка ОГРНИП без цифр"


def test_ooo_carrier_keeps_kpp_in_1_2(generator, work_file):
    """У ООО КПП в п. 1.2 на месте — он есть в данных."""
    doc = _render(generator, OOO_CARRIER, "ООО (с НДС)",
                  work_file("ogrn_ooo_kpp.docx"))

    clause = _clause_text(doc, "1.2. Перевозчик:")
    assert "КПП 770101001," in clause


def test_ip_carrier_no_kpp_in_1_2(generator, work_file):
    """У ИП КПП не существует: ни значения, ни «КПП ,» в п. 1.2."""
    doc = _render(generator, IP_CARRIER, "ИП без НДС",
                  work_file("ogrn_ip_kpp.docx"))

    clause = _clause_text(doc, "1.2. Перевозчик:")
    assert "КПП" not in clause
    assert "КПП ," not in _document_text(doc)


def test_ooo_carrier_kpp_absent_prints_no_comma(generator, work_file):
    """
    ООО без КПП: строка КПП исчезает вместе с запятой.

    Именно так выглядела «дырка» в договоре: «ИНН 7701234567, КПП , ОГРН …».
    """
    carrier = dict(OOO_CARRIER, kpp="")
    doc = _render(generator, carrier, "ООО (с НДС)",
                  work_file("ogrn_ooo_no_kpp.docx"))

    clause = _clause_text(doc, "1.2. Перевозчик:")
    assert "КПП" not in clause
    assert "ИНН 7701234567, ОГРН 1027700132195" in clause
    assert "КПП ," not in _document_text(doc)


# ─────────────────────────────────────────────────────────────
# П. 1.2: «в лице» только у ООО
# ─────────────────────────────────────────────────────────────

def test_ip_carrier_acts_himself(generator, work_file):
    """ИП действует сам: «в лице» и «именуемое» в его ветке не печатаются."""
    doc = _render(generator, IP_CARRIER, "ИП без НДС",
                  work_file("ogrn_ip_acting.docx"))

    clause = _clause_text(doc, "1.2. Перевозчик:")
    assert "действующий на основании свидетельства о государственной регистрации" in clause
    assert "именуемый в дальнейшем «Перевозчик»" in clause
    assert "в лице" not in clause


def test_ooo_carrier_acts_through_director(generator, work_file):
    """ООО действует через директора: формулировка прежняя."""
    doc = _render(generator, OOO_CARRIER, "ООО (с НДС)",
                  work_file("ogrn_ooo_acting.docx"))

    clause = _clause_text(doc, "1.2. Перевозчик:")
    assert "в лице директора Сидоров Сидор Сидорович, действующего" in clause
    assert "именуемое в дальнейшем «Перевозчик»" in clause


def test_ip_carrier_signature_has_his_name(generator, work_file):
    """П. 9: ИП подписывает сам — в подписи его ФИО, а не пустые «//»."""
    doc = _render(generator, IP_CARRIER, "ИП без НДС",
                  work_file("ogrn_ip_sign.docx"))

    text = _document_text(doc)
    assert (
        "Индивидуальный предприниматель ________ "
        "/Смирнов Алексей Николаевич/" in text
    )
    assert "//" not in text, "пустое ФИО в подписи"
