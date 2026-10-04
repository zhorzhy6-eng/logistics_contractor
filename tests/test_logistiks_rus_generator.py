#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты генератора заявки «Логистикс Рус» (ЭТАП 3.1.C.A.3).

Проверяют: тип больше не заглушка; оба бланка (ООО и ИП) рендерятся в
изолированную папку; все плейсхолдеры получают значения; в ООО-варианте три
суммы (без НДС, НДС по ставке, итого), в ИП-варианте одна сумма «Без НДС»
без плейсхолдеров НДС; таблица автомобилей рассчитана на переменное число
машин (лишние строки удаляются постобработкой); блоки грузоотправителей и
грузополучателей сверх фактического числа точек удаляются; тягач и прицеп в
груз не попадают; логи идут в канал «core.contract_generator» и не содержат
персональных данных.

Точки маршрута в тестовых данных вложены в contract["loadings"] /
contract["unloadings"]: только этот путь сохраняет название
грузоотправителя/грузополучателя — при приведении данных через ContractData
поле name у точек отбрасывается (core.contract_data._as_point_list), см.
докстринг core/contracts/logistiks_rus/generator.py. Поведение с точками
верхнего уровня описано отдельным тестом.

Шаги постобработки по отдельности (на программно собранном Document())
проверяются в tests/test_logistiks_rus_postprocess.py; здесь остаются
интеграционные проверки — тот же результат в готовом документе после
generate().

Все данные синтетические, реальных ПДн нет.
"""

import gc
import logging
import re
from pathlib import Path

import pytest
from docx import Document

from core.contracts.base_generator import ConvertNewlinesStep
from core.contracts.factory import GeneratorFactory
from core.contracts.logistiks_rus.generator import (
    DEFAULT_SPECIAL_CONDITIONS,
    LogistiksRusGenerator,
)
from core.contracts.logistiks_rus.postprocess import (
    RemoveEmptyShipperConsigneeBlocksStep,
    RemoveEmptyVehicleRowsStep,
)
from core.contracts.registry import ContractTypeRegistry

CARGO_HEADERS = ("№", "Марка, модель", "VIN-номер")

PLACEHOLDER_RE = re.compile(r"\{\{[^{}]*\}\}")

#: Варианты бланка: ключ — carrier_type, значение — файл шаблона.
CARRIER_TYPES = {
    "ООО": "ООО (с НДС)",
    "ИП": "ИП без НДС",
}


@pytest.fixture(autouse=True)
def _release_documents():
    """Word-файлы на Windows освобождаем до удаления временных копий."""
    yield
    gc.collect()


@pytest.fixture(scope="module", autouse=True)
def _builtin_types_loaded():
    """Регистрация типов: без неё фабрика не знает про logistiks_rus."""
    ContractTypeRegistry.load_builtin()


@pytest.fixture
def generator(templates_dir) -> LogistiksRusGenerator:
    return LogistiksRusGenerator(templates_dir=str(templates_dir))


# ─────────────────────────────────────────────────────────────
# Тестовые данные
# ─────────────────────────────────────────────────────────────

def _vehicle(number: int) -> dict:
    """Синтетическая перевозимая машина."""
    return {
        "vin": f"TESTVIN000000000{number:02d}",
        "brand_model": f"МОДЕЛЬ {number}",
        "vehicle_type": "Легковой автомобиль",
    }


def _point(kind: str, number: int, name=None, address=None) -> dict:
    """
    Точка маршрута: грузоотправитель (kind="shipper") или грузополучатель.

    name/address можно подменить (в том числе пустой строкой) — так
    проверяются частично заполненные блоки.
    """
    labels = {"shipper": "Грузоотправитель", "consignee": "Грузополучатель"}
    return {
        "name": f"ООО «{labels[kind]} {number}»" if name is None else name,
        "address": f"Адрес {kind} {number}" if address is None else address,
        "date": "2026-09-26" if kind == "shipper" else "2026-10-01",
        "time_window": "08:00-20:00",
    }


def _payload(
    carrier_type: str = "ООО (с НДС)",
    cars: int = 3,
    shippers: int = 2,
    consignees: int = 2,
    price=221099.18,
    loadings=None,
    unloadings=None,
    contract_extra=None,
) -> dict:
    """Синтетические данные заявки «Логистикс Рус»."""
    contract = {
        "number": "ЛР-2026-1",
        "date": "2026-09-24",
        "carrier_type": carrier_type,
        "vat_rate_num": 22,
        "price_without_vat": price,
        "loading_date": "2026-09-26",
        "loading_time_from": "08:00",
        "loading_time_to": "20:00",
        "unloading_date": "2026-10-01",
        "unloading_time_from": "08:00",
        "unloading_time_to": "20:00",
        "loadings": (
            [_point("shipper", n) for n in range(1, shippers + 1)]
            if loadings is None else loadings
        ),
        "unloadings": (
            [_point("consignee", n) for n in range(1, consignees + 1)]
            if unloadings is None else unloadings
        ),
    }
    if contract_extra:
        contract.update(contract_extra)

    return {
        "driver": {"full_name": "Иванов Иван Иванович"},
        "customer": {"full_name": "ООО «Заказчик Тест»"},
        "vehicles": [_vehicle(number) for number in range(1, cars + 1)] + [
            # Тягач и полуприцеп лежат в справочнике машин, но грузом не являются.
            {"brand_model": "Тягач-Модель", "plate_number": "А001АА01",
             "vehicle_type": "Тягач"},
            {"brand_model": "Прицеп-Модель", "plate_number": "Б002ББ02",
             "vehicle_type": "Полуприцеп"},
        ],
        "tractor": {"brand_model": "Тягач-Модель", "plate_number": "А001АА01"},
        "trailer": {"brand_model": "Прицеп-Модель", "plate_number": "Б002ББ02"},
        "contract": contract,
    }


def _document_text(doc) -> str:
    """Весь текст документа: абзацы и таблицы (включая вложенные)."""
    parts = [p.text for p in doc.paragraphs]

    def table_text(table):
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
                for nested in cell.tables:
                    table_text(nested)

    for table in doc.tables:
        table_text(table)

    return "\n".join(parts)


def _body_texts(doc):
    """Тексты абзацев верхнего уровня без крайних пробелов."""
    return [p.text.strip() for p in doc.paragraphs]


def _cargo_table(doc):
    """Таблица автомобилей — по заголовкам колонок, а не по индексу."""
    for table in doc.tables:
        if not table.rows:
            continue
        headers = tuple(cell.text.strip() for cell in table.rows[0].cells)
        if headers == CARGO_HEADERS:
            return table
    return None


def _count_lines(doc, prefix: str) -> int:
    """Сколько абзацев начинается с указанной метки."""
    return sum(1 for text in _body_texts(doc) if text.startswith(prefix))


def _generate(generator, payload, output_dir) -> str:
    path = generator.generate(payload, output_dir=str(output_dir))
    assert Path(path).exists(), f"файл не создан: {path}"
    return path


# ─────────────────────────────────────────────────────────────
# Тип зарегистрирован и больше не заглушка
# ─────────────────────────────────────────────────────────────

def test_generator_is_registered_and_not_stub():
    spec = ContractTypeRegistry.get("logistiks_rus", strict=True)
    assert spec.generator_class is LogistiksRusGenerator
    assert isinstance(
        GeneratorFactory.get_generator("logistiks_rus", strict=True),
        LogistiksRusGenerator,
    )


def test_type_metadata_matches_templates(generator):
    assert LogistiksRusGenerator.CONTRACT_TYPE == "logistiks_rus"
    assert LogistiksRusGenerator.FILE_PREFIX == "Заявка_Логистикс_Рус"
    assert LogistiksRusGenerator.TEMPLATE_NAMES == {
        "ООО": "shablon_logistiks_rus_ooo.docx",
        "ИП": "shablon_logistiks_rus_ip.docx",
    }
    assert LogistiksRusGenerator.DEFAULT_VAT_RATE == 22.0
    assert generator.templates["ООО"].endswith("shablon_logistiks_rus_ooo.docx")
    assert generator.templates["ИП"].endswith("shablon_logistiks_rus_ip.docx")


def test_generate_does_not_raise_not_implemented(generator, work_dir):
    path = _generate(generator, _payload(cars=1), work_dir)
    assert path.endswith(".docx")


def test_filename_uses_prefix_number_and_date(generator, work_dir):
    path = Path(_generate(generator, _payload(cars=1), work_dir))
    assert path.name.startswith(f"{LogistiksRusGenerator.FILE_PREFIX}_ЛР-2026-1_")


def test_default_output_dir_is_output():
    assert Path(LogistiksRusGenerator.default_output_dir()).name == "output"


# ─────────────────────────────────────────────────────────────
# Выбор шаблона и конвейер постобработки
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize(
    "carrier_type, filename",
    [
        ("ООО (с НДС)", "shablon_logistiks_rus_ooo.docx"),
        ("ООО", "shablon_logistiks_rus_ooo.docx"),
        ("", "shablon_logistiks_rus_ooo.docx"),
        ("ИП без НДС", "shablon_logistiks_rus_ip.docx"),
        ("ИП", "shablon_logistiks_rus_ip.docx"),
        ("ИП Хейгетян Е.В.", "shablon_logistiks_rus_ip.docx"),
    ],
)
def test_template_path_selects_variant(generator, carrier_type, filename):
    assert generator._get_template_path(carrier_type).endswith(filename)


def test_postprocess_steps_order(generator):
    steps = generator.postprocess_steps(None)
    assert [type(step) for step in steps] == [
        ConvertNewlinesStep,
        RemoveEmptyVehicleRowsStep,
        RemoveEmptyShipperConsigneeBlocksStep,
    ]
    assert [step.name for step in steps] == [
        "convert_newlines",
        "remove_empty_vehicle_rows",
        "remove_empty_shipper_consignee_blocks",
    ]


def test_insert_route_tables_is_noop(generator):
    """Таблиц по точкам маршрута в бланке нет — метод обязан быть безопасным."""
    doc = Document()
    doc.add_paragraph("текст")
    assert generator._insert_route_tables(doc, None) is None
    assert len(doc.paragraphs) == 1


# ─────────────────────────────────────────────────────────────
# Рендер обоих шаблонов
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("variant", list(CARRIER_TYPES))
def test_both_templates_render_to_isolated_dir(generator, work_dir, variant):
    """ООО- и ИП-бланк рендерятся в изолированную папку без исключений."""
    output_dir = work_dir / f"lr_{variant}"
    output_dir.mkdir(parents=True, exist_ok=True)

    path = Path(_generate(generator, _payload(CARRIER_TYPES[variant], cars=2), output_dir))
    assert path.parent == output_dir, "файл ушёл мимо изолированной папки"

    text = _document_text(Document(path))
    assert "{{" not in text and "}}" not in text

    # Экспедитор — свой у каждого варианта (подставляется шаблоном).
    if variant == "ООО":
        assert "Экспедитор: ООО «ТЕХНОЛОГИСТИКА»" in text
        assert "ИП Хейгетян" not in text
    else:
        assert "Экспедитор: ИП Хейгетян Е.В." in text
        assert "ООО «ТЕХНОЛОГИСТИКА»" not in text

    assert "________________ / Гао Фанфан /" in text


@pytest.mark.parametrize("variant", list(CARRIER_TYPES))
def test_all_fields_are_substituted(generator, work_dir, variant):
    """Каждое поле данных попадает в документ своего варианта."""
    output_dir = work_dir / f"lr_fields_{variant}"
    output_dir.mkdir(parents=True, exist_ok=True)
    path = _generate(generator, _payload(CARRIER_TYPES[variant], cars=2), output_dir)

    doc = Document(path)
    text = _document_text(doc)
    texts = _body_texts(doc)

    # Шапка и заказчик.
    assert "ЗАЯВКА № ЛР-2026-1" in text
    assert "Дата Заявки: «24» сентября 2026 года." in text
    assert "Заказчик: ООО «Заказчик Тест»" in text

    # Грузоотправители и грузополучатели.
    assert "Грузоотправитель: ООО «Грузоотправитель 1»" in texts
    assert "Адрес погрузки: Адрес shipper 1" in texts
    assert "Грузополучатель №1: ООО «Грузополучатель 1»" in texts
    assert "Адрес выгрузки: Адрес consignee 1" in texts

    # План погрузки и выгрузки.
    assert "Дата / время погрузки: 26.09.2026 г. Время с 08:00 по 20:00" in texts
    assert ("Плановая дата / время завершения выгрузки: 01.10.2026 г. "
            "Время с 08:00 по 20:00") in texts

    # Груз и автовоз.
    assert "Общее количество: 2 шт." in texts
    assert "Тягач: Тягач-Модель гос. №: А001АА01" in texts
    assert "Прицеп: Прицеп-Модель гос. №: Б002ББ02" in texts
    assert "Водитель: Иванов Иван Иванович" in texts

    table = _cargo_table(doc)
    assert [cell.text.strip() for cell in table.rows[1].cells] == [
        "1", "МОДЕЛЬ 1", "TESTVIN00000000001"
    ]
    assert [cell.text.strip() for cell in table.rows[2].cells] == [
        "2", "МОДЕЛЬ 2", "TESTVIN00000000002"
    ]


def test_special_conditions_default_matches_sample(generator):
    """Пустые особые условия заменяются отпиской образца."""
    replacements = generator._build_replacements_map(_payload())
    assert replacements["special_conditions"] == DEFAULT_SPECIAL_CONDITIONS


def test_special_conditions_from_data_are_used(generator, work_dir):
    payload = _payload(contract_extra={"special_conditions": "Погрузка круглосуточно."})
    replacements = generator._build_replacements_map(payload)
    assert replacements["special_conditions"] == "Погрузка круглосуточно."

    path = _generate(generator, payload, work_dir)
    assert "Погрузка круглосуточно." in _document_text(Document(path))


# ─────────────────────────────────────────────────────────────
# Стоимость: три суммы для ООО, одна для ИП
# ─────────────────────────────────────────────────────────────

def test_ooo_cost_has_three_sums(generator, work_dir):
    """ООО: без НДС + НДС по ставке = итого (как в образце 221 099,18)."""
    output_dir = work_dir / "lr_cost_ooo"
    output_dir.mkdir(parents=True, exist_ok=True)
    path = _generate(generator, _payload("ООО (с НДС)"), output_dir)
    texts = _body_texts(Document(path))

    assert ("Стоимость услуг: 221\u00a0099,18 руб. "
            "(Двести двадцать одна тысяча девяносто девять рублей "
            "восемнадцать копеек)") in texts
    assert ("НДС 22%: 48\u00a0641,82 руб. "
            "(Сорок восемь тысяч шестьсот сорок один рубль восемьдесят "
            "две копейки)") in texts
    assert ("Итого: 269\u00a0741,00 руб. "
            "(Двести шестьдесят девять тысяч семьсот сорок один рубль "
            "00 копеек)") in texts


def test_ooo_cost_values_in_replacements(generator):
    replacements = generator._build_replacements_map(_payload("ООО (с НДС)"))

    assert replacements["vat_rate"] == "22%"
    assert replacements["sum_wo_vat"] == "221\u00a0099,18"
    assert replacements["sum_vat"] == "48\u00a0641,82"
    assert replacements["sum_total"] == "269\u00a0741,00"
    assert replacements["sum_wo_vat_words"].startswith("Двести двадцать одна тысяча")
    assert replacements["sum_vat_words"].startswith("Сорок восемь тысяч")
    assert replacements["sum_total_words"].startswith("Двести шестьдесят девять тысяч")


def test_ip_cost_is_single_sum_without_vat(generator, work_dir):
    """ИП: одна сумма «Без НДС», плейсхолдеров НДС в карте замен нет."""
    payload = _payload("ИП без НДС")
    replacements = generator._build_replacements_map(payload)

    assert replacements["sum_total"] == "221\u00a0099,18"
    assert replacements["sum_total_words"].startswith("Двести двадцать одна тысяча")
    for name in ("sum_wo_vat", "sum_vat", "vat_rate",
                 "sum_wo_vat_words", "sum_vat_words"):
        assert name not in replacements, f"в ИП-вариант положен лишний ключ {name!r}"

    output_dir = work_dir / "lr_cost_ip"
    output_dir.mkdir(parents=True, exist_ok=True)
    texts = _body_texts(Document(_generate(generator, payload, output_dir)))

    assert texts.count("Стоимость услуг: 221\u00a0099,18 руб. "
                       "(Двести двадцать одна тысяча девяносто девять рублей "
                       "восемнадцать копеек) Без НДС") == 1
    assert not [t for t in texts if t.startswith("Итого:")]
    assert not [t for t in texts if t.startswith("НДС ")]


def test_ip_vat_rate_is_zero_even_if_data_has_rate(generator):
    """Ставка в данных ИП не важна: вариант «Без НДС» ставку обнуляет."""
    payload = _payload("ИП без НДС", contract_extra={"vat_rate_num": 22})
    replacements = generator._build_replacements_map(payload)
    assert replacements["sum_total"] == "221\u00a0099,18"
    assert "vat_rate" not in replacements


def test_cost_falls_back_to_recognized_price_input(generator):
    """Без price_without_vat берётся распознанная сумма price_input."""
    payload = _payload()
    payload["contract"].pop("price_without_vat")
    payload["contract"]["price_input"] = "221 099,18"

    replacements = generator._build_replacements_map(payload)

    assert replacements["sum_wo_vat"] == "221\u00a0099,18"
    assert replacements["sum_vat"] == "48\u00a0641,82"
    assert replacements["sum_total"] == "269\u00a0741,00"


def test_cost_missing_is_zero_not_invented(generator):
    payload = _payload()
    payload["contract"].pop("price_without_vat")

    replacements = generator._build_replacements_map(payload)

    assert replacements["sum_wo_vat"] == "0,00"
    assert replacements["sum_vat"] == "0,00"
    assert replacements["sum_total"] == "0,00"
    assert replacements["vat_rate"] == "22%"


def test_vat_rate_can_come_as_text(generator):
    """Ставка строкой «22%» принимается, если числа в данных нет."""
    payload = _payload(contract_extra={"vat_rate": "22%"})
    payload["contract"].pop("vat_rate_num")

    assert generator._build_replacements_map(payload)["vat_rate"] == "22%"


def test_explicit_zero_vat_rate_is_respected(generator):
    """Явный ноль у ООО — это «без НДС», а не отсутствие ставки."""
    replacements = generator._build_replacements_map(
        _payload(contract_extra={"vat_rate_num": 0})
    )

    assert replacements["vat_rate"] == "0%"
    assert replacements["sum_vat"] == "0,00"
    assert replacements["sum_total"] == "221\u00a0099,18"


# ─────────────────────────────────────────────────────────────
# Карта замен покрывает шаблон полностью
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("variant", list(CARRIER_TYPES))
def test_replacements_map_covers_all_template_placeholders(
    generator, templates_dir, variant
):
    """Каждый плейсхолдер бланка получает значение из карты замен."""
    name = LogistiksRusGenerator.TEMPLATE_NAMES[variant]
    template_text = _document_text(Document(str(templates_dir / name)))
    template_names = {
        match.strip("{} ") for match in PLACEHOLDER_RE.findall(template_text)
    }
    assert template_names, "в шаблоне не найдено ни одного плейсхолдера"

    replacements = generator._build_replacements_map(
        _payload(CARRIER_TYPES[variant])
    )
    missing = template_names - set(replacements)

    assert not missing, f"нет значений для плейсхолдеров ({variant}): {sorted(missing)}"


@pytest.mark.parametrize("variant", list(CARRIER_TYPES))
def test_replacements_map_has_no_extra_placeholders(
    generator, templates_dir, variant
):
    """Лишних ключей карта не содержит — только плейсхолдеры своего бланка."""
    name = LogistiksRusGenerator.TEMPLATE_NAMES[variant]
    template_text = _document_text(Document(str(templates_dir / name)))
    template_names = {
        match.strip("{} ") for match in PLACEHOLDER_RE.findall(template_text)
    }

    replacements = generator._build_replacements_map(
        _payload(CARRIER_TYPES[variant])
    )

    assert set(replacements) == template_names


def test_generated_document_has_no_placeholders_left(generator, work_dir):
    for variant, carrier_type in CARRIER_TYPES.items():
        output_dir = work_dir / f"lr_leftovers_{variant}"
        output_dir.mkdir(parents=True, exist_ok=True)
        path = _generate(generator, _payload(carrier_type), output_dir)
        text = _document_text(Document(path))

        assert "{{" not in text, f"в документе {variant} остался плейсхолдер"
        assert "}}" not in text


def test_generate_with_empty_data_does_not_crash(generator, work_dir):
    """Пустые данные — пустой бланк без исключений (заглушки больше нет)."""
    path = _generate(generator, {"contract": {}}, work_dir)
    text = _document_text(Document(path))

    assert "ЗАЯВКА №" in text
    assert "{{" not in text
    # Заказчика нет — пустая строка вместо выдуманного названия.
    assert "Заказчик:" in text
    assert "«Заказчик Тест»" not in text


def test_missing_template_raises_file_not_found(work_dir):
    generator = LogistiksRusGenerator(templates_dir=str(work_dir / "нет-такого"))
    with pytest.raises(FileNotFoundError):
        generator.generate(_payload(cars=1), output_dir=str(work_dir))


# ─────────────────────────────────────────────────────────────
# Таблица автомобилей: переменное число машин
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("cars", [1, 2, 3, 12])
def test_cargo_table_keeps_only_filled_rows(generator, work_dir, cars):
    output_dir = work_dir / f"lr_cars_{cars}"
    output_dir.mkdir(parents=True, exist_ok=True)
    path = _generate(generator, _payload(cars=cars), output_dir)
    table = _cargo_table(Document(path))

    assert table is not None, "таблица автомобилей не найдена в документе"
    assert len(table.rows) == cars + 1, (
        f"машин {cars}, а строк в таблице {len(table.rows)} (ожидалось {cars + 1})"
    )
    for number in range(1, cars + 1):
        cells = [cell.text.strip() for cell in table.rows[number].cells]
        assert cells == [
            str(number),
            f"МОДЕЛЬ {number}",
            f"TESTVIN000000000{number:02d}",
        ]


def test_generate_without_vehicles_keeps_header_only(generator, work_dir):
    """Ни одной машины — в документе остаётся только шапка таблицы."""
    path = _generate(generator, _payload(cars=0), work_dir)
    table = _cargo_table(Document(path))

    assert len(table.rows) == 1
    assert generator._build_replacements_map(_payload(cars=0))["cargo_count"] == "0"


def test_tractor_and_trailer_are_not_in_cargo_table(generator, work_dir):
    path = _generate(generator, _payload(cars=2), work_dir)
    table = _cargo_table(Document(path))
    text = "\n".join(cell.text for row in table.rows for cell in row.cells)

    assert "Тягач-Модель" not in text, "тягач попал в таблицу груза"
    assert "Прицеп-Модель" not in text, "прицеп попал в таблицу груза"


def test_cargo_count_matches_vehicles(generator):
    for cars in (0, 1, 5, 12):
        replacements = generator._build_replacements_map(_payload(cars=cars))
        assert replacements["cargo_count"] == str(cars)


def test_more_than_twelve_cars_are_truncated(generator, caplog):
    with caplog.at_level(logging.WARNING, logger="core.contract_generator"):
        replacements = generator._build_replacements_map(_payload(cars=14))

    assert replacements["cargo_count"] == "14"
    assert replacements["car_12_brand"] == "МОДЕЛЬ 12"
    assert "car_13_brand" not in replacements
    assert "в бланк помещается 12" in caplog.text


# ─────────────────────────────────────────────────────────────
# Интеграция: постобработка в готовом документе
# ─────────────────────────────────────────────────────────────

def test_empty_blocks_removed_in_generated_document(generator, work_dir):
    """В готовом документе остаются только фактически заполненные блоки."""
    output_dir = work_dir / "lr_points"
    output_dir.mkdir(parents=True, exist_ok=True)
    path = _generate(
        generator, _payload(cars=1, shippers=2, consignees=3), output_dir
    )
    doc = Document(path)

    assert _count_lines(doc, "Грузоотправитель:") == 2
    assert _count_lines(doc, "Адрес погрузки:") == 2
    assert _count_lines(doc, "Грузополучатель №") == 3
    assert _count_lines(doc, "Адрес выгрузки:") == 3

    texts = _body_texts(doc)
    assert "Грузополучатель №3: ООО «Грузополучатель 3»" in texts
    assert not [t for t in texts if t.startswith("Грузополучатель №4")]
    # Строки дат остаются на месте: удаляются только блоки точек.
    assert [t for t in texts if t.startswith("Дата / время погрузки:")]
    assert [t for t in texts if t.startswith("Плановая дата / время завершения выгрузки:")]


def test_points_from_top_level_still_render_addresses(generator, work_dir):
    """
    Точки верхнего уровня: адреса доходят до бланка.

    Название точки при этом теряется — ContractData хранит точки как
    {address, date, time_window} (core.contract_data._as_point_list), и по
    пути generate() исходный словарь с name до генератора не доживает.
    Чтобы имя доживало и здесь, нужно расширить _as_point_list; на этом шаге
    core/contract_data.py не трогаем и фиксируем фактическое поведение:
    блок с одним адресом остаётся.
    """
    output_dir = work_dir / "lr_top_level"
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = _payload(cars=1)
    loadings = payload["contract"].pop("loadings")
    payload["loadings"] = loadings

    path = _generate(generator, payload, output_dir)
    doc = Document(path)

    assert _count_lines(doc, "Грузоотправитель:") == 2
    assert "Адрес погрузки: Адрес shipper 1" in _body_texts(doc)


def test_route_points_use_names_from_raw_data(generator):
    """При прямом вызове карты замен название берётся из исходных точек."""
    payload = _payload(shippers=1, consignees=1)
    loadings = payload["contract"].pop("loadings")
    unloadings = payload["contract"].pop("unloadings")
    payload["loadings"] = loadings
    payload["unloadings"] = unloadings

    replacements = generator._build_replacements_map(payload)

    assert replacements["shipper_1_name"] == "ООО «Грузоотправитель 1»"
    assert replacements["shipper_1_address"] == "Адрес shipper 1"
    assert replacements["consignee_1_name"] == "ООО «Грузополучатель 1»"
    assert replacements["consignee_1_address"] == "Адрес consignee 1"


def test_more_than_ten_points_are_truncated(generator, caplog):
    with caplog.at_level(logging.WARNING, logger="core.contract_generator"):
        replacements = generator._build_replacements_map(
            _payload(shippers=12, consignees=12)
        )

    assert replacements["shipper_10_name"] == "ООО «Грузоотправитель 10»"
    assert "shipper_11_name" not in replacements
    assert "в бланк помещается 10" in caplog.text


def test_route_points_fall_back_to_legacy_contract_addresses(generator):
    """Исторические поля contract.loading_address/unloading_address_1/2."""
    payload = _payload()
    payload["contract"].pop("loadings")
    payload["contract"].pop("unloadings")
    payload["contract"]["loading_address"] = "Легаси адрес погрузки"
    payload["contract"]["unloading_address_1"] = "Легаси адрес выгрузки 1"
    payload["contract"]["unloading_address_2"] = "Легаси адрес выгрузки 2"

    replacements = generator._build_replacements_map(payload)

    assert replacements["shipper_1_name"] == ""
    assert replacements["shipper_1_address"] == "Легаси адрес погрузки"
    assert replacements["consignee_1_address"] == "Легаси адрес выгрузки 1"
    assert replacements["consignee_2_address"] == "Легаси адрес выгрузки 2"
    assert replacements["consignee_3_address"] == ""


def test_route_time_falls_back_to_point_window(generator):
    """Пустые поля времени берутся из окна времени точки."""
    payload = _payload(contract_extra={
        "loading_time_from": "", "loading_time_to": "",
        "unloading_time_from": "", "unloading_time_to": "",
    })

    replacements = generator._build_replacements_map(payload)

    assert replacements["loading_time_from"] == "08:00"
    assert replacements["loading_time_to"] == "20:00"
    assert replacements["unloading_time_from"] == "08:00"
    assert replacements["unloading_time_to"] == "20:00"


def test_route_date_falls_back_to_point_date(generator):
    payload = _payload(contract_extra={"loading_date": "", "unloading_date": ""})

    replacements = generator._build_replacements_map(payload)

    assert replacements["loading_date"] == "26.09.2026"
    assert replacements["unloading_date"] == "01.10.2026"


def test_contract_year_comes_from_contract_date(generator):
    replacements = generator._build_replacements_map(_payload())

    assert replacements["contract_date_day"] == "24"
    assert replacements["contract_date_month"] == "сентября"
    assert replacements["contract_date_year"] == "2026"


def test_empty_customer_gives_empty_placeholder(generator):
    """Пустой заказчик — пустая строка, а не выдуманное название."""
    payload = _payload()
    payload["customer"] = {}

    assert generator._build_replacements_map(payload)["customer_name"] == ""


# ─────────────────────────────────────────────────────────────
# Логи
# ─────────────────────────────────────────────────────────────

def test_generation_logs_to_contract_generator(caplog, generator, work_dir):
    with caplog.at_level(logging.INFO, logger="core.contract_generator"):
        _generate(generator, _payload(cars=4, shippers=2, consignees=3), work_dir)

    records = [r for r in caplog.records if r.name == "core.contract_generator"]
    assert records, "генерация ничего не записала в core.contract_generator"

    messages = "\n".join(r.getMessage() for r in records)
    assert "машин в заявке — 4" in messages
    assert "удалено пустых строк таблицы груза: 8" in messages
    assert "удалено пустых блоков грузоотправителей и грузополучателей: 15" in messages
    assert "Логистикс Рус [ООО]: в бланк подставлено" in messages


def test_logs_have_no_personal_data(caplog, generator, work_dir):
    """В логи не попадают ФИО, VIN, адреса, марки машин и названия сторон."""
    with caplog.at_level(logging.DEBUG, logger="core.contract_generator"):
        _generate(generator, _payload(cars=4, shippers=2, consignees=2), work_dir)

    messages = "\n".join(
        r.getMessage() for r in caplog.records
        if r.name == "core.contract_generator"
    )
    for fragment in ("Иванов", "TESTVIN00000000001", "Заказчик Тест",
                     "Адрес shipper 1", "Адрес consignee 1",
                     "Тягач-Модель", "Прицеп-Модель", "МОДЕЛЬ 1",
                     "Грузоотправитель 1", "Грузополучатель 1"):
        assert fragment not in messages, f"в логе есть «{fragment}»"
