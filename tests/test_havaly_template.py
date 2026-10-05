#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты эталонного пустого бланка Хавалов (ЭТАП 3.1.E.A.1).

Проверяют, что templates/shablon_havaly.xlsx — именно ПУСТОЙ БЛАНК,
собранный по образцу templates/Хавалы_образец.xlsx, а не копия файла
заказчика с данными:

  * файл существует и открывается openpyxl;
  * лист РОВНО ОДИН и называется TDSheet, скрытых листов нет;
  * строка 5 содержит подпись 'Дата заявки:' (ищется по ТЕКСТУ, а не по
    номеру строки — так же её будет искать импорт), ячейка даты B5 пустая;
  * строка 6 содержит ровно 32 заголовка A..AF в эталонном порядке;
  * строк данных ровно 10 (R7-R16), все пустые, с рамкой и числовыми
    форматами колонок времени и ставки;
  * нижний блок сторон: 'Заказчик'/'Перевозчик' в колонках C и J,
    'Сюрлогистик' и 'ООО "ТЕХНОЛОГИСТИКА"', затем две пустые строки
    под подписи и 'ФИО, подпись, печать' — ровно как в образце;
  * зелёной обводки вокруг 'Город погрузки' НЕТ: рамка E6 — обычная
    тонкая чёрная, заливок на листе нет вообще, условного форматирования,
    фигур и линий (xl/drawings) в файле нет;
  * данные образца в бланк не попали (в том числе пример даты
    '16.09.2026') — полный список непустых ячеек сверяется поимённо;
  * ширины колонок, высоты строк и формат листа совпадают с образцом;
  * шаблон закреплён по SHA256, и сборщик даёт побайтово тот же файл
    (см. tools/make_havaly_template.py и TEMPLATE_SHA256);
  * сборщик идемпотентен: повторный запуск не перезаписывает актуальный
    бланк и не оставляет временных файлов.

Образец templates/Хавалы_образец.xlsx отслеживается в git (ПДн третьих лиц
в нём нет: только шапка, пример даты и блок сторон), поэтому проверки
«образец не изменён» и «структура совпадает с образцом» работают всегда;
если файла всё же нет — они пропускаются, а не падают.
"""

import hashlib
import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────────────────────
# Эталон: имена файлов, лист, тексты
# ─────────────────────────────────────────────────────────────

TEMPLATE_NAME = "shablon_havaly.xlsx"
SAMPLE_NAME = "Хавалы_образец.xlsx"

#: Лист всегда один и всегда с этим именем.
SHEET_NAME = "TDSheet"

#: SHA256 собранного бланка (ЭТАП 3.1.E.A.1). Если бланк пересобрали
#: осознанно — обнови константу; если правки не было, значит файл
#: отредактировали руками мимо сборщика.
TEMPLATE_SHA256 = (
    "7a4e9716f536bf4b854fb191f1477bc4f201837cf3bccef7c56ccf960f2ae130"
)

#: SHA256 образца-источника: он только читается, из него собирается бланк.
SAMPLE_SHA256 = (
    "106da4ec323d8b850dbdcef2e089043e7de0f889d41f1a466137258c796d1280"
)

#: Заголовок бланка — две строки в одной ячейке B2 (объединение B2:C2).
TITLE = "ЗАЯВКА\n на перевозку автомобилей"

#: Подпись поля даты заявки.
DATE_LABEL = "Дата заявки:"

#: Пример даты из образца: в бланк попасть НЕ должен.
SAMPLE_DATE_VALUE = "16.09.2026"

#: Шапка таблицы (строка 6): ровно 32 колонки A..AF, порядок эталонный.
HEADERS = (
    "Номер Лота",                           # A
    "VIN",                                  # B
    "Марка",                                # C
    "Модель",                               # D
    "Город погрузки",                       # E
    "Пункт погрузки",                       # F
    "Город доставки",                       # G
    "Пункт разгрузки",                      # H
    "Дилер",                                # I
    "Код дилера",                           # J
    "Наименование транспортной компании",   # K
    "Марка Автовоза",                       # L
    "Цвет кабины",                          # M
    "Номер автовоза",                       # N
    "Марка прицепа",                        # O
    "Номер Прицепа",                        # P
    "Фамилия",                              # Q
    "Имя",                                  # R
    "Отчество",                             # S
    "Номер В/У",                            # T
    "Дата выдачи В/У",                      # U
    "Серия Паспорта",                       # V
    "Номер Паспорта",                       # W
    "Кем выдан паспорт",                    # X
    "Дата выдачи паспорта",                 # Y
    "Гражданство",                          # Z
    "Дата Рождения",                        # AA
    "Прописка",                             # AB
    "Планируемая дата\n погрузки",          # AC — перенос строки как в образце
    "Время погрузки",                       # AD
    "№ телефона водителя",                  # AE
    "Ставка с НДС",                         # AF
)

#: Ячейка шапки с переносом текста — как в образце.
HEADER_WRAP_CELLS = ("E6", "G6", "AC6", "AE6")

#: Строки бланка.
TITLE_ROW = 2
DATE_ROW = 5
DATE_CELL = "B5"
HEADER_ROW = 6
FIRST_DATA_ROW = 7
DATA_ROWS = 10
LAST_DATA_ROW = FIRST_DATA_ROW + DATA_ROWS - 1

#: Нижний блок: две пустые строки, стороны, две пустые, подпись.
BLOCK_SPACER_ROWS = (17, 18, 21, 22)
PARTY_ROW = 19
PARTY_NAME_ROW = 20
SIGNATURE_ROW = 23
LAST_ROW = SIGNATURE_ROW

#: Колонки сторон — как в образце (в ТЗ ошибочно указана I).
PARTY_COLUMNS = ("C", "J")

PARTY_LABELS = ("Заказчик", "Перевозчик")
CUSTOMER = "Сюрлогистик"
CARRIER = 'ООО "ТЕХНОЛОГИСТИКА"'
SIGNATURE_TEXT = "ФИО, подпись, печать"

#: Ширины колонок бланка (= ws.column_dimensions образца).
COLUMN_WIDTHS = {
    "A": 18.83203125, "B": 19.33203125, "C": 15.0, "D": 14.33203125,
    "E": 19.1640625, "F": 19.1640625, "G": 19.1640625, "H": 66.33203125,
    "I": 23.0, "J": 18.0, "K": 44.33203125, "L": 22.1640625,
    "M": 19.6640625, "N": 19.6640625, "O": 21.83203125, "P": 21.6640625,
    "Q": 10.5, "R": 10.5, "S": 10.5, "T": 18.33203125, "U": 21.0,
    "V": 26.5, "W": 18.0, "X": 27.1640625, "Y": 24.1640625,
    "Z": 28.1640625, "AA": 26.83203125, "AB": 25.6640625,
    "AC": 28.1640625, "AD": 26.33203125, "AE": 23.33203125,
    "AF": 23.33203125,
}

#: Колонки образца с bestFit=1 (Excel подгоняет ширину под содержимое).
COLUMN_BEST_FIT = ("B", "H", "I")

#: Высоты строк бланка (= ws.row_dimensions образца).
ROW_HEIGHTS = {
    1: 14.25,
    TITLE_ROW: 23.25,
    3: 23.25,
    DATE_ROW: 15.0,
    HEADER_ROW: 23.25,
    **{row: 11.25 for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1)},
    17: 11.25,
    18: 12.75,
    PARTY_ROW: 12.75,
    PARTY_NAME_ROW: 15.75,
    21: 11.25,
    22: 11.25,
    SIGNATURE_ROW: 11.25,
}

#: Числовые форматы колонок данных: время погрузки и ставка с НДС.
TIME_FORMAT = "h:mm"
MONEY_FORMAT = '#,##0" р."'

#: Оформление, унаследованное от образца.
MAIN_FONT_FAMILY = "Arial"
DATA_FONT_SIZE = 8.0
HEADER_FONT_SIZE = 10.0
VALUE_FONT_FAMILY = "Calibri"
VALUE_FONT_COLOR = "FF000000"

#: Единственный цвет, который допустим в применённых стилях бланка.
BLACK_ARGB = "FF000000"

#: Подпись 'Дата заявки:' встречается в листе ровно один раз.
DATE_LABEL_OCCURRENCES = 1

#: Полный список непустых ячеек бланка: шапка, подпись даты и блок сторон.
#: Всё, что не входит в этот набор, — данные образца, которых в бланке
#: быть не должно (персональные данные, VIN, телефоны, адреса, суммы).
EXPECTED_NON_EMPTY_CELLS = 40

#: В образце на одну непустую ячейку больше — заполненная дата заявки.
SAMPLE_NON_EMPTY_CELLS = EXPECTED_NON_EMPTY_CELLS + 1


# ─────────────────────────────────────────────────────────────
# Фикстуры и вспомогательные функции
# ─────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def template_path(templates_dir) -> Path:
    """Путь бланка в templates/."""
    return templates_dir / TEMPLATE_NAME


@pytest.fixture(scope="module")
def sample_path(templates_dir) -> Path:
    """Путь образца в templates/."""
    return templates_dir / SAMPLE_NAME


@pytest.fixture
def template_ws(template_path):
    """
    Лист TDSheet собранного бланка.

    Фикстура не кэшируется на модуль намеренно: тесты читают ячейки за
    границами листа (например, колонку AG), а openpyxl при таком чтении
    создаёт ячейки в памяти и меняет размеры листа — общий объект между
    тестами привёл бы к зависимости результатов от порядка запуска.
    """
    workbook = load_workbook(template_path)
    try:
        yield workbook[SHEET_NAME]
    finally:
        workbook.close()


@pytest.fixture
def sample_ws(sample_path):
    """Лист TDSheet образца — только для сравнения структуры."""
    if not sample_path.exists():
        pytest.skip(f"образец {SAMPLE_NAME} не найден в templates/")
    workbook = load_workbook(sample_path)
    try:
        yield workbook[SHEET_NAME]
    finally:
        workbook.close()


def _sha256(path: Path) -> str:
    """SHA256 файла в нижнем регистре."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _cell_texts(ws) -> dict:
    """Координата → значение для всех непустых ячеек листа."""
    return {cell.coordinate: cell.value
            for row in ws.iter_rows() for cell in row
            if cell.value is not None}


def _row_values(ws, row: int, columns: int = len(HEADERS)) -> list:
    """Значения строки по колонкам 1..columns (None для пустых)."""
    return [ws.cell(row=row, column=index).value
            for index in range(1, columns + 1)]


def _row_is_empty(ws, row: int) -> bool:
    """В строке нет ни одного значения во всех 32 колонках."""
    return all(value is None for value in _row_values(ws, row))


def _effective_widths(ws) -> dict:
    """
    Ширина каждой колонки A..AF с учётом групповых записей <col min max>.

    openpyxl хранит одну запись на группу колонок (например, E, F и G —
    это одна запись min=5 max=7), поэтому обращение по букве F вернуло бы
    ширину по умолчанию. Здесь группа разворачивается по буквам.
    """
    widths = {}
    for dimension in ws.column_dimensions.values():
        if dimension.width is None:
            continue
        for index in range(dimension.min, dimension.max + 1):
            widths[get_column_letter(index)] = dimension.width
    return widths


def _is_green(value) -> bool:
    """
    Похож ли цвет на зелёный.

    Принимает значение вида 'FF00B050' (ARGB), '00B050' (RGB) или None.
    Зелёный — тот, у которого зелёная составляющая строго больше красной
    и синей: именно так выглядела бы «зелёная обводка» из ТЗ.
    """
    if value is None:
        return False
    digits = str(value)
    if digits in ("00000000", "FFFFFFFF"):
        return False
    if len(digits) == 8:
        digits = digits[2:]
    if len(digits) != 6:
        return False
    try:
        red, green, blue = (int(digits[index:index + 2], 16)
                            for index in (0, 2, 4))
    except ValueError:
        return False
    return green > red and green > blue


def _applied_style_colors(path: Path) -> set:
    """
    RGB-цвета, реально применённые в книге (шрифты, заливки, рамки, dxf).

    Легаси-палитра <colors><indexedColors> не учитывается: openpyxl пишет
    в неё стандартные 64 цвета Excel, включая зелёные, но они не
    используются ни одной ячейкой.
    """
    with zipfile.ZipFile(path) as archive:
        styles = archive.read("xl/styles.xml").decode("utf-8")

    colors = set()
    for section in ("fonts", "fills", "borders", "dxfs"):
        match = re.search(rf"<{section}[ >].*?</{section}>", styles, re.S)
        if match:
            colors.update(re.findall(r'rgb="([^"]+)"', match.group(0)))
    return colors


def _archive_parts(path: Path) -> list:
    """Список частей архива xlsx."""
    with zipfile.ZipFile(path) as archive:
        return archive.namelist()


# ─────────────────────────────────────────────────────────────
# Файл, лист, дата заявки
# ─────────────────────────────────────────────────────────────

def test_template_file_exists_and_opens(template_path):
    """Бланк есть, он не пустой и открывается openpyxl."""
    if not template_path.exists():
        pytest.skip(f"бланк {TEMPLATE_NAME} ещё не собран")

    assert template_path.stat().st_size > 0
    workbook = load_workbook(template_path)
    try:
        assert workbook.sheetnames == [SHEET_NAME]
    finally:
        workbook.close()


def test_single_visible_sheet_named_tdsheet(template_ws):
    """Лист ровно один, он видимый и называется TDSheet."""
    workbook = template_ws.parent

    assert len(workbook.worksheets) == 1
    assert workbook.sheetnames == [SHEET_NAME]
    assert template_ws.title == SHEET_NAME
    assert template_ws.sheet_state == "visible"
    assert workbook.active.title == SHEET_NAME


def test_no_hidden_rows_or_columns(template_ws):
    """Скрытых строк и колонок в бланке нет."""
    hidden_rows = [row for row, dimension in template_ws.row_dimensions.items()
                   if dimension.hidden]
    hidden_columns = [column for column, dimension
                      in template_ws.column_dimensions.items()
                      if dimension.hidden]

    assert hidden_rows == []
    assert hidden_columns == []


def test_date_label_is_found_by_text(template_ws):
    """
    Подпись 'Дата заявки:' есть в листе ровно один раз, в колонке A.

    Импорт ищет её по тексту, а не по номеру строки (требование ТЗ),
    поэтому проверяется и сам текст, и его единственность.
    """
    matches = [(cell.coordinate, cell.value)
               for row in template_ws.iter_rows() for cell in row
               if isinstance(cell.value, str) and DATE_LABEL in cell.value]

    assert len(matches) == DATE_LABEL_OCCURRENCES, (
        f"подпись {DATE_LABEL!r} должна встречаться один раз: {matches}"
    )

    coordinate, value = matches[0]
    assert value == DATE_LABEL
    assert coordinate == f"A{DATE_ROW}"
    assert template_ws.cell(row=DATE_ROW, column=1).value == DATE_LABEL


def test_date_cell_is_empty_and_styled(template_ws):
    """
    Ячейка даты пустая, но сохраняет оформление образца.

    В образце в B5 стоит '16.09.2026' — это данные заказчика, в бланк
    они не переносятся.
    """
    cell = template_ws[DATE_CELL]

    assert cell.value is None
    assert cell.font.name == MAIN_FONT_FAMILY
    assert cell.font.sz == DATA_FONT_SIZE


def test_title_cell_is_merged_and_bold(template_ws):
    """Заголовок 'ЗАЯВКА...' стоит в B2 в объединении B2:C2, жирным."""
    assert template_ws["B2"].value == TITLE
    assert str(template_ws.merged_cells.ranges) != ""
    assert "B2:C2" in {str(rng) for rng in template_ws.merged_cells.ranges}

    font = template_ws["B2"].font
    assert font.name == MAIN_FONT_FAMILY
    assert font.bold is True
    assert template_ws["B2"].alignment.horizontal == "center"


def test_rows_before_and_after_title_are_empty(template_ws):
    """Строки 1, 3 и 4 бланка пустые — как в образце."""
    for row in (1, 3, 4):
        assert _row_is_empty(template_ws, row), f"строка {row} должна быть пустой"


def test_merged_cells_are_only_the_title(template_ws):
    """
    Объединение в бланке ровно одно — B2:C2.

    В образце есть ещё AF7:AF12 (ставка, объединённая на 6 машин), но это
    след заполнения файла, а не часть пустого бланка.
    """
    assert {str(rng) for rng in template_ws.merged_cells.ranges} == {"B2:C2"}


# ─────────────────────────────────────────────────────────────
# Шапка таблицы (строка 6)
# ─────────────────────────────────────────────────────────────

def test_header_row_has_32_columns_in_expected_order(template_ws):
    """Строка 6 — ровно 32 заголовка A..AF в эталонном порядке."""
    assert len(HEADERS) == 32
    assert _row_values(template_ws, HEADER_ROW) == list(HEADERS)
    assert template_ws.max_column == len(HEADERS)


def test_header_row_has_nothing_beyond_column_af(template_ws):
    """За колонкой AF (32-й) в шапке ничего нет."""
    beyond = [template_ws.cell(row=HEADER_ROW, column=index).value
              for index in range(len(HEADERS) + 1, len(HEADERS) + 4)]

    assert beyond == [None, None, None]


def test_header_cells_are_bold_arial_centered(template_ws):
    """Каждая ячейка шапки: Arial 10, жирный, по центру."""
    for index in range(1, len(HEADERS) + 1):
        cell = template_ws.cell(row=HEADER_ROW, column=index)
        coordinate = get_column_letter(index)

        assert cell.font.name == MAIN_FONT_FAMILY, f"{coordinate}6: шрифт"
        assert cell.font.sz == HEADER_FONT_SIZE, f"{coordinate}6: кегль"
        assert cell.font.bold is True, f"{coordinate}6: не жирный"
        assert cell.alignment.horizontal == "center", f"{coordinate}6"
        assert cell.alignment.vertical == "center", f"{coordinate}6"


def test_header_wrap_matches_sample(template_ws):
    """Перенос текста включён ровно у E6, G6, AC6 и AE6 — как в образце."""
    wrapped = [f"{get_column_letter(index)}6"
               for index in range(1, len(HEADERS) + 1)
               if template_ws.cell(row=HEADER_ROW, column=index)
               .alignment.wrap_text]

    assert wrapped == list(HEADER_WRAP_CELLS)
    assert "\n" in template_ws["AC6"].value


def test_header_has_thin_black_borders(template_ws):
    """Рамка шапки — тонкая чёрная со всех четырёх сторон."""
    for index in range(1, len(HEADERS) + 1):
        cell = template_ws.cell(row=HEADER_ROW, column=index)
        for edge in ("left", "right", "top", "bottom"):
            side = getattr(cell.border, edge)
            assert side.style == "thin", f"{cell.coordinate}.{edge}"
            assert side.color.rgb == BLACK_ARGB, f"{cell.coordinate}.{edge}"


# ─────────────────────────────────────────────────────────────
# Строки данных (R7-R16)
# ─────────────────────────────────────────────────────────────

def test_there_are_exactly_ten_data_rows(template_ws):
    """Строк данных ровно 10, последняя — R16."""
    assert DATA_ROWS == 10
    assert LAST_DATA_ROW == 16
    assert template_ws.max_row == LAST_ROW


def test_data_rows_are_empty(template_ws):
    """Все 10 строк данных пустые: ни VIN, ни ФИО, ни ставок."""
    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        assert _row_is_empty(template_ws, row), f"строка {row} не пустая"


def test_data_rows_are_bordered_and_styled(template_ws):
    """Строки данных — с рамкой, шрифтом Arial 8 и в тех же 32 колонках."""
    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        for index in range(1, len(HEADERS) + 1):
            cell = template_ws.cell(row=row, column=index)

            assert cell.border.left.style == "thin", cell.coordinate
            assert cell.border.right.style == "thin", cell.coordinate
            assert cell.border.top.style == "thin", cell.coordinate
            assert cell.border.bottom.style == "thin", cell.coordinate
            assert cell.font.name == MAIN_FONT_FAMILY, cell.coordinate
            assert cell.font.sz == DATA_FONT_SIZE, cell.coordinate


def test_number_formats_of_data_columns(template_ws):
    """Форматы колонок: AD — время ('h:mm'), AF — ставка ('#,##0 " р."')."""
    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        assert template_ws.cell(row=row, column=30).number_format == TIME_FORMAT
        assert template_ws.cell(row=row, column=32).number_format == MONEY_FORMAT

    assert template_ws["A7"].number_format == "General"


def test_spacer_rows_around_parties_block_are_empty(template_ws):
    """Строки 17, 18, 21 и 22 пустые — место перед блоком и под подписи."""
    assert BLOCK_SPACER_ROWS == (17, 18, 21, 22)
    for row in BLOCK_SPACER_ROWS:
        assert _row_is_empty(template_ws, row), f"строка {row} должна быть пустой"


def test_spacer_rows_have_no_borders_or_shapes(template_ws):
    """Пустые строки блока сторон рамки не имеют (в образце её тоже нет)."""
    for row in BLOCK_SPACER_ROWS:
        cell = template_ws.cell(row=row, column=3)
        assert cell.border.left.style is None, cell.coordinate
        assert cell.border.bottom.style is None, cell.coordinate


# ─────────────────────────────────────────────────────────────
# Нижний блок: Заказчик / Перевозчик
# ─────────────────────────────────────────────────────────────

def test_parties_block_columns_and_labels(template_ws):
    """'Заказчик' и 'Перевозчик' — в колонках C и J, как в образце."""
    left, right = PARTY_COLUMNS

    assert (left, right) == ("C", "J")
    assert template_ws[f"{left}{PARTY_ROW}"].value == PARTY_LABELS[0]
    assert template_ws[f"{right}{PARTY_ROW}"].value == PARTY_LABELS[1]


def test_parties_block_fixed_parties(template_ws):
    """Стороны зафиксированы: Сюрлогистик (Заказчик) и ООО ТЕХНОЛОГИСТИКА."""
    left, right = PARTY_COLUMNS

    assert template_ws[f"{left}{PARTY_NAME_ROW}"].value == CUSTOMER
    assert template_ws[f"{right}{PARTY_NAME_ROW}"].value == CARRIER
    assert CUSTOMER == "Сюрлогистик"
    assert CARRIER == 'ООО "ТЕХНОЛОГИСТИКА"'


def test_signature_row(template_ws):
    """Строка подписи: 'ФИО, подпись, печать' в обеих колонках блока."""
    left, right = PARTY_COLUMNS

    assert template_ws[f"{left}{SIGNATURE_ROW}"].value == SIGNATURE_TEXT
    assert template_ws[f"{right}{SIGNATURE_ROW}"].value == SIGNATURE_TEXT


def test_parties_block_uses_only_columns_c_and_j(template_ws):
    """В строках блока заняты только колонки C и J: колонка I пустая."""
    for row in (PARTY_ROW, PARTY_NAME_ROW, SIGNATURE_ROW):
        values = [(get_column_letter(index),
                   template_ws.cell(row=row, column=index).value)
                  for index in range(1, len(HEADERS) + 1)
                  if template_ws.cell(row=row, column=index).value is not None]

        assert [column for column, _ in values] == ["C", "J"], f"строка {row}"


def test_customer_is_on_the_left_and_carrier_on_the_right(template_ws):
    """Сюрлогистик — слева (Заказчик), ООО ТЕХНОЛОГИСТИКА — справа."""
    assert template_ws[f"{PARTY_COLUMNS[0]}{PARTY_NAME_ROW}"].value == CUSTOMER
    assert template_ws[f"{PARTY_COLUMNS[1]}{PARTY_NAME_ROW}"].value == CARRIER


# ─────────────────────────────────────────────────────────────
# Оформление: обводка «Город погрузки», заливки, фигуры
# ─────────────────────────────────────────────────────────────

def test_no_green_outline_around_loading_city(template_ws):
    """
    Вокруг 'Город погрузки' нет зелёной обводки.

    В ТЗ это ручное выделение в файле заказчика; воспроизводить его
    нельзя. Проверяются обе половины: рамка E6 — обычная тонкая чёрная,
    и ни у одной ячейки листа нет зелёного цвета в рамке, заливке
    или шрифте.
    """
    assert template_ws["E6"].value == "Город погрузки"

    for edge in ("left", "right", "top", "bottom"):
        side = getattr(template_ws["E6"].border, edge)
        assert side.style == "thin", f"E6.{edge}"
        assert side.color.rgb == BLACK_ARGB, f"E6.{edge}: {side.color.rgb}"

    for row in template_ws.iter_rows():
        for cell in row:
            for edge in ("left", "right", "top", "bottom"):
                side = getattr(cell.border, edge)
                assert not _is_green(side.color.rgb if side.color else None), (
                    f"{cell.coordinate}.{edge}: зелёная обводка"
                )
            assert not _is_green(cell.fill.fgColor.rgb), cell.coordinate
            assert not _is_green(cell.font.color.rgb
                                 if cell.font.color else None), cell.coordinate


def test_no_green_colors_in_applied_styles(template_path):
    """В применённых стилях книги нет ни одного зелёного цвета."""
    colors = _applied_style_colors(template_path)

    assert colors == {BLACK_ARGB}, f"неожиданные цвета в стилях: {colors}"
    assert not any(_is_green(color) for color in colors)


def test_no_fills_on_the_sheet(template_ws):
    """
    Заливок на листе нет вообще — как в образце.

    В xl/styles.xml образца только patternFill none и gray125, ни одной
    цветной заливки; ТЗ предполагало заливку шапки, но в файле её нет.
    """
    for row in template_ws.iter_rows():
        for cell in row:
            assert cell.fill.patternType is None, cell.coordinate


def test_no_drawings_or_shapes_in_archive(template_path):
    """
    В бланке нет фигур и линий (xl/drawings).

    В образце есть xl/drawings/drawing1.xml — две чёрные линии подписи
    под ячейками C19/J19. openpyxl фигуры создавать не умеет, поэтому
    линий в бланке нет: это осознанное решение (доложено в отчёте).
    """
    parts = _archive_parts(template_path)

    assert not [part for part in parts if "drawing" in part], parts
    assert not [part for part in parts if "media" in part], parts


def test_no_conditional_formatting_or_data_validation(template_ws):
    """Условного форматирования и проверок данных нет — скрытых выделений тоже."""
    assert list(template_ws.conditional_formatting) == []
    assert list(template_ws.data_validations.dataValidation) == []


def test_no_comments_hyperlinks_or_tables(template_ws):
    """Примечаний, гиперссылок и объектов-таблиц в бланке нет."""
    assert list(template_ws._hyperlinks) == []
    assert list(getattr(template_ws, "tables", {})) == []
    assert [cell.coordinate for row in template_ws.iter_rows()
            for cell in row if cell.comment] == []


def test_no_freeze_panes_or_autofilter(template_ws):
    """Закрепления областей и автофильтра нет — как в образце."""
    assert template_ws.freeze_panes is None
    assert template_ws.auto_filter.ref is None


# ─────────────────────────────────────────────────────────────
# Данных образца в бланке нет
# ─────────────────────────────────────────────────────────────

def test_blank_contains_only_headers_and_parties(template_ws):
    """
    Полный список непустых ячеек бланка — ровно шапка, дата и стороны.

    Это главная проверка на отсутствие данных: если в бланк случайно
    попадут VIN, ФИО, телефоны, адреса или суммы образца, количество
    или состав непустых ячеек изменится.
    """
    texts = _cell_texts(template_ws)

    assert len(texts) == EXPECTED_NON_EMPTY_CELLS, sorted(texts.items())

    expected = {"B2": TITLE, f"A{DATE_ROW}": DATE_LABEL}
    expected.update({f"{get_column_letter(index)}{HEADER_ROW}": title
                     for index, title in enumerate(HEADERS, start=1)})
    expected[f"{PARTY_COLUMNS[0]}{PARTY_ROW}"] = PARTY_LABELS[0]
    expected[f"{PARTY_COLUMNS[1]}{PARTY_ROW}"] = PARTY_LABELS[1]
    expected[f"{PARTY_COLUMNS[0]}{PARTY_NAME_ROW}"] = CUSTOMER
    expected[f"{PARTY_COLUMNS[1]}{PARTY_NAME_ROW}"] = CARRIER
    expected[f"{PARTY_COLUMNS[0]}{SIGNATURE_ROW}"] = SIGNATURE_TEXT
    expected[f"{PARTY_COLUMNS[1]}{SIGNATURE_ROW}"] = SIGNATURE_TEXT

    assert texts == expected


def test_sample_date_is_not_copied(template_ws):
    """Пример даты из образца в бланк не перенесён."""
    assert template_ws[DATE_CELL].value is None
    assert SAMPLE_DATE_VALUE not in _cell_texts(template_ws).values()


def test_no_data_like_values_in_blank(template_ws):
    """В бланке нет значений, похожих на данные: VIN, госномеров, телефонов."""
    patterns = {
        "VIN (17 символов)": r"\b[A-HJ-NPR-Z0-9]{17}\b",
        "госномер": r"\b[АВЕКМНОРСТУХ]\d{3}[АВЕКМНОРСТУХ]{2}\d{2,3}\b",
        "телефон": r"(\+7|8)[\s\-(]*\d{3}",
        "паспорт": r"\b\d{4}\s?\d{6}\b",
        "ИНН": r"\b\d{10,12}\b",
        "дата цифрами": r"\b\d{2}\.\d{2}\.\d{4}\b",
    }
    texts = [str(value) for value in _cell_texts(template_ws).values()]

    for label, pattern in patterns.items():
        hits = [text for text in texts if re.search(pattern, text)]
        assert hits == [], f"{label}: {hits}"


# ─────────────────────────────────────────────────────────────
# Геометрия: ширины, высоты, формат листа
# ─────────────────────────────────────────────────────────────

def test_column_widths_match_reference(template_ws):
    """Ширины колонок A..AF совпадают с эталонными (из образца)."""
    widths = _effective_widths(template_ws)

    assert set(widths) == set(COLUMN_WIDTHS)
    for column, expected in COLUMN_WIDTHS.items():
        assert widths[column] == pytest.approx(expected, abs=1e-9), column


def test_best_fit_columns(template_ws):
    """bestFit=1 стоит у тех же колонок, что в образце: B, H и I."""
    best_fit = sorted(column for column, dimension
                      in template_ws.column_dimensions.items()
                      if dimension.bestFit)

    assert best_fit == sorted(COLUMN_BEST_FIT)


def test_row_heights_match_reference(template_ws):
    """Высоты строк совпадают с эталонными (из образца)."""
    for row, expected in ROW_HEIGHTS.items():
        dimension = template_ws.row_dimensions.get(row)
        actual = dimension.height if dimension is not None else None

        assert actual == pytest.approx(expected, abs=1e-9), f"строка {row}"


def test_row_four_keeps_default_height(template_ws):
    """У строки 4 высота не задана — в образце её тоже нет."""
    assert 4 not in ROW_HEIGHTS
    assert template_ws.row_dimensions.get(4) is None


def test_sheet_format_defaults(template_ws):
    """Формат листа по умолчанию совпадает с образцом."""
    sheet_format = template_ws.sheet_format

    assert sheet_format.baseColWidth == 8
    assert sheet_format.defaultColWidth == pytest.approx(16.83203125)
    assert sheet_format.defaultRowHeight == pytest.approx(15.0)
    assert sheet_format.customHeight is True


# ─────────────────────────────────────────────────────────────
# Сверка с образцом
# ─────────────────────────────────────────────────────────────

def test_sample_is_committed_and_untouched(sample_path):
    """
    Образец — источник структуры, только чтение.

    Если этот тест упал, образец отредактировали: он нужен как есть,
    из него собирается бланк и по нему сверяется структура.
    """
    if not sample_path.exists():
        pytest.skip(f"образец {SAMPLE_NAME} не найден в templates/")

    digest = _sha256(sample_path)
    assert digest == SAMPLE_SHA256, (
        f"образец {SAMPLE_NAME!r} изменён: ожидался {SAMPLE_SHA256}, "
        f"получен {digest}"
    )


def test_sample_has_same_sheet_name(template_ws, sample_ws):
    """Имя листа бланка совпадает с образцом: TDSheet."""
    assert sample_ws.title == template_ws.title == SHEET_NAME


def test_sample_headers_match_blank_headers(template_ws, sample_ws):
    """Шапка бланка посимвольно совпадает с шапкой образца."""
    assert _row_values(sample_ws, HEADER_ROW) == _row_values(template_ws,
                                                             HEADER_ROW)


def test_column_widths_and_best_fit_match_sample(template_ws, sample_ws):
    """Ширины колонок и bestFit бланка совпадают с образцом."""
    sample_widths = _effective_widths(sample_ws)

    assert sample_widths == _effective_widths(template_ws)
    assert {column for column, dimension in sample_ws.column_dimensions.items()
            if dimension.bestFit} == {
        column for column, dimension in template_ws.column_dimensions.items()
        if dimension.bestFit}


def test_parties_block_matches_sample_layout(template_ws, sample_ws):
    """
    Блок сторон бланка — та же структура, что в образце, только ниже.

    В образце таблица кончается на R12, и дальше идут R13-R14 пусто,
    R15 'Заказчик/Перевозчик', R16 'Сюрлогистик/ООО ТЕХНОЛОГИСТИКА',
    R17-R18 пусто, R19 'ФИО, подпись, печать'. В бланке на 10 строк
    данных тот же блок сдвинут на 4 строки вниз (R17-R23), а высоты
    строк повторяются один в один.
    """
    shift = LAST_DATA_ROW - 12
    assert shift == 4

    sample_heights = [sample_ws.row_dimensions[row].height
                      for row in range(13, 20)]
    blank_heights = [template_ws.row_dimensions[row].height
                     for row in range(13 + shift, 20 + shift)]

    assert sample_heights == blank_heights

    sample_values = [sample_ws.cell(row=row, column=3).value
                     for row in range(13, 20)]
    blank_values = [template_ws.cell(row=row, column=3).value
                    for row in range(13 + shift, 20 + shift)]

    assert sample_values == blank_values
    assert sample_values == [None, None, "Заказчик", "Сюрлогистик",
                             None, None, SIGNATURE_TEXT]

def test_lower_block_columns_match_sample(sample_ws):
    """В образце блок сторон стоит в колонках C и J — тест фиксирует факт."""
    filled = [get_column_letter(cell.column)
              for row in sample_ws.iter_rows(min_row=13, max_row=19)
              for cell in row if cell.value is not None]

    assert filled == ["C", "J", "C", "J", "C", "J"]


def test_sheet_format_matches_sample(template_ws, sample_ws):
    """Формат листа по умолчанию совпадает с образцом."""
    assert (template_ws.sheet_format.baseColWidth
            == sample_ws.sheet_format.baseColWidth)
    assert (template_ws.sheet_format.defaultColWidth
            == pytest.approx(sample_ws.sheet_format.defaultColWidth))
    assert (template_ws.sheet_format.defaultRowHeight
            == pytest.approx(sample_ws.sheet_format.defaultRowHeight))
    assert (template_ws.sheet_format.customHeight
            == sample_ws.sheet_format.customHeight)


def test_sample_has_no_third_party_personal_data(sample_ws):
    """
    В образце нет персональных данных третьих лиц.

    Проверяется по ячейкам: только шапка, пример даты и блок сторон.
    Именно поэтому образец можно хранить в git (в отличие от образца
    разовой аренды, который из-за ПДн в git не хранится).
    """
    texts = _cell_texts(sample_ws)

    assert len(texts) == SAMPLE_NON_EMPTY_CELLS, sorted(texts.items())
    assert texts[DATE_CELL] == SAMPLE_DATE_VALUE


# ─────────────────────────────────────────────────────────────
# SHA256 и воспроизводимость сборки
# ─────────────────────────────────────────────────────────────

def test_template_matches_pinned_sha(template_path):
    """
    Бланк на диске — ровно то, что собрал сборщик.

    Если тест упал после осознанной пересборки — обнови TEMPLATE_SHA256;
    если правки не было — файл отредактировали вручную.
    """
    if not template_path.exists():
        pytest.skip(f"бланк {TEMPLATE_NAME} ещё не собран")

    digest = _sha256(template_path)
    assert digest == TEMPLATE_SHA256, (
        f"бланк {TEMPLATE_NAME!r} не совпадает со сборкой: ожидался "
        f"{TEMPLATE_SHA256}, получен {digest}"
    )


def test_archive_timestamps_are_fixed(template_path):
    """
    В архиве бланка фиксированные метки времени.

    Без этого SHA256 «плыл» бы от запуска к запуску: openpyxl пишет
    текущее время и в zip, и в docProps/core.xml.
    """
    with zipfile.ZipFile(template_path) as archive:
        stamps = {info.date_time for info in archive.infolist()}
        core = archive.read("docProps/core.xml").decode("utf-8")

    assert stamps == {(1980, 1, 1, 0, 0, 0)}, stamps
    assert core.count("2026-01-01T00:00:00Z") == 2
    assert re.search(r"<dcterms:created[^>]*>\d{4}-\d{2}-\d{2}", core)
    assert re.search(r"<dcterms:modified[^>]*>\d{4}-\d{2}-\d{2}", core)


def test_builder_reproduces_template_bytes(template_path, work_dir):
    """Сборщик даёт побайтово тот же файл, что лежит в templates/."""
    import tools.make_havaly_template as builder

    rebuilt = builder.build_template(work_dir / "havaly_rebuild.xlsx")

    assert rebuilt.read_bytes() == template_path.read_bytes(), (
        "сборщик и файл в templates/ разъехались: пересобери бланк "
        "(python tools/make_havaly_template.py) и обнови TEMPLATE_SHA256"
    )


def test_builder_reproduces_same_layout(template_path, work_dir):
    """Структура пересобранного бланка совпадает с лежащим в templates/."""
    import tools.make_havaly_template as builder

    rebuilt = builder.build_template(work_dir / "havaly_rebuild_layout.xlsx")
    expected = load_workbook(template_path)
    actual = load_workbook(rebuilt)

    try:
        assert actual.sheetnames == expected.sheetnames
        expected_ws = expected[SHEET_NAME]
        actual_ws = actual[SHEET_NAME]

        assert _cell_texts(actual_ws) == _cell_texts(expected_ws)
        assert _effective_widths(actual_ws) == _effective_widths(expected_ws)
        assert ({row: dimension.height
                 for row, dimension in actual_ws.row_dimensions.items()}
                == {row: dimension.height
                    for row, dimension in expected_ws.row_dimensions.items()})
        assert ({str(rng) for rng in actual_ws.merged_cells.ranges}
                == {str(rng) for rng in expected_ws.merged_cells.ranges})
    finally:
        expected.close()
        actual.close()


def test_builder_is_idempotent(work_dir):
    """Повторная сборка не меняет уже собранный файл."""
    import tools.make_havaly_template as builder

    target = work_dir / "havaly_idempotent.xlsx"
    builder.build_template(target)
    first_bytes = target.read_bytes()
    first_mtime = target.stat().st_mtime_ns

    builder.build_template(target)

    assert target.read_bytes() == first_bytes
    assert target.stat().st_mtime_ns == first_mtime, (
        "актуальный бланк был перезаписан: идемпотентность нарушена"
    )


def test_builder_leaves_no_temporary_files(work_dir):
    """После сборки в папке не остаётся служебных файлов сборщика."""
    import tools.make_havaly_template as builder

    folder = work_dir / "havaly_tmp_check"
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / "havaly.xlsx"

    builder.build_template(target)
    builder.build_template(target)

    assert sorted(path.name for path in folder.iterdir()) == [target.name]


def test_builder_constants_match_test_reference():
    """Константы сборщика не разъехались с эталоном теста."""
    import tools.make_havaly_template as builder

    assert builder.TEMPLATE_NAME == TEMPLATE_NAME
    assert builder.SAMPLE_NAME == SAMPLE_NAME
    assert builder.SHEET_NAME == SHEET_NAME
    assert builder.TITLE == TITLE
    assert builder.DATE_LABEL == DATE_LABEL
    assert tuple(builder.HEADERS) == HEADERS
    assert builder.DATA_ROWS == DATA_ROWS
    assert builder.FIRST_DATA_ROW == FIRST_DATA_ROW
    assert builder.HEADER_ROW == HEADER_ROW
    assert builder.DATE_ROW == DATE_ROW
    assert builder.PARTY_ROW == PARTY_ROW
    assert builder.PARTY_NAME_ROW == PARTY_NAME_ROW
    assert builder.SIGNATURE_ROW == SIGNATURE_ROW
    assert builder.PARTY_COLUMNS == PARTY_COLUMNS
    assert builder.CUSTOMER == CUSTOMER
    assert builder.CARRIER == CARRIER
    assert builder.SIGNATURE_TEXT == SIGNATURE_TEXT
    assert builder.COLUMN_WIDTHS == COLUMN_WIDTHS
    assert set(builder.ROW_HEIGHTS) == set(ROW_HEIGHTS)
    assert builder.TIME_FORMAT == TIME_FORMAT
    assert builder.MONEY_FORMAT == MONEY_FORMAT
    assert tuple(builder.HEADER_WRAP_COLUMNS) == ("E", "G", "AC", "AE")


# ─────────────────────────────────────────────────────────────
# Самопроверка: детектор зелёного цвета действительно работает
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("value", [
    "FF00B050",     # зелёный акцент Office
    "0092D050",     # светло-зелёный
    "FF00FF00",     # чистый зелёный
    "00339966",     # зелёный из легаси-палитры
])
def test_green_detector_recognizes_green(value):
    """Зелёные цвета определяются — иначе проверка обводки была бы пустой."""
    assert _is_green(value)


@pytest.mark.parametrize("value", [
    None,           # цвета нет
    "FF000000",     # чёрный (рамка и текст бланка)
    "00000000",     # «пустой» цвет openpyxl
    "FFFFFFFF",     # белый
    "FFFF0000",     # красный
    "FF0000FF",     # синий
])
def test_green_detector_ignores_non_green(value):
    """Не зелёные цвета не считаются обводкой."""
    assert not _is_green(value)
