#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты генератора договора-заявки «Формика» (ЭТАП 3.1.A.3).

Проверяют: тип больше не заглушка; плейсхолдеры шаблона заменяются;
таблица груза рассчитана на переменное число машин (лишние строки
удаляются постобработкой); тягач и полуприцеп в груз не попадают;
стоимость выводится одной суммой с НДС; срок оплаты берётся из поля
«Срок оплаты (дней)» и печатается с прописью в родительном падеже
(ШАГ FIX-2.4: константы «3 (трех) банковских дней» в бланке больше нет);
логи идут в канал «core.contract_generator» и не содержат персональных
данных; готовый документ совпадает с золотым эталоном
tests/data/golden/formika_sample.* (4 машины).

Все данные синтетические, реальных ПДн нет.
"""

import gc
import json
import logging
import re
import shutil
from pathlib import Path

import pytest
from docx import Document

from core.contracts.base_generator import ConvertNewlinesStep
from core.contracts.factory import GeneratorFactory
from core.contracts.formika.generator import FormikaGenerator
from core.contracts.formika.postprocess import RemoveEmptyVehicleRowsStep
from core.contracts.registry import ContractTypeRegistry
from tools.make_golden import (
    FORMIKA_SCENARIOS,
    GOLDEN_DIR,
    fingerprint,
    strip_embedded_fonts,
)

PROJECT_ROOT = Path(__file__).resolve().parent.parent

CARGO_HEADERS = ("№", "Марка, модель", "VIN-номер")

PLACEHOLDER_RE = re.compile(r"\{\{[^{}]*\}\}")


@pytest.fixture(autouse=True)
def _release_documents():
    """Word-файлы на Windows освобождаем до удаления временных копий."""
    yield
    gc.collect()


@pytest.fixture(scope="module", autouse=True)
def _builtin_types_loaded():
    """Регистрация типов: без неё фабрика не знает про formika."""
    ContractTypeRegistry.load_builtin()


def _vehicle(number: int) -> dict:
    return {
        "vin": f"EC3TEUMB0T000{number:04d}",
        "brand_model": f"МОДЕЛЬ {number}",
        "vehicle_type": "Легковой автомобиль",
    }


def _payload(cars: int = 4) -> dict:
    """Синтетические данные договора-заявки Формики."""
    return {
        "driver": {
            "full_name": "Иванов Иван Иванович",
            "birth_date": "1980-01-01",
            "passport_series": "18 22",
            "passport_number": "926830",
            "passport_issue_date": "2023-01-30",
            "passport_issuer": "Отделом УФМС России по г. Москве",
            "registration_address": "г. Москва, ул. Тестовая, д. 1",
            "license_series": "99 36",
            "license_number": "123456",
            "phone": "+7 (999) 123-45-67",
        },
        "carrier": {"full_name": "ООО «Ромашка»", "entity_type": "ООО"},
        "customer": {"full_name": "ООО «Заказчик»"},
        "vehicles": [_vehicle(number) for number in range(1, cars + 1)] + [
            # Тягач и полуприцеп лежат в справочнике машин, но грузом не являются.
            {"brand_model": "Foton Auman", "plate_number": "O844XY196",
             "vehicle_type": "Тягач"},
            {"brand_model": "YANGMINDA", "plate_number": "71ABF18",
             "vehicle_type": "Полуприцеп"},
        ],
        "tractor": {"brand_model": "Foton Auman", "plate_number": "O844XY196",
                    "color": "Белый", "year": 2023},
        "trailer": {"brand_model": "YANGMINDA", "plate_number": "71ABF18",
                    "color": "Серый", "year": 2020},
        "contract": {
            "number": "ФМ-2026-1",
            "date": "2026-07-24",
            "route": "г. Воронеж - г. Москва",
            "carrier_type": "ООО (с НДС)",
            "vat_rate": "22%",
            "vat_rate_num": 22,
            "price_without_vat": 180300.0,
            "price_with_vat": 219966.0,
            "loading_plan_date": "2026-07-27",
        },
        "loadings": [{"address": "г. Воронеж, ул. Остужева 52Б",
                      "date": "2026-07-27", "time_window": "09:00-15:00"}],
        "unloadings": [{"address": "г. Москва, Перерва 19 стр 3",
                        "date": "2026-07-30", "time_window": ""}],
    }


@pytest.fixture
def generator(templates_dir) -> FormikaGenerator:
    return FormikaGenerator(templates_dir=str(templates_dir))


def _document_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join(parts)


def _cargo_table(doc):
    """Таблица груза — по заголовкам колонок, а не по индексу."""
    for table in doc.tables:
        if not table.rows:
            continue
        headers = tuple(cell.text.strip() for cell in table.rows[0].cells)
        if headers == CARGO_HEADERS:
            return table
    return None


def _generate(generator, payload, output_dir) -> str:
    path = generator.generate(payload, output_dir=str(output_dir))
    assert Path(path).exists(), f"файл не создан: {path}"
    return path


# ─────────────────────────────────────────────────────────────
# Тип зарегистрирован и больше не заглушка
# ─────────────────────────────────────────────────────────────

def test_generator_is_registered_and_not_stub():
    spec = ContractTypeRegistry.get("formika", strict=True)
    assert spec.generator_class is FormikaGenerator
    assert isinstance(
        GeneratorFactory.get_generator("formika", strict=True), FormikaGenerator
    )


def test_generate_does_not_raise_not_implemented(generator, work_dir):
    path = _generate(generator, _payload(1), work_dir)
    assert path.endswith(".docx")


def test_filename_uses_prefix_number_and_date(generator, work_dir):
    path = Path(_generate(generator, _payload(1), work_dir))
    assert path.name.startswith(f"{FormikaGenerator.FILE_PREFIX}_ФМ-2026-1_")


def test_template_path_ignores_carrier_type(generator):
    for carrier_type in ("ООО (с НДС)", "ИП без НДС", ""):
        assert generator._get_template_path(carrier_type).endswith(
            "shablon_formika.docx"
        )


def test_postprocess_steps_are_newlines_and_empty_rows(generator):
    steps = generator.postprocess_steps(None)
    assert [type(step) for step in steps] == [
        ConvertNewlinesStep, RemoveEmptyVehicleRowsStep
    ]


def test_insert_route_tables_is_noop(generator):
    """У Формики нет таблиц по погрузкам — метод обязан быть безопасным."""
    doc = Document()
    doc.add_paragraph("текст")
    assert generator._insert_route_tables(doc, None) is None
    assert len(doc.paragraphs) == 1


# ─────────────────────────────────────────────────────────────
# Таблица груза: переменное число машин
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("cars", [1, 2, 4, 12])
def test_cargo_table_keeps_only_filled_rows(generator, work_dir, cars):
    path = _generate(generator, _payload(cars), work_dir)
    table = _cargo_table(Document(path))

    assert table is not None, "таблица груза не найдена в готовом документе"
    assert len(table.rows) == cars + 1, (
        f"машин {cars}, а строк в таблице {len(table.rows)} (ожидалось {cars + 1})"
    )
    for number in range(1, cars + 1):
        cells = [cell.text.strip() for cell in table.rows[number].cells]
        assert cells == [str(number), f"МОДЕЛЬ {number}",
                         f"EC3TEUMB0T000{number:04d}"]


def test_tractor_and_trailer_are_not_in_cargo_table(generator, work_dir):
    path = _generate(generator, _payload(2), work_dir)
    table = _cargo_table(Document(path))
    text = "\n".join(cell.text for row in table.rows for cell in row.cells)

    assert "Foton Auman" not in text, "тягач попал в таблицу груза"
    assert "YANGMINDA" not in text, "прицеп попал в таблицу груза"


def test_cargo_count_matches_vehicles(generator):
    for cars in (1, 5, 12):
        replacements = generator._build_replacements_map(_payload(cars))
        assert replacements["cargo_count"] == str(cars)


def test_replacements_map_covers_all_template_placeholders(generator, templates_dir):
    """Каждый плейсхолдер бланка получает значение из карты замен."""
    template_text = _document_text(Document(str(templates_dir / "shablon_formika.docx")))
    template_names = {
        match.strip("{} ") for match in PLACEHOLDER_RE.findall(template_text)
    }

    replacements = generator._build_replacements_map(_payload(4))
    missing = template_names - set(replacements)
    assert not missing, f"нет значений для плейсхолдеров: {sorted(missing)}"


def test_generated_document_has_no_placeholders_left(generator, work_dir):
    path = _generate(generator, _payload(4), work_dir)
    text = _document_text(Document(path))
    assert "{{" not in text
    assert "}}" not in text


# ─────────────────────────────────────────────────────────────
# Содержимое готового документа
# ─────────────────────────────────────────────────────────────

def test_document_content_matches_sample_structure(generator, work_dir):
    path = _generate(generator, _payload(4), work_dir)
    text = _document_text(Document(path))

    assert "ДОГОВОР-ЗАЯВКА № ФМ-2026-1" in text
    assert "«24» июля 2026 г." in text
    assert "2. МАРШРУТ ПЕРЕВОЗКИ: г. Воронеж - г. Москва" in text
    assert "Пункт погрузки: г. Воронеж, ул. Остужева 52Б" in text
    assert "Пункт выгрузки: г. Москва, Перерва 19 стр 3" in text
    assert "Дата и время: 27.07.2026 г. с 09:00 до 15:00" in text
    assert "Срок доставки: 3 календарных дня с момента погрузки" in text
    assert "ФИО: Иванов Иван Иванович" in text
    assert "Паспорт: 18 22 926830" in text
    assert "Тягач: Foton Auman гос. номер: O844XY196" in text
    assert "Прицеп: YANGMINDA гос. номер: 71ABF18" in text
    assert "8. ПОДПИСИ СТОРОН" in text
    assert "ООО «Формика»" in text


def test_cost_is_single_amount_with_vat(generator, work_dir):
    """Сумма в бланке — с НДС, в формате образца «219 966,00»."""
    path = _generate(generator, _payload(4), work_dir)
    text = _document_text(Document(path))

    assert "Стоимость перевозки составляет: 219\u00a0966,00 руб." in text
    assert "включая НДС 22%." in text


def test_cost_is_calculated_from_price_without_vat_if_needed(generator):
    payload = _payload(1)
    payload["contract"].pop("price_with_vat")

    replacements = generator._build_replacements_map(payload)

    assert replacements["sum_total"] == "219\u00a0966,00"
    assert replacements["sum_total_words"].startswith("Двести девятнадцать тысяч")


def test_cost_missing_is_zero_not_invented(generator):
    payload = _payload(1)
    payload["contract"].pop("price_with_vat")
    payload["contract"].pop("price_without_vat")

    replacements = generator._build_replacements_map(payload)

    assert replacements["sum_total"] == "0,00"
    assert replacements["vat_rate"] == "22%"


def test_plan_time_is_taken_from_loading_window(generator, work_dir):
    """Время погрузки берётся из окна времени точки, если поля пусты."""
    path = _generate(generator, _payload(1), work_dir)
    assert "с 09:00 до 15:00" in _document_text(Document(path))


def test_contract_year_comes_from_contract_date(generator):
    replacements = generator._build_replacements_map(_payload(1))
    assert replacements["contract_year"] == "2026"
    assert replacements["contract_month"] == "июля"
    assert replacements["contract_date"] == "24"


# ─────────────────────────────────────────────────────────────
# Срок оплаты (ШАГ FIX-2.4)
# ─────────────────────────────────────────────────────────────

#: Константа, стоявшая в бланке до шага FIX-2.4. Её не должно быть ни в
#: одном сгенерированном документе — ни при каком сроке оплаты.
OLD_PAYMENT_CONSTANT = "3 (трех) банковских дней"


def _payment_clause(text: str) -> str:
    """Строка «Порядок оплаты: …» из готового документа."""
    return next(line for line in text.splitlines()
                if line.startswith("Порядок оплаты:"))


def _payload_with_payment_days(payment_days):
    """Данные договора; payment_days=None — поле не задано вовсе."""
    payload = _payload(1)
    if payment_days is None:
        payload["contract"].pop("payment_days", None)
    else:
        payload["contract"]["payment_days"] = payment_days
    return payload


def test_payment_days_forty_five(generator, work_dir):
    """payment_days=45 → «45 (сорока пяти) банковских дней»."""
    path = _generate(generator, _payload_with_payment_days(45), work_dir)
    clause = _payment_clause(_document_text(Document(path)))

    assert "в течение 45 (сорока пяти) банковских дней" in clause
    assert OLD_PAYMENT_CONSTANT not in clause


def test_payment_days_thirty(generator, work_dir):
    """payment_days=30 → «30 (тридцати) банковских дней»."""
    path = _generate(generator, _payload_with_payment_days(30), work_dir)
    clause = _payment_clause(_document_text(Document(path)))

    assert "в течение 30 (тридцати) банковских дней" in clause


def test_payment_days_ten_matches_default_field(generator, work_dir):
    """Значение по умолчанию вкладки «Стоимость» (10) доходит до бланка."""
    path = _generate(generator, _payload_with_payment_days(10), work_dir)
    clause = _payment_clause(_document_text(Document(path)))

    assert "в течение 10 (десяти) банковских дней" in clause


def test_payment_days_empty_leaves_blank_not_three(generator, work_dir):
    """
    Незаданный срок оплаты — пустое место, а не «3 (трех)».

    docxtpl съедает пробелы вокруг пустого плейсхолдера, поэтому в бланке
    остаётся «в течение  () банковских дней» — цифр там нет. Число
    генератор не выдумывает: о незаполненном сроке скажет валидатор.
    """
    path = _generate(generator, _payload_with_payment_days(None), work_dir)
    text = _document_text(Document(path))
    clause = _payment_clause(text)

    assert OLD_PAYMENT_CONSTANT not in text
    assert "в течение  () банковских дней" in clause, clause
    assert "3" not in clause.split("после получения")[0]


def test_payment_days_empty_in_replacements(generator):
    """Пустой payment_days → обе замены пустые (в бланке пробел, не «3»)."""
    replacements = generator._build_replacements_map(_payload_with_payment_days(None))

    assert replacements["payment_days"] == ""
    assert replacements["payment_days_words"] == ""


def test_payment_days_replacements_are_numeric_and_words(generator):
    """Карта замен: цифры числом, скобки — родительным падежом прописью."""
    replacements = generator._build_replacements_map(_payload_with_payment_days(45))

    assert replacements["payment_days"] == "45"
    assert replacements["payment_days_words"] == "сорока пяти"


@pytest.mark.parametrize("cars", [1, 4])
def test_old_payment_constant_absent_from_every_document(generator, work_dir, cars):
    """«3 (трех) банковских дней» больше не встречается ни в одном документе."""
    for payment_days in (None, 0, 5, 30, 45, 365):
        payload = _payload(cars)
        if payment_days is None:
            payload["contract"].pop("payment_days", None)
        else:
            payload["contract"]["payment_days"] = payment_days

        path = _generate(generator, payload, work_dir)
        assert OLD_PAYMENT_CONSTANT not in _document_text(Document(path)), (
            f"константа вернулась при payment_days={payment_days!r}"
        )


# ─────────────────────────────────────────────────────────────
# Постобработка: шаг удаления пустых строк
# ─────────────────────────────────────────────────────────────

def _cargo_doc_with_empty_rows() -> Document:
    """Документ с таблицей груза, где заполнены только две строки из трёх."""
    doc = Document()
    table = doc.add_table(rows=4, cols=3)
    for cell, title in zip(table.rows[0].cells, CARGO_HEADERS):
        cell.text = title
    rows = [
        ("1", "МОДЕЛЬ 1", "VIN1"),
        ("2", "", ""),
        ("3", "МОДЕЛЬ 3", "VIN3"),
    ]
    for index, values in enumerate(rows, 1):
        for cell, value in zip(table.rows[index].cells, values):
            cell.text = value
    return doc


def test_remove_empty_rows_step_deletes_only_empty_rows(tmp_path):
    generator = FormikaGenerator(templates_dir=str(tmp_path))
    doc = _cargo_doc_with_empty_rows()

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    values = [[cell.text for cell in row.cells] for row in doc.tables[0].rows]
    assert values == [
        ["№", "Марка, модель", "VIN-номер"],
        ["1", "МОДЕЛЬ 1", "VIN1"],
        ["3", "МОДЕЛЬ 3", "VIN3"],
    ]


def test_remove_empty_rows_keeps_row_with_only_brand(tmp_path):
    """Строка с маркой, но без VIN, не удаляется: данные не теряем."""
    generator = FormikaGenerator(templates_dir=str(tmp_path))
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    for cell, title in zip(table.rows[0].cells, CARGO_HEADERS):
        cell.text = title
    table.rows[1].cells[1].text = "МОДЕЛЬ БЕЗ VIN"

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 2


def test_remove_empty_rows_ignores_foreign_tables(tmp_path):
    """Таблица подписей (1×2) не трогается: заголовки не те груза."""
    generator = FormikaGenerator(templates_dir=str(tmp_path))
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Заказчик:"
    table.rows[0].cells[1].text = "Экспедитор:"

    RemoveEmptyVehicleRowsStep(generator).apply(doc, None)

    assert len(doc.tables[0].rows) == 1


def test_generate_without_vehicles_keeps_header_only(generator, work_dir):
    """Ни одной машины — в документе остаётся только шапка таблицы."""
    path = _generate(generator, _payload(0), work_dir)
    table = _cargo_table(Document(path))

    assert len(table.rows) == 1
    assert generator._build_replacements_map(_payload(0))["cargo_count"] == "0"


def test_generate_with_empty_data_does_not_crash(generator, work_dir):
    """Пустые данные — пустой бланк без исключений (заглушки больше нет)."""
    path = _generate(generator, {"contract": {}}, work_dir)
    text = _document_text(Document(path))

    assert "ДОГОВОР-ЗАЯВКА №" in text
    assert "{{" not in text


def test_missing_template_raises_file_not_found(work_dir):
    generator = FormikaGenerator(templates_dir=str(work_dir / "нет-такого"))
    with pytest.raises(FileNotFoundError):
        generator.generate(_payload(1), output_dir=str(work_dir))


# ─────────────────────────────────────────────────────────────
# Логи
# ─────────────────────────────────────────────────────────────

def test_generation_logs_to_contract_generator(caplog, generator, work_dir):
    with caplog.at_level(logging.INFO, logger="core.contract_generator"):
        _generate(generator, _payload(4), work_dir)

    records = [r for r in caplog.records if r.name == "core.contract_generator"]
    assert records, "генерация ничего не записала в core.contract_generator"
    messages = "\n".join(r.getMessage() for r in records)
    assert "машин в договоре — 4" in messages
    assert "удалено пустых строк таблицы груза: 8" in messages


def test_logs_have_no_personal_data(caplog, generator, work_dir):
    """В логи не попадают ФИО, паспорт, VIN, адреса и марки машин."""
    with caplog.at_level(logging.DEBUG, logger="core.contract_generator"):
        _generate(generator, _payload(4), work_dir)

    messages = "\n".join(
        r.getMessage() for r in caplog.records
        if r.name == "core.contract_generator"
    )
    for fragment in ("Иванов", "926830", "EC3TEUMB0T0000001", "Остужева",
                     "Foton", "YANGMINDA", "МОДЕЛЬ"):
        assert fragment not in messages, f"в логе есть «{fragment}»"


def test_default_output_dir_is_output():
    assert Path(FormikaGenerator.default_output_dir()).name == "output"


# ─────────────────────────────────────────────────────────────
# Золотой эталон (4 машины)
# ─────────────────────────────────────────────────────────────

def test_golden_formika_scenario_exists():
    for scenario in FORMIKA_SCENARIOS:
        assert (GOLDEN_DIR / f"{scenario}.docx").exists(), scenario
        assert (GOLDEN_DIR / f"{scenario}.json").exists(), scenario


def test_golden_formika_sample_matches_baseline(work_dir):
    """
    Готовый договор Формики структурно совпадает с эталоном
    tests/data/golden/formika_sample.* (4 машины, пустые строки удалены).
    """
    template_name, payload = FORMIKA_SCENARIOS["formika_sample"]
    expected = json.loads(
        (GOLDEN_DIR / "formika_sample.json").read_text(encoding="utf-8")
    )

    templates_dir = work_dir / "golden_tpl_formika_sample"
    shutil.rmtree(templates_dir, ignore_errors=True)
    templates_dir.mkdir(parents=True, exist_ok=True)
    strip_embedded_fonts(
        PROJECT_ROOT / "templates" / template_name,
        templates_dir / template_name,
    )

    output = work_dir / "golden_out_formika_sample.docx"
    try:
        generator = GeneratorFactory.get_generator(
            "formika", templates_dir=str(templates_dir), strict=True
        )
        generator.generate_docx(payload, str(output))

        actual = fingerprint(output)

        assert actual["paragraphs"] == expected["paragraphs"]
        assert actual["tables"] == expected["tables"]
        assert actual["body"] == expected["body"]
        assert actual["placeholders_left"] == []
        assert actual["sha256_text"] == expected["sha256_text"]
    finally:
        shutil.rmtree(templates_dir, ignore_errors=True)
        output.unlink(missing_ok=True)
