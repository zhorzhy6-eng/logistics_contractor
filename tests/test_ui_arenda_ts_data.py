#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты сборки данных «Разовой аренды» из вкладок (ЭТАП 3.1.D.B.1, дополнены
на ШАГЕ FIX-1).

Проверяется ui/windows/arenda_ts/data.py::collect_arenda_ts_data: раскладка
полей семи вкладок по ContractData, пять маппингов этого шага (корневые поля
распознавания → в contract, точки с time_from / time_to, суммы по виду
Арендатора, разбор паспорта и удостоверения, carrier_type из entity_type
и ставки НДС) и устойчивость сборки (нет вкладки, нет get_data(), get_data()
упал, вернул не словарь).

Отдельно проверяются два поля ШАГА FIX-1: три РАЗНЫЕ даты договора
(lease_start_date / lease_end_date / planned_completion_date — п. 2.5 и
п. 3.3.2 бланка, автоподстановки между ними нет) и срок оплаты
contract.payment_days.

Qt не нужен: вкладки подменяются простыми объектами-заглушками со своим
get_data(). Отдельные тесты проверяют, что собранных данных достаточно
валидатору и генератору типа, что генератор читает время подачи ТС именно из
собранных точек и что данные доходят до готового документа. Модуль сборки
при этом не тянет PyQt5.

Все данные синтетические, реальных ПДн нет.
"""

import logging
import re
from pathlib import Path

import pytest

from core.contract_data import ContractData
from ui.windows.arenda_ts.data import (
    BASIS_IP,
    BASIS_OOO,
    CARRIER_TYPE_IP_WITHOUT_VAT,
    CARRIER_TYPE_IP_WITH_VAT,
    CARRIER_TYPE_OOO,
    DEFAULT_CARRIER_TYPE,
    DEFAULT_VAT_RATE_NUM,
    MAX_CARS,
    MAX_POINTS,
    OGRNIP_LABEL,
    OGRN_LABEL,
    build,
    collect_arenda_ts_data,
)

# ─────────────────────────────────────────────────────────────
# Константы тестовых данных
# ─────────────────────────────────────────────────────────────

#: Шапка договора и срок аренды (п. 2.5).
CONTRACT_NUMBER = "ТЛ-574"
CONTRACT_DATE = "2026-09-19"
LEASE_START = "2026-09-21"
LEASE_END = "2026-09-27"
ROUTE = "г. Москва — г. Калуга"

#: Арендатор-ООО — наша сторона в варианте по умолчанию.
LESSEE_OOO_NAME = "ООО «Арендатор-Тест»"
LESSEE_OOO_SHORT = "ООО «АТ»"
LESSEE_OOO_INN = "7701234567"
LESSEE_OOO_KPP = "770101001"
LESSEE_OOO_OGRN = "1027700132195"
LESSEE_OOO_ADDRESS = "г. Москва, ул. Арендаторская, д. 1"
LESSEE_OOO_DIRECTOR = "Петров Пётр Петрович"

#: Арендатор-ИП: КПП у индивидуального предпринимателя не бывает.
LESSEE_IP_NAME = "Индивидуальный предприниматель Смирнов Сергей Сергеевич"
LESSEE_IP_SHORT = "ИП Смирнов С.С."
LESSEE_IP_INN = "770123456789"
LESSEE_IP_OGRNIP = "321770000123456"
LESSEE_IP_ADDRESS = "г. Москва, ул. Предпринимательская, д. 7"
LESSEE_IP_DIRECTOR = "Смирнов Сергей Сергеевич"

#: Арендодатель — вторая сторона; во всех бланках это ООО.
LESSOR_NAME = "ООО «Арендодатель-Тест»"
LESSOR_SHORT = "ООО «АДТ»"
LESSOR_INN = "7709876543"
LESSOR_OGRN = "1027700132196"
LESSOR_ADDRESS = "г. Москва, ул. Арендодательская, д. 2"
LESSOR_DIRECTOR = "Сидоров Сидор Сидорович"

#: Объект аренды (п. 2.1) — тягач и прицеп.
TRACTOR_BRAND = "Тягач-Модель 5440"
TRACTOR_PLATE = "А001АА01"
TRACTOR_TYPE = "грузовой тягач седельный"
TRAILER_BRAND = "Прицеп-Модель 9"
TRAILER_PLATE = "Б002ББ02"

#: Перевозимые машины (таблица п. 3.1): VIN — 17 символов без букв I, O, Q.
VIN_1 = "XTC651150N0001001"
VIN_2 = "XTC651150N0001002"

#: Точки маршрута (п. 3.2 и 3.3).
LOADING_NAME_1 = "ООО «Склад Север»"
LOADING_ADDRESS_1 = "г. Москва, ул. Складская, д. 1"
LOADING_NAME_2 = "ООО «Склад Юг»"
LOADING_ADDRESS_2 = "г. Калуга, ул. Промышленная, д. 5"
UNLOADING_NAME_1 = "ООО «Приёмка»"
UNLOADING_ADDRESS_1 = "г. Чехов, ул. Приёмная, д. 9"
UNLOADING_NAME_2 = "ООО «Возврат»"
UNLOADING_ADDRESS_2 = "г. Калуга, ул. Возвратная, д. 4"
LOADING_DATE_1 = "2026-09-21"
LOADING_DATE_2 = "2026-09-22"
LOADING_TIME_FROM = "08:00"
LOADING_TIME_TO = "18:00"
UNLOADING_DATE_1 = "2026-09-27"
#: Дата ВТОРОЙ точки выгрузки — своя, не равная дате завершения рейса: с шага
#: FIX-1-T п. 3.3.2 печатает {{planned_completion_date}}, а не дату точки.
UNLOADING_DATE_2 = "2026-09-25"

#: Планируемая дата завершения рейса (п. 3.3.2 бланка). Это ТРЕТЬЯ, отдельная
#: дата договора: в образце ТЛ-574 она не равна окончанию аренды.
COMPLETION_DATE = "2026-09-26"

#: Экипаж (п. 3.5): документы приходят одной строкой — как их отдаёт промпт.
DRIVER_NAME = "Иванов Иван Иванович"
DRIVER_BIRTH = "1980-01-01"
DRIVER_PASSPORT = "18 22 926830"
DRIVER_PASSPORT_SERIES = "18 22"
DRIVER_PASSPORT_NUMBER = "926830"
DRIVER_PASSPORT_ISSUER = "Отделом УФМС России по г. Москве"
DRIVER_PASSPORT_DATE = "2023-01-30"
DRIVER_LICENSE = "99 36 123456"
DRIVER_LICENSE_SERIES = "99 36"
DRIVER_LICENSE_NUMBER = "123456"
DRIVER_LICENSE_DATE = "2020-01-01"
DRIVER_ADDRESS = "г. Москва, ул. Водительская, д. 3"
DRIVER_PHONE = "+7 (999) 123-45-67"

#: Суммы. ООО / ИП с НДС: 221 099,18 + 22% (48 641,82) = 269 741,00.
OOO_BASE_SUM = 221099.18
OOO_VAT_SUM = 48641.82
OOO_TOTAL_SUM = 269741.00
VAT_RATE_NUM = 22.0
VAT_RATE_TEXT = "22%"

#: ИП без НДС: одна сумма, НДС не облагается.
IP_SUM = 135833.00
IP_VAT_RATE_TEXT = "0%"

SPECIAL_CONDITIONS = "Простой не более 24 часов"

#: Варианты бланка — те же значения, что у генератора и валидатора.
VARIANTS = (CARRIER_TYPE_OOO, CARRIER_TYPE_IP_WITH_VAT, CARRIER_TYPE_IP_WITHOUT_VAT)

#: Имена файлов бланков по виду Арендатора (ArendaTsGenerator.TEMPLATE_NAMES).
TEMPLATE_NAMES = {
    CARRIER_TYPE_OOO: "shablon_arenda_ts_ooo.docx",
    CARRIER_TYPE_IP_WITH_VAT: "shablon_arenda_ts_ip_with_vat.docx",
    CARRIER_TYPE_IP_WITHOUT_VAT: "shablon_arenda_ts_ip_without_vat.docx",
}

#: Незаменённый плейсхолдер бланка («{{car_12_brand}}») в готовом документе.
PLACEHOLDER_RE = re.compile(r"\{\{[^{}]*\}\}")


class StubTab:
    """Вкладка-заглушка: отдаёт заранее заданный словарь."""

    def __init__(self, data):
        self._data = data
        self.calls = 0

    def get_data(self):
        self.calls += 1
        return self._data


class BrokenTab:
    """Вкладка, чей get_data() падает."""

    def get_data(self):
        raise RuntimeError("вкладка сломана")


class NoGetDataTab:
    """Вкладка без get_data() (например, ещё не написанная)."""


class NotADictTab:
    """Вкладка, вернувшая не словарь."""

    def get_data(self):
        return ["не", "словарь"]


#: Плейсхолдер Jinja, оставшийся в готовом документе.
PLACEHOLDER_RE = re.compile(r"\{\{[^{}]*\}\}")


def _flatten(text: str) -> str:
    """Текст одной строкой: любые пробелы (в том числе неразрывные) — по одному."""
    return re.sub(r"\s+", " ", text.replace("\u00a0", " ")).strip()


def _format_date(value: str) -> str:
    """ISO-дата в формате бланка: «2026-09-26» → «26.09.2026»."""
    year, month, day = str(value).split("-")
    return f"{day}.{month}.{year}"


# ─────────────────────────────────────────────────────────────
# Данные вкладок: раскладка полей, зафиксированная на ЭТАПЕ 3.1.D.B
# ─────────────────────────────────────────────────────────────

def _lessee_tab(variant: str = CARRIER_TYPE_OOO) -> dict:
    """Вкладка «Арендатор»: вид стороны выбирает пользователь (carrier_type)."""
    if variant == CARRIER_TYPE_OOO:
        return {
            "carrier_type": variant,
            "entity_type": "ООО",
            "full_name": LESSEE_OOO_NAME,
            "short_name": LESSEE_OOO_SHORT,
            "inn": LESSEE_OOO_INN,
            "kpp": LESSEE_OOO_KPP,
            "ogrn": LESSEE_OOO_OGRN,
            "address": LESSEE_OOO_ADDRESS,
            "actual_address": LESSEE_OOO_ADDRESS,
            "account": "40702810000000000001",
            "bank": "ПАО Сбербанк",
            "bik": "044525225",
            "corr_account": "30101810400000000225",
            "email": "arenda@example.ru",
            "edo": "2AE-7F31-4C50",
            "director_position": "Генеральный директор",
            "director_name": LESSEE_OOO_DIRECTOR,
        }

    return {
        "carrier_type": variant,
        "entity_type": "ИП",
        "full_name": LESSEE_IP_NAME,
        "short_name": LESSEE_IP_SHORT,
        "inn": LESSEE_IP_INN,
        # У индивидуального предпринимателя КПП не бывает: поле пустое.
        "kpp": "",
        "ogrn": LESSEE_IP_OGRNIP,
        "address": LESSEE_IP_ADDRESS,
        "actual_address": LESSEE_IP_ADDRESS,
        "account": "40802810000000000011",
        "bank": "АО «Банк Тест»",
        "bik": "044525227",
        "corr_account": "30101810400000000227",
        "email": "ip@example.ru",
        "edo": "1BC-6E20-3B40",
        "director_position": "Индивидуальный предприниматель",
        "director_name": LESSEE_IP_DIRECTOR,
    }


def _lessor_tab() -> dict:
    """Вкладка «Арендодатель» — вторая сторона с полными реквизитами."""
    return {
        "full_name": LESSOR_NAME,
        "short_name": LESSOR_SHORT,
        "inn": LESSOR_INN,
        "ogrn": LESSOR_OGRN,
        "address": LESSOR_ADDRESS,
        "actual_address": LESSOR_ADDRESS,
        "account": "40702810000000000002",
        "bank": "АО «Банк Второй»",
        "bik": "044525226",
        "corr_account": "30101810400000000226",
        "email": "lessor@example.ru",
        "edo": "3CD-8A42-5D60",
        "director_position": "Директор",
        "director_name": LESSOR_DIRECTOR,
    }


def _vehicle_tab() -> dict:
    """Вкладка «ТС»: шапка договора, три даты рейса, тягач и прицеп."""
    return {
        "contract_number": CONTRACT_NUMBER,
        "contract_date": CONTRACT_DATE,
        "lease_start_date": LEASE_START,
        "lease_end_date": LEASE_END,
        "planned_completion_date": COMPLETION_DATE,
        "tractor_brand": TRACTOR_BRAND,
        "tractor_plate": TRACTOR_PLATE,
        "tractor_type": TRACTOR_TYPE,
        "trailer_brand": TRAILER_BRAND,
        "trailer_plate": TRAILER_PLATE,
    }


def _route_tab() -> dict:
    """Вкладка «Маршрут»: направление и точки с датой и временем подачи ТС."""
    return {
        "route": ROUTE,
        "loadings": [
            {"name": LOADING_NAME_1, "address": LOADING_ADDRESS_1,
             "date": LOADING_DATE_1, "time_from": LOADING_TIME_FROM,
             "time_to": LOADING_TIME_TO},
            {"name": LOADING_NAME_2, "address": LOADING_ADDRESS_2,
             "date": LOADING_DATE_2, "time_from": LOADING_TIME_FROM,
             "time_to": LOADING_TIME_TO},
        ],
        "unloadings": [
            {"name": UNLOADING_NAME_1, "address": UNLOADING_ADDRESS_1,
             "date": UNLOADING_DATE_1},
        ],
    }


def _vehicle(number: int) -> dict:
    """Строка таблицы машин: марка, VIN и точки погрузки/выгрузки в строке."""
    return {
        "brand_model": f"МОДЕЛЬ {number}",
        "vin": f"XTC651150N0001{number:03d}",
        "loading_point": LOADING_ADDRESS_1,
        "unloading_point": UNLOADING_ADDRESS_1,
    }


def _cargo_tab(count: int = 2) -> dict:
    """Вкладка «Груз»: перевозимые автомобили таблицы п. 3.1."""
    return {
        "cargo_count": count,
        "vehicles": [_vehicle(n) for n in range(1, count + 1)],
    }


def _crew_tab() -> dict:
    """Вкладка «Экипаж»: паспорт и удостоверение — одной строкой."""
    return {
        "driver_full_name": DRIVER_NAME,
        "driver_birth_date": DRIVER_BIRTH,
        "driver_passport": DRIVER_PASSPORT,
        "driver_passport_issuer": DRIVER_PASSPORT_ISSUER,
        "driver_passport_issue_date": DRIVER_PASSPORT_DATE,
        "driver_license": DRIVER_LICENSE,
        "driver_license_issue_date": DRIVER_LICENSE_DATE,
        "driver_registration_address": DRIVER_ADDRESS,
        "driver_phone": DRIVER_PHONE,
    }


def _price_tab(variant: str = CARRIER_TYPE_OOO) -> dict:
    """Вкладка «Стоимость»: у ИП без НДС сумма одна, НДС не облагается."""
    if variant == CARRIER_TYPE_IP_WITHOUT_VAT:
        return {
            "sum_wo_vat": 0.0,
            "sum_vat": 0.0,
            "sum_total": IP_SUM,
            "vat_rate": IP_VAT_RATE_TEXT,
            "vat_rate_num": 0.0,
            "special_conditions": SPECIAL_CONDITIONS,
        }

    return {
        "sum_wo_vat": OOO_BASE_SUM,
        "sum_vat": OOO_VAT_SUM,
        "sum_total": OOO_TOTAL_SUM,
        "vat_rate": VAT_RATE_TEXT,
        "vat_rate_num": VAT_RATE_NUM,
        "special_conditions": SPECIAL_CONDITIONS,
    }


def _tabs_of_variant(variant: str = CARRIER_TYPE_OOO) -> dict:
    """Все семь вкладок окна, заполненные как в жизни (обычные dict)."""
    return {
        "lessee": _lessee_tab(variant),
        "lessor": _lessor_tab(),
        "vehicle": _vehicle_tab(),
        "route": _route_tab(),
        "cargo": _cargo_tab(),
        "crew": _crew_tab(),
        "price": _price_tab(variant),
    }


@pytest.fixture
def full_tabs() -> dict:
    """Те же семь вкладок, но вкладки — объекты со своим get_data()."""
    return {key: StubTab(data) for key, data in _tabs_of_variant().items()}


# ─────────────────────────────────────────────────────────────
# Пустой вход и устойчивость
# ─────────────────────────────────────────────────────────────

def test_empty_input_gives_empty_contract_data():
    cd = collect_arenda_ts_data({})

    assert isinstance(cd, ContractData)
    assert cd.driver == {}
    assert cd.customer == {}
    assert cd.carrier == {}
    assert cd.vehicles == []
    assert cd.tractor == {}
    assert cd.trailer == {}
    assert cd.loadings == []
    assert cd.unloadings == []
    # Массивы точек остаются в contract всегда (генератор и валидатор читают
    # именно их), но пустой вход не добавляет ни одного значения.
    assert cd.contract == {"loadings": [], "unloadings": []}


def test_empty_input_keeps_city_empty():
    """Город в сборе не задаём: ContractData.resolved_city() подставит «Москва»."""
    cd = collect_arenda_ts_data({})

    assert cd.city == ""
    assert cd.resolved_city() == "Москва"


def test_empty_input_has_no_carrier_type():
    """
    Вид Арендатора не выдумывается: у пустого входа поля нет.

    Генератор и валидатор и без него берут вариант по умолчанию (ООО), но
    в данных интерфейса пустое должно остаться пустым.
    """
    cd = collect_arenda_ts_data({})

    assert "carrier_type" not in cd.contract
    assert "vat_rate" not in cd.contract
    assert "vat_rate_num" not in cd.contract


def test_partial_tabs_do_not_break_collection():
    cd = collect_arenda_ts_data({
        "lessee": {"full_name": LESSEE_OOO_NAME},
        "vehicle": {"contract_number": CONTRACT_NUMBER},
    })

    assert cd.contract["number"] == CONTRACT_NUMBER
    assert cd.contract["lessee"]["full_name"] == LESSEE_OOO_NAME
    assert cd.vehicles == []
    assert cd.driver == {}


def test_tab_object_with_get_data_is_used():
    tab = StubTab({"contract_number": f"  {CONTRACT_NUMBER}  "})

    cd = collect_arenda_ts_data({"vehicle": tab})

    assert tab.calls == 1
    assert cd.contract["number"] == CONTRACT_NUMBER


def test_every_tab_is_read_only_once(full_tabs):
    """
    Каждая вкладка читается ровно один раз.

    Вид Арендатора выводится из вкладок «Арендатор» и «Стоимость», а от него
    зависит сборка сторон и сумм: если бы разделы читали вкладки сами, get_data()
    вызывался бы по несколько раз (и, например, перечитывал справочник).
    """
    collect_arenda_ts_data(full_tabs)

    for key, tab in full_tabs.items():
        assert tab.calls == 1, f"вкладка «{key}» прочитана {tab.calls} раз(а)"


def test_failing_get_data_does_not_break_collection():
    tabs = {
        "lessee": BrokenTab(),
        "route": {"route": ROUTE},
    }

    cd = collect_arenda_ts_data(tabs)

    assert "lessee" not in cd.contract
    assert cd.contract["route"] == ROUTE


def test_tab_without_get_data_is_skipped():
    cd = collect_arenda_ts_data({"crew": NoGetDataTab()})

    assert cd.driver == {}


def test_tab_returning_not_a_dict_is_skipped():
    cd = collect_arenda_ts_data({"cargo": NotADictTab()})

    assert cd.vehicles == []


def test_tabs_not_a_mapping_gives_empty_result():
    """Даже совсем неверный вход не роняет сборку."""
    cd = collect_arenda_ts_data(None)

    assert isinstance(cd, ContractData)
    assert cd.contract == {"loadings": [], "unloadings": []}
    assert cd.vehicles == []


def test_data_module_does_not_import_qt():
    """Сборка данных не зависит от интерфейса — PyQt5 в модуле не нужен."""
    import ui.windows.arenda_ts.data as data_module

    assert not hasattr(data_module, "QWidget")
    assert "PyQt5" not in getattr(data_module, "__dict__", {})


# ─────────────────────────────────────────────────────────────
# Маппинг 1: корневые поля распознавания → в contract
# ─────────────────────────────────────────────────────────────

def test_lessee_block_goes_to_contract(full_tabs):
    cd = collect_arenda_ts_data(full_tabs)

    assert cd.contract["lessee"]["full_name"] == LESSEE_OOO_NAME
    assert cd.contract["lessee"]["inn"] == LESSEE_OOO_INN


def test_lessor_block_goes_to_contract(full_tabs):
    cd = collect_arenda_ts_data(full_tabs)

    assert cd.contract["lessor"]["full_name"] == LESSOR_NAME
    assert cd.contract["lessor"]["inn"] == LESSOR_INN


def test_route_and_lease_dates_go_to_contract(full_tabs):
    cd = collect_arenda_ts_data(full_tabs)

    assert cd.contract["route"] == ROUTE
    assert cd.contract["lease_start_date"] == LEASE_START
    assert cd.contract["lease_end_date"] == LEASE_END


def test_root_fields_of_recognition_are_lost_by_coerce(full_tabs):
    """
    Маппинг 1: ContractData.coerce корневые ключи не хранит.

    Промпт отдаёт lessee / lessor / route / lease_start_date / lease_end_date
    в КОРНЕ ответа, а генератор и валидатор читают их из contract. Сборщик
    обязан положить их сразу в contract — на _hoist_contract_fields надежды нет:
    окно передаёт генератору готовый ContractData, а тот не-Mapping.
    """
    root_only = ContractData.coerce({
        "lessee": {"full_name": LESSEE_OOO_NAME},
        "route": ROUTE,
        "lease_start_date": LEASE_START,
    })
    assert "lessee" not in root_only.contract
    assert "route" not in root_only.contract
    assert "lease_start_date" not in root_only.contract

    cd = collect_arenda_ts_data(full_tabs)

    assert cd.contract["lessee"]["full_name"] == LESSEE_OOO_NAME
    assert cd.contract["lessor"]["full_name"] == LESSOR_NAME
    assert cd.contract["route"] == ROUTE
    assert cd.contract["lease_start_date"] == LEASE_START
    assert cd.contract["lease_end_date"] == LEASE_END

    # Повторное приведение (генератор делает его сам) данные не теряет.
    coerced = ContractData.coerce(cd)
    assert coerced.contract["lessee"]["short_name"] == LESSEE_OOO_SHORT
    assert coerced.contract["lease_end_date"] == LEASE_END


def test_empty_party_tab_leaves_contract_without_block():
    """Незаполненная вкладка не оставляет за собой пустой блок стороны."""
    cd = collect_arenda_ts_data({"lessee": {"full_name": "  "}, "lessor": {}})

    assert "lessee" not in cd.contract
    assert "lessor" not in cd.contract
    assert cd.customer == {}
    assert cd.carrier == {}


def test_customer_and_carrier_mirror_party_names(full_tabs):
    """
    Арендатор дублируется в customer, Арендодатель — в carrier.

    Так их читает валидатор, если блока стороны в contract не окажется
    (ArendaTsValidator._party_block), и так же устроены данные других типов.
    """
    cd = collect_arenda_ts_data(full_tabs)

    assert cd.customer == {
        "full_name": LESSEE_OOO_NAME,
        "short_name": LESSEE_OOO_SHORT,
    }
    assert cd.carrier == {
        "full_name": LESSOR_NAME,
        "short_name": LESSOR_SHORT,
    }


# ─────────────────────────────────────────────────────────────
# Маппинг 2: точки маршрута с time_from / time_to
# ─────────────────────────────────────────────────────────────

def test_points_are_stored_twice(full_tabs):
    """
    Точки лежат и в contract (полный набор), и в ContractData (приведённые).

    ContractData._as_point_list оставляет у точки только
    {address, date, time_window}: если бы точки клались лишь наверх, генератор
    не увидел бы ни названий, ни времени подачи ТС.
    """
    cd = collect_arenda_ts_data(full_tabs)

    assert len(cd.contract["loadings"]) == 2
    assert len(cd.loadings) == 2
    assert set(cd.contract["loadings"][0]) == {
        "name", "address", "date", "time_from", "time_to", "time_window",
    }
    assert set(cd.loadings[0]) == {"address", "date", "time_window"}


def test_loading_points_keep_time_from_and_time_to(full_tabs):
    cd = collect_arenda_ts_data(full_tabs)

    point = cd.contract["loadings"][0]
    assert point["name"] == LOADING_NAME_1
    assert point["address"] == LOADING_ADDRESS_1
    assert point["date"] == LOADING_DATE_1
    assert point["time_from"] == LOADING_TIME_FROM
    assert point["time_to"] == LOADING_TIME_TO

    # В приведённой точке время подачи ТС остаётся только окном одной строкой.
    assert cd.loadings[0]["time_window"] == f"{LOADING_TIME_FROM}-{LOADING_TIME_TO}"
    assert "name" not in cd.loadings[0]
    assert "time_from" not in cd.loadings[0]


def test_point_time_window_is_built_from_times():
    cd = collect_arenda_ts_data({"route": {"loadings": [
        {"address": LOADING_ADDRESS_1, "time_from": "07:30", "time_to": "12:45"},
    ]}})

    assert cd.contract["loadings"][0]["time_window"] == "07:30-12:45"
    assert cd.loadings[0]["time_window"] == "07:30-12:45"


def test_point_with_single_time_boundary_keeps_it():
    """Заполнена одна граница — в окно попадает она, вторая не выдумывается."""
    cd = collect_arenda_ts_data({"route": {"loadings": [
        {"address": LOADING_ADDRESS_1, "time_from": "08:00", "time_to": ""},
    ]}})

    assert cd.contract["loadings"][0]["time_window"] == "08:00"


def test_unloadings_keep_name_address_and_date(full_tabs):
    cd = collect_arenda_ts_data(full_tabs)

    point = cd.contract["unloadings"][0]
    assert point["name"] == UNLOADING_NAME_1
    assert point["address"] == UNLOADING_ADDRESS_1
    assert point["date"] == UNLOADING_DATE_1
    # У точки выгрузки времени подачи ТС в бланке нет — поля пустые, а не выдуманные.
    assert point["time_from"] == ""
    assert point["time_to"] == ""
    assert point["time_window"] == ""
    assert cd.unloadings[0]["address"] == UNLOADING_ADDRESS_1


def test_points_without_name_and_address_are_dropped():
    cd = collect_arenda_ts_data({"route": {
        "loadings": [
            {"name": "", "address": "", "date": LOADING_DATE_1},
            {"name": "  ", "address": "  ", "time_from": "08:00"},
            {"name": LOADING_NAME_1, "address": LOADING_ADDRESS_1},
        ],
    }})

    assert len(cd.contract["loadings"]) == 1
    assert cd.contract["loadings"][0]["address"] == LOADING_ADDRESS_1


def test_point_with_only_name_is_kept():
    """Заполнено название, адрес допишут позже — строку не выбрасываем."""
    cd = collect_arenda_ts_data({"route": {"loadings": [
        {"name": LOADING_NAME_1, "address": ""},
    ]}})

    assert cd.contract["loadings"][0]["name"] == LOADING_NAME_1
    assert cd.contract["loadings"][0]["address"] == ""


def test_points_are_limited_to_max_points():
    rows = [
        {"name": f"ПОСТ {n}", "address": f"г. Москва, ул. Складская, д. {n}"}
        for n in range(1, MAX_POINTS + 3)
    ]

    cd = collect_arenda_ts_data({"route": {"loadings": rows, "unloadings": rows}})

    assert MAX_POINTS == 10
    assert len(cd.contract["loadings"]) == MAX_POINTS
    assert len(cd.contract["unloadings"]) == MAX_POINTS
    assert cd.contract["loadings"][-1]["address"] == (
        f"г. Москва, ул. Складская, д. {MAX_POINTS}"
    )


def test_points_not_a_list_are_ignored():
    cd = collect_arenda_ts_data({"route": {"loadings": "мусор", "unloadings": None}})

    assert cd.contract["loadings"] == []
    assert cd.contract["unloadings"] == []
    assert cd.loadings == []


def test_points_not_mappings_are_ignored():
    cd = collect_arenda_ts_data({"route": {"loadings": ["строка", 42, None]}})

    assert cd.contract["loadings"] == []


def test_generator_reads_point_times_from_collected_data(full_tabs, templates_dir):
    """
    Время подачи ТС доходит до карты замен генератора.

    Генератор читает точку двумя путями (ContractData.loadings и
    contract["loadings"]); время берётся из второго — значит, точки обязаны
    лежать там полным набором.
    """
    from core.contracts.arenda_ts.generator import ArendaTsGenerator

    generator = ArendaTsGenerator(templates_dir=str(templates_dir))
    replacements = generator.build_replacements(collect_arenda_ts_data(full_tabs))

    assert replacements["loading_1_address"] == LOADING_ADDRESS_1
    assert replacements["loading_1_time_from"] == LOADING_TIME_FROM
    assert replacements["loading_1_time_to"] == LOADING_TIME_TO
    assert replacements["loading_1_date"] == "21.09.2026"
    assert replacements["loading_2_address"] == LOADING_ADDRESS_2
    assert replacements["unloading_1_address"] == UNLOADING_ADDRESS_1
    assert replacements["unloading_1_date"] == "27.09.2026"


def test_generator_gets_point_time_from_time_window_when_times_are_absent(
    templates_dir,
):
    """Если вкладка отдала только окно одной строкой, генератор разберёт его сам."""
    from core.contracts.arenda_ts.generator import ArendaTsGenerator

    cd = collect_arenda_ts_data({"route": {"loadings": [
        {"address": LOADING_ADDRESS_1, "date": LOADING_DATE_1,
         "time_from": "", "time_to": "", "time_window": "09:00-15:00"},
    ]}})

    # Готовое окно вкладки не теряется: границ нет, значит берётся оно.
    assert cd.contract["loadings"][0]["time_window"] == "09:00-15:00"
    assert cd.loadings[0]["time_window"] == "09:00-15:00"

    generator = ArendaTsGenerator(templates_dir=str(templates_dir))
    replacements = generator.build_replacements(cd)

    assert replacements["loading_1_time_from"] == "09:00"
    assert replacements["loading_1_time_to"] == "15:00"


# ─────────────────────────────────────────────────────────────
# Маппинг 3: суммы по виду Арендатора
# ─────────────────────────────────────────────────────────────

def test_ooo_sums_are_mapped():
    cd = collect_arenda_ts_data({"price": _price_tab(CARRIER_TYPE_OOO)})

    assert cd.contract["price_without_vat"] == OOO_BASE_SUM
    assert cd.contract["price_with_vat"] == OOO_TOTAL_SUM
    assert cd.contract["vat_amount"] == OOO_VAT_SUM
    assert cd.contract["vat_rate_num"] == VAT_RATE_NUM
    assert cd.contract["vat_rate"] == VAT_RATE_TEXT
    # Суммы документа остаются рядом: их читает генератор (_base_price).
    assert cd.contract["sum_wo_vat"] == OOO_BASE_SUM
    assert cd.contract["sum_vat"] == OOO_VAT_SUM
    assert cd.contract["sum_total"] == OOO_TOTAL_SUM


def test_ip_with_vat_sums_are_mapped():
    tabs = {"lessee": _lessee_tab(CARRIER_TYPE_IP_WITH_VAT),
            "price": _price_tab(CARRIER_TYPE_IP_WITH_VAT)}

    cd = collect_arenda_ts_data(tabs)

    assert cd.contract["carrier_type"] == CARRIER_TYPE_IP_WITH_VAT
    assert cd.contract["price_without_vat"] == OOO_BASE_SUM
    assert cd.contract["price_with_vat"] == OOO_TOTAL_SUM
    assert cd.contract["vat_rate_num"] == VAT_RATE_NUM


def test_ip_without_vat_single_sum_is_mapped():
    tabs = {"lessee": _lessee_tab(CARRIER_TYPE_IP_WITHOUT_VAT),
            "price": _price_tab(CARRIER_TYPE_IP_WITHOUT_VAT)}

    cd = collect_arenda_ts_data(tabs)

    assert cd.contract["carrier_type"] == CARRIER_TYPE_IP_WITHOUT_VAT
    # Единственная сумма документа — без НДС, ставка не применяется.
    assert cd.contract["price_without_vat"] == IP_SUM
    assert cd.contract["sum_total"] == IP_SUM
    assert cd.contract["vat_rate_num"] == 0.0
    assert cd.contract["vat_rate"] == IP_VAT_RATE_TEXT
    # Плейсхолдеров сумм НДС в варианте без НДС нет — и ключей тоже.
    assert "sum_wo_vat" not in cd.contract
    assert "sum_vat" not in cd.contract
    assert "price_with_vat" not in cd.contract


def test_ooo_with_only_total_uses_total_without_vat():
    """
    У ООО в документе только итог: он и становится суммой без НДС.

    Ставка при этом не выдумывается (0.0): считать НДС от чужой суммы нельзя,
    о расхождении скажет валидатор.
    """
    cd = collect_arenda_ts_data({
        "lessee": _lessee_tab(CARRIER_TYPE_OOO),
        "price": {"sum_total": OOO_TOTAL_SUM},
    })

    assert cd.contract["price_without_vat"] == OOO_TOTAL_SUM
    assert cd.contract["vat_rate_num"] == 0.0
    assert cd.contract["vat_rate"] == "0%"


def test_vat_rate_string_is_parsed():
    """Ставка строкой «22%» разбирается, если числа в поле нет."""
    cd = collect_arenda_ts_data({"price": {
        "sum_wo_vat": OOO_BASE_SUM,
        "vat_rate": "22%",
    }})

    assert cd.contract["vat_rate_num"] == VAT_RATE_NUM
    assert cd.contract["vat_rate"] == VAT_RATE_TEXT


def test_zero_vat_rate_string_means_without_vat():
    cd = collect_arenda_ts_data({"price": {
        "sum_wo_vat": OOO_BASE_SUM,
        "vat_rate": "Без НДС",
    }})

    assert cd.contract["vat_rate_num"] == 0.0
    assert cd.contract["vat_rate"] == "0%"


def test_vat_rate_defaults_when_absent():
    """Ставки нет ни числом, ни строкой — берём ставку по умолчанию, как генератор."""
    cd = collect_arenda_ts_data({"price": {"sum_wo_vat": OOO_BASE_SUM}})

    assert cd.contract["vat_rate_num"] == DEFAULT_VAT_RATE_NUM
    assert cd.contract["vat_rate"] == VAT_RATE_TEXT


def test_special_conditions_are_mapped():
    cd = collect_arenda_ts_data({"price": {"special_conditions": SPECIAL_CONDITIONS}})

    assert cd.contract["special_conditions"] == SPECIAL_CONDITIONS


def test_empty_price_tab_adds_no_sums():
    cd = collect_arenda_ts_data({"lessee": _lessee_tab(CARRIER_TYPE_OOO)})

    assert "price_without_vat" not in cd.contract
    assert "sum_total" not in cd.contract
    assert "vat_rate" not in cd.contract


def test_price_tab_returning_recognition_contract_block_is_read():
    """Вкладка может отдать блок contract распознавания как есть."""
    cd = collect_arenda_ts_data({"price": {"contract": {
        "sum_wo_vat": OOO_BASE_SUM,
        "sum_vat": OOO_VAT_SUM,
        "sum_total": OOO_TOTAL_SUM,
        "vat_rate": VAT_RATE_TEXT,
    }}})

    assert cd.contract["price_without_vat"] == OOO_BASE_SUM
    assert cd.contract["price_with_vat"] == OOO_TOTAL_SUM
    assert cd.contract["vat_rate_num"] == VAT_RATE_NUM


# ─────────────────────────────────────────────────────────────
# Маппинг 4: экипаж — разбор строк
# ─────────────────────────────────────────────────────────────

def test_passport_string_is_split():
    cd = collect_arenda_ts_data({"crew": {"driver_passport": DRIVER_PASSPORT}})

    assert cd.driver["passport_series"] == DRIVER_PASSPORT_SERIES
    assert cd.driver["passport_number"] == DRIVER_PASSPORT_NUMBER
    # Строкой паспорт не дублируется: генератор соберёт её из серии и номера.
    assert "passport" not in cd.driver


def test_license_string_is_split():
    cd = collect_arenda_ts_data({"crew": {"driver_license": DRIVER_LICENSE}})

    assert cd.driver["license_series"] == DRIVER_LICENSE_SERIES
    assert cd.driver["license_number"] == DRIVER_LICENSE_NUMBER
    assert "license" not in cd.driver


def test_document_with_other_separators_is_split():
    """Разделители в документе бывают любыми — цифры важнее."""
    cd = collect_arenda_ts_data({"crew": {
        "driver_passport": "18-22 № 926830",
        "driver_license": "99/36 123456",
    }})

    assert cd.driver["passport_series"] == DRIVER_PASSPORT_SERIES
    assert cd.driver["passport_number"] == DRIVER_PASSPORT_NUMBER
    assert cd.driver["license_series"] == DRIVER_LICENSE_SERIES
    assert cd.driver["license_number"] == DRIVER_LICENSE_NUMBER


def test_already_split_documents_are_kept():
    """Справочник водителя отдаёт серию и номер отдельно — не трогаем."""
    cd = collect_arenda_ts_data({"crew": {
        "driver_passport_series": DRIVER_PASSPORT_SERIES,
        "driver_passport_number": DRIVER_PASSPORT_NUMBER,
        "driver_license_series": DRIVER_LICENSE_SERIES,
        "driver_license_number": DRIVER_LICENSE_NUMBER,
    }})

    assert cd.driver["passport_series"] == DRIVER_PASSPORT_SERIES
    assert cd.driver["passport_number"] == DRIVER_PASSPORT_NUMBER
    assert cd.driver["license_series"] == DRIVER_LICENSE_SERIES
    assert cd.driver["license_number"] == DRIVER_LICENSE_NUMBER
    assert "passport" not in cd.driver
    assert "license" not in cd.driver


def test_unparsable_document_is_kept_as_is():
    """
    Нестандартная строка документа остаётся строкой.

    Генератор (_fill_driver) и валидатор (_document) принимают и такой вид:
    потерять данные хуже, чем напечатать строку как в документе.
    """
    cd = collect_arenda_ts_data({"crew": {"driver_passport": "18 22 926830 5"}})

    assert cd.driver["passport"] == "18 22 926830 5"
    assert "passport_series" not in cd.driver
    assert "passport_number" not in cd.driver


def test_document_without_digits_is_kept_as_is():
    cd = collect_arenda_ts_data({"crew": {"driver_license": "б/н"}})

    assert cd.driver["license"] == "б/н"


def test_license_issue_date_is_extracted_from_license_line():
    """
    Дата выдачи из строки удостоверения уходит в своё поле.

    Промпт предупреждает, что дата может стоять в той же строке, что и номер
    (core/prompts/arenda_ts.py): если её не отделить, цифры даты испортят разбор.
    """
    cd = collect_arenda_ts_data({"crew": {
        "driver_license": f"{DRIVER_LICENSE} от 30.01.2020",
    }})

    assert cd.driver["license_series"] == DRIVER_LICENSE_SERIES
    assert cd.driver["license_number"] == DRIVER_LICENSE_NUMBER
    assert cd.driver["license_issue_date"] == "30.01.2020"


def test_explicit_license_date_wins_over_date_in_line():
    cd = collect_arenda_ts_data({"crew": {
        "driver_license": f"{DRIVER_LICENSE} от 30.01.2020",
        "driver_license_issue_date": DRIVER_LICENSE_DATE,
    }})

    assert cd.driver["license_issue_date"] == DRIVER_LICENSE_DATE


def test_all_driver_fields_are_collected():
    """Экипаж: девять полей бланка — ФИО, даты, документы, адрес и телефон."""
    cd = collect_arenda_ts_data({"crew": _crew_tab()})

    assert cd.driver == {
        "full_name": DRIVER_NAME,
        "birth_date": DRIVER_BIRTH,
        "passport_series": DRIVER_PASSPORT_SERIES,
        "passport_number": DRIVER_PASSPORT_NUMBER,
        "passport_issuer": DRIVER_PASSPORT_ISSUER,
        "passport_issue_date": DRIVER_PASSPORT_DATE,
        "license_series": DRIVER_LICENSE_SERIES,
        "license_number": DRIVER_LICENSE_NUMBER,
        "license_issue_date": DRIVER_LICENSE_DATE,
        "registration_address": DRIVER_ADDRESS,
        "phone": DRIVER_PHONE,
    }


def test_empty_crew_tab_leaves_driver_empty():
    cd = collect_arenda_ts_data({"crew": {"driver_full_name": "   "}})

    assert cd.driver == {}


def test_recognition_driver_block_is_read():
    """
    Вкладка может отдать блок driver распознавания как есть.

    Имя поля адреса у промпта — address; генератору и валидатору нужен
    registration_address.
    """
    cd = collect_arenda_ts_data({"crew": {"driver": {
        "full_name": DRIVER_NAME,
        "birth_date": DRIVER_BIRTH,
        "passport": DRIVER_PASSPORT,
        "passport_issuer": DRIVER_PASSPORT_ISSUER,
        "passport_issue_date": DRIVER_PASSPORT_DATE,
        "license": DRIVER_LICENSE,
        "license_issue_date": DRIVER_LICENSE_DATE,
        "address": DRIVER_ADDRESS,
        "phone": DRIVER_PHONE,
    }}})

    assert cd.driver["full_name"] == DRIVER_NAME
    assert cd.driver["passport_series"] == DRIVER_PASSPORT_SERIES
    assert cd.driver["license_number"] == DRIVER_LICENSE_NUMBER
    assert cd.driver["registration_address"] == DRIVER_ADDRESS
    assert cd.driver["phone"] == DRIVER_PHONE


def test_generator_assembles_documents_from_split_fields(full_tabs, templates_dir):
    """Разобранные серия и номер доходят до бланка одной строкой."""
    from core.contracts.arenda_ts.generator import ArendaTsGenerator

    generator = ArendaTsGenerator(templates_dir=str(templates_dir))
    replacements = generator.build_replacements(collect_arenda_ts_data(full_tabs))

    assert replacements["driver_passport"] == DRIVER_PASSPORT
    assert replacements["driver_license"] == DRIVER_LICENSE
    assert replacements["driver_full_name"] == DRIVER_NAME
    assert replacements["driver_address"] == DRIVER_ADDRESS


# ─────────────────────────────────────────────────────────────
# Маппинг 5: carrier_type из entity_type и ставки НДС
# ─────────────────────────────────────────────────────────────

def test_explicit_carrier_type_wins():
    """Выбор пользователя на вкладке «Арендатор» важнее догадки по entity_type."""
    cd = collect_arenda_ts_data({
        "lessee": {"carrier_type": CARRIER_TYPE_IP_WITHOUT_VAT,
                   "entity_type": "ООО",
                   "full_name": LESSEE_IP_NAME},
    })

    assert cd.contract["carrier_type"] == CARRIER_TYPE_IP_WITHOUT_VAT


def test_ip_with_zero_rate_gives_ip_without_vat():
    cd = collect_arenda_ts_data({
        "lessee": {"entity_type": "ИП", "full_name": LESSEE_IP_NAME},
        "price": {"sum_total": IP_SUM, "vat_rate": IP_VAT_RATE_TEXT},
    })

    assert cd.contract["carrier_type"] == CARRIER_TYPE_IP_WITHOUT_VAT
    assert cd.contract["price_without_vat"] == IP_SUM


def test_ip_with_rate_gives_ip_with_vat():
    cd = collect_arenda_ts_data({
        "lessee": {"entity_type": "ИП", "full_name": LESSEE_IP_NAME},
        "price": {"sum_wo_vat": OOO_BASE_SUM, "vat_rate": VAT_RATE_TEXT},
    })

    assert cd.contract["carrier_type"] == CARRIER_TYPE_IP_WITH_VAT
    assert cd.contract["price_without_vat"] == OOO_BASE_SUM


def test_ip_rate_number_without_string_is_enough():
    """Ставка числом (vat_rate_num) работает так же, как строка «22%»."""
    cd = collect_arenda_ts_data({
        "lessee": {"entity_type": "ИП", "full_name": LESSEE_IP_NAME},
        "price": {"sum_wo_vat": OOO_BASE_SUM, "vat_rate_num": VAT_RATE_NUM},
    })

    assert cd.contract["carrier_type"] == CARRIER_TYPE_IP_WITH_VAT


def test_entity_type_ooo_gives_ooo():
    cd = collect_arenda_ts_data({
        "lessee": {"entity_type": "ООО", "full_name": LESSEE_OOO_NAME},
    })

    assert cd.contract["carrier_type"] == CARRIER_TYPE_OOO


def test_entity_type_is_read_from_recognition_block():
    """entity_type ищется и во вложенном блоке lessee (ответ распознавания)."""
    cd = collect_arenda_ts_data({
        "lessee": {"lessee": {"entity_type": "ИП", "full_name": LESSEE_IP_NAME},
                   "contract": {"vat_rate": "0%"}},
        "price": {"sum_total": IP_SUM, "vat_rate": IP_VAT_RATE_TEXT},
    })

    assert cd.contract["carrier_type"] == CARRIER_TYPE_IP_WITHOUT_VAT


def test_short_ip_value_uses_rate():
    """«ИП» без уточнения: вариант выбирает ставка НДС."""
    cd = collect_arenda_ts_data({
        "lessee": {"carrier_type": "ИП", "full_name": LESSEE_IP_NAME},
        "price": {"sum_total": IP_SUM, "vat_rate": IP_VAT_RATE_TEXT},
    })

    assert cd.contract["carrier_type"] == CARRIER_TYPE_IP_WITHOUT_VAT


def test_unknown_entity_type_defaults_to_ooo():
    cd = collect_arenda_ts_data({
        "lessee": {"full_name": LESSEE_OOO_NAME},
        "price": {"sum_wo_vat": OOO_BASE_SUM, "vat_rate": VAT_RATE_TEXT},
    })

    assert cd.contract["carrier_type"] == DEFAULT_CARRIER_TYPE == CARRIER_TYPE_OOO


def test_carrier_type_selects_template_variant(templates_dir):
    """Вид Арендатора из собранных данных выбирает файл бланка."""
    from core.contracts.arenda_ts.generator import ArendaTsGenerator

    generator = ArendaTsGenerator(templates_dir=str(templates_dir))

    for variant in VARIANTS:
        cd = collect_arenda_ts_data(_tabs_of_variant(variant))
        path = generator._get_template_path(cd.contract["carrier_type"])
        assert path.endswith(TEMPLATE_NAMES[variant]), f"вариант {variant}"


# ─────────────────────────────────────────────────────────────
# Арендатор и Арендодатель: реквизиты
# ─────────────────────────────────────────────────────────────

def test_lessee_requisites_are_translated(full_tabs):
    """Поля вкладки переводятся в имена блока lessee, которые читает генератор."""
    cd = collect_arenda_ts_data(full_tabs)
    lessee = cd.contract["lessee"]

    assert lessee["full_name"] == LESSEE_OOO_NAME
    assert lessee["short_name"] == LESSEE_OOO_SHORT
    assert lessee["inn"] == LESSEE_OOO_INN
    assert lessee["kpp"] == LESSEE_OOO_KPP
    assert lessee["ogrn"] == LESSEE_OOO_OGRN
    assert lessee["legal_address"] == LESSEE_OOO_ADDRESS
    assert lessee["actual_address"] == LESSEE_OOO_ADDRESS
    assert lessee["bank_account"] == "40702810000000000001"
    assert lessee["bank_name"] == "ПАО Сбербанк"
    assert lessee["bik"] == "044525225"
    assert lessee["corr_account"] == "30101810400000000225"
    assert lessee["email"] == "arenda@example.ru"
    assert lessee["edo"] == "2AE-7F31-4C50"
    assert lessee["director_position"] == "Генеральный директор"
    assert lessee["director_name"] == LESSEE_OOO_DIRECTOR


def test_lessee_address_is_not_duplicated_as_plain_address(full_tabs):
    """Адрес вкладки — это юридический адрес бланка (legal_address)."""
    lessee = collect_arenda_ts_data(full_tabs).contract["lessee"]

    assert "address" not in lessee
    assert "account" not in lessee
    assert "bank" not in lessee


def test_lessee_kpp_only_for_ooo():
    """У индивидуального предпринимателя КПП не бывает — ключ не кладём."""
    cd = collect_arenda_ts_data({"lessee": {
        "carrier_type": CARRIER_TYPE_IP_WITHOUT_VAT,
        "full_name": LESSEE_IP_NAME,
        "kpp": LESSEE_OOO_KPP,
    }})

    assert "kpp" not in cd.contract["lessee"]

    cd_ooo = collect_arenda_ts_data({"lessee": {
        "carrier_type": CARRIER_TYPE_OOO,
        "full_name": LESSEE_OOO_NAME,
        "kpp": LESSEE_OOO_KPP,
    }})

    assert cd_ooo.contract["lessee"]["kpp"] == LESSEE_OOO_KPP


def test_ooo_labels_and_basis():
    cd = collect_arenda_ts_data({"lessee": _lessee_tab(CARRIER_TYPE_OOO)})

    assert cd.contract["lessee"]["ogrn_label"] == OGRN_LABEL
    assert cd.contract["lessee"]["basis"] == BASIS_OOO
    assert cd.contract["lessee"]["entity_type"] == "ООО"


def test_ip_labels_and_basis():
    """ИП: метка «ОГРНИП» и основание — свидетельство о государственной регистрации."""
    cd = collect_arenda_ts_data({"lessee": _lessee_tab(CARRIER_TYPE_IP_WITHOUT_VAT)})

    assert cd.contract["lessee"]["ogrn_label"] == OGRNIP_LABEL
    assert cd.contract["lessee"]["basis"] == BASIS_IP
    assert cd.contract["lessee"]["entity_type"] == "ИП"
    assert "kpp" not in cd.contract["lessee"]


def test_tab_basis_wins_over_variant_default():
    cd = collect_arenda_ts_data({"lessee": {
        "carrier_type": CARRIER_TYPE_OOO,
        "full_name": LESSEE_OOO_NAME,
        "basis": "Устава и договора купли-продажи",
    }})

    assert cd.contract["lessee"]["basis"] == "Устава и договора купли-продажи"


def test_lessor_requisites_are_translated(full_tabs):
    cd = collect_arenda_ts_data(full_tabs)
    lessor = cd.contract["lessor"]

    assert lessor["full_name"] == LESSOR_NAME
    assert lessor["short_name"] == LESSOR_SHORT
    assert lessor["inn"] == LESSOR_INN
    assert lessor["ogrn"] == LESSOR_OGRN
    assert lessor["legal_address"] == LESSOR_ADDRESS
    assert lessor["bank_account"] == "40702810000000000002"
    assert lessor["bank_name"] == "АО «Банк Второй»"
    assert lessor["bik"] == "044525226"
    assert lessor["corr_account"] == "30101810400000000226"
    assert lessor["email"] == "lessor@example.ru"
    assert lessor["edo"] == "3CD-8A42-5D60"
    assert lessor["director_position"] == "Директор"
    assert lessor["director_name"] == LESSOR_DIRECTOR


def test_lessor_stays_ooo_in_ip_variants():
    """
    Бланки рассчитаны на Арендодателя-ООО: у Арендатора-ИП метка и основание
    Арендодателя не меняются — те же константы, что подставляет генератор.
    """
    for variant in (CARRIER_TYPE_IP_WITH_VAT, CARRIER_TYPE_IP_WITHOUT_VAT):
        cd = collect_arenda_ts_data({
            "lessee": _lessee_tab(variant),
            "lessor": _lessor_tab(),
        })
        lessor = cd.contract["lessor"]

        assert lessor["ogrn_label"] == OGRN_LABEL
        assert lessor["basis"] == BASIS_OOO
        assert "kpp" not in lessor


# ─────────────────────────────────────────────────────────────
# ТС: договор, срок аренды, тягач и прицеп
# ─────────────────────────────────────────────────────────────

def test_contract_number_date_and_lease_dates(full_tabs):
    cd = collect_arenda_ts_data(full_tabs)

    assert cd.contract["number"] == CONTRACT_NUMBER
    assert cd.contract["date"] == CONTRACT_DATE
    assert cd.contract["lease_start_date"] == LEASE_START
    assert cd.contract["lease_end_date"] == LEASE_END


# ─────────────────────────────────────────────────────────────
# Три даты рейса (FIX-1): п. 2.5 и п. 3.3.2 — разные плейсхолдеры
# ─────────────────────────────────────────────────────────────

def test_planned_completion_date_is_collected(full_tabs):
    """Планируемая дата завершения рейса (п. 3.3.2) доходит до contract."""
    cd = collect_arenda_ts_data(full_tabs)

    assert cd.contract["planned_completion_date"] == COMPLETION_DATE


def test_planned_completion_date_from_contract_block():
    """Дата может прийти и внутри блока contract вкладки «ТС»."""
    cd = collect_arenda_ts_data({"vehicle": {
        "contract": {"planned_completion_date": COMPLETION_DATE},
    }})

    assert cd.contract["planned_completion_date"] == COMPLETION_DATE


def test_three_dates_are_collected_independently(full_tabs):
    """Сборщик не связывает даты: каждая читается своим ключом."""
    cd = collect_arenda_ts_data(full_tabs)
    only_start = collect_arenda_ts_data({"vehicle": {
        "lease_start_date": LEASE_START,
    }})
    only_end = collect_arenda_ts_data({"vehicle": {
        "lease_end_date": LEASE_END,
    }})
    only_completion = collect_arenda_ts_data({"vehicle": {
        "planned_completion_date": COMPLETION_DATE,
    }})

    assert cd.contract["lease_start_date"] == LEASE_START
    assert cd.contract["lease_end_date"] == LEASE_END
    assert cd.contract["planned_completion_date"] == COMPLETION_DATE

    # Только своя дата: остальных ключей в contract нет. Пустые массивы точек
    # сборщик кладёт всегда («точек нет» и «раздел не собирался» — разное).
    assert only_start.contract["lease_start_date"] == LEASE_START
    assert only_end.contract["lease_end_date"] == LEASE_END
    assert only_completion.contract["planned_completion_date"] == COMPLETION_DATE
    assert set(only_start.contract) == {"lease_start_date", "loadings", "unloadings"}
    assert set(only_end.contract) == {"lease_end_date", "loadings", "unloadings"}
    assert set(only_completion.contract) == {
        "planned_completion_date", "loadings", "unloadings",
    }


def test_equal_and_different_dates_are_kept_as_entered():
    """Совпали даты или нет — в contract уходит то, что ввели."""
    same = collect_arenda_ts_data({"vehicle": {
        "lease_end_date": COMPLETION_DATE,
        "planned_completion_date": COMPLETION_DATE,
    }})
    different = collect_arenda_ts_data({"vehicle": {
        "lease_end_date": LEASE_END,
        "planned_completion_date": COMPLETION_DATE,
    }})

    assert same.contract["lease_end_date"] == COMPLETION_DATE
    assert same.contract["planned_completion_date"] == COMPLETION_DATE
    assert different.contract["lease_end_date"] == LEASE_END
    assert different.contract["planned_completion_date"] == COMPLETION_DATE


def test_empty_completion_date_is_not_invented():
    """Дату рейса не выводим из срока аренды: нет ввода — нет ключа."""
    cd = collect_arenda_ts_data({"vehicle": {
        "lease_start_date": LEASE_START,
        "lease_end_date": LEASE_END,
    }})

    assert "planned_completion_date" not in cd.contract


def test_dates_reach_generated_document_in_their_clauses(
        full_tabs, templates_dir, work_dir):
    """
    Три даты стоят в договоре каждая на своём месте и не путаются.

    П. 2.5 — плановый период аренды (начало и конец), п. 3.3.2 — планируемая
    дата завершения рейса. В образце ТЛ-574 это разные даты (28.09.2026 и
    26.09.2026), поэтому проверяются обе формулировки.

    Поле «Планируемая дата завершения рейса» с шага FIX-1-T печатает бланк
    САМ (плейсхолдер {{planned_completion_date}}), а не дата второй точки
    выгрузки: дата точки в п. 3.3.2 документа больше не участвует.
    """
    from docx import Document

    from core.contracts.arenda_ts.generator import ArendaTsGenerator

    data = _tabs_of_variant()
    data["vehicle"]["lease_start_date"] = LEASE_START
    data["vehicle"]["lease_end_date"] = LEASE_END
    data["vehicle"]["planned_completion_date"] = COMPLETION_DATE
    # Даты точек выгрузки — свои и обе отличны от даты завершения рейса:
    # если бы бланк печатал дату точки, в договоре стояла бы не та дата.
    data["route"]["unloadings"] = [
        {"name": UNLOADING_NAME_1, "address": UNLOADING_ADDRESS_1,
         "date": UNLOADING_DATE_1},
        {"name": UNLOADING_NAME_2, "address": UNLOADING_ADDRESS_2,
         "date": UNLOADING_DATE_2},
    ]

    generator = ArendaTsGenerator(templates_dir=str(templates_dir))
    path = generator.generate(
        collect_arenda_ts_data(data), output_dir=str(work_dir)
    )

    doc = Document(path)
    text = _flatten("\n".join(
        [p.text for p in doc.paragraphs]
        + [cell.text for t in doc.tables for r in t.rows for cell in r.cells]
    ))
    start_text = _format_date(LEASE_START)
    end_text = _format_date(LEASE_END)
    completion_text = _format_date(COMPLETION_DATE)
    second_point_text = _format_date(UNLOADING_DATE_2)

    assert start_text != end_text != completion_text
    assert completion_text != second_point_text
    assert f"с {start_text} г. по {end_text} г. включительно" in text
    assert f"Плановая дата завершения: {completion_text} г." in text
    # Дата второй точки выгрузки в п. 3.3.2 не печатается — там дата рейса.
    assert (f"{UNLOADING_ADDRESS_2}. Плановая дата завершения: "
            f"{second_point_text} г.") not in text


def test_completion_date_is_not_printed_as_lease_end(
        templates_dir, work_dir):
    """Дата завершения рейса не подменяет собой окончание аренды."""
    from docx import Document

    from core.contracts.arenda_ts.generator import ArendaTsGenerator

    data = _tabs_of_variant()
    data["vehicle"]["lease_end_date"] = LEASE_END
    data["vehicle"]["planned_completion_date"] = COMPLETION_DATE

    generator = ArendaTsGenerator(templates_dir=str(templates_dir))
    path = generator.generate(
        collect_arenda_ts_data(data), output_dir=str(work_dir / "clause_2_5")
    )

    doc = Document(path)
    lines = [p.text for p in doc.paragraphs]
    period = next(line for line in lines if line.strip().startswith("2.5."))

    assert _format_date(LEASE_END) in period
    assert _format_date(COMPLETION_DATE) not in period


def test_vehicle_tractor_and_trailer_are_mapped(full_tabs):
    cd = collect_arenda_ts_data(full_tabs)

    assert cd.tractor == {
        "brand_model": TRACTOR_BRAND,
        "plate_number": TRACTOR_PLATE,
        "vehicle_type": TRACTOR_TYPE,
    }
    assert cd.trailer == {
        "brand_model": TRAILER_BRAND,
        "plate_number": TRAILER_PLATE,
    }


def test_empty_vehicle_tab_gives_empty_vehicle():
    cd = collect_arenda_ts_data({"vehicle": {"tractor_brand": "  "}})

    assert cd.tractor == {}
    assert cd.trailer == {}
    assert "number" not in cd.contract
    assert "lease_start_date" not in cd.contract


def test_recognition_vehicle_blocks_are_read():
    """Вкладка «ТС» может отдать блоки tractor / trailer распознавания как есть."""
    cd = collect_arenda_ts_data({"vehicle": {
        "contract": {"number": CONTRACT_NUMBER, "date": CONTRACT_DATE},
        "lease_start_date": LEASE_START,
        "lease_end_date": LEASE_END,
        "tractor": {"brand_model": TRACTOR_BRAND, "plate_number": TRACTOR_PLATE,
                    "vehicle_type": TRACTOR_TYPE},
        "trailer": {"brand_model": TRAILER_BRAND, "plate_number": TRAILER_PLATE},
    }})

    assert cd.contract["number"] == CONTRACT_NUMBER
    assert cd.contract["date"] == CONTRACT_DATE
    assert cd.contract["lease_start_date"] == LEASE_START
    assert cd.tractor["vehicle_type"] == TRACTOR_TYPE
    assert cd.trailer["plate_number"] == TRAILER_PLATE


# ─────────────────────────────────────────────────────────────
# Груз: перевозимые машины
# ─────────────────────────────────────────────────────────────

def test_cargo_vehicles_are_collected():
    cd = collect_arenda_ts_data({"cargo": {"vehicles": [
        {"brand_model": "МОДЕЛЬ 1", "vin": VIN_1,
         "loading_point": LOADING_ADDRESS_1,
         "unloading_point": UNLOADING_ADDRESS_1},
    ]}})

    assert cd.vehicles == [{
        "brand_model": "МОДЕЛЬ 1",
        "vin": VIN_1,
        "loading_point": LOADING_ADDRESS_1,
        "unloading_point": UNLOADING_ADDRESS_1,
    }]


def test_empty_cargo_rows_are_dropped():
    cd = collect_arenda_ts_data({"cargo": {"vehicles": [
        {"brand_model": "", "vin": ""},
        {"brand_model": "  ", "vin": "  ", "loading_point": LOADING_ADDRESS_1},
        {"brand_model": "МОДЕЛЬ 1", "vin": VIN_1},
    ]}})

    assert len(cd.vehicles) == 1
    assert cd.vehicles[0]["vin"] == VIN_1


def test_vehicle_with_only_brand_or_only_vin_is_kept():
    """Заполнена хотя бы одна колонка — строку не выбрасываем (скажет валидатор)."""
    cd = collect_arenda_ts_data({"cargo": {"vehicles": [
        {"brand_model": "МОДЕЛЬ 1", "vin": ""},
        {"brand_model": "", "vin": VIN_1},
    ]}})

    assert len(cd.vehicles) == 2
    assert cd.vehicles[0]["brand_model"] == "МОДЕЛЬ 1"
    assert cd.vehicles[1]["vin"] == VIN_1


def test_cargo_is_limited_to_max_cars():
    cd = collect_arenda_ts_data({"cargo": _cargo_tab(MAX_CARS + 3)})

    assert MAX_CARS == 12
    assert len(cd.vehicles) == MAX_CARS
    assert cd.vehicles[-1]["brand_model"] == f"МОДЕЛЬ {MAX_CARS}"


def test_cargo_not_a_list_is_ignored():
    cd = collect_arenda_ts_data({"cargo": {"vehicles": "мусор"}})

    assert cd.vehicles == []


def test_cargo_count_is_not_invented():
    """
    cargo_count в сборке не выдумывается: его считает генератор по списку.

    Выдуманное значение спорило бы с таблицей 3.1 и давало замечание валидатора.
    """
    cd = collect_arenda_ts_data({"cargo": _cargo_tab()})

    assert "cargo_count" not in cd.contract


def test_vehicle_points_reach_the_car_table(full_tabs, templates_dir):
    """Точки погрузки и выгрузки машин доходят до карты замен таблицы п. 3.1."""
    from core.contracts.arenda_ts.generator import ArendaTsGenerator

    generator = ArendaTsGenerator(templates_dir=str(templates_dir))
    replacements = generator.build_replacements(collect_arenda_ts_data(full_tabs))

    assert replacements["car_1_brand"] == "МОДЕЛЬ 1"
    assert replacements["car_1_vin"] == VIN_1
    assert replacements["car_1_loading_point"] == LOADING_ADDRESS_1
    assert replacements["car_1_unloading_point"] == UNLOADING_ADDRESS_1
    assert replacements["car_2_brand"] == "МОДЕЛЬ 2"
    assert replacements["car_3_brand"] == ""


# ─────────────────────────────────────────────────────────────
# Полный сценарий: все семь вкладок вместе
# ─────────────────────────────────────────────────────────────

def test_full_scenario_collects_everything(full_tabs):
    cd = collect_arenda_ts_data(full_tabs)

    # Договор и срок аренды.
    assert cd.contract["number"] == CONTRACT_NUMBER
    assert cd.contract["date"] == CONTRACT_DATE
    assert cd.contract["lease_start_date"] == LEASE_START
    assert cd.contract["lease_end_date"] == LEASE_END
    assert cd.contract["route"] == ROUTE

    # Стороны.
    assert cd.contract["carrier_type"] == CARRIER_TYPE_OOO
    assert cd.customer["full_name"] == LESSEE_OOO_NAME
    assert cd.carrier["full_name"] == LESSOR_NAME

    # Автопоезд, машины, маршрут и экипаж.
    assert cd.tractor["plate_number"] == TRACTOR_PLATE
    assert cd.trailer["plate_number"] == TRAILER_PLATE
    assert len(cd.vehicles) == 2
    assert len(cd.contract["loadings"]) == 2
    assert len(cd.contract["unloadings"]) == 1
    assert cd.driver["full_name"] == DRIVER_NAME

    # Стоимость.
    assert cd.contract["price_without_vat"] == OOO_BASE_SUM
    assert cd.contract["vat_rate_num"] == VAT_RATE_NUM
    assert cd.contract["special_conditions"] == SPECIAL_CONDITIONS


def test_full_scenario_summary_is_safe_for_logs(full_tabs):
    """Сводка для логов не падает и не содержит самих данных."""
    cd = collect_arenda_ts_data(full_tabs)
    summary = cd.summary()

    assert "vehicles=2" in summary
    assert "loadings=2" in summary
    assert LESSEE_OOO_NAME not in summary
    assert DRIVER_NAME not in summary


def test_logs_do_not_contain_personal_data(full_tabs, caplog):
    """
    В INFO сборки нет персональных данных: только вид Арендатора и количества.

    Наименования сторон, ФИО, адреса, VIN, госномера и суммы в лог не попадают.
    """
    with caplog.at_level(logging.INFO, logger="ui.windows.arenda_ts.data"):
        cd = collect_arenda_ts_data(full_tabs)
        assert cd.contract["number"] == CONTRACT_NUMBER

    messages = "\n".join(
        record.getMessage() for record in caplog.records
        if record.name == "ui.windows.arenda_ts.data"
    )
    assert messages, "сборка ничего не записала в свой лог"

    for fragment in (
        "Арендатор-Тест",   # наименование Арендатора
        "Арендодатель-Тест",  # наименование Арендодателя
        "Иванов",           # ФИО водителя
        "Склад Север",      # название точки маршрута
        "Складская",        # адрес точки
        VIN_1,              # VIN
        "МОДЕЛЬ",           # марка машины
        "Тягач-Модель",     # тягач
        "Б002ББ02",         # госномер прицепа
        "926830",           # паспорт
    ):
        assert fragment not in messages, f"в логе есть «{fragment}»"

    # Вид Арендатора и количества — можно.
    assert "вариант=ООО" in messages
    assert "машин=2" in messages
    assert "точек погрузки=2" in messages
    assert "точек выгрузки=1" in messages


@pytest.mark.parametrize("variant", VARIANTS)
def test_full_scenario_passes_arenda_ts_validator(variant):
    """
    Собранных данных достаточно, чтобы валидатор типа не нашёл ни ошибок,
    ни замечаний — во всех трёх вариантах бланка (ООО / ИП с НДС / ИП без НДС).
    """
    from core.contracts.arenda_ts.validator import ArendaTsValidator

    report = ArendaTsValidator().check(
        collect_arenda_ts_data(_tabs_of_variant(variant))
    )

    assert report.errors == [], f"{variant}: {report.errors}"
    assert report.warnings == [], f"{variant}: {report.warnings}"


def test_full_scenario_generator_replacements(full_tabs, templates_dir):
    """Генератор читает из собранных данных ровно то, что ждёт."""
    from core.contracts.arenda_ts.generator import ArendaTsGenerator

    generator = ArendaTsGenerator(templates_dir=str(templates_dir))
    replacements = generator.build_replacements(collect_arenda_ts_data(full_tabs))

    assert replacements["contract_number"] == CONTRACT_NUMBER
    assert replacements["contract_date"] == "19.09.2026"
    assert replacements["lease_start_date"] == "21.09.2026"
    assert replacements["lease_end_date"] == "27.09.2026"
    assert replacements["lessee_full_name"] == LESSEE_OOO_NAME
    assert replacements["lessee_inn"] == LESSEE_OOO_INN
    assert replacements["lessee_kpp"] == LESSEE_OOO_KPP
    assert replacements["lessee_ogrn_label"] == OGRN_LABEL
    assert replacements["lessee_basis"] == BASIS_OOO
    assert replacements["lessor_full_name"] == LESSOR_NAME
    assert replacements["lessor_basis"] == BASIS_OOO
    assert replacements["tractor_brand"] == TRACTOR_BRAND
    assert replacements["tractor_plate"] == TRACTOR_PLATE
    assert replacements["tractor_type"] == TRACTOR_TYPE
    assert replacements["trailer_plate"] == TRAILER_PLATE
    assert replacements["route"] == ROUTE
    assert replacements["cargo_count"] == "2"

    # Стоимость: база и ставка из формы, НДС и итог генератор считает сам.
    assert replacements["sum_wo_vat"] == "221\u00a0099,18"
    assert replacements["sum_vat"] == "48\u00a0641,82"
    assert replacements["sum_total"] == "269\u00a0741,00"
    assert replacements["vat_rate"] == VAT_RATE_TEXT


def test_full_scenario_renders_docx(full_tabs, templates_dir, work_dir):
    """
    Собранные данные доходят до ГОТОВОГО документа.

    Это проверка стыка «вкладки → сборщик → генератор → бланк»: время подачи
    ТС, разобранные паспорт и удостоверение, реквизиты сторон и суммы видны
    в тексте договора.
    """
    from docx import Document

    from core.contracts.arenda_ts.generator import ArendaTsGenerator

    output_dir = work_dir / "ui_arenda_ts_data_full"
    output_dir.mkdir(parents=True, exist_ok=True)

    generator = ArendaTsGenerator(templates_dir=str(templates_dir))
    path = None
    try:
        path = Path(generator.generate(
            collect_arenda_ts_data(full_tabs), output_dir=str(output_dir)
        ))
        assert path.exists(), f"файл не создан: {path}"

        doc = Document(str(path))
        texts = [p.text for p in doc.paragraphs]
        parts = list(texts)

        def table_text(table):
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
                    for nested in cell.tables:
                        table_text(nested)

        for table in doc.tables:
            table_text(table)
        text = "\n".join(parts)

        assert PLACEHOLDER_RE.search(text) is None, "в документе остался плейсхолдер"

        # Точка погрузки с временем подачи ТС — маппинг 2.
        assert ("3.2.1. Точка погрузки № 1 — г. Москва, ул. Складская, д. 1. "
                "Плановая дата и время подачи ТС: 21.09.2026 г., "
                "с 08:00 до 18:00.") in texts
        assert f"3.4. Согласованный маршрут: {ROUTE}." in texts

        # Срок аренды, автопоезд и реквизиты сторон.
        assert any(
            t.startswith("2.5. Плановый период аренды: с 21.09.2026 г. "
                         "по 27.09.2026 г. включительно")
            for t in texts
        )
        assert (f"– тягач: {TRACTOR_BRAND}, государственный регистрационный "
                f"знак {TRACTOR_PLATE}, тип ТС — {TRACTOR_TYPE};") in texts
        assert (f"– прицеп/полуприцеп: {TRAILER_BRAND}, государственный "
                f"регистрационный знак {TRAILER_PLATE}.") in texts
        assert LESSEE_OOO_NAME in text
        assert LESSOR_NAME in text

        # Экипаж: разобранные паспорт и удостоверение — маппинг 4.
        assert DRIVER_NAME in text
        assert f"Паспорт: {DRIVER_PASSPORT}" in text
        assert f"Водительское удостоверение: {DRIVER_LICENSE}" in text
        assert f"Адрес регистрации: {DRIVER_ADDRESS}" in text

        # Арендная плата: 221 099,18 + 22% (48 641,82) = 269 741,00.
        assert "221\u00a0099,18" in text
        assert "48\u00a0641,82" in text
        assert "269\u00a0741,00" in text
    finally:
        if path is not None:
            path.unlink(missing_ok=True)


def test_recognition_payload_shape_is_accepted():
    """
    Вкладки могут отдать ответ распознавания как есть — сборка его понимает.

    Раскладка ответа (core/prompts/arenda_ts.py): блоки lessee / lessor /
    tractor / trailer / driver, суммы в contract, срок аренды и route в корне.
    """
    tabs = {
        "lessee": {"lessee": {
            "entity_type": "ИП",
            "full_name": LESSEE_IP_NAME,
            "short_name": LESSEE_IP_SHORT,
            "inn": LESSEE_IP_INN,
            "ogrn": LESSEE_IP_OGRNIP,
            "legal_address": LESSEE_IP_ADDRESS,
            "director_name": LESSEE_IP_DIRECTOR,
            "director_position": "Индивидуальный предприниматель",
        }},
        "lessor": {"lessor": {
            "full_name": LESSOR_NAME,
            "short_name": LESSOR_SHORT,
            "inn": LESSOR_INN,
            "ogrn": LESSOR_OGRN,
            "legal_address": LESSOR_ADDRESS,
            "director_name": LESSOR_DIRECTOR,
        }},
        "vehicle": {
            "tractor": {"brand_model": TRACTOR_BRAND, "plate_number": TRACTOR_PLATE,
                        "vehicle_type": TRACTOR_TYPE},
            "trailer": {"brand_model": TRAILER_BRAND, "plate_number": TRAILER_PLATE},
            "lease_start_date": LEASE_START,
            "lease_end_date": LEASE_END,
            "contract": {"number": CONTRACT_NUMBER, "date": CONTRACT_DATE},
        },
        "route": {"route": ROUTE, "loadings": [
            {"address": LOADING_ADDRESS_1, "date": LOADING_DATE_1,
             "time_from": LOADING_TIME_FROM, "time_to": LOADING_TIME_TO},
        ], "unloadings": [
            {"address": UNLOADING_ADDRESS_1, "date": UNLOADING_DATE_1},
        ]},
        "cargo": {"vehicles": [_vehicle(1)]},
        "crew": {"driver": {
            "full_name": DRIVER_NAME,
            "birth_date": DRIVER_BIRTH,
            "passport": DRIVER_PASSPORT,
            "passport_issuer": DRIVER_PASSPORT_ISSUER,
            "passport_issue_date": DRIVER_PASSPORT_DATE,
            "license": DRIVER_LICENSE,
            "license_issue_date": DRIVER_LICENSE_DATE,
            "address": DRIVER_ADDRESS,
            "phone": DRIVER_PHONE,
        }},
        "price": {"contract": {
            "sum_total": IP_SUM, "vat_rate": IP_VAT_RATE_TEXT,
        }},
    }

    cd = collect_arenda_ts_data(tabs)

    assert cd.contract["carrier_type"] == CARRIER_TYPE_IP_WITHOUT_VAT
    assert cd.contract["lessee"]["full_name"] == LESSEE_IP_NAME
    assert cd.contract["lessor"]["full_name"] == LESSOR_NAME
    assert cd.contract["number"] == CONTRACT_NUMBER
    assert cd.contract["lease_start_date"] == LEASE_START
    assert cd.tractor["vehicle_type"] == TRACTOR_TYPE
    assert cd.vehicles[0]["vin"] == VIN_1
    assert cd.driver["passport_number"] == DRIVER_PASSPORT_NUMBER
    assert cd.contract["loadings"][0]["time_from"] == LOADING_TIME_FROM
    assert cd.contract["price_without_vat"] == IP_SUM


# ─────────────────────────────────────────────────────────────
# Точка входа build()
# ─────────────────────────────────────────────────────────────

def test_build_uses_same_layout_as_collect(full_tabs):
    """build() — та же сборка, только вкладки переданы по именам."""
    cd = build(
        full_tabs["lessee"],
        full_tabs["lessor"],
        full_tabs["vehicle"],
        full_tabs["route"],
        full_tabs["cargo"],
        full_tabs["crew"],
        full_tabs["price"],
    )

    assert cd.contract["number"] == CONTRACT_NUMBER
    assert cd.contract["carrier_type"] == CARRIER_TYPE_OOO
    assert cd.contract["lessee"]["full_name"] == LESSEE_OOO_NAME
    assert cd.customer["full_name"] == LESSEE_OOO_NAME
    assert cd.carrier["full_name"] == LESSOR_NAME
    assert cd.driver["full_name"] == DRIVER_NAME
    assert [p["address"] for p in cd.contract["loadings"]] == [
        LOADING_ADDRESS_1, LOADING_ADDRESS_2,
    ]
    assert len(cd.vehicles) == 2


def test_build_accepts_none_tabs():
    """Вкладок может ещё не быть — сборка не падает."""
    cd = build(None, None, None, None, None, None, None)

    assert isinstance(cd, ContractData)
    assert cd.contract == {"loadings": [], "unloadings": []}
    assert cd.driver == {}
    assert cd.vehicles == []


def test_two_entry_points_give_the_same_result(full_tabs):
    """Обе точки входа используют одну раскладку полей."""
    from_dict = collect_arenda_ts_data(full_tabs)
    from_names = build(*[full_tabs[key] for key in (
        "lessee", "lessor", "vehicle", "route", "cargo", "crew", "price",
    )])

    assert from_dict.contract == from_names.contract
    assert from_dict.driver == from_names.driver
    assert from_dict.customer == from_names.customer
    assert from_dict.carrier == from_names.carrier
