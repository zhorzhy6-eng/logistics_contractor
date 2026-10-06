#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты шаблонов заявки «Логистикс Рус» (ЭТАП 3.1.C.A.1).

Проверяют, что оба шаблона —

    templates/shablon_logistiks_rus_ooo.docx  (ООО «ТЕХНОЛОГИСТИКА»)
    templates/shablon_logistiks_rus_ip.docx   (ИП Хейгетян Е.В.)

— именно ПУСТЫЕ БЛАНКИ с плейсхолдерами docxtpl, а не копии образцов
с данными:

  * файл существует и открывается python-docx;
  * все ключевые плейсхолдеры на месте и ни один не разорван между runs;
  * раздел 1 «Погрузка»: ОДИН грузоотправитель ({{shipper_name}}) и 10
    нумерованных адресов погрузки ({{shipper_N_address}}); раздел 2
    «Выгрузка»: по-прежнему 10 пар «грузополучатель + адрес»;
  * таблица автомобилей: шапка «№ / Марка, модель / VIN-номер» и ровно
    12 строк с плейсхолдерами марки и VIN;
  * в ООО-шаблоне три суммы (sum_wo_vat / vat_rate+sum_vat / sum_total)
    и все три суммы прописью (sum_wo_vat_words / sum_vat_words /
    sum_total_words), в ИП-шаблоне — одна сумма (sum_total, «Без НДС»)
    и только sum_total_words;
  * в шаблоне нет данных образцов (марки, VIN, ФИО, госномера, адреса,
    суммы, номера заявок) и нет незакрытых «{{»;
  * вариант экспедитора и номер генерального договора — свои у каждого
    шаблона (ТЭ0909/01 у ООО, ТЭ0909/02 у ИП);
  * оформление повторяет образец: A4, поля 2,54 см, Arial 12 pt;
  * образцы-источники не изменены (сверка по SHA256);
  * docxtpl рендерит шаблон без остатка плейсхолдеров;
  * сборщик tools/make_logistiks_rus_template.py даёт ровно тот же текст.

Образцы (Заявка_600_ООО …, Заявка_359_ИП_ …) — только источник структуры,
формулировок и оформления: тест test_sample_documents_untouched намеренно
падает, если их отредактировали или пересобрали.
"""

import gc
import hashlib
import re

import pytest
from docx import Document

#: Варианты экспедитора и их файлы — те же ключи, что у генератора.
TEMPLATE_NAMES = {
    "ООО": "shablon_logistiks_rus_ooo.docx",
    "ИП": "shablon_logistiks_rus_ip.docx",
}

CARRIER_TYPES = tuple(TEMPLATE_NAMES)

#: SHA256 собранных шаблонов (ШАГ FIX-2.2: грузоотправитель ОДИН,
#: адреса погрузки нумеруются, {{shipper_N_name}} из бланка убран).
#: Если шаблон пересобрали осознанно (например, поменяли формулировку),
#: значения нужно обновить — тест ловит ручную правку .docx мимо сборщика.
TEMPLATE_SHA256 = {
    "ООО": "d78fa27185991af236a9f6697e8a4e5fa88859ce9b62e70ca248c0bc52d3a826",
    "ИП": "fb05c6bbe308668238ce3b254fbb7678842ac2efde020f5ddae66febf3c69b3c",
}

#: Образцы-источники и их SHA256 на момент сборки шаблонов.
SAMPLES = {
    "ООО": (
        "Приложение № 1 Заявка_600_ООО «ДжейСиСиТиЭс Интернейшнл "
        "Логистикс Рус».docx",
        "55c98140b6ae55ee2885c4aa7284354fad3120f7b187d0738253744b4de2117e",
    ),
    "ИП": (
        "Приложение № 1 Заявка_359_ИП_ «ДжейСиСиТиЭс Интернейшнл "
        "Логистикс Рус».docx",
        "3b8f7406b9d0978edf940243df74a6caf3fa017914f5b829cb3cad69c06254b5",
    ),
}

#: Заголовки таблицы автомобилей — ровно как в образце (запятая, не слэш).
CARGO_HEADERS = ("№", "Марка, модель", "VIN-номер")

CARGO_ROWS = 12
MAX_POINTS = 10

#: Плейсхолдеры, общие для обоих вариантов.
COMMON_PLACEHOLDERS = (
    "contract_number",
    "contract_date_day",
    "contract_date_month",
    "contract_date_year",
    "customer_name",
    "loading_date",
    "loading_time_from",
    "loading_time_to",
    "unloading_date",
    "unloading_time_from",
    "unloading_time_to",
    "cargo_count",
    "tractor_brand",
    "tractor_plate",
    "trailer_brand",
    "trailer_plate",
    "driver_name",
    "sum_total",
    "sum_total_words",
    "special_conditions",
)

#: Плейсхолдеры ООО-варианта: расчёт НДС по ставке (в ИП-шаблоне их нет).
OOO_PLACEHOLDERS = (
    "sum_wo_vat", "sum_wo_vat_words",
    "vat_rate",
    "sum_vat", "sum_vat_words",
)

#: Плейсхолдеры сумм прописью, которых в ИП-шаблоне быть не должно.
OOO_WORDS_ONLY = ("sum_wo_vat_words", "sum_vat_words")

#: Номер генерального договора: свой у каждого варианта.
CONTRACT_NUMBERS = {
    "ООО": ("№ТЭ0909/01", "№ ТЭ 0909/01"),
    "ИП": ("№ТЭ0909/02", "№ ТЭ 0909/02"),
}

#: Данные образцов, которых в бланках быть не должно.
FORBIDDEN_FRAGMENTS = (
    # марки и модели перевозимых машин
    "JETOUR", "DASHING", "X70PLUS", "Престиж", "Комфорт",
    # VIN из образцов
    "EC3DLUFD", "EC3DCUFD", "EC37CUSM",
    # тягачи и прицепы
    "DAF XF", "LUXUDA", "Dongfeng", "BLACKSMITH",
    "М342СА761", "МУ582323", "Т5630Х761", "СТО98761",
    # ФИО водителей
    "Соин", "Скрынник",
    # грузоотправители и грузополучатели
    "ВОТУР", "ЮГ-АВТО", "ТЕМП АВТО", "Р-МОТОРС", "НОВОКАР", "АВТОРИТЭЙЛ",
    # адреса
    "Клин", "Белый Раст", "Тахтамукай", "Краснодар", "Новороссийск",
    "Мысхакское",
    # суммы и даты образцов
    "221 099", "48 641", "269 741", "135 833",
    "24.09.2026", "25.09.2026", "26.09.2026", "01.10.2026",
    # номера заявок и время из образцов
    "600", "359", "08:00", "20:00",
)

PLACEHOLDER_RE = re.compile(r"\{\{[^{}]*\}\}")


@pytest.fixture(autouse=True)
def _release_documents():
    """Word-файлы на Windows освобождаем до удаления временных копий."""
    yield
    gc.collect()


@pytest.fixture(params=CARRIER_TYPES)
def carrier_type(request) -> str:
    """Вариант экспедитора: «ООО» или «ИП» — тесты идут по обоим."""
    return request.param


@pytest.fixture
def template_path(templates_dir, carrier_type):
    path = templates_dir / TEMPLATE_NAMES[carrier_type]
    assert path.exists(), f"нет шаблона Логистикс Рус: {path}"
    return path


@pytest.fixture(scope="module")
def templates(templates_dir):
    """Оба шаблона: {«ООО»: Document, «ИП»: Document}."""
    return {
        key: Document(str(templates_dir / name))
        for key, name in TEMPLATE_NAMES.items()
    }


@pytest.fixture
def template_doc(templates, carrier_type):
    return templates[carrier_type]


# ─────────────────────────────────────────────────────────────
# Вспомогательные функции
# ─────────────────────────────────────────────────────────────

def _expected_placeholders(carrier_type: str) -> set:
    """Полный набор имён плейсхолдеров шаблона этого варианта."""
    names = set(COMMON_PLACEHOLDERS)
    # Раздел 1: ОДИН грузоотправитель + до 10 нумерованных адресов погрузки.
    names.add("shipper_name")
    names |= {
        f"shipper_{number}_address"
        for number in range(1, MAX_POINTS + 1)
    }
    names |= {
        f"consignee_{number}_{field}"
        for number in range(1, MAX_POINTS + 1)
        for field in ("name", "address")
    }
    names |= {
        f"car_{number}_{field}"
        for number in range(1, CARGO_ROWS + 1)
        for field in ("brand", "vin")
    }
    if carrier_type == "ООО":
        names |= set(OOO_PLACEHOLDERS)
    return names


def _document_text(doc) -> str:
    """Весь текст документа: абзацы, таблицы (включая вложенные), колонтитулы."""
    parts = [p.text for p in doc.paragraphs]

    def table_text(table):
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
                for nested in cell.tables:
                    table_text(nested)

    for table in doc.tables:
        table_text(table)

    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)

    return "\n".join(parts)


def _all_paragraphs(doc):
    """Все абзацы документа, включая абзацы таблиц и колонтитулов."""
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _cargo_table(doc):
    """Таблица автомобилей — по заголовкам колонок, а не по индексу."""
    for table in doc.tables:
        if not table.rows:
            continue
        headers = tuple(cell.text.strip() for cell in table.rows[0].cells)
        if headers == CARGO_HEADERS:
            return table
    return None


def _body_texts(doc):
    """Тексты абзацев верхнего уровня (без таблиц), без крайних пробелов."""
    return [p.text.strip() for p in doc.paragraphs]


# ─────────────────────────────────────────────────────────────
# Файл и плейсхолдеры
# ─────────────────────────────────────────────────────────────

def test_template_file_exists_and_opens(template_path, template_doc):
    assert template_path.stat().st_size > 0
    assert template_doc.tables, "в бланке нет ни одной таблицы"


def test_common_placeholders_present(template_doc, carrier_type):
    text = _document_text(template_doc)
    missing = [name for name in COMMON_PLACEHOLDERS
               if "{{" + name + "}}" not in text]
    assert not missing, f"в шаблоне {carrier_type} нет плейсхолдеров: {missing}"


def test_point_placeholders_present(template_doc, carrier_type):
    """Адреса погрузки (10 нумерованных) и блоки выгрузки (10 пар) на месте."""
    text = _document_text(template_doc)
    missing = [
        f"{{{{shipper_{number}_address}}}}"
        for number in range(1, MAX_POINTS + 1)
        if f"{{{{shipper_{number}_address}}}}" not in text
    ]
    missing += [
        f"{{{{{prefix}_{number}_{field}}}}}"
        for prefix in ("consignee",)
        for number in range(1, MAX_POINTS + 1)
        for field in ("name", "address")
        if f"{{{{{prefix}_{number}_{field}}}}}" not in text
    ]
    assert "{{shipper_name}}" in text, "нет плейсхолдера грузоотправителя"
    assert not missing, f"нет плейсхолдеров точек маршрута: {missing}"


def test_all_car_placeholders_present(template_doc, carrier_type):
    """Все 12 машин: и марка, и VIN."""
    text = _document_text(template_doc)
    missing = [
        f"{{{{car_{number}_{field}}}}}"
        for number in range(1, CARGO_ROWS + 1)
        for field in ("brand", "vin")
        if f"{{{{car_{number}_{field}}}}}" not in text
    ]
    assert not missing, f"нет плейсхолдеров машин: {missing}"


def test_placeholders_are_not_split_across_runs(template_doc, carrier_type):
    """
    Плейсхолдер обязан лежать целиком в одном run.

    Если Word разорвал «{{contract_number}}» на несколько runs, docxtpl
    (Jinja) его не найдёт и в документе останется сырой текст.
    """
    joined_placeholders = set()
    for paragraph in _all_paragraphs(template_doc):
        joined = "".join(run.text for run in paragraph.runs)
        for placeholder in PLACEHOLDER_RE.findall(joined):
            if not any(placeholder in run.text for run in paragraph.runs):
                joined_placeholders.add(placeholder)

    assert not joined_placeholders, (
        f"плейсхолдеры разорваны между runs: {sorted(joined_placeholders)}"
    )


def test_no_unclosed_placeholders(template_doc, carrier_type):
    text = _document_text(template_doc)
    opening = text.count("{{")
    closing = text.count("}}")

    assert opening == closing, (
        f"несбалансированные скобки: «{{{{» = {opening}, «}}}}» = {closing}"
    )
    assert len(PLACEHOLDER_RE.findall(text)) == opening, (
        "есть «{{» вне корректно закрытого плейсхолдера"
    )


def test_placeholder_names_are_known(template_doc, carrier_type):
    """В бланке нет опечаток в именах: каждое имя — из ожидаемого набора."""
    expected = _expected_placeholders(carrier_type)
    found = {
        match.strip("{} ")
        for match in PLACEHOLDER_RE.findall(_document_text(template_doc))
    }
    assert found == expected, (
        f"лишние/недостающие плейсхолдеры ({carrier_type}): {found ^ expected}"
    )


def test_ooo_template_has_vat_placeholders(templates):
    """ООО-вариант: три суммы — без НДС, НДС по ставке, итого, каждая с прописью."""
    texts = _body_texts(templates["ООО"])
    assert "Стоимость услуг: {{sum_wo_vat}} руб. ({{sum_wo_vat_words}})" in texts
    assert "НДС {{vat_rate}}: {{sum_vat}} руб. ({{sum_vat_words}})" in texts
    assert "Итого: {{sum_total}} руб. ({{sum_total_words}})" in texts


def test_ip_template_has_single_sum_without_vat(templates):
    """ИП-вариант: одна сумма «Без НДС» с прописью, плейсхолдеров НДС нет."""
    doc = templates["ИП"]
    texts = _body_texts(doc)
    assert ("Стоимость услуг: {{sum_total}} руб. ({{sum_total_words}}) "
            "Без НДС") in texts

    text = _document_text(doc)
    for name in OOO_PLACEHOLDERS:
        assert "{{" + name + "}}" not in text, (
            f"в ИП-шаблоне лишний плейсхолдер {name!r}"
        )
    assert text.count("{{sum_total}}") == 1


def test_words_placeholders_present(templates, carrier_type):
    """Суммы прописью: у ООО — три, у ИП — одна (ИП без НДС)."""
    text = _document_text(templates[carrier_type])

    assert "{{sum_total_words}}" in text, "нет суммы итого прописью"

    if carrier_type == "ООО":
        assert "{{sum_wo_vat_words}}" in text
        assert "{{sum_vat_words}}" in text
    else:
        for name in OOO_WORDS_ONLY:
            assert "{{" + name + "}}" not in text, (
                f"в ИП-шаблоне лишний плейсхолдер {name!r}"
            )


# ─────────────────────────────────────────────────────────────
# Таблица автомобилей
# ─────────────────────────────────────────────────────────────

def test_cargo_table_has_header_and_12_data_rows(template_doc, carrier_type):
    table = _cargo_table(template_doc)
    assert table is not None, "таблица автомобилей не найдена по заголовкам"
    assert len(table.rows) == 1 + CARGO_ROWS, (
        f"строк в таблице {len(table.rows)}, ожидалось {1 + CARGO_ROWS}"
    )
    assert len(table.columns) == 3


def test_cargo_rows_are_numbered_and_placeholder_driven(template_doc, carrier_type):
    table = _cargo_table(template_doc)
    for number in range(1, CARGO_ROWS + 1):
        cells = [cell.text.strip() for cell in table.rows[number].cells]
        assert cells == [
            str(number),
            f"{{{{car_{number}_brand}}}}",
            f"{{{{car_{number}_vin}}}}",
        ], f"строка {number} таблицы автомобилей: {cells!r}"


def test_cargo_table_has_no_sample_brands_or_vins(template_doc, carrier_type):
    table = _cargo_table(template_doc)
    text = "\n".join(cell.text for row in table.rows for cell in row.cells)
    for fragment in ("JETOUR", "DASHING", "X70PLUS", "EC3", "Престиж", "Комфорт"):
        assert fragment not in text, f"в таблице автомобилей осталось «{fragment}»"


def test_cargo_header_is_bold_and_centered(template_doc, carrier_type):
    table = _cargo_table(template_doc)
    for cell in table.rows[0].cells:
        paragraph = cell.paragraphs[0]
        assert str(paragraph.alignment) == "CENTER (1)"
        assert all(run.bold for run in paragraph.runs), "шапка не полужирная"


def test_cargo_column_widths_match_sample(template_doc, carrier_type):
    """Ширины колонок таблицы автомобилей — как в образце (twips)."""
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    table = _cargo_table(template_doc)
    widths = [column.get(ns + "w") for column in table._tbl.tblGrid]
    assert widths == ["1000", "5091", "3932"]


# ─────────────────────────────────────────────────────────────
# Раздел 1 (погрузка) и раздел 2 (выгрузка)
# ─────────────────────────────────────────────────────────────

def test_shipper_name_appears_once(template_doc, carrier_type):
    """Грузоотправитель в бланке ОДИН: метка «Грузоотправитель:» одна."""
    texts = _body_texts(template_doc)
    labels = [t for t in texts if t.startswith("Грузоотправитель:")]

    assert len(labels) == 1, f"строк «Грузоотправитель:» — {len(labels)}"
    assert labels[0] == "Грузоотправитель: {{shipper_name}}"


def test_shipper_addresses_are_numbered(template_doc, carrier_type):
    """Адреса погрузки пронумерованы: «Адрес погрузки №N: {{shipper_N_address}}»."""
    texts = _body_texts(template_doc)
    for number in range(1, MAX_POINTS + 1):
        assert (f"Адрес погрузки №{number}: "
                f"{{{{shipper_{number}_address}}}}") in texts


def test_no_old_shipper_name_placeholders(template_doc, carrier_type):
    """Прежних {{shipper_N_name}} в бланке нет: грузоотправитель один."""
    text = _document_text(template_doc)
    leftovers = [
        f"{{{{shipper_{number}_name}}}}"
        for number in range(1, MAX_POINTS + 1)
        if f"{{{{shipper_{number}_name}}}}" in text
    ]
    assert not leftovers, f"в бланке остались старые плейсхолдеры: {leftovers}"
    assert "{{shipper_name}}" in text, "нет нового плейсхолдера грузоотправителя"


def test_consignee_blocks_are_still_numbered(template_doc, carrier_type):
    """Раздел 2 не изменился: 10 пронумерованных пар грузополучателей."""
    texts = _body_texts(template_doc)
    for number in range(1, MAX_POINTS + 1):
        assert f"Грузополучатель №{number}: {{{{consignee_{number}_name}}}}" in texts
        assert f"Адрес выгрузки: {{{{consignee_{number}_address}}}}" in texts


def test_date_lines_are_present_once_per_section(template_doc, carrier_type):
    """Дата/время погрузки и выгрузки — по одной строке на раздел."""
    texts = _body_texts(template_doc)
    loading = [t for t in texts if t.startswith("Дата / время погрузки:")]
    unloading = [t for t in texts
                 if t.startswith("Плановая дата / время завершения выгрузки:")]

    assert len(loading) == 1, f"строк погрузки: {len(loading)}"
    assert len(unloading) == 1, f"строк выгрузки: {len(unloading)}"

    assert "{{loading_date}}" in loading[0]
    assert "{{loading_time_from}}" in loading[0]
    assert "{{loading_time_to}}" in loading[0]
    assert "{{unloading_date}}" in unloading[0]
    assert "{{unloading_time_from}}" in unloading[0]
    assert "{{unloading_time_to}}" in unloading[0]


# ─────────────────────────────────────────────────────────────
# Вариант: экспедитор, договор, подписи
# ─────────────────────────────────────────────────────────────

def test_header_points_to_own_general_contract(templates, carrier_type):
    """В шапке — номер генерального договора своего варианта."""
    header_number, reference = CONTRACT_NUMBERS[carrier_type]
    other = "ИП" if carrier_type == "ООО" else "ООО"
    text = _document_text(templates[carrier_type])

    assert header_number in text
    assert reference in text, "нет ссылки на генеральный договор в разделе 6"
    assert CONTRACT_NUMBERS[other][0] not in text


def test_expeditor_of_variant_is_fixed(templates, carrier_type):
    """Экспедитор подставляется шаблоном — в промпте его не извлекаем."""
    text = _document_text(templates[carrier_type])

    if carrier_type == "ООО":
        assert "Экспедитор: ООО «ТЕХНОЛОГИСТИКА»" in text
        assert "________________ / Т.А. Ахмедов /" in text
        assert "Хейгетян" not in text, "в ООО-шаблоне чужой экспедитор"
    else:
        assert "Экспедитор: ИП Хейгетян Е.В." in text
        assert "ИП Хейгетян Елена Валентиновна" in text
        assert "________________ / Е.В.Хейгетян /" in text
        assert "ТЕХНОЛОГИСТИКА" not in text, "в ИП-шаблоне чужой экспедитор"


def test_customer_is_placeholder_with_fixed_signer(templates, carrier_type):
    """Заказчик вводится данными, подписант заказчика фиксирован."""
    text = _document_text(templates[carrier_type])
    assert "Заказчик: {{customer_name}}" in text
    assert "{{customer_name}}" in text.split("ЗАКАЗЧИК")[-1]
    assert "________________ / Гао Фанфан /" in text


# ─────────────────────────────────────────────────────────────
# Отсутствие данных образцов
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("fragment", FORBIDDEN_FRAGMENTS)
def test_template_has_no_sample_data(templates, carrier_type, fragment):
    assert fragment not in _document_text(templates[carrier_type]), (
        f"в шаблоне {carrier_type} найдены данные образца: «{fragment}»"
    )


# ─────────────────────────────────────────────────────────────
# Оформление
# ─────────────────────────────────────────────────────────────

def test_title_is_arial_bold_centered(template_doc, carrier_type):
    """«ЗАЯВКА № {{contract_number}}» — по центру, полужирным, Arial 12 pt."""
    title = next(
        p for p in template_doc.paragraphs
        if p.text.strip().startswith("ЗАЯВКА № ")
    )
    assert title.text.strip() == "ЗАЯВКА № {{contract_number}}"
    assert str(title.alignment) == "CENTER (1)"

    for run in title.runs:
        assert run.font.name == "Arial"
        assert run.font.size.pt == 12
        assert run.bold is True


def test_page_geometry_matches_sample(template_doc, carrier_type):
    """A4 и поля образца: 2,54 см со всех сторон."""
    section = template_doc.sections[0]
    assert round(section.page_width.cm, 2) == 21.01
    assert round(section.page_height.cm, 2) == 29.69
    assert round(section.left_margin.cm, 2) == 2.54
    assert round(section.right_margin.cm, 2) == 2.54
    assert round(section.top_margin.cm, 2) == 2.54
    assert round(section.bottom_margin.cm, 2) == 2.54


def test_body_font_is_arial(template_doc, carrier_type):
    """Основной текст — Arial 12 pt, ровно как в образцах."""
    normal = template_doc.styles["Normal"]
    assert normal.font.name == "Arial"
    assert normal.font.size.pt == 12

    paragraph = next(
        p for p in template_doc.paragraphs if p.text.startswith("Тягач:")
    )
    for run in paragraph.runs:
        assert run.font.name == "Arial"
        assert run.font.size.pt == 12


def test_all_runs_are_arial_twelve(template_doc, carrier_type):
    """Во всём бланке (абзацы и таблица) нет ни одного run другого шрифта."""
    wrong = []
    for paragraph in _all_paragraphs(template_doc):
        for run in paragraph.runs:
            if run.text and (run.font.name != "Arial" or run.font.size.pt != 12):
                wrong.append((run.text[:30], run.font.name,
                              run.font.size.pt if run.font.size else None))

    assert not wrong, f"runs не Arial 12 pt: {wrong[:5]}"


def test_sections_go_in_order(template_doc, carrier_type):
    """Разделы 1–6 и подписи идут в том же порядке, что в образце."""
    texts = _body_texts(template_doc)
    expected = [
        "Приложение № 1",
        "1. ПОГРУЗКА",
        "2. ВЫГРУЗКА",
        "3. ПЕРЕВОЗИМЫЕ АВТОМОБИЛИ",
        "4. АВТОВОЗ И ВОДИТЕЛЬ",
        "5. СТОИМОСТЬ",
        "6. ОСОБЫЕ УСЛОВИЯ",
        "ЗАКАЗЧИК",
        "ЭКСПЕДИТОР",
    ]
    positions = []
    for marker in expected:
        assert marker in texts, f"в шаблоне нет раздела «{marker}»"
        positions.append(texts.index(marker))

    assert positions == sorted(positions), (
        f"разделы идут не по порядку: {list(zip(expected, positions))}"
    )


# ─────────────────────────────────────────────────────────────
# Образцы не изменены и шаблон рендерится
# ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("carrier_type", CARRIER_TYPES)
def test_sample_documents_untouched(templates_dir, carrier_type):
    """
    Образцы — источник структуры, только чтение.

    Если этот тест упал, значит образец отредактировали (или пересобрали).
    Вернуть файл из git:
        git checkout -- "templates/Приложение № 1 Заявка_600_ООО …docx"
    """
    name, expected_digest = SAMPLES[carrier_type]
    sample = templates_dir / name
    assert sample.exists(), f"пропал образец Логистикс Рус: {sample}"

    digest = hashlib.sha256(sample.read_bytes()).hexdigest()
    assert digest == expected_digest, (
        f"образец {name!r} изменён: ожидался {expected_digest}, получен {digest}"
    )


@pytest.mark.parametrize("carrier_type", CARRIER_TYPES)
def test_templates_match_pinned_sha(templates_dir, carrier_type):
    """
    Шаблон на диске — ровно то, что собрал сборщик (ЭТАП 3.1.C.A.1-fix).

    Хеши зафиксированы после пересборки под Arial и суммы прописью. Если
    этот тест упал после осознанной правки шаблона — пересобери его
    (python tools/make_logistiks_rus_template.py) и обнови TEMPLATE_SHA256;
    если правки не было — .docx кто-то отредактировал вручную.
    """
    path = templates_dir / TEMPLATE_NAMES[carrier_type]
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    assert digest == TEMPLATE_SHA256[carrier_type], (
        f"шаблон {TEMPLATE_NAMES[carrier_type]!r} не совпадает со сборкой: "
        f"ожидался {TEMPLATE_SHA256[carrier_type]}, получен {digest}"
    )


def test_docxtpl_renders_without_leftovers(template_path, carrier_type, work_file):
    """Шаблон совместим с docxtpl: после рендера плейсхолдеров не остаётся."""
    docxtpl = pytest.importorskip("docxtpl")

    text_before = _document_text(Document(str(template_path)))
    values = {
        name.strip("{} "): "ЗНАЧЕНИЕ"
        for name in PLACEHOLDER_RE.findall(text_before)
    }
    assert set(values) == _expected_placeholders(carrier_type)

    template = docxtpl.DocxTemplate(str(template_path))
    template.render(values, autoescape=True)

    output = work_file(f"logistiks_rus_{carrier_type}_rendered.docx")
    template.save(str(output))

    text_after = _document_text(Document(str(output)))
    assert "{{" not in text_after
    assert "}}" not in text_after
    assert "ЗНАЧЕНИЕ" in text_after


def test_builder_reproduces_same_document(templates_dir, carrier_type, work_dir,
                                          monkeypatch):
    """
    Сборщик и лежащий в templates/ шаблон не разъехались.

    Байты сравнивать нельзя: python-docx пишет в zip текущее время. Поэтому
    шаблон пересобирается во временную папку и сравнивается текст документа.
    """
    import tools.make_logistiks_rus_template as builder

    monkeypatch.setattr(builder, "TEMPLATES_DIR", work_dir)
    rebuilt = builder.build_template(carrier_type)

    expected = Document(str(templates_dir / TEMPLATE_NAMES[carrier_type]))
    actual = Document(str(rebuilt))

    assert _document_text(actual) == _document_text(expected)
    assert len(actual.paragraphs) == len(expected.paragraphs)
    assert len(actual.tables[0].rows) == len(expected.tables[0].rows)
