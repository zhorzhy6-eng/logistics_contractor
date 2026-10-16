#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Форматирование DOCX: копирование стилей, обход абзацев, ширины колонок.

Сюда переехали методы `BaseContractGenerator`, которые работают с XML
python-docx и ничего не знают о типе договора (Этап 4, 16.10.2026).

Класс `Formatting` — не объект с состоянием: все методы статические, класс
здесь только как пространство имён (как `docx.shared`). Экземпляр
`self.formatting` в генераторе создаётся ради единообразия с `renderer`
и `tables`.

Логика методов не менялась — перенос из `core/contracts/base_generator.py`
один в один.
"""

import re
from copy import deepcopy
from typing import Any, Dict

#: Пространство имён WordprocessingML.
NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

#: Полное имя тега разрыва строки Word — по нему находим абзацы, которые
#: косметика пробелов не трогает (см. `postprocess.normalize_document_spaces`).
W_BR_TAG = NS + "br"


class Formatting:
    """Утилиты оформления документа: теги, стили, абзацы, ячейки, колонки."""

    # ─────────────────────────────────────────────────────────
    # Имена элементов и однострочные значения
    # ─────────────────────────────────────────────────────────

    @staticmethod
    def qname(tag: str) -> str:
        """Имя элемента WordprocessingML с пространством имён."""
        return NS + tag

    @staticmethod
    def single_line(value: Any) -> str:
        """
        Значение одной строкой: переносы строк, табуляции и лишние пробелы —
        в один обычный пробел.
        """
        return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()

    @staticmethod
    def normalized_point_address(point: Dict[str, Any]) -> str:
        """Адрес точки одной строкой: переносы строк в заголовке не нужны."""
        return Formatting.single_line(point.get("address", "") or "")

    # ─────────────────────────────────────────────────────────
    # Копирование форматирования (без новых зависимостей)
    # ─────────────────────────────────────────────────────────

    @staticmethod
    def copy_paragraph_format(sample, paragraph) -> None:
        """
        Копирует форматирование абзаца-образца (если образец найден).

        Существующий pPr заменяется: двух pPr в одном абзаце быть не должно,
        иначе Word берёт первый и наш выравнивание/отступы теряются.
        """
        if sample is None:
            return

        props = sample._p.find(Formatting.qname("pPr"))
        if props is None:
            return

        existing = paragraph._p.find(Formatting.qname("pPr"))
        if existing is not None:
            paragraph._p.remove(existing)

        paragraph._p.insert(0, deepcopy(props))

    @staticmethod
    def copy_run_format(sample_run, run) -> None:
        """Копирует шрифт (имя, размер, начертание) из образца."""
        if sample_run is None:
            return

        props = sample_run.find(Formatting.qname("rPr"))
        if props is None:
            return

        run._r.insert(0, deepcopy(props))

    @staticmethod
    def apply_sample_cell_format(cell, sample_cell) -> None:
        """
        Копирует tcPr образца (заливка, границы, отступы) в нашу ячейку.

        Ширина и объединения не переносятся: у наших колонок свои размеры,
        а gridSpan/vMerge из образца сломали бы сетку таблицы.
        """
        sample_props = sample_cell._tc.find(Formatting.qname("tcPr"))
        if sample_props is None:
            return

        copied = deepcopy(sample_props)
        for tag in ("tcW", "gridSpan", "vMerge", "hMerge"):
            node = copied.find(Formatting.qname(tag))
            if node is not None:
                copied.remove(node)

        existing = cell._tc.find(Formatting.qname("tcPr"))
        if existing is not None:
            cell._tc.remove(existing)
        cell._tc.insert(0, copied)

    @staticmethod
    def set_cell_text(paragraph, text: str, sample_run, bold: bool = False) -> None:
        """Пишет текст в абзац ячейки, копируя шрифт образца."""
        run = paragraph.add_run(text)
        Formatting.copy_run_format(sample_run, run)
        run.bold = bold
        if run.font.size is None and sample_run is None:
            run.font.name = "Times New Roman"

    @staticmethod
    def sample_cell(sample_cells, index: int):
        """Ячейка-образец с нужным номером колонки (или None)."""
        if sample_cells and index < len(sample_cells):
            return sample_cells[index]
        return None

    @staticmethod
    def fill_table_cell(cell, text: str, sample_cell, bold: bool) -> None:
        """
        Пишет текст в ячейку, повторяя оформление ячейки-образца.

        Из образца переносятся границы, заливка и отступы (tcPr без ширины)
        плюс формат абзаца и шрифта. Ширину колонки задаёт генератор.
        """
        if sample_cell is not None:
            Formatting.apply_sample_cell_format(cell, sample_cell)

        paragraph = cell.paragraphs[0]
        sample_paragraph = (
            sample_cell.paragraphs[0]
            if (sample_cell is not None and sample_cell.paragraphs)
            else None
        )
        # Сначала формат абзаца из образца (интервалы, отступы), потом своё
        # выравнивание: иначе копия pPr перетёрла бы центр.
        Formatting.copy_paragraph_format(sample_paragraph, paragraph)
        paragraph.alignment = 1  # WD_ALIGN_PARAGRAPH.CENTER

        sample_run = (
            sample_paragraph.runs[0]._r
            if (sample_paragraph is not None and sample_paragraph.runs)
            else None
        )
        Formatting.set_cell_text(paragraph, text, sample_run, bold=bold)

    @staticmethod
    def set_grid_widths(table, widths_cm) -> None:
        """
        Фиксирует ширину колонок таблицы: tblGrid + tblW + tblLayout.

        Одного cell.width мало: Word считает колонки по tblGrid, а его
        python-docx при создании таблицы делает равномерным.
        """
        from docx.shared import Cm, Emu

        emu_widths = [Cm(width) for width in widths_cm]
        # Cm() возвращает Length (подкласс int), а sum() даёт обычный int —
        # приводим обратно, чтобы пользоваться .twips.
        total = Emu(sum(emu_widths))

        props = table._tbl.tblPr
        layout = props.get_or_add_tblLayout()
        layout.type = "fixed"

        tbl_width = props.find(Formatting.qname("tblW"))
        if tbl_width is None:
            tbl_width = props.makeelement(Formatting.qname("tblW"), {})
            props.insert(0, tbl_width)
        tbl_width.set(Formatting.qname("type"), "dxa")
        tbl_width.set(Formatting.qname("w"), str(int(total.twips)))

        for grid_col, width in zip(table._tbl.tblGrid, emu_widths):
            grid_col.set(Formatting.qname("w"), str(int(Emu(width).twips)))

        for row in table.rows:
            for cell, width in zip(row.cells, emu_widths):
                cell.width = width

    # ─────────────────────────────────────────────────────────
    # Обход документа
    # ─────────────────────────────────────────────────────────

    @staticmethod
    def iter_paragraphs(doc):
        """Все абзацы документа, включая таблицы (в т.ч. вложенные) и колонтитулы."""
        for paragraph in doc.paragraphs:
            yield paragraph

        for table in doc.tables:
            yield from Formatting.iter_table_paragraphs(table)

        for section in doc.sections:
            for paragraph in section.header.paragraphs:
                yield paragraph
            for paragraph in section.footer.paragraphs:
                yield paragraph

    @staticmethod
    def iter_table_paragraphs(table):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
                for nested in cell.tables:
                    yield from Formatting.iter_table_paragraphs(nested)
